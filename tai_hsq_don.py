# -*- coding: utf-8 -*-
"""
Tool dòng lệnh (chưa có UI): kiểm tra đơn có những hồ sơ quét nào và tải toàn bộ về.

Luồng:
  1. Mở Chrome (giống dohs_1.py: có thư mục tải + bật log mạng), đăng nhập MPLIS.
  2. Với mỗi tinhHinhDangKyId: gọi GetThongTinDangKyByTinhHinhDangKyIds (getHoSoQuet=True)
     → liệt kê ListHoSoQuet / ListFileHoSoQuet.
  3. Tải từng file ngay trong trình duyệt (fetch bằng session của Chrome)
     về <thư mục ra>/<tinhHinhDangKyId>/.

Endpoint tải file: /dc/Handlers/FileHandler.ashx?DocId=<nodeId>&MimeType=application/pdf
Nếu MPLIS đổi endpoint, chạy --hoc-lai:
tool bắt request khi bạn mở xem 1 file (Chrome) hoặc cho dán URL từ F12 (token/cookie)
→ suy ra mẫu URL → lưu vào hsq_url_mau.json để các lần chạy sau dùng lại.

Cách dùng:
  python tai_hsq_don.py 13456477 12913166          # kiểm tra + tải
  python tai_hsq_don.py 13456477 --chi-kiem-tra     # chỉ liệt kê, không tải
  python tai_hsq_don.py --excel ds_don.xlsx         # lấy ID đơn từ cột A (từ dòng 2)
  python tai_hsq_don.py 13456477 --hoc-lai          # bỏ mẫu URL đã lưu, học lại
  python tai_hsq_don.py 13456477 --token XXX --cookie "a=1; b=2"   # khỏi đăng nhập
  python tai_hsq_don.py 13456477 --chrome           # đăng nhập bằng Chrome luôn

Khi chạy không có --token/--cookie/--chrome, tool hỏi chọn: dán token + cookie hoặc mở Chrome.
Chế độ token/cookie không có Chrome để bắt request, nên nếu chưa có hsq_url_mau.json
tool sẽ yêu cầu dán URL xem file (copy từ F12) để suy ra mẫu.
"""

import argparse
import base64
import getpass
import json
import os
import re
import sys
import time
from typing import Any
from urllib.parse import parse_qs, unquote, urlparse

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from webdriver_manager.chrome import ChromeDriverManager

import cap_nhat_don_hsq as core

# Chạy dạng .exe (PyInstaller): __file__ nằm trong thư mục giải nén tạm → lưu cạnh file .exe.
THU_MUC_APP = (os.path.dirname(sys.executable) if getattr(sys, "frozen", False)
               else os.path.dirname(os.path.abspath(__file__)))
FILE_MAU_URL = os.path.join(THU_MUC_APP, "hsq_url_mau.json")
KHOA_ID_FILE = ("nodeId", "thanhPhanHoSoQuetNId", "thanhPhanHoSoQuetId")
THOI_GIAN_CHO_HOC = 300  # giây chờ người dùng mở xem 1 file

# Endpoint tải file HSQ của MPLIS (đã xác nhận: DocId = nodeId của file).
MAU_MAC_DINH = [
    {
        "method": "GET",
        "url": core.BASE_URL + "/dc/Handlers/FileHandler.ashx?DocId={nodeId}&MimeType=application/pdf",
        "body": "",
        "content_type": "",
    }
]


_log_handler = None  # UI (Flet) gán hàm nhận log vào đây; None = in ra console.


def dat_ham_log(fn) -> None:
    global _log_handler
    _log_handler = fn


def log(msg: str) -> None:
    if _log_handler is not None:
        _log_handler(msg)
        return
    print(f"{time.strftime('%H:%M:%S')}  {msg}", flush=True)


def tach_danh_sach_id(raw_values) -> list[int]:
    """Nhận list chuỗi (hoặc 1 chuỗi nhiều dòng/dấu phẩy) → list tinhHinhDangKyId không trùng."""
    if isinstance(raw_values, str):
        raw_values = re.split(r"[\r\n,;\t ]+", raw_values)
    ids: list[int] = []
    for raw in raw_values:
        if not str(raw or "").strip():
            continue
        id_text, _ = core.tach_thong_tin_tu_chuoi_madon(raw)
        if id_text and int(id_text) not in ids:
            ids.append(int(id_text))
        elif not id_text:
            log(f"⚠ Bỏ qua giá trị không nhận ra mã đơn: {raw!r}")
    return ids


def ten_file_an_toan(s: str, limit: int = 120) -> str:
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", str(s or "")).strip(" .")
    return s[:limit] or "file"


# ============================ CHROME (theo dohs_1.py) ============================

def mo_chrome(download_dir: str) -> webdriver.Chrome:
    options = Options()
    options.add_argument("--start-maximized")
    options.add_experimental_option(
        "prefs",
        {
            "download.default_directory": download_dir,
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "plugins.always_open_pdf_externally": True,
        },
    )
    # Bật log mạng để bắt request tải file khi người dùng mở xem hồ sơ quét.
    options.set_capability("goog:loggingPrefs", {"performance": "ALL"})

    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)
    try:
        driver.execute_cdp_cmd("Page.setDownloadBehavior", {"behavior": "allow", "downloadPath": download_dir})
    except Exception:
        pass
    driver.set_script_timeout(core.TIMEOUT)
    return driver


def dang_nhap(driver: webdriver.Chrome, username: str, password: str) -> None:
    driver.get(core.REFERER_LOGIN)
    time.sleep(2)
    try:
        user_box = pass_box = None
        for inp in driver.find_elements(By.CSS_SELECTOR, "input"):
            typ = (inp.get_attribute("type") or "").lower()
            if user_box is None and typ in {"text", "email"}:
                user_box = inp
            if pass_box is None and typ == "password":
                pass_box = inp
        if user_box and pass_box:
            user_box.clear()
            user_box.send_keys(username)
            pass_box.clear()
            pass_box.send_keys(password)
            pass_box.send_keys(Keys.ENTER)
            log("Đã điền thông tin đăng nhập, chờ trang tải...")
        else:
            log("Không nhận dạng được form đăng nhập, hãy đăng nhập tay trên Chrome.")
    except Exception as exc:
        log(f"Không tự điền được form đăng nhập ({exc}), hãy đăng nhập tay.")


# ============================ ĐỌC HỒ SƠ QUÉT ============================

def liet_ke_file_hsq(detail: dict[str, Any]) -> list[dict[str, Any]]:
    """Trả về danh sách phẳng các file HSQ của đơn, kèm hoSoQuetId cha (bỏ trùng nodeId)."""
    ket_qua: list[dict[str, Any]] = []
    da_co: set[str] = set()
    for hoso in detail.get("ListHoSoQuet") or []:
        if not isinstance(hoso, dict):
            continue
        wrapper = hoso.get("ListFileHoSoQuet") or {}
        files = wrapper.get("ListFileHoSoQuet") if isinstance(wrapper, dict) else wrapper
        for vi_tri, f in enumerate(files or [], 1):
            if not isinstance(f, dict):
                continue
            node_id = str(f.get("nodeId") or "")
            if node_id and node_id in da_co:
                continue
            da_co.add(node_id)
            item = dict(f)
            item.setdefault("hoSoQuetId", hoso.get("hoSoQuetId"))
            item["_vi_tri"] = vi_tri  # vị trí file trong HSQ (chỉ để ghi Excel, không gửi lên)
            ket_qua.append(item)
    return ket_qua


def node_goc(f: dict[str, Any]) -> dict[str, Any]:
    """Bỏ các khóa nội bộ (bắt đầu bằng '_') → đúng node MPLIS trả về."""
    return {k: v for k, v in f.items() if not k.startswith("_")}


# ============================ EXCEL KẾT QUẢ ============================

# Excel kết quả tải = đầu vào cho bước đẩy file ký số (day_hsq_ky_so.py).
# Mỗi dòng = 1 file HSQ. node_json lưu nguyên node gốc để đẩy lại y hệt, chỉ thay file.
EXCEL_HEADERS = [
    "stt", "madon", "hosoquetid", "vitri", "thanhphanhosoquetid", "thanhphanhosoquetnid",
    "nodeid", "mota", "loaihosoquet", "giaychungnhanid", "file_tai", "trang_thai_tai",
    "ghi_chu_tai", "file_ky_so", "trang_thai_day", "ghi_chu_day", "thoi_gian_day",
    "nodeid_moi", "node_json",
]
EXCEL_WIDTHS = {
    "stt": 6, "madon": 12, "hosoquetid": 11, "vitri": 6, "thanhphanhosoquetid": 13,
    "thanhphanhosoquetnid": 20, "nodeid": 20, "mota": 40, "loaihosoquet": 16,
    "giaychungnhanid": 12, "file_tai": 50, "trang_thai_tai": 14, "ghi_chu_tai": 30,
    "file_ky_so": 40, "trang_thai_day": 16, "ghi_chu_day": 40, "thoi_gian_day": 18,
    "nodeid_moi": 20, "node_json": 12,
}


def dong_excel_tu_file(id_don: int, f: dict[str, Any], file_tai: str, trang_thai: str, ghi_chu: str) -> dict:
    loai = core.safe_int(f.get("loaiHoSoQuet"), -1)
    return {
        "madon": id_don,
        "hosoquetid": f.get("hoSoQuetId"),
        "vitri": f.get("_vi_tri"),
        "thanhphanhosoquetid": f.get("thanhPhanHoSoQuetId"),
        "thanhphanhosoquetnid": f.get("thanhPhanHoSoQuetNId"),
        "nodeid": f.get("nodeId"),
        "mota": f.get("moTa") or f.get("tenGiayTo") or "",
        "loaihosoquet": core.LOAI_HO_SO_QUET_OPTIONS.get(loai, f.get("loaiHoSoQuet")),
        "giaychungnhanid": f.get("giayChungNhanId"),
        "file_tai": file_tai,
        "trang_thai_tai": trang_thai,
        "ghi_chu_tai": ghi_chu,
        "node_json": json.dumps(node_goc(f), ensure_ascii=False),
    }


def ghi_excel_ket_qua(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    mau_trang_thai = {"Đã tải": "C6EFCE", "Lỗi": "F4CCCC", "Không có HSQ": "FFF2CC"}
    col_tt = EXCEL_HEADERS.index("trang_thai_tai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in EXCEL_HEADERS])
        mau = mau_trang_thai.get(str(row.get("trang_thai_tai") or ""))
        if mau:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=mau)

    for i, h in enumerate(EXCEL_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.column_dimensions[get_column_letter(EXCEL_HEADERS.index("node_json") + 1)].hidden = True
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}1"

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def in_danh_sach(id_don: int, detail: dict[str, Any], files: list[dict[str, Any]]) -> None:
    thdk = detail.get("TinhHinhDangKy") or {}
    log(
        f"Đơn {id_don} | tờ {thdk.get('soHieuToBanDo', '')} thửa {thdk.get('soThuTuThua', '')} "
        f"| xã {thdk.get('xaId', '')} | {len(detail.get('ListHoSoQuet') or [])} HSQ, {len(files)} file"
    )
    for i, f in enumerate(files, 1):
        loai = core.LOAI_HO_SO_QUET_OPTIONS.get(core.safe_int(f.get("loaiHoSoQuet"), -1), f.get("loaiHoSoQuet"))
        log(
            f"   {i:>2}. [{loai}] {f.get('moTa') or f.get('tenGiayTo') or '(không mô tả)'} "
            f"| hoSoQuetId={f.get('hoSoQuetId')} | nodeId={f.get('nodeId') or '(trống)'}"
        )


# ============================ HỌC ENDPOINT TẢI FILE ============================

def _thay_id_bang_placeholder(text: str, files: list[dict[str, Any]]) -> str | None:
    """Nếu text chứa id của 1 file thì thay id đó bằng {khoa}. Không khớp → None."""
    for f in files:
        for khoa in KHOA_ID_FILE:
            gia_tri = str(f.get(khoa) or "").strip()
            if len(gia_tri) >= 6 and gia_tri in text:
                return text.replace(gia_tri, "{" + khoa + "}")
    return None


def _cac_url_ung_vien(url: str) -> list[str]:
    """URL gốc + URL file lồng trong viewer pdf.js (viewer.html?file=...)."""
    ds = [url]
    for v in parse_qs(urlparse(url).query).get("file", []):
        inner = unquote(v)
        if inner.startswith("/"):
            inner = core.BASE_URL + inner
        ds.insert(0, inner)
    return ds


def tao_mau(url: str, method: str, post: str, ctype: str, files: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Từ 1 request (url/body) chứa id file → mẫu tải file. Không khớp id nào → None."""
    for url_ung_vien in _cac_url_ung_vien(url):
        la_url_goc = url_ung_vien == url
        url_mau = _thay_id_bang_placeholder(url_ung_vien, files)
        post_mau = _thay_id_bang_placeholder(post, files) if post and la_url_goc else None
        if url_mau is None and post_mau is None:
            continue
        return {
            "method": method if la_url_goc else "GET",
            "url": url_mau or url_ung_vien,
            "body": (post_mau or post) if la_url_goc else "",
            "content_type": ctype if la_url_goc else "",
        }
    return None


def hoc_mau_tu_url_dan(files: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Chế độ token/cookie (không có Chrome): người dùng dán URL xem file copy từ F12."""
    log("🔎 Chưa biết endpoint tải file hồ sơ quét.")
    log("   → Trên MPLIS mở F12 > Network, bấm XEM 1 file hồ sơ quét của đơn này,")
    log("     chuột phải request tải file > Copy > Copy URL rồi dán vào đây.")
    while True:
        url = input("URL (Enter trống để bỏ qua): ").strip()
        if not url:
            return None
        post = input("Request là POST? Dán Payload (Enter trống nếu là GET): ").strip()
        mau = tao_mau(
            url, "POST" if post else "GET", post,
            "application/x-www-form-urlencoded; charset=UTF-8" if post and not post.startswith("{")
            else ("application/json; charset=UTF-8" if post else ""),
            files,
        )
        if mau:
            log(f"✅ Mẫu tải file: {mau['method']} {mau['url']}")
            return mau
        log("❌ URL/payload không chứa nodeId hay id file nào của đơn này, thử lại.")


def hoc_mau_url(driver: webdriver.Chrome, files: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Chờ người dùng mở xem 1 file HSQ trên Chrome, bắt request mạng chứa id file."""
    driver.get_log("performance")  # xả log cũ
    log("🔎 Chưa biết endpoint tải file hồ sơ quét.")
    log("   → Trên Chrome, mở đơn đang xử lý và bấm XEM 1 file hồ sơ quét bất kỳ.")
    log(f"   → Tool đang theo dõi request mạng (tối đa {THOI_GIAN_CHO_HOC}s)...")

    het_han = time.time() + THOI_GIAN_CHO_HOC
    while time.time() < het_han:
        for entry in driver.get_log("performance"):
            try:
                msg = json.loads(entry["message"])["message"]
            except Exception:
                continue
            if msg.get("method") != "Network.requestWillBeSent":
                continue
            req = msg.get("params", {}).get("request", {})
            url = req.get("url") or ""
            if not url.startswith("http") or "/GetThongTinDangKy" in url:
                continue
            mau = tao_mau(
                url,
                (req.get("method") or "GET").upper(),
                req.get("postData") or "",
                (req.get("headers") or {}).get("Content-Type", ""),
                files,
            )
            if mau:
                log(f"✅ Đã bắt được endpoint: {mau['method']} {mau['url']}")
                if mau["body"]:
                    log(f"   body: {mau['body'][:200]}")
                return mau
        time.sleep(0.5)

    return None


def doc_mau_da_luu() -> dict[str, Any] | None:
    try:
        with open(FILE_MAU_URL, encoding="utf-8") as fp:
            return json.load(fp)
    except Exception:
        return None


def luu_mau(mau: dict[str, Any]) -> None:
    with open(FILE_MAU_URL, "w", encoding="utf-8") as fp:
        json.dump(mau, fp, ensure_ascii=False, indent=2)
    log(f"   Đã lưu mẫu vào {FILE_MAU_URL}")


# ============================ TẢI FILE TRONG TRÌNH DUYỆT ============================

JS_FETCH = """
const [url, method, body, ctype, token, done] = arguments;
const headers = {'X-Requested-With': 'XMLHttpRequest'};
if (token) headers['__requestverificationtoken'] = token;
if (body && ctype) headers['Content-Type'] = ctype;
fetch(url, {method, body: body || undefined, headers, credentials: 'include'})
  .then(async r => {
    const b = new Uint8Array(await r.arrayBuffer());
    let s = '';
    for (let i = 0; i < b.length; i += 0x8000) s += String.fromCharCode.apply(null, b.subarray(i, i + 0x8000));
    done({status: r.status, ctype: r.headers.get('content-type') || '',
          cd: r.headers.get('content-disposition') || '', b64: btoa(s)});
  })
  .catch(e => done({error: String(e)}));
"""


class Downloader:
    """Tải trong Chrome (driver) hoặc bằng requests (session từ token/cookie dán tay)."""

    def __init__(
        self,
        driver: webdriver.Chrome | None,
        session,
        token: str,
        mau: dict[str, Any] | None,
        ung_vien: list[dict[str, Any]] | None = None,
    ):
        self.driver = driver
        self.session = session
        self.token = token
        self.mau = mau
        # Các mẫu thử thêm khi self.mau chưa được xác nhận; mẫu nào tải được thì chốt luôn.
        self.ung_vien = list(ung_vien or [])

    def hoc_mau(self, files: list[dict[str, Any]]) -> dict[str, Any] | None:
        if self.driver is not None:
            return hoc_mau_url(self.driver, files)
        return hoc_mau_tu_url_dan(files)

    def _gui(self, url: str, method: str, body: str, ctype: str) -> dict[str, Any]:
        if self.driver is not None:
            res = self.driver.execute_async_script(JS_FETCH, url, method, body, ctype, self.token)
            if not res or res.get("error"):
                raise RuntimeError(f"fetch lỗi: {res and res.get('error')}")
            return {**res, "data": base64.b64decode(res.get("b64") or "")}

        headers = {"Accept": "*/*"}
        if body and ctype:
            headers["Content-Type"] = ctype
        r = self.session.request(
            method, url, data=body.encode("utf-8") if body else None,
            headers=headers, timeout=core.TIMEOUT, allow_redirects=False,
        )
        return {
            "status": r.status_code,
            "ctype": r.headers.get("Content-Type", ""),
            "cd": r.headers.get("Content-Disposition", ""),
            "data": r.content,
        }

    def tai(self, f: dict[str, Any]) -> tuple[bytes, str]:
        """Trả về (bytes, tên file từ server hoặc ''). Thử mẫu hiện tại rồi tới các ứng viên."""
        ds_mau = ([self.mau] if self.mau else []) + [m for m in self.ung_vien if m != self.mau]
        if not ds_mau:
            raise RuntimeError("Chưa có mẫu URL tải file")

        loi_cuoi: Exception | None = None
        for mau in ds_mau:
            try:
                ket_qua = self._tai_theo_mau(mau, f)
            except Exception as exc:
                loi_cuoi = exc
                continue
            if self.ung_vien:
                log(f"   ℹ Dùng endpoint: {mau['url']}")
                self.ung_vien = []
            self.mau = mau
            return ket_qua
        raise loi_cuoi  # type: ignore[misc]

    def _tai_theo_mau(self, mau: dict[str, Any], f: dict[str, Any]) -> tuple[bytes, str]:
        def dien(text: str) -> str:
            # Không dùng str.format vì body JSON có sẵn dấu { }.
            for k in KHOA_ID_FILE:
                text = text.replace("{" + k + "}", str(f.get(k) or ""))
            return text

        url = dien(mau["url"])
        body = dien(mau.get("body") or "")

        res = self._gui(url, mau.get("method") or "GET", body, mau.get("content_type") or "")
        data = res["data"]
        ctype = (res.get("ctype") or "").lower()
        if res.get("status") in {301, 302, 303, 307, 308}:
            raise RuntimeError(f"Bị chuyển hướng (HTTP {res.get('status')}) — token/cookie có thể đã hết hạn.")
        if res.get("status") != 200 or not data or "html" in ctype:
            raise RuntimeError(f"HTTP {res.get('status')} {ctype}: {data[:150]!r}")

        # Có hệ thống trả JSON chứa base64 của file.
        if "json" in ctype:
            data = self._lay_file_tu_json(data)

        ten_server = ""
        m = re.search(r"filename\*?=(?:UTF-8'')?\"?([^\";]+)", res.get("cd") or "", re.IGNORECASE)
        if m:
            ten_server = unquote(m.group(1))
        return data, ten_server

    @staticmethod
    def _lay_file_tu_json(raw: bytes) -> bytes:
        js = json.loads(raw.decode("utf-8", errors="replace"))

        def tim(o):
            if isinstance(o, str) and len(o) > 200:
                try:
                    b = base64.b64decode(o.split(",", 1)[-1], validate=True)
                    if b[:4] == b"%PDF" or b[:3] == b"\xff\xd8\xff" or b[:4] in (b"II*\x00", b"MM\x00*"):
                        return b
                except Exception:
                    return None
            if isinstance(o, dict):
                o = list(o.values())
            if isinstance(o, list):
                for v in o:
                    r = tim(v)
                    if r:
                        return r
            return None

        data = tim(js)
        if not data:
            raise RuntimeError(f"API trả JSON không chứa file: {raw[:200]!r}")
        return data


def doan_duoi_file(data: bytes, ten_server: str) -> str:
    if ten_server and "." in ten_server:
        return os.path.splitext(ten_server)[1]
    if data[:4] == b"%PDF":
        return ".pdf"
    if data[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if data[:4] in (b"II*\x00", b"MM\x00*"):
        return ".tif"
    return ".bin"


# ============================ LUỒNG CHÍNH ============================

def xu_ly_don(
    client,
    downloader: Downloader,
    id_don: int,
    out_dir: str,
    chi_kiem_tra: bool,
    nen_dung=lambda: False,
) -> dict[str, Any]:
    # kq["dong"]: các dòng Excel kết quả (mỗi file HSQ 1 dòng; đơn lỗi/không HSQ cũng có 1 dòng).
    kq = {"id_don": id_don, "so_file": 0, "da_tai": 0, "loi": [], "dong": []}

    detail = client.lay_chi_tiet_don(id_don)
    if detail is None:
        kq["loi"].append("Không lấy được chi tiết đơn")
        kq["dong"].append({"madon": id_don, "trang_thai_tai": "Lỗi", "ghi_chu_tai": "Không lấy được chi tiết đơn"})
        log(f"❌ Đơn {id_don}: không lấy được chi tiết")
        return kq

    files = liet_ke_file_hsq(detail)
    kq["so_file"] = len(files)
    in_danh_sach(id_don, detail, files)

    if not files:
        kq["dong"].append({"madon": id_don, "trang_thai_tai": "Không có HSQ"})
        return kq
    if chi_kiem_tra:
        kq["dong"] = [dong_excel_tu_file(id_don, f, "", "Chưa tải", "") for f in files]
        return kq

    if downloader.mau is None and not downloader.ung_vien:
        mau = downloader.hoc_mau(files)
        if mau is None:
            raise RuntimeError("Chưa có mẫu URL tải file hồ sơ quét (hết thời gian chờ hoặc bỏ qua).")
        downloader.mau = mau

    thu_muc_don = os.path.join(out_dir, str(id_don))
    os.makedirs(thu_muc_don, exist_ok=True)

    with open(os.path.join(thu_muc_don, "hosoquet.json"), "w", encoding="utf-8") as fp:
        json.dump(detail.get("ListHoSoQuet"), fp, ensure_ascii=False, indent=2, default=str)

    for i, f in enumerate(files, 1):
        if nen_dung():
            log("⏹ Đã dừng theo yêu cầu.")
            break
        mo_ta = f.get("moTa") or f.get("tenGiayTo") or f"file_{i}"
        try:
            data, ten_server = downloader.tai(f)
        except Exception as exc:
            kq["loi"].append(f"{mo_ta}: {exc}")
            kq["dong"].append(dong_excel_tu_file(id_don, f, "", "Lỗi", core.rut_gon_text(exc, 300)))
            log(f"   ❌ {i}. {mo_ta}: {core.rut_gon_text(exc, 300)}")
            continue

        ten = f"{i:02d}_{ten_file_an_toan(mo_ta)}{doan_duoi_file(data, ten_server)}"
        duong_dan = os.path.join(thu_muc_don, ten)
        with open(duong_dan, "wb") as fp:
            fp.write(data)
        kq["da_tai"] += 1
        kq["dong"].append(dong_excel_tu_file(id_don, f, duong_dan, "Đã tải", ""))
        log(f"   ✅ {i}. {ten} ({len(data) / 1024:.0f} KB)")
        time.sleep(core.REQUEST_DELAY_SECONDS)

    return kq


def doc_id_tu_excel(path: str) -> list[str]:
    """Lấy ID đơn từ cột 'madon' (hoặc 'tinhhinhdangkyid'...), không có header thì lấy cột A."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(v or "").strip().lower() for v in rows[0]]
    col = next((i for i, h in enumerate(headers) if h in core.ALIASES_ID_DON), None)
    if col is None:
        col = 0
        log("⚠ Excel không có cột 'madon' → lấy cột A.")
    return [str(r[col]) for r in rows[1:] if len(r) > col and r[col] not in (None, "")]


def ten_excel_ket_qua(out_dir: str) -> str:
    return os.path.join(out_dir, f"ket_qua_tai_hsq_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="Kiểm tra và tải toàn bộ hồ sơ quét của đơn MPLIS")
    ap.add_argument("ids", nargs="*", help="tinhHinhDangKyId (hoặc chuỗi chứa mã đơn)")
    ap.add_argument("--excel", help="File Excel có cột madon (tinhHinhDangKyId)")
    ap.add_argument("--out", default="hsq_tai_ve", help="Thư mục lưu (mặc định: hsq_tai_ve)")
    ap.add_argument("--chi-kiem-tra", action="store_true", help="Chỉ liệt kê HSQ, không tải")
    ap.add_argument("--hoc-lai", action="store_true", help="Bỏ mẫu URL đã lưu, học lại từ Chrome")
    ap.add_argument("--user", help="Username đăng nhập MPLIS")
    ap.add_argument("--token", help="__RequestVerificationToken copy từ F12 (bỏ qua đăng nhập)")
    ap.add_argument("--cookie", help="Chuỗi Cookie copy từ F12 (dùng cùng --token)")
    ap.add_argument("--chrome", action="store_true", help="Đăng nhập bằng Chrome, không hỏi token/cookie")
    args = ap.parse_args()

    raw_ids = list(args.ids)
    if args.excel:
        raw_ids += doc_id_tu_excel(args.excel)

    ids = tach_danh_sach_id(raw_ids)
    if not ids:
        ap.error("Chưa có ID đơn nào (truyền ids hoặc --excel).")

    out_dir = os.path.abspath(args.out)
    os.makedirs(out_dir, exist_ok=True)

    token, cookie = (args.token or "").strip(), (args.cookie or "").strip()
    if not args.chrome and not (token and cookie):
        print("Chọn cách đăng nhập:")
        print("  1. Dán token + cookie (copy từ F12 > Network > Request Headers)")
        print("  2. Mở Chrome đăng nhập bằng username/password")
        if (input("Chọn [1]: ").strip() or "1") == "1":
            token = token or input("__RequestVerificationToken: ").strip()
            cookie = cookie or input("Cookie: ").strip()

    client = core.MplisClient(log)
    try:
        if token and cookie:
            client.build_session_from_manual(token, cookie)
        else:
            username = args.user or input("Username: ").strip()
            password = getpass.getpass("Password: ")
            client.driver = mo_chrome(out_dir)
            dang_nhap(client.driver, username, password)
            input(">>> Đăng nhập xong trên Chrome (OTP/captcha nếu có) rồi nhấn Enter...")
            client.build_session_from_browser()

        mau = None if args.hoc_lai else doc_mau_da_luu()
        if mau:
            log(f"Dùng mẫu URL đã lưu: {mau.get('method')} {mau.get('url')}")
        token = client.session.headers.get("__requestverificationtoken", "")
        # --hoc-lai: bỏ cả mẫu mặc định FileHandler.ashx, bắt buộc học từ request thật.
        ung_vien = [] if args.hoc_lai else MAU_MAC_DINH
        downloader = Downloader(client.driver, client.session, token, mau, ung_vien)

        tong = []
        excel_kq = ten_excel_ket_qua(out_dir)
        for n, id_don in enumerate(ids, 1):
            log(f"--- [{n}/{len(ids)}] Đơn {id_don} ---")
            da_co_mau = downloader.mau is not None
            try:
                kq = xu_ly_don(client, downloader, id_don, out_dir, args.chi_kiem_tra)
            except Exception as exc:
                log(f"❌ Đơn {id_don}: {core.rut_gon_text(exc, 300)}")
                kq = {"id_don": id_don, "so_file": 0, "da_tai": 0, "loi": [str(exc)],
                      "dong": [{"madon": id_don, "trang_thai_tai": "Lỗi", "ghi_chu_tai": str(exc)}]}
            tong.append(kq)
            ghi_excel_ket_qua(excel_kq, [d for k in tong for d in k["dong"]])

            # Mẫu vừa học mà tải được file → lưu lại cho lần sau.
            if not da_co_mau and downloader.mau is not None and downloader.mau not in MAU_MAC_DINH:
                if kq["da_tai"] > 0:
                    luu_mau(downloader.mau)
                else:
                    log("⚠ Mẫu URL vừa bắt được không tải được file nào → bỏ, sẽ học lại ở đơn sau.")
                    downloader.mau = None

        log("=" * 60)
        log(f"Tổng: {len(tong)} đơn | {sum(k['so_file'] for k in tong)} file HSQ"
            + ("" if args.chi_kiem_tra else f" | đã tải {sum(k['da_tai'] for k in tong)} → {out_dir}"))
        khong_hsq = [str(k["id_don"]) for k in tong if k["so_file"] == 0]
        if khong_hsq:
            log(f"Đơn không có HSQ: {', '.join(khong_hsq)}")
        co_loi = [k for k in tong if k["loi"]]
        if co_loi:
            log(f"Đơn có lỗi: {len(co_loi)}")
        log(f"Excel kết quả: {excel_kq}")
    finally:
        client.close_browser()


if __name__ == "__main__":
    main()
