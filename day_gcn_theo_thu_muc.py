# -*- coding: utf-8 -*-
"""
Đẩy file Giấy chứng nhận theo thư mục mã đơn — tìm file theo SỐ GCN của đơn, ghi nhận là thành phần GCN.

Cấu trúc thư mục:
    <thư mục gốc>/
        12567890/                              ← tên thư mục bắt đầu bằng mã đơn (tinhHinhDangKyId)
            DG 690247.pdf                      ← tên file = số phát hành GCN (ưu tiên 1)
            24232_12_40_GCN_DG 690247.pdf      ← hoặc có chữ "GCN" + số phát hành (ưu tiên 2)
            PC_NGUYEN DANH CHUNG.pdf           ← file khác: không đụng tới

Với mỗi đơn:
  1. Lấy chi tiết đơn → GCN hiện hành (bỏ isLastest = False) + HSQ.
  2. Mỗi GCN tìm 1 file trong thư mục đơn (không tính đuôi ký số _signed..., tiền tố "NN_"):
       ưu tiên 1: tên file đúng bằng số GCN ("DG 690247", "GCN DG 690247", "Giấy chứng nhận DG 690247");
       ưu tiên 2: tên file có chữ "GCN"/"Giấy chứng nhận" và chứa số GCN.
     File chỉ có số mà không có chữ GCN (vd "GT_DG 690247") KHÔNG lấy. Cùng mức có nhiều file → lấy file
     đã ký số nếu chỉ 1 file ký, còn lại bỏ qua GCN đó. 1 file chứa nhiều số GCN → dùng chung cho các GCN.
  3. Đặt file vào HSQ (luôn ghi nhận loại Giấy chứng nhận + gắn đúng GCN):
       a. thành phần đã gắn GCN đó          → thay file;
       b. thành phần có mô tả trùng tên file (hoặc tên file chứa nguyên mô tả ≥ 8 ký tự,
          vd 'Ly Kim Hoang 24221_3_96_GCN_DI 923410') → thay file + ghi nhận GCN;
       c. không có                           → thêm thành phần mới "Giấy chứng nhận <số>" vào HSQ đầu tiên;
       d. đơn chưa có HSQ                    → tạo HSQ mới gồm các thành phần GCN.
     Mỗi HSQ gửi 1 request kèm ĐỦ thành phần cũ (giữ nguyên), như day_hsq_ky_so / bo_sung_gcn_hsq.
  4. Xác nhận: số thành phần không giảm (giảm → DỪNG TOÀN BỘ), mỗi GCN có thành phần loại Giấy chứng nhận
     gắn đúng GCN, file thay có nodeId mới.
"""

from __future__ import annotations

import contextlib
import json
import os
import time
from typing import Any, Callable

import bo_sung_gcn_hsq as bs
import cap_nhat_don_hsq as core
import nhom1_mplis as nhom1
import day_hsq_ky_so as day
import day_hsq_theo_mota as dm
import tai_hsq_don as thd

log = thd.log

EXCEL_HEADERS = ["stt", "madon", "so_gcn", "gcn_id", "file", "cach", "hosoquetid", "mota_thanh_phan",
                 "da_ky_so", "trang_thai", "ghi_chu"]
EXCEL_WIDTHS = {"stt": 6, "madon": 12, "so_gcn": 14, "gcn_id": 10, "file": 45, "cach": 24, "hosoquetid": 11,
                "mota_thanh_phan": 40, "da_ky_so": 8, "trang_thai": 13, "ghi_chu": 60}
MAU_TRANG_THAI = {"Đã đẩy": "C6EFCE", "Lỗi": "F4CCCC", "Bỏ qua": "FFF2CC", "Chạy thử": "D9EAF7"}
CACH = {"a": "Thay file thành phần GCN", "b": "Thay file + ghi nhận GCN", "c": "Thêm thành phần GCN mới",
        "d": "Tạo HSQ mới"}


# ============================ TÌM FILE THEO SỐ GCN ============================

def _co_chu_gcn(khoa: str) -> bool:
    return "gcn" in khoa or "giaychungnhan" in khoa


def tim_file_gcn(pdfs: list[str], gcn: dict[str, Any]) -> tuple[str | None, str]:
    """File cho 1 GCN. Trả về (đường dẫn, '') hoặc (None, lý do)."""
    so = dm.chuan_hoa(gcn.get("soPhatHanh"))
    if len(so) < 5:
        return None, f"GCN {gcn.get('giayChungNhanId')} không có số phát hành để tìm file"
    muc1, muc2 = [], []
    for p in pdfs:
        khoa = dm.chuan_hoa(dm.tach_ten_file(p)[1])
        if khoa in (so, "gcn" + so, "giaychungnhan" + so):
            muc1.append(p)
        elif so in khoa and _co_chu_gcn(khoa):
            muc2.append(p)
    for ung_vien in (muc1, muc2):
        if len(ung_vien) == 1:
            return ung_vien[0], ""
        if len(ung_vien) > 1:
            ky = [p for p in ung_vien if day.pdf_da_ky_so(p)]
            if len(ky) == 1:
                return ky[0], ""
            return None, (f"{len(ung_vien)} file cùng khớp GCN {gcn.get('soPhatHanh')}: "
                          + ", ".join(os.path.basename(p) for p in ung_vien) + " — giữ lại 1 file")
    return None, (f"Không có file cho GCN {gcn.get('soPhatHanh')} (đặt tên '{gcn.get('soPhatHanh')}.pdf' "
                  "hoặc tên có chữ GCN + số phát hành)")


# ============================ KẾ HOẠCH 1 ĐƠN ============================

def info_gcn_moi(gcn: dict[str, Any], da_ky: bool) -> dict[str, Any]:
    """Thành phần GCN tạo mới — đúng mẫu bo_sung_gcn_hsq.info_thanh_phan_moi (đã chạy thật)."""
    return bs.info_thanh_phan_moi({"daKySo": da_ky}, gcn)


def lap_ke_hoach(detail: dict[str, Any], ds_gcn: list[dict], file_gcn: dict[int, tuple[str, bool]]) -> list[dict]:
    """
    Mỗi GCN có file → {"gcn", "path", "da_ky", "cach", "hoso", "vi_tri"}.
    cach a/b: thay thành phần (hoso, vi_tri); c: thêm vào hoso; d: tạo HSQ mới (hoso None).
    """
    hsq_list = [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]
    da_dung: set[tuple[int, int]] = set()
    ke_hoach = []
    for gcn in ds_gcn:
        gid = core.safe_int(gcn.get("giayChungNhanId"))
        if gid not in file_gcn:
            continue
        path, da_ky = file_gcn[gid]
        muc = {"gcn": gcn, "path": path, "da_ky": da_ky, "hoso": None, "vi_tri": None}
        if not hsq_list:
            ke_hoach.append({**muc, "cach": "d"})
            continue
        khoa_file = dm.chuan_hoa(dm.tach_ten_file(path)[1])
        tim = None
        for cach, dk in (("a", lambda n: core.safe_int(n.get("giayChungNhanId")) == gid),
                         ("b", lambda n: dm.chuan_hoa(n.get("moTa")) == khoa_file),
                         ("b", lambda n: dm.mo_ta_nam_trong(khoa_file, dm.chuan_hoa(n.get("moTa"))))):
            for h_i, h in enumerate(hsq_list):
                for i, n in enumerate(day.cac_node(h)):
                    if (h_i, i) not in da_dung and dk(n):
                        tim = (cach, h_i, h, i)
                        break
                if tim:
                    break
            if tim:
                break
        if tim:
            da_dung.add((tim[1], tim[3]))
            ke_hoach.append({**muc, "cach": tim[0], "hoso": tim[2], "vi_tri": tim[3]})
        else:
            ke_hoach.append({**muc, "cach": "c", "hoso": hsq_list[0]})
    return ke_hoach


# ============================ GỬI ============================

def gui_tao_moi(client: core.MplisClient, thong_tin_ho_so_id: int, items: list[tuple[dict, str]]) -> None:
    """Tạo HSQ mới nhiều thành phần trong 1 request (mẫu khoi_phuc_hsq — đã chạy thật)."""
    session = client._require_session()
    data = {
        "hoSoQuet": json.dumps({"thongTinHoSoId": thong_tin_ho_so_id, "TuiHoSo": None, "tuiHoSoId": 0,
                                "hoSoQuetId": 0}, ensure_ascii=False),
        "count": str(len(items)),
        "isLuuKhoHoSoQuet": "false",
    }
    headers = dict(session.headers)
    headers.pop("Content-Type", None)
    with contextlib.ExitStack() as stack:
        files = {}
        for i, (info, p) in enumerate(items, 1):
            data[f"infoHoSoQuet_{i}"] = json.dumps(info, ensure_ascii=False)
            files[f"fileHoSoQuet_{i}"] = (os.path.basename(p), stack.enter_context(open(p, "rb")), "application/pdf")
        response = session.post(core.URL_UPDATE_HOSOQUET, data=data, files=files, headers=headers,
                                timeout=core.TIMEOUT, allow_redirects=False)
    result = client._response_json(response, "UpdateHoSoQuetExistFile (tạo mới)")
    if not isinstance(result, dict) or not result.get("success"):
        raise RuntimeError(f"Tạo HSQ không thành công: {core.rut_gon_text(result)}")


def items_hsq_co_san(hoso: dict[str, Any], ke_hoach_hsq: list[dict]) -> list[tuple[dict, str | None]]:
    """ĐỦ thành phần hiện có (thay hoặc giữ nguyên) + thành phần GCN thêm mới, theo thứ tự."""
    hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
    thay = {k["vi_tri"]: k for k in ke_hoach_hsq if k["cach"] in ("a", "b")}
    items: list[tuple[dict, str | None]] = []
    for i, node in enumerate(day.cac_node(hoso)):
        if i in thay:
            k = thay[i]
            info = day.info_thay_file(node, hsq_id, k["da_ky"])
            dm.gan_loai_gcn(info, k["gcn"])
            items.append((info, k["path"]))
        else:
            items.append((day.info_giu_nguyen(node), None))
    items += [(info_gcn_moi(k["gcn"], k["da_ky"]), k["path"]) for k in ke_hoach_hsq if k["cach"] == "c"]
    return items


# ============================ XỬ LÝ 1 ĐƠN ============================

def _dong(madon: int, k: dict | None = None, **kw) -> dict[str, Any]:
    k = k or {}
    gcn = k.get("gcn") or {}
    node = {}
    if k.get("hoso") is not None and k.get("vi_tri") is not None:
        node = day.cac_node(k["hoso"])[k["vi_tri"]]
    return {
        "madon": madon,
        "so_gcn": gcn.get("soPhatHanh", ""),
        "gcn_id": gcn.get("giayChungNhanId", ""),
        "file": os.path.basename(k["path"]) if k.get("path") else "",
        "cach": CACH.get(k.get("cach"), ""),
        "hosoquetid": (k.get("hoso") or {}).get("hoSoQuetId", "") if k.get("hoso") else "",
        "mota_thanh_phan": node.get("moTa") or (bs.mo_ta_gcn(gcn) if k.get("cach") in ("c", "d") else ""),
        "da_ky_so": "Có" if k.get("da_ky") else "",
        **kw,
    }


def xu_ly_don(client: core.MplisClient, madon: int, pdfs: list[str], chay_thu: bool,
              bo_qua_nhom_1: bool = True) -> list[dict[str, Any]]:
    detail = client.lay_chi_tiet_don(madon)
    if detail is None:
        return [_dong(madon, trang_thai="Lỗi", ghi_chu=dm.LOI_CHI_TIET)]
    ds_gcn = bs.gcn_hien_hanh(detail)
    if not ds_gcn:
        return [_dong(madon, trang_thai="Bỏ qua", ghi_chu="Đơn chưa có GCN hiện hành")]
    if bo_qua_nhom_1 and nhom1.da_dat(client, madon, detail):
        return [_dong(madon, trang_thai="Bỏ qua", ghi_chu=nhom1.LY_DO_DA_DAT)]

    rows: list[dict[str, Any]] = []
    file_gcn: dict[int, tuple[str, bool]] = {}
    for gcn in ds_gcn:
        path, ly_do = tim_file_gcn(pdfs, gcn)
        if path is None:
            rows.append(_dong(madon, {"gcn": gcn}, trang_thai="Bỏ qua", ghi_chu=ly_do))
            log(f"   ⚠ {ly_do}")
        else:
            file_gcn[core.safe_int(gcn.get("giayChungNhanId"))] = (path, day.pdf_da_ky_so(path))
    if not file_gcn:
        return rows

    ke_hoach = lap_ke_hoach(detail, ds_gcn, file_gcn)
    for k in ke_hoach:
        log(f"   GCN {k['gcn'].get('soPhatHanh')} ← {os.path.basename(k['path'])}"
            f"{' (đã ký số)' if k['da_ky'] else ''} | {CACH[k['cach']]}"
            + (f" '{day.cac_node(k['hoso'])[k['vi_tri']].get('moTa')}'" if k["cach"] in ("a", "b") else ""))
    if chay_thu:
        return rows + [_dong(madon, k, trang_thai="Chạy thử", ghi_chu=f"Sẽ {CACH[k['cach']].lower()}")
                       for k in ke_hoach]

    # ---- gửi: đơn chưa có HSQ → tạo mới; có HSQ → mỗi HSQ 1 request đủ thành phần
    nhom: list[tuple[dict | None, list[dict]]] = []
    tao_moi = [k for k in ke_hoach if k["cach"] == "d"]
    if tao_moi:
        nhom.append((None, tao_moi))
    for h in [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]:
        ds = [k for k in ke_hoach if k["hoso"] is h]
        if ds:
            nhom.append((h, ds))

    for hoso, ds in nhom:
        so_tp_truoc = len(day.cac_node(hoso)) if hoso else 0
        node_id_truoc = {k["vi_tri"]: str(day.cac_node(hoso)[k["vi_tri"]].get("nodeId") or "")
                         for k in ds if k["cach"] in ("a", "b")}
        try:
            if hoso is None:
                tths_id = core.safe_int(core.lay_thong_tin_ho_so_id(detail))
                if not tths_id:
                    raise RuntimeError("Không tìm được thongTinHoSoId của đơn để tạo HSQ mới")
                gui_tao_moi(client, tths_id, [(info_gcn_moi(k["gcn"], k["da_ky"]), k["path"]) for k in ds])
            else:
                bs.gui_hsq(client, hoso, items_hsq_co_san(hoso, ds))
        except core.HetPhien:
            raise
        except Exception as exc:
            rows += [_dong(madon, k, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for k in ds]
            log(f"   ❌ {core.rut_gon_text(exc, 300)}")
            continue

        # ---- xác nhận
        detail_moi = client.lay_chi_tiet_don(madon) or {}
        if hoso is not None:
            hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
            hoso_moi = next((h for h in detail_moi.get("ListHoSoQuet") or []
                             if isinstance(h, dict) and core.safe_int(h.get("hoSoQuetId")) == hsq_id), None)
            nodes_moi = day.cac_node(hoso_moi) if hoso_moi else []
            if len(nodes_moi) < so_tp_truoc:
                rows += [_dong(madon, k, trang_thai="Lỗi", ghi_chu=f"MPLIS còn {len(nodes_moi)}/{so_tp_truoc} "
                                                                     "thành phần!") for k in ds]
                raise day.DungKhanCap(f"HSQ {hsq_id} (đơn {madon}) giảm từ {so_tp_truoc} còn {len(nodes_moi)} "
                                      "thành phần sau khi đẩy → DỪNG TOÀN BỘ.")
        else:
            nodes_moi = [n for h in detail_moi.get("ListHoSoQuet") or [] if isinstance(h, dict)
                         for n in day.cac_node(h)]
        for k in ds:
            gid = core.safe_int(k["gcn"].get("giayChungNhanId"))
            gcn_nodes = [n for n in nodes_moi if core.safe_int(n.get("giayChungNhanId")) == gid
                         and core.safe_int(n.get("loaiHoSoQuet"), -1) == 1]
            loi = "" if gcn_nodes else "chưa có thành phần loại Giấy chứng nhận gắn GCN này"
            if not loi and k["cach"] in ("a", "b") and all(
                    str(n.get("nodeId") or "") == node_id_truoc[k["vi_tri"]] for n in gcn_nodes):
                loi = "file chưa đổi"
            if loi:
                rows.append(_dong(madon, k, trang_thai="Lỗi", ghi_chu=f"Server báo OK nhưng {loi} (kiểm tra MPLIS)"))
                log(f"   ❌ GCN {k['gcn'].get('soPhatHanh')}: server báo OK nhưng {loi}")
            else:
                rows.append(_dong(madon, k, trang_thai="Đã đẩy", ghi_chu=time.strftime("%d/%m/%Y %H:%M:%S")))
                log(f"   ✅ GCN {k['gcn'].get('soPhatHanh')}: {CACH[k['cach']].lower()}")
        time.sleep(core.REQUEST_DELAY_SECONDS)
    return rows


# ============================ EXCEL + LUỒNG CHÍNH ============================

def ghi_excel(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_tt = EXCEL_HEADERS.index("trang_thai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in EXCEL_HEADERS])
        fill = MAU_TRANG_THAI.get(str(row.get("trang_thai") or ""))
        if fill:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=fill)
    for i, h in enumerate(EXCEL_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}{max(ws.max_row, 1)}"
    wb.save(path)


def _ghi_an_toan(path: str, rows: list[dict[str, Any]]) -> str:
    try:
        ghi_excel(path, rows)
        return path
    except PermissionError:
        alt = os.path.splitext(path)[0] + "_capnhat.xlsx"
        ghi_excel(alt, rows)
        log(f"⚠ Excel đang mở, đã lưu sang: {alt}")
        return alt


def chay(
    client: core.MplisClient,
    goc: str,
    chay_thu: bool = True,
    nen_dung: Callable[[], bool] = lambda: False,
    bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None,
) -> str:
    """Chạy cho mọi thư mục mã đơn trong `goc`; trả về đường dẫn Excel kết quả (ghi trong `goc`)."""
    theo_don, khong_nhan = dm.quet_thu_muc_goc(goc)
    if khong_nhan:
        log(f"⚠ Bỏ qua {len(khong_nhan)} thư mục có PDF nhưng không nằm trong thư mục mã đơn nào: {', '.join(khong_nhan[:10])}"
            + (" ..." if len(khong_nhan) > 10 else ""))
    theo_don = {m: p for m, p in theo_don.items() if p}
    if not theo_don:
        raise RuntimeError(f"Không có thư mục mã đơn nào chứa file PDF trong {goc}")
    excel = os.path.join(goc, f"ket_qua_day_gcn_{'thu_' if chay_thu else ''}{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
    log(f"{'[CHẠY THỬ] ' if chay_thu else ''}{len(theo_don)} đơn | {sum(len(v) for v in theo_don.values())} file PDF")

    rows: list[dict[str, Any]] = []
    loi_lien_tiep = 0
    try:
        for n, (madon, pdfs) in enumerate(theo_don.items(), 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(n - 1, len(theo_don), f"Đơn {madon} ({n}/{len(theo_don)})")
            log(f"--- [{n}/{len(theo_don)}] Đơn {madon}: {len(pdfs)} file ---")
            try:
                dong = xu_ly_don(client, madon, pdfs, chay_thu)
                rows += dong
                loi_lien_tiep = loi_lien_tiep + 1 if dong and dong[0].get("ghi_chu") == dm.LOI_CHI_TIET else 0
                if loi_lien_tiep >= dm.DUNG_SAU_SO_DON_LOI:
                    log(f"❌ {loi_lien_tiep} đơn liên tiếp không đọc được từ MPLIS — có thể đã hết phiên đăng nhập. "
                        "ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            except day.DungKhanCap:
                rows.append(_dong(madon, trang_thai="Lỗi", ghi_chu="MPLIS giảm thành phần sau khi đẩy — ĐÃ DỪNG"))
                raise
            except Exception as exc:
                rows.append(_dong(madon, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)))
                log(f"   ❌ {core.rut_gon_text(exc, 300)}")
                if isinstance(exc, core.HetPhien):
                    log(f"❌ {exc} — ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            excel = _ghi_an_toan(excel, rows)
            core.nghi_giua_don()
    finally:
        if rows:
            excel = _ghi_an_toan(excel, rows)

    dem: dict[str, int] = {}
    for r in rows:
        dem[r["trang_thai"]] = dem.get(r["trang_thai"], 0) + 1
    bao_tien_do(len(theo_don), len(theo_don), "Hoàn tất")
    log("=" * 50)
    log(("[CHẠY THỬ — chưa gửi gì] " if chay_thu else "✅ ") + "GCN: " + " | ".join(f"{k}: {v}" for k, v in dem.items()))
    log(f"Excel kết quả: {excel}")
    return excel
