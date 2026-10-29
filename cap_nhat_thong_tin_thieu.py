# -*- coding: utf-8 -*-
"""
Cập nhật thông tin thiếu lên MPLIS từ Excel ket_qua_thong_tin_thieu_*.xlsx (do tai_tach_hsq_don /
lay_thong_tin_thieu_hsq sinh ra, ĐÃ DUYỆT).

Chỉ ghi các dòng có cột duyet = 'x' (tool đọc điền sẵn cho Đã đọc / Mặc định; dòng Cần kiểm tra bạn tự đánh x
sau khi kiểm tra). Giá trị ghi = gia_tri_sua nếu có, không thì gia_tri_ai.

Mỗi đơn:
  1. Lấy bản ghi làm sạch của đơn trên MPLIS (GetThongKePhanLoaiThuaDatChiTiet) — mục nào MPLIS không còn báo
     thiếu thì bỏ qua (không ghi đè dữ liệu đã có).
  2. Tìm khoá delta của từng mục trong thongTinDangKyChuaDapUngNhom1 của bản ghi (đúng khoá MPLIS dùng, vd
     TINHHINHDANGKY.<id>|GIAYCHUNGNHAN.<id>_<ver>|ngayVaoSo); không có danh sách khoá thì dựng theo đối tượng.
  3. Chuẩn hoá giá trị theo kiểu trường (ngày → ISO UTC như MPLIS, diện tích → số, mã loại GCN...).
  4. Gửi delta làm sạch (LamSachDuLieuAjax/CapNhatThongKePhanLoaiThuaDatChiTiet — cùng API lamsach_delta), mỗi
     bản ghi 1 request. Mặc định CHẠY THỬ: chỉ ghi kế hoạch vào Excel.
  5. Đọc lại bản ghi làm sạch: mục hết bị báo thiếu → "Đã cập nhật"; còn → "Đã gửi — MPLIS vẫn báo thiếu".
Trạng thái ghi ngược vào chính file Excel (cột trang_thai_cap_nhat / ghi_chu_cap_nhat / thoi_gian_cap_nhat),
chạy lại thì dòng "Đã cập nhật" được bỏ qua.
"""

from __future__ import annotations

import json
import re
import time
from typing import Any, Callable

import cap_nhat_don_hsq as core
import lay_thong_tin_thieu_hsq as ltt
import tai_hsq_don as thd

log = thd.log

URL_DELTA_LAM_SACH = "https://dla.mplis.gov.vn/dc/LamSachDuLieuAjax/CapNhatThongKePhanLoaiThuaDatChiTiet"
REFERER = "https://dla.mplis.gov.vn/dc/DonDangKy/KeKhaiDangKyV2"

# loại đối tượng trong thông báo → tiền tố trong khoá delta
TAG = {
    "thửa đất": "THUADAT", "mục đích sử dụng": "MUCDICHSUDUNG", "nguồn gốc sử dụng": "NGUONGOCSUDUNGDAT",
    "tài liệu đo đạc": "TAILIEUDODAC", "giấy chứng nhận": "GIAYCHUNGNHAN", "nội dung thay đổi": "NOIDUNGTHAYDOI",
    "hộ gia đình": "HOGIADINH", "vợ chồng": "VOCHONG", "cá nhân": "CANHAN", "giấy tờ tùy thân": "GIAYTOTUYTHAN",
    "tổ chức": "TOCHUC", "cộng đồng dân cư": "CONGDONGDANCU",
}
# tên trường trong thông báo → tên trường trong khoá delta / chi tiết đơn
TRUONG_MPLIS = {
    "thời hạn sử dụng": "thoiHanSuDung", "số giấy tờ": "soGiayTo", "ngayVaoSo": "ngayVaoSo",
    "loaiGiayChungNhanId": "loaiGiayChungNhanId", "số phát hành": "soPhatHanh", "địa chỉ": "diaChi",
    "diện tích": "dienTich", "loaiBanDoDiaChinhId": "loaiBanDoDiaChinhId", "năm sinh": "namSinh",
    "ngayXacNhan": "ngayXacNhan",
}
# Mã loaiBanDoDiaChinhId theo tên loại bản đồ (ô chọn ddlloaiBanDoDiaChinhId trên MPLIS). None = chưa biết mã
# → không ghi trường này (báo lỗi rõ) cho tới khi điền mã đúng.
MA_LOAI_BAN_DO: dict[str, int | None] = {"Bản đồ địa chính (VN2000)": None}

COT_TRANG_THAI = ("trang_thai_cap_nhat", "ghi_chu_cap_nhat", "thoi_gian_cap_nhat")
GIA_TRI_DUYET = {"x", "có", "co", "1", "y", "yes", "ok", "✓", "v"}
MAU = {"Đã cập nhật": "C6EFCE", "Chạy thử": "D9EAF7", "Bỏ qua": "EDEDED", "Lỗi": "F4CCCC",
       "Đã gửi — MPLIS vẫn báo thiếu": "FFF2CC"}


# ============================ CHUẨN HOÁ GIÁ TRỊ ============================

def chuan_hoa_gia_tri(truong: str, text: str) -> Any:
    """Giá trị đọc được (chữ) → giá trị gửi MPLIS. Sai định dạng → ValueError (lý do)."""
    t = str(text or "").strip()
    if not t:
        raise ValueError("giá trị rỗng")
    if truong in ("ngayVaoSo", "ngayXacNhan"):
        m = re.search(r"(\d{1,2})\s*[/.-]\s*(\d{1,2})\s*[/.-]\s*(\d{4})", t)
        if not m:
            raise ValueError(f"ngày '{t}' không phải dạng DD/MM/YYYY")
        try:
            return core.ddmmyyyy_to_iso_utc_start_of_day_vn(f"{int(m[1]):02d}/{int(m[2]):02d}/{m[3]}")
        except ValueError as exc:
            raise ValueError(f"ngày '{t}' không hợp lệ") from exc
    if truong == "năm sinh":
        m = re.search(r"\b(19\d\d|20\d\d)\b", t)
        if not m or int(m[1]) > time.localtime().tm_year:
            raise ValueError(f"năm sinh '{t}' không hợp lệ")
        return m[1]
    if truong == "diện tích":
        m = re.search(r"\d+(?:[.,]\d+)?", t.replace(" ", ""))
        if not m:
            raise ValueError(f"diện tích '{t}' không phải số")
        so = float(m[0].replace(",", "."))
        if so <= 0:
            raise ValueError(f"diện tích '{t}' phải > 0")
        return f"{so:.4f}".rstrip("0").rstrip(".")
    if truong == "loaiGiayChungNhanId":
        ma = re.sub(r"\D", "", t)
        if ma not in ltt.LOAI_GCN:
            raise ValueError(f"mã loại GCN '{t}' không có trong bảng")
        return ma
    if truong == "loaiBanDoDiaChinhId":
        ma = MA_LOAI_BAN_DO.get(t)
        if ma is None:
            raise ValueError(f"chưa có mã loaiBanDoDiaChinhId cho '{t}' (điền MA_LOAI_BAN_DO)")
        return str(ma)
    if truong == "số phát hành":
        m = re.fullmatch(r"([A-ZĐ]{1,2})\s*[-_]?\s*(\d{6,8})", core.remove_accents(t).upper().replace("Đ", "D"))
        if not m:
            raise ValueError(f"số phát hành '{t}' không đúng dạng 2 chữ cái + 6 số")
        return f"{m[1]} {m[2]}"
    if truong == "số giấy tờ":
        so = re.sub(r"[\s.]", "", t)
        if not re.fullmatch(r"[A-Za-z0-9]{6,15}", so):
            raise ValueError(f"số giấy tờ '{t}' không hợp lệ")
        return so
    return t  # thời hạn sử dụng, địa chỉ: chữ


# ============================ KHOÁ DELTA ============================

def _ds_khoa_thieu(ban_ghi: dict) -> list[str]:
    v = ban_ghi.get("thongTinDangKyChuaDapUngNhom1")
    if isinstance(v, str):
        try:
            v = json.loads(v)
        except ValueError:
            v = [x for x in re.split(r"[,;\s]+", v) if "|" in x]
    return [str(x) for x in v] if isinstance(v, list) else []


def _id_trong_khoa(khoa: str) -> list[int]:
    return [int(m[1]) for m in re.finditer(r"\|[A-Z]+\.(\d+)", khoa)]


def tim_khoa(muc: ltt.MucThieu, ds_ban_ghi: list[dict], da_dung: set[str]) -> tuple[str, dict | None, str]:
    """→ (khoá delta, bản ghi làm sạch chứa khoá, ghi chú). Không có danh sách khoá thì dựng theo đối tượng."""
    ten = TRUONG_MPLIS[muc.truong].lower()
    ids = [so for _, so in muc.chuoi]
    for bg in ds_ban_ghi:
        ung_vien = [k for k in _ds_khoa_thieu(bg) if k not in da_dung and ids
                    and ids[-1] in _id_trong_khoa(k) and set(ids[-1:]) <= set(_id_trong_khoa(k))]
        dung_ten = [k for k in ung_vien if k.rsplit("|", 1)[-1].lower() == ten]
        if len(dung_ten) == 1:
            return dung_ten[0], bg, "khoá theo MPLIS"
        if len(dung_ten) > 1:
            khop_het = [k for k in dung_ten if set(ids) <= set(_id_trong_khoa(k))]
            if len(khop_het) == 1:
                return khop_het[0], bg, "khoá theo MPLIS"
    # Dựng khoá: cá nhân trong hộ / vợ chồng đứng thẳng dưới TINHHINHDANGKY (như delta mã định danh đã gửi).
    chuoi = list(muc.chuoi)
    if any(loai == "cá nhân" for loai, _ in chuoi):
        while chuoi and chuoi[0][0] in ("hộ gia đình", "vợ chồng"):
            chuoi.pop(0)
    phan = [f"{TAG[loai]}.{so}" + (f"_{muc.phien_ban[so]}" if so in muc.phien_ban else "") for loai, so in chuoi]
    khoa = "|".join([f"TINHHINHDANGKY.{muc.id_don}", *phan, TRUONG_MPLIS[muc.truong]])
    bg = next((b for b in ds_ban_ghi if ids and str(ids[0]) in ltt.thong_bao_cua([b])), None)
    co_ds = any(_ds_khoa_thieu(b) for b in ds_ban_ghi)
    return khoa, bg or (ds_ban_ghi[0] if ds_ban_ghi else None), (
        "khoá tự dựng — không thấy trong danh sách khoá thiếu của MPLIS" if co_ds
        else "khoá tự dựng — MPLIS không trả danh sách khoá thiếu")


# ============================ EXCEL ============================

def doc_excel(path: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    hdr = [str(h or "").strip() for h in rows[0]]
    for can in ("madon", "doi_tuong", "truong_thieu", "gia_tri_ai", "duyet"):
        if can not in hdr:
            raise RuntimeError(f"Excel thiếu cột '{can}' — cần file ket_qua_thong_tin_thieu_*.xlsx bản mới")
    return [{**{h: r[i] if i < len(r) else None for i, h in enumerate(hdr) if h}, "_dong": n}
            for n, r in enumerate(rows[1:], 2) if any(v not in (None, "") for v in r)]


def ghi_trang_thai(path: str, ket_qua: dict[int, tuple[str, str]]) -> None:
    """{số dòng Excel: (trạng thái, ghi chú)} → ghi vào chính file. File đang mở → PermissionError."""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    wb = load_workbook(path)
    ws = wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    cot = {}
    for ten in COT_TRANG_THAI:
        if ten not in hdr:
            hdr.append(ten)
            ws.cell(row=1, column=len(hdr), value=ten)
        cot[ten] = hdr.index(ten) + 1
    luc = time.strftime("%d/%m/%Y %H:%M:%S")
    for dong, (tt, ghi_chu) in ket_qua.items():
        ws.cell(row=dong, column=cot["trang_thai_cap_nhat"], value=tt)
        ws.cell(row=dong, column=cot["ghi_chu_cap_nhat"], value=ghi_chu)
        ws.cell(row=dong, column=cot["thoi_gian_cap_nhat"], value=luc)
        if tt in MAU:
            ws.cell(row=dong, column=cot["trang_thai_cap_nhat"]).fill = PatternFill("solid", fgColor=MAU[tt])
    wb.save(path)


# ============================ 1 ĐƠN ============================

def _gui_delta(client, payload: dict) -> None:
    session = client._require_session()
    headers = {"Accept": "application/json, text/javascript, */*; q=0.01",
               "Content-Type": "application/json; charset=UTF-8", "X-Requested-With": "XMLHttpRequest",
               "Referer": REFERER, "Origin": "https://dla.mplis.gov.vn"}
    resp = session.post(URL_DELTA_LAM_SACH, data=json.dumps(payload, ensure_ascii=False), headers=headers,
                        timeout=core.TIMEOUT, allow_redirects=False)  # ghi: không retry
    js = client._response_json(resp, "CapNhatThongKePhanLoaiThuaDatChiTiet")
    if isinstance(js, dict) and (js.get("success") is False or js.get("Success") is False):
        raise RuntimeError(f"MPLIS từ chối: {core.rut_gon_text(js, 300)}")


def cap_nhat_don(client, id_don: int, dong: list[dict[str, Any]], chay_thu: bool) -> dict[int, tuple[str, str]]:
    """Các dòng Excel (đã duyệt) của 1 đơn → {số dòng: (trạng thái, ghi chú)}. Hết phiên → core.HetPhien."""
    kq: dict[int, tuple[str, str]] = {}
    detail = client.lay_chi_tiet_don(id_don)
    if detail is None:
        return {d["_dong"]: ("Lỗi", "Không lấy được chi tiết đơn") for d in dong}
    da_dap_ung, ds_bg = ltt.loc_chua_dap_ung(ltt.lay_ban_ghi_lam_sach(client, id_don, detail))
    if da_dap_ung:
        return {d["_dong"]: ("Bỏ qua", "MPLIS: thongTinDangKyDapUngNhom1 = true — đơn đã đáp ứng, không ghi")
                for d in dong}
    con_thieu = {m.khoa for m in ltt.doc_thong_bao(ltt.thong_bao_cua(ds_bg))[0]}
    if not ds_bg:
        return {d["_dong"]: ("Bỏ qua", "MPLIS không còn bản ghi làm sạch nào của đơn (đã đủ thông tin?)")
                for d in dong}

    ke_hoach: dict[str, dict] = {}  # id bản ghi → {"bg", "data", "dong": [(số dòng, muc, khoá, ghi chú)]}
    da_dung: set[str] = set()
    for d in dong:
        muc = ltt.tu_doi_tuong(id_don, d.get("doi_tuong"), str(d.get("truong_thieu") or ""))
        if muc.truong not in TRUONG_MPLIS or not muc.chuoi:
            kq[d["_dong"]] = ("Lỗi", f"Không nhận ra trường/đối tượng: {muc.truong} / {d.get('doi_tuong')}")
            continue
        if muc.khoa not in con_thieu:
            kq[d["_dong"]] = ("Bỏ qua", "MPLIS không còn báo thiếu mục này — không ghi đè")
            continue
        gia_tri_goc = d.get("gia_tri_sua") if str(d.get("gia_tri_sua") or "").strip() else d.get("gia_tri_ai")
        try:
            gia_tri = chuan_hoa_gia_tri(muc.truong, str(gia_tri_goc or ""))
        except ValueError as exc:
            kq[d["_dong"]] = ("Lỗi", f"Giá trị không hợp lệ: {exc}")
            continue
        khoa, bg, ghi_chu = tim_khoa(muc, ds_bg, da_dung)
        if bg is None or not bg.get("id"):
            kq[d["_dong"]] = ("Lỗi", "Không có id bản ghi làm sạch để gửi delta")
            continue
        da_dung.add(khoa)
        kh = ke_hoach.setdefault(str(bg["id"]), {"bg": bg, "data": {}, "dong": []})
        kh["data"][khoa] = gia_tri
        kh["dong"].append((d["_dong"], muc, khoa, f"{khoa} = {gia_tri!r} ({ghi_chu})"))
        log(f"   {muc.truong} ({muc.doi_tuong}): {khoa} = {gia_tri!r} [{ghi_chu}]")

    if chay_thu:
        for kh in ke_hoach.values():
            for so_dong, _, _, mo_ta in kh["dong"]:
                kq[so_dong] = ("Chạy thử", f"Sẽ ghi {mo_ta}")
        return kq

    for pl_id, kh in ke_hoach.items():
        payload = {"id": kh["bg"]["id"], "tinhHinhDangKyId": id_don,
                   "data": json.dumps(kh["data"], ensure_ascii=False)}
        log(f"   Gửi delta bản ghi {pl_id}: {len(kh['data'])} trường...")
        try:
            _gui_delta(client, payload)
        except core.HetPhien:
            raise
        except Exception as exc:
            log(f"   ❌ {core.rut_gon_text(exc, 300)}")
            for so_dong, _, _, mo_ta in kh["dong"]:
                kq[so_dong] = ("Lỗi", f"Gửi lỗi: {core.rut_gon_text(exc, 200)} — {mo_ta}")
            kh["loi"] = True
        time.sleep(core.REQUEST_DELAY_SECONDS)

    gui_ok = [x for kh in ke_hoach.values() if not kh.get("loi") for x in kh["dong"]]
    if gui_ok:
        con_lai = {m.khoa for m in ltt.doc_thong_bao(ltt.thong_bao_cua(
            ltt.lay_ban_ghi_lam_sach(client, id_don, detail)))[0]}
        for so_dong, muc, _, mo_ta in gui_ok:
            if muc.khoa in con_lai:
                kq[so_dong] = ("Đã gửi — MPLIS vẫn báo thiếu", f"Đã ghi {mo_ta} nhưng MPLIS vẫn báo thiếu — kiểm tra")
                log(f"   ⚠ {muc.truong} ({muc.doi_tuong}): MPLIS vẫn báo thiếu")
            else:
                kq[so_dong] = ("Đã cập nhật", f"Đã ghi {mo_ta}")
                log(f"   ✅ {muc.truong} ({muc.doi_tuong}): đã cập nhật")
    return kq


# ============================ CHẠY ============================

def la_duyet(v: Any) -> bool:
    return str(v or "").strip().lower() in GIA_TRI_DUYET


def la_excel_da_duyet(path: str) -> bool:
    """Excel ket_qua_thong_tin_thieu (có cột gia_tri_ai + duyet) → True; Excel chỉ có madon → False."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        hdr = {str(c or "").strip() for c in next(wb.active.iter_rows(max_row=1, values_only=True), ())}
    finally:
        wb.close()
    return {"gia_tri_ai", "duyet", "doi_tuong"} <= hdr


# ============================ TỰ ĐỘNG: CHỈ CÓ MADON ============================

def doc_dau_vao_madon(path: str) -> tuple[list[int], list[ltt.MucThieu]]:
    """
    Excel cột madon: ô là câu thông báo thiếu ('Tình hình đăng ký ... không có ...') → mục thiếu đọc thẳng;
    ô chỉ là mã đơn → id (tool hỏi MPLIS đơn đó thiếu gì). → (id chỉ có mã, mục thiếu từ câu thông báo).
    """
    muc, _ = ltt.doc_thong_bao(ltt.doc_file_dau_vao(path))
    gia_tri = [v for v in thd.doc_id_tu_excel(path) if "tình hình đăng ký" not in str(v).lower()]
    co_muc = {m.id_don for m in muc}
    return [i for i in thd.tach_danh_sach_id(gia_tri) if i not in co_muc], muc


def chay_tu_madon(client, ids: list[int], out_dir: str, api_key: str, model: str | None = None, rpm: float = 30,
                  chay_thu: bool = True, ghi_ca_can_kiem_tra: bool = False, dung_file_da_tai: bool = True,
                  muc_dau_vao: list | None = None,
                  nen_dung: Callable[[], bool] = lambda: False,
                  bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None) -> str:
    """
    Mỗi đơn: mục thiếu lấy từ muc_dau_vao (câu thông báo trong Excel), đơn không có thì hỏi MPLIS → tải HSQ (tới đâu đọc tới đó) → AI đọc đúng các mục thiếu → ghi lên MPLIS
    các mục Đã đọc / Mặc định (+ Cần kiểm tra nếu ghi_ca_can_kiem_tra) → đọc lại xác nhận.
    → Excel ket_qua_cap_nhat_thieu_*.xlsx (đúng mẫu ket_qua_thong_tin_thieu + trạng thái cập nhật: duyệt thêm
    dòng Cần kiểm tra rồi nạp lại file này vào tool để ghi tiếp).
    """
    import os
    from pathlib import Path

    import tach_gcn_theo_don as tg
    from tach_gcn.gemini import DEFAULT_MODEL, GeminiAnalyzer, Usage
    from tach_gcn.report import load_cache, save_cache

    if not api_key:
        raise RuntimeError("Chưa có Gemini API key (GEMINI_API_KEY)")
    model = model or DEFAULT_MODEL
    analyzer = GeminiAnalyzer(api_key, model, rpm=rpm)
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    token = client.session.headers.get("__requestverificationtoken", "")
    downloader = thd.Downloader(None, client.session, token, thd.doc_mau_da_luu(), thd.MAU_MAC_DINH)
    cache_path = Path(out_dir) / tg.FILE_CACHE
    cache = load_cache(cache_path)
    usage = Usage()
    excel = os.path.join(out_dir, f"ket_qua_cap_nhat_thieu_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
    dong_excel: list[dict[str, Any]] = []  # dòng Excel (dict) theo thứ tự
    co_san: dict[int, list] = {}
    for m in muc_dau_vao or []:
        co_san.setdefault(m.id_don, []).append(m)
    ids = list(dict.fromkeys([*co_san, *ids]))

    def luu() -> None:
        try:
            ltt.ghi_excel(excel, dong_excel, COT_TRANG_THAI)
        except PermissionError:
            log("⚠ Không ghi được Excel kết quả (file đang mở?)")
        save_cache(cache_path, cache)

    log(f"{'CHẠY THỬ — ' if chay_thu else ''}Tự động lấy & cập nhật thông tin thiếu cho {len(ids)} đơn"
        f" | ghi: Đã đọc + Mặc định{' + Cần kiểm tra' if ghi_ca_can_kiem_tra else ''} | model {model}")
    tt = tg.TaiTruoc(client, downloader, out_dir, dung_file_da_tai)
    tt.dang_ky(ids)
    ghi_duoc = {"Đã đọc", "Mặc định"} | ({"Cần kiểm tra"} if ghi_ca_can_kiem_tra else set())
    try:
        for n, id_don in enumerate(ids, 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(n - 1, len(ids), f"Đơn {id_don} ({n}/{len(ids)})")
            log(f"--- [{n}/{len(ids)}] Đơn {id_don} ---")
            muc: list[ltt.MucThieu] = []
            try:
                detail = tt.chi_tiet(id_don)
                if detail is None:
                    log("   ❌ Không lấy được chi tiết đơn")
                    dong_excel.append({"madon": id_don, "trang_thai": "Lỗi", "ghi_chu": "Không lấy được chi tiết đơn"})
                    continue
                # thông tin thiếu lấy từ response GetThongKePhanLoaiThuaDatChiTiet; MPLIS không trả gì mới dùng câu trong Excel
                with tt.khoa:
                    muc = ltt.muc_thieu_tu_ban_ghi(id_don, ltt.lay_ban_ghi_lam_sach(client, id_don, detail))
                if muc is None:  # đơn đã đáp ứng nhóm 1 → bỏ qua, không tải HSQ / gọi AI
                    log("   MPLIS: đơn đã đáp ứng nhóm 1 (thongTinDangKyDapUngNhom1 = true) — bỏ qua đơn")
                    dong_excel.append({"madon": id_don, "trang_thai": "Bỏ qua",
                                       "ghi_chu": "MPLIS: thongTinDangKyDapUngNhom1 = true — đơn đã đáp ứng"})
                    continue
                nguon = "MPLIS"
                if not muc and id_don in co_san:
                    muc = co_san[id_don]
                    nguon = "Excel (MPLIS không trả thông tin thiếu)"
                if not muc:
                    log("   MPLIS: đơn không thiếu thông tin nào — bỏ qua")
                    dong_excel.append({"madon": id_don, "trang_thai": "Bỏ qua",
                                       "ghi_chu": "MPLIS không báo thiếu thông tin"})
                    continue
                log(f"   {nguon}: đơn thiếu {len(muc)} mục: " + ", ".join(dict.fromkeys(m.truong for m in muc)))
                ltt.xu_ly_don(analyzer, id_don, muc, cache, usage, tt, model, nen_dung)
                dong = [{**ltt.dong_excel(m, True), "_dong": len(dong_excel) + i} for i, m in enumerate(muc)]
                dong_excel.extend(dong)
                can_ghi = [d for d in dong if d["trang_thai"] in ghi_duoc and d["gia_tri_ai"]]
                for d in dong:
                    if d not in can_ghi:
                        d["trang_thai_cap_nhat"] = "Bỏ qua"
                        d["ghi_chu_cap_nhat"] = ("Chưa ghi — cần duyệt (đánh x cột duyet rồi nạp lại file này)"
                                                 if d["gia_tri_ai"] else "Không có giá trị để ghi")
                if can_ghi and not nen_dung():
                    with tt.khoa:
                        kq = cap_nhat_don(client, id_don, can_ghi, chay_thu)
                    for d in can_ghi:
                        d["trang_thai_cap_nhat"], d["ghi_chu_cap_nhat"] = kq.get(d["_dong"], ("Lỗi", "Không có kết quả"))
                        d["thoi_gian_cap_nhat"] = time.strftime("%d/%m/%Y %H:%M:%S")
            except core.HetPhien as exc:
                log(f"❌ {exc}")
                log("⛔ HẾT PHIÊN — đã dừng. Lấy session mới rồi chạy lại (kết quả AI lấy từ cache, không tốn thêm).")
                dong_excel.append({"madon": id_don, "trang_thai": "Lỗi", "ghi_chu": f"HẾT PHIÊN — ĐÃ DỪNG: {exc}"})
                break
            except Exception as exc:
                log(f"❌ Đơn {id_don}: {core.rut_gon_text(exc, 300)}")
                dong_excel.append({"madon": id_don, "trang_thai": "Lỗi", "ghi_chu": core.rut_gon_text(exc, 300)})
            finally:
                luu()
            core.nghi_giua_don()
    finally:
        tt.dong()
        luu()
    bao_tien_do(len(ids), len(ids), "Xong")
    dem: dict[str, int] = {}
    for d in dong_excel:
        k = d.get("trang_thai_cap_nhat") or d.get("trang_thai") or ""
        dem[k] = dem.get(k, 0) + 1
    log("=" * 60)
    log("Kết quả: " + ", ".join(f"{k} {v}" for k, v in dem.items() if k))
    log(f"Chi phí Gemini ước tính: {usage.estimated_cost_usd:.4f} USD")
    log(f"Excel kết quả: {excel}")
    return excel


def chay(client, path: str, chay_thu: bool = True, nen_dung: Callable[[], bool] = lambda: False,
         bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None) -> dict[str, int]:
    """Cập nhật mọi dòng đã duyệt trong Excel → đếm theo trạng thái. Hết phiên → dừng (đã ghi Excel)."""
    rows = doc_excel(path)
    duyet = [r for r in rows if la_duyet(r.get("duyet")) and r.get("trang_thai_cap_nhat") != "Đã cập nhật"]
    theo_don: dict[int, list] = {}
    for r in duyet:
        if core.safe_int(r.get("madon")):
            theo_don.setdefault(core.safe_int(r.get("madon")), []).append(r)
    log(f"{'CHẠY THỬ — ' if chay_thu else ''}Cập nhật thông tin thiếu: {len(duyet)}/{len(rows)} dòng đã duyệt, "
        f"{len(theo_don)} đơn")
    tong: dict[int, tuple[str, str]] = {}
    ids = list(theo_don)

    def luu() -> None:
        if tong:
            ghi_trang_thai(path, tong)

    try:
        for n, id_don in enumerate(ids, 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(n - 1, len(ids), f"Đơn {id_don} ({n}/{len(ids)})")
            log(f"--- [{n}/{len(ids)}] Đơn {id_don}: {len(theo_don[id_don])} mục ---")
            try:
                tong.update(cap_nhat_don(client, id_don, theo_don[id_don], chay_thu))
            except core.HetPhien as exc:
                log(f"❌ {exc}")
                log("⛔ HẾT PHIÊN — đã dừng. Lấy session mới rồi chạy lại (dòng đã cập nhật được bỏ qua).")
                for r in theo_don[id_don]:
                    tong.setdefault(r["_dong"], ("Lỗi", f"HẾT PHIÊN — ĐÃ DỪNG: {exc}"))
                break
            except Exception as exc:
                log(f"❌ Đơn {id_don}: {core.rut_gon_text(exc, 300)}")
                for r in theo_don[id_don]:
                    tong.setdefault(r["_dong"], ("Lỗi", core.rut_gon_text(exc, 300)))
            if n % 10 == 0:
                luu()
            core.nghi_giua_don()
    finally:
        luu()
    bao_tien_do(len(ids), len(ids), "Xong")
    dem: dict[str, int] = {}
    for tt, _ in tong.values():
        dem[tt] = dem.get(tt, 0) + 1
    log("=" * 60)
    log("Kết quả: " + (", ".join(f"{k} {v}" for k, v in dem.items()) or "không có dòng nào") + f" → {path}")
    return dem
