# -*- coding: utf-8 -*-
"""
Khôi phục hồ sơ quét của 1 đơn bằng cách ĐẨY MỚI toàn bộ thành phần trong 1 request.

Dữ liệu dùng:
  - <thư mục tải>/<madon>/hosoquet.json : metadata HSQ gốc lúc tải (thongTinHoSoId, các thành phần).
  - Excel kết quả tải (ket_qua_tai_hsq_*.xlsx): file tương ứng từng thành phần
    (ưu tiên file_ky_so nếu có, không thì file_tai).

Payload theo mẫu "tạo mới" của cap_nhat_don_hsq.api_update_hosoquet_exist_file
(hoSoQuetId = 0 + thongTinHoSoId), nhưng count = số thành phần, mỗi thành phần
infoHoSoQuet_i + fileHoSoQuet_i.

Mặc định CHỈ IN KẾ HOẠCH. Thêm --thuc-hien mới gửi lên MPLIS.

  venv\\Scripts\\python.exe khoi_phuc_hsq.py 12830794
  venv\\Scripts\\python.exe khoi_phuc_hsq.py 12830794 --thuc-hien
"""

from __future__ import annotations

import argparse
import contextlib
import glob
import json
import os
import sys
import uuid
from typing import Any

import cap_nhat_don_hsq as core
import day_hsq_ky_so as day
import tai_hsq_don as thd

log = thd.log


def tim_excel_ket_qua(thu_muc: str, madon: str) -> str | None:
    """Excel kết quả mới nhất trong thư mục có chứa madon."""
    from openpyxl import load_workbook

    for path in sorted(glob.glob(os.path.join(thu_muc, "ket_qua_tai_hsq_*.xlsx")), reverse=True):
        wb = load_workbook(path, read_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
        if not rows:
            continue
        h = [str(x or "").strip().lower() for x in rows[0]]
        if "madon" in h and any(str(r[h.index("madon")]) == madon for r in rows[1:]):
            return path
    return None


def file_theo_thanh_phan(excel_path: str, madon: str) -> dict[int, str]:
    """thanhPhanHoSoQuetId → đường dẫn file sẽ đẩy (file_ky_so nếu tồn tại, không thì file_tai)."""
    wb, _ws, _col, rows = day.doc_excel(excel_path)
    wb.close()
    kq: dict[int, str] = {}
    for r in rows:
        if str(r.get("madon")) != madon:
            continue
        tp = core.safe_int(r.get("thanhphanhosoquetid"))
        for key in ("file_ky_so", "file_tai"):
            p = str(r.get(key) or "").strip()
            if p and os.path.isfile(p):
                kq[tp] = p
                break
    return kq


def info_tao_moi(node: dict[str, Any], da_ky_so: bool) -> dict[str, Any]:
    """Mẫu infoHoSoQuet tạo mới (như cap_nhat_don_hsq), giá trị lấy từ thành phần gốc."""
    gcn_id = node.get("giayChungNhanId")
    return {
        "loaiHoSoQuet": core.safe_int(node.get("loaiHoSoQuet"), 0),
        "laGiayToVeNguonGoc": bool(node.get("laGiayToVeNguonGoc")),
        "giayChungNhanId": str(gcn_id) if gcn_id else "",
        "moTa": node.get("moTa") or "",
        "tenGiayTo": node.get("tenGiayTo") or "",
        "trichYeu": node.get("trichYeu") or "",
        "versionGiayChungNhan": node.get("versionGiayChungNhan"),
        "laGiayChungNhan": bool(node.get("laGiayChungNhan")),
        "daKySo": bool(node.get("daKySo")) or da_ky_so,
        "__id": str(uuid.uuid4()),
        "files": None,
    }


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="Khôi phục HSQ của đơn bằng cách đẩy mới đủ thành phần")
    ap.add_argument("madon", help="tinhHinhDangKyId")
    ap.add_argument("--thu-muc", default="hsq_tai_ve", help="Thư mục tải (mặc định hsq_tai_ve)")
    ap.add_argument("--excel", help="Excel kết quả tải (mặc định: file mới nhất có madon)")
    ap.add_argument("--token")
    ap.add_argument("--cookie")
    ap.add_argument("--thuc-hien", action="store_true", help="Gửi thật lên MPLIS (không có = chỉ in kế hoạch)")
    ap.add_argument("--van-tao-moi", action="store_true",
                    help="Vẫn tạo mới dù đơn hiện còn HSQ có file (sẽ thành 2 HSQ)")
    args = ap.parse_args()

    madon = str(core.safe_int(args.madon))
    thu_muc = os.path.abspath(args.thu_muc)

    # ---- dữ liệu gốc
    json_goc = os.path.join(thu_muc, madon, "hosoquet.json")
    with open(json_goc, encoding="utf-8") as fp:
        hsq_goc = json.load(fp) or []
    excel = args.excel or tim_excel_ket_qua(thu_muc, madon)
    if not excel:
        sys.exit(f"Không tìm thấy Excel kết quả có madon {madon} trong {thu_muc}")
    file_map = file_theo_thanh_phan(excel, madon)
    log(f"HSQ gốc: {json_goc}")
    log(f"Excel: {excel}")

    # ---- session
    token = (args.token or "").strip() or input("__RequestVerificationToken: ").strip()
    cookie = (args.cookie or "").strip() or input("Cookie: ").strip()
    client = core.MplisClient(log)
    client.build_session_from_manual(token, cookie)

    # ---- hiện trạng trên MPLIS
    detail = client.lay_chi_tiet_don(int(madon))
    if detail is None:
        sys.exit("Không lấy được chi tiết đơn.")
    hien_tai = detail.get("ListHoSoQuet") or []
    so_file_hien_tai = len(thd.liet_ke_file_hsq(detail))
    log(f"Hiện trạng MPLIS: {len(hien_tai)} HSQ, {so_file_hien_tai} file")
    for hs in hien_tai:
        for f in thd.liet_ke_file_hsq({"ListHoSoQuet": [hs]}):
            log(f"   HSQ {hs.get('hoSoQuetId')}: {f.get('moTa')} | nodeId={f.get('nodeId')}")
    if so_file_hien_tai and not args.van_tao_moi:
        sys.exit("Đơn hiện còn HSQ có file → không tạo mới (thêm --van-tao-moi nếu chắc chắn).")

    # ---- kế hoạch: mỗi HSQ gốc → 1 request tạo mới đủ thành phần
    ke_hoach = []
    for hs in hsq_goc:
        nodes = (hs.get("ListFileHoSoQuet") or {}).get("ListFileHoSoQuet") or []
        tths_id = core.safe_int(hs.get("thongTinHoSoId")) or core.safe_int(core.lay_thong_tin_ho_so_id(detail))
        items = []
        for n in nodes:
            p = file_map.get(core.safe_int(n.get("thanhPhanHoSoQuetId")))
            if not p:
                sys.exit(f"Thiếu file cho thành phần '{n.get('moTa')}' ({n.get('thanhPhanHoSoQuetId')}).")
            items.append((info_tao_moi(n, day.pdf_da_ky_so(p)), p))
        ke_hoach.append(({"thongTinHoSoId": tths_id, "TuiHoSo": None, "tuiHoSoId": 0, "hoSoQuetId": 0}, items))

    for ho_so_quet, items in ke_hoach:
        log("=" * 60)
        log(f"Tạo mới HSQ: {json.dumps(ho_so_quet, ensure_ascii=False)} | count={len(items)}")
        for i, (info, p) in enumerate(items, 1):
            log(f"  infoHoSoQuet_{i}: loai={info['loaiHoSoQuet']} GCN={info['laGiayChungNhan']}"
                f"({info['giayChungNhanId']}) kySo={info['daKySo']} | {info['moTa'][:50]}")
            log(f"  fileHoSoQuet_{i}: {p} ({os.path.getsize(p) // 1024} KB)")

    if not args.thuc_hien:
        log("CHẠY THỬ — chưa gửi gì. Kiểm tra kế hoạch trên rồi chạy lại với --thuc-hien.")
        return

    # ---- gửi
    session = client._require_session()
    headers = dict(session.headers)
    headers.pop("Content-Type", None)
    for ho_so_quet, items in ke_hoach:
        data = {
            "hoSoQuet": json.dumps(ho_so_quet, ensure_ascii=False),
            "count": str(len(items)),
            "isLuuKhoHoSoQuet": "false",
        }
        with contextlib.ExitStack() as stack:
            files = {}
            for i, (info, p) in enumerate(items, 1):
                data[f"infoHoSoQuet_{i}"] = json.dumps(info, ensure_ascii=False)
                files[f"fileHoSoQuet_{i}"] = (os.path.basename(p), stack.enter_context(open(p, "rb")),
                                              "application/pdf")
            response = session.post(core.URL_UPDATE_HOSOQUET, data=data, files=files, headers=headers,
                                    timeout=core.TIMEOUT, allow_redirects=False)
        result = client._response_json(response, "UpdateHoSoQuetExistFile (khôi phục)")
        log(f"Kết quả: {core.rut_gon_text(result, 500)}")
        if not isinstance(result, dict) or not result.get("success"):
            sys.exit("❌ Khôi phục không thành công.")

    # ---- xác nhận
    detail = client.lay_chi_tiet_don(int(madon)) or {}
    files_moi = thd.liet_ke_file_hsq(detail)
    log(f"Sau khôi phục: {len(detail.get('ListHoSoQuet') or [])} HSQ, {len(files_moi)} file")
    for f in files_moi:
        log(f"   {f.get('moTa')} | loai={f.get('loaiHoSoQuet')} | nodeId={f.get('nodeId')}")
    tong_goc = sum(len(it) for _, it in ke_hoach)
    log("✅ Đủ thành phần." if len(files_moi) >= tong_goc
        else f"❌ Mới có {len(files_moi)}/{tong_goc} thành phần — gửi log cho mình.")


if __name__ == "__main__":
    main()
