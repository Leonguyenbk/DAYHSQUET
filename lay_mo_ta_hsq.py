# -*- coding: utf-8 -*-
"""
Lấy mô tả (moTa) của tất cả file hồ sơ quét trong các đơn đăng ký → xuất Excel.

Với mỗi tinhHinhDangKyId: gọi GetThongTinDangKyByTinhHinhDangKyIds (getHoSoQuet=True),
duyệt ListHoSoQuet / ListFileHoSoQuet, ghi mỗi file 1 dòng. Không tải file.

Excel kết quả có 2 sheet:
  - ChiTiet : mỗi file HSQ 1 dòng (madon, tờ, thửa, xã, GCN của đơn, hosoquetid, vị trí, mô tả, loại...)
  - TongHop : các mô tả khác nhau, số lần xuất hiện, số đơn, danh sách mã đơn

Cách dùng:
  python lay_mo_ta_hsq.py 13456477 12913166
  python lay_mo_ta_hsq.py --excel ds_don.xlsx      # cột madon/tinhhinhdangkyid, không có thì cột A
  python lay_mo_ta_hsq.py --txt ds_don.txt         # mỗi dòng 1 ID (hoặc cách nhau dấu phẩy)
  python lay_mo_ta_hsq.py --token XXX --cookie "a=1; b=2"
  python lay_mo_ta_hsq.py --chrome --user abc
Chạy không tham số (python lay_mo_ta_hsq.py) → mở giao diện lay_mo_ta_hsq_flet.py.
"""

import argparse
import getpass
from datetime import datetime, timedelta, timezone
import os
import re
import sys
import time
from collections import OrderedDict
from typing import Any

import cap_nhat_don_hsq as core
import tai_hsq_don as tai

log = tai.log

HEADERS = ["stt", "madon", "soto", "sothua", "xaid",
           "gcn_sophathanh", "gcn_sovaoso", "gcn_mavach", "gcn_ngayvaoso", "hosoquetid", "vitri", "mota",
           "loaihosoquet", "giaychungnhanid", "dakyso", "nodeid", "ghichu"]
WIDTHS = {"stt": 6, "madon": 12, "soto": 7, "sothua": 7, "xaid": 8,
          "gcn_sophathanh": 20, "gcn_sovaoso": 16, "gcn_mavach": 18, "gcn_ngayvaoso": 14, "hosoquetid": 11, "vitri": 6,
          "mota": 55, "loaihosoquet": 22, "giaychungnhanid": 13, "dakyso": 8, "nodeid": 38, "ghichu": 40}


def ngay_mplis(value: Any) -> str:
    """/Date(1543942800000)/ hoặc ISO → DD/MM/YYYY (giờ Việt Nam)."""
    s = str(value or "").strip()
    m = re.fullmatch(r"/Date\((-?\d+)(?:[+-]\d{4})?\)/", s)
    try:
        if m:
            dt = datetime.fromtimestamp(int(m.group(1)) / 1000, tz=timezone.utc)
        else:
            dt = datetime.fromisoformat(s[:-1] + "+00:00" if s.endswith("Z") else s)
    except ValueError:
        return s
    if dt.tzinfo is not None:
        dt = dt.astimezone(timezone(timedelta(hours=7)))
    return dt.strftime("%d/%m/%Y")


def thong_tin_gcn(detail: dict[str, Any]) -> dict[str, str]:
    """GCN hiện hành của đơn (bỏ isLastest = False); nhiều GCN thì nối bằng '; '."""
    ds = [g for g in detail.get("ListGiayChungNhan") or []
          if isinstance(g, dict) and g.get("isLastest") is not False]

    def noi(lay) -> str:
        return "; ".join(str(lay(g) or "").strip() for g in ds) if ds else ""

    return {
        "gcn_sophathanh": noi(lambda g: g.get("soPhatHanh")),
        "gcn_sovaoso": noi(lambda g: g.get("soVaoSo")),
        "gcn_mavach": noi(lambda g: g.get("maVach")),
        "gcn_ngayvaoso": noi(lambda g: ngay_mplis(g.get("ngayVaoSo"))),
    }


def dong_tu_don(id_don: int, detail: dict[str, Any]) -> list[dict[str, Any]]:
    thdk = detail.get("TinhHinhDangKy") or {}
    chung = {"madon": id_don, "soto": thdk.get("soHieuToBanDo", ""),
             "sothua": thdk.get("soThuTuThua", ""), "xaid": thdk.get("xaId", ""),
             **thong_tin_gcn(detail)}
    files = tai.liet_ke_file_hsq(detail)
    if not files:
        return [{**chung, "ghichu": "Không có HSQ"}]
    rows = []
    for f in files:
        loai = core.safe_int(f.get("loaiHoSoQuet"), -1)
        rows.append({
            **chung,
            "hosoquetid": f.get("hoSoQuetId"),
            "vitri": f.get("_vi_tri"),
            "mota": f.get("moTa") or f.get("tenGiayTo") or "",
            "loaihosoquet": core.LOAI_HO_SO_QUET_OPTIONS.get(loai, f.get("loaiHoSoQuet")),
            "giaychungnhanid": f.get("giayChungNhanId"),
            "dakyso": "Có" if f.get("daKySo") else "",
            "nodeid": f.get("nodeId") or "",
        })
    return rows


def tong_hop_mo_ta(rows: list[dict[str, Any]]) -> list[list[Any]]:
    nhom: "OrderedDict[str, dict]" = OrderedDict()
    for r in rows:
        if r.get("ghichu"):
            continue
        mo_ta = str(r.get("mota") or "").strip() or "(trống)"
        g = nhom.setdefault(mo_ta, {"so_lan": 0, "don": []})
        g["so_lan"] += 1
        if r["madon"] not in g["don"]:
            g["don"].append(r["madon"])
    ket_qua = [[mt, g["so_lan"], len(g["don"]), ", ".join(map(str, g["don"]))] for mt, g in nhom.items()]
    return sorted(ket_qua, key=lambda x: (-x[1], x[0].lower()))


def ghi_excel(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def tieu_de(ws):
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "ChiTiet"
    ws.append(HEADERS)
    tieu_de(ws)
    vang = PatternFill("solid", fgColor="FFF2CC")
    for stt, r in enumerate(rows, 1):
        ws.append([stt if h == "stt" else r.get(h, "") for h in HEADERS])
        if r.get("ghichu"):
            ws.cell(row=ws.max_row, column=HEADERS.index("ghichu") + 1).fill = vang
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 12)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 1)}"

    ws2 = wb.create_sheet("TongHop")
    ws2.append(["mota", "so_file", "so_don", "ds_madon"])
    tieu_de(ws2)
    for dong in tong_hop_mo_ta(rows):
        ws2.append(dong)
    for col, w in zip("ABCD", (55, 9, 9, 80)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:D{max(ws2.max_row, 1)}"

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def ten_excel_mac_dinh(thu_muc: str = tai.THU_MUC_APP) -> str:
    return os.path.join(thu_muc, f"mo_ta_hsq_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")


def chay(client, ids: list[int], out_path: str, nen_dung=lambda: False, bao_tien_do=None) -> str:
    """Lấy mô tả HSQ của từng đơn → ghi Excel out_path (lưu dần mỗi 20 đơn). Trả về out_path."""
    log(f"POST {core.URL_CHI_TIET}")
    log('   payload: {"getHoSoQuet": true, "tinhHinhDangKyIds": [<id>]}')
    rows: list[dict[str, Any]] = []
    try:
        for n, id_don in enumerate(ids, 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            if bao_tien_do:
                bao_tien_do(n - 1, len(ids), f"Đơn {id_don} ({n}/{len(ids)})")
            try:
                detail = client.lay_chi_tiet_don(id_don)
                if detail is None:
                    dong = [{"madon": id_don, "ghichu": "Không lấy được chi tiết đơn"}]
                else:
                    dong = dong_tu_don(id_don, detail)
            except core.HetPhien as exc:
                rows.append({"madon": id_don, "ghichu": f"Lỗi: {core.rut_gon_text(exc, 300)}"})
                log(f"❌ {exc} — ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                break
            except Exception as exc:
                dong = [{"madon": id_don, "ghichu": f"Lỗi: {core.rut_gon_text(exc, 300)}"}]
            rows += dong
            so_file = sum(1 for d in dong if not d.get("ghichu"))
            ghi_chu = dong[0].get("ghichu")
            log(("❌ " if ghi_chu and ghi_chu != "Không có HSQ" else "⚠ " if ghi_chu else "")
                + f"[{n}/{len(ids)}] Đơn {id_don}: {so_file} file" + (f" — {ghi_chu}" if ghi_chu else ""))
            for d in dong:
                if not d.get("ghichu"):
                    log(f"     {d['vitri']:>2}. [{d['loaihosoquet']}] {d['mota'] or '(trống)'}")
            if n % 20 == 0:
                _ghi_an_toan(out_path, rows)  # lưu dần, lỡ dừng giữa chừng vẫn có dữ liệu
            time.sleep(core.REQUEST_DELAY_SECONDS)
    finally:
        if rows:
            _ghi_an_toan(out_path, rows)

    if bao_tien_do:
        bao_tien_do(len({r["madon"] for r in rows}), len(ids), "Hoàn tất")
    so_file = sum(1 for r in rows if not r.get("ghichu"))
    log("=" * 50)
    log(f"✅ {len({r['madon'] for r in rows})} đơn | {so_file} file HSQ "
        f"| {len(tong_hop_mo_ta(rows))} mô tả khác nhau")
    khong_hsq = [str(r["madon"]) for r in rows if r.get("ghichu") == "Không có HSQ"]
    if khong_hsq:
        log(f"⚠ Đơn không có HSQ: {', '.join(khong_hsq)}")
    loi = [str(r["madon"]) for r in rows if r.get("ghichu") and r.get("ghichu") != "Không có HSQ"]
    if loi:
        log(f"❌ Đơn lỗi: {', '.join(loi)}")
    if rows:
        log(f"Excel kết quả: {out_path}")
    return out_path


def _ghi_an_toan(path: str, rows: list[dict[str, Any]]) -> None:
    try:
        ghi_excel(path, rows)
    except PermissionError:
        log(f"⚠ Không ghi được Excel (file đang mở?): {path}")


def doc_id_tu_txt(path: str) -> list[str]:
    with open(path, encoding="utf-8-sig") as fp:
        return [fp.read()]


def nhap_id_tu_ban_phim() -> list[str]:
    print("Dán danh sách tinhHinhDangKyId (mỗi dòng 1 ID hoặc cách nhau dấu phẩy), dòng trống để kết thúc:")
    dong = []
    while True:
        s = input().strip()
        if not s:
            break
        dong.append(s)
    return ["\n".join(dong)]


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="Lấy mô tả tất cả hồ sơ quét của các đơn MPLIS ra Excel")
    ap.add_argument("ids", nargs="*", help="tinhHinhDangKyId (hoặc chuỗi chứa mã đơn)")
    ap.add_argument("--excel", help="File Excel có cột madon/tinhhinhdangkyid")
    ap.add_argument("--txt", help="File text chứa danh sách ID")
    ap.add_argument("--out", help="File Excel kết quả (mặc định: mo_ta_hsq_<thời gian>.xlsx)")
    ap.add_argument("--user", help="Username đăng nhập MPLIS")
    ap.add_argument("--token", help="__RequestVerificationToken copy từ F12")
    ap.add_argument("--cookie", help="Chuỗi Cookie copy từ F12 (dùng cùng --token)")
    ap.add_argument("--chrome", action="store_true", help="Đăng nhập bằng Chrome, không hỏi token/cookie")
    args = ap.parse_args()

    raw_ids = list(args.ids)
    if args.excel:
        raw_ids += tai.doc_id_tu_excel(args.excel)
    if args.txt:
        raw_ids += doc_id_tu_txt(args.txt)
    if not raw_ids:
        raw_ids = nhap_id_tu_ban_phim()
    ids = tai.tach_danh_sach_id([x for s in raw_ids for x in re.split(r"[\r\n,;\t ]+", str(s))])
    if not ids:
        ap.error("Chưa có ID đơn nào.")
    log(f"Có {len(ids)} đơn cần lấy mô tả HSQ.")

    out_path = os.path.abspath(args.out or ten_excel_mac_dinh())

    token, cookie = (args.token or "").strip(), (args.cookie or "").strip()
    if not args.chrome and not (token and cookie):
        print("Chọn cách đăng nhập:")
        print("  1. Dán token + cookie (copy từ F12 > Network > Request Headers)")
        print("  2. Mở Chrome đăng nhập bằng username/password")
        if (input("Chọn [1]: ").strip() or "1") == "1":
            token = token or input("__RequestVerificationToken: ").strip()
            cookie = cookie or input("Cookie: ").strip()

    client = core.MplisClient(log)
    try:
        if token and cookie:
            client.build_session_from_manual(token, cookie)
        else:
            username = args.user or input("Username: ").strip()
            password = getpass.getpass("Password: ")
            client.driver = tai.mo_chrome(os.path.dirname(out_path))
            tai.dang_nhap(client.driver, username, password)
            input(">>> Đăng nhập xong trên Chrome (OTP/captcha nếu có) rồi nhấn Enter...")
            client.build_session_from_browser()

        chay(client, ids, out_path)
    except KeyboardInterrupt:
        log("⏹ Dừng theo yêu cầu.")
    finally:
        client.close_browser()


if __name__ == "__main__":
    if len(sys.argv) == 1:
        # Không truyền tham số → mở giao diện (nhập token/cookie, ID trên UI).
        import lay_mo_ta_hsq_flet

        lay_mo_ta_hsq_flet.main()
    else:
        main()
