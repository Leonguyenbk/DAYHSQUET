# -*- coding: utf-8 -*-
"""
Đẩy HSQ theo mô tả, mỗi đơn 1 thư mục — file đuôi "-GCN" gán cho GCN có số nằm trong tên file.

Cấu trúc thư mục:
    <thư mục gốc>/
        13456477/                          ← tên thư mục bắt đầu bằng mã đơn (tinhHinhDangKyId)
            Đơn đăng ký.pdf                ← tên file = mô tả (moTa) thành phần HSQ → thay file thành phần đó
            03_Trích lục thửa đất.pdf      ← file tải bằng TaiDayHSQ ("NN_<mô tả>") cũng được
            PC_A_DH 770214_DH 770215-GCN.pdf ← đuôi "-GCN": gán cho GCN có số nằm trong tên (ở đây 2 GCN)
            Ban sao-GCN DH 770214.pdf      ← "-GCN <số phát hành>": chỉ gán cho GCN đó

Với mỗi đơn:
  1. Lấy chi tiết đơn → thành phần HSQ + GCN hiện hành (bỏ isLastest = False).
  2. File thường: khớp mô tả như day_hsq_theo_mota (bỏ dấu / ký tự đặc biệt / đuôi ký số, tiền tố "NN_",
     tên file chứa nguyên mô tả ≥ 8 ký tự) → thay file, giữ nguyên loại HSQ / GCN.
  3. File "-GCN": chỉ gán cho GCN hiện hành có SỐ PHÁT HÀNH nằm trong tên file (tên không có số GCN nào
     của đơn → không đẩy; cùng 1 GCN nhiều file → lấy file đã ký số nếu chỉ 1 file ký):
       a. thành phần đã gắn GCN đó                         → thay file (ghi nhận lại loại GCN);
       b. thành phần có mô tả trùng tên file (bỏ "-GCN")   → thay file + ghi nhận Giấy chứng nhận + gắn GCN;
       c. không có                                          → thêm thành phần "Giấy chứng nhận <số>" vào HSQ
                                                             chứa thành phần ở (b), không có thì HSQ đầu tiên;
       d. đơn chưa có HSQ                                   → tạo HSQ mới gồm các thành phần GCN.
     Thành phần đã nhận file "-GCN" thì file thường cùng mô tả bị bỏ qua.
  4. Mỗi HSQ gửi 1 request kèm ĐỦ thành phần (giữ nguyên thành phần không đổi). HSQ không có thành phần loại
     Giấy chứng nhận / Đơn đăng ký / Thông báo xác nhận → không gửi (MPLIS sẽ từ chối), báo trong Excel.
     Tuỳ chọn giay_to: HSQ có đẩy GCN → mọi thành phần khác của HSQ đó chuyển loại Giấy tờ (giữ thành phần
     GCN gắn GCN hiện hành chưa có file lần này); HSQ không đẩy GCN → không đụng.
  5. Xác nhận: số thành phần không giảm (giảm → DỪNG TOÀN BỘ), thành phần thay có nodeId mới, mỗi GCN có
     thành phần loại Giấy chứng nhận gắn đúng GCN, thành phần chuyển Giấy tờ đã đúng loại.
"""

from __future__ import annotations

import copy
import json
import os
import time
from typing import Any, Callable

import bo_sung_gcn_hsq as bs
import cap_nhat_don_hsq as core
import nhom1_mplis as nhom1
import day_gcn_theo_thu_muc as dg
import day_hsq_ky_so as day
import day_hsq_theo_mota as dm
import tai_hsq_don as thd

log = thd.log

EXCEL_HEADERS = ["stt", "madon", "thu_muc", "file", "hosoquetid", "vitri", "mota", "loaihosoquet", "so_gcn",
                 "viec", "da_ky_so", "trang_thai", "ghi_chu"]
EXCEL_WIDTHS = {"stt": 6, "madon": 12, "thu_muc": 40, "file": 45, "hosoquetid": 11, "vitri": 6, "mota": 45,
                "loaihosoquet": 18, "so_gcn": 14, "viec": 30, "da_ky_so": 8, "trang_thai": 13, "ghi_chu": 60}
MAU_TRANG_THAI = {"Đã đẩy": "C6EFCE", "Lỗi": "F4CCCC", "Bỏ qua": "FFF2CC", "Chạy thử": "D9EAF7",
                  "Giữ nguyên": "EDEDED"}


# ============================ PHÂN LOẠI FILE ============================

def tach_file(pdfs: list[str]) -> tuple[list[str], list[tuple[str, str, str, int]]]:
    """→ (file thường, [(file "-GCN", tên bỏ hậu tố, số ghi trong tên, tiền tố stt)])."""
    thuong, gcn = [], []
    for p in pdfs:
        stt, ten = dm.tach_ten_file(p)
        hau_to = dm.tach_hau_to_gcn(ten)
        if hau_to is None:
            thuong.append(p)
        else:
            gcn.append((p, hau_to[0], hau_to[1], stt))
    return thuong, gcn


def _mot_file(files: list[str]) -> tuple[str | None, str]:
    if len(files) == 1:
        return files[0], ""
    ky = [p for p in files if day.pdf_da_ky_so(p)]
    if len(ky) == 1:
        return ky[0], ""
    return None, f"{len(files)} file cùng khớp ({', '.join(os.path.basename(p) for p in files)}) — giữ lại 1 file"


def chia_file_gcn(ds_gcn: list[dict], file_gcn: list[tuple[str, str, str, int]],
                  ) -> tuple[dict[int, tuple[str, str, int]], list[tuple[str, str]]]:
    """
    {gcnId: (file, tên bỏ hậu tố, stt)} — file "-GCN" chỉ gán cho GCN có SỐ PHÁT HÀNH nằm trong tên file
    (vd 'PC_A_DH 770214-GCN', 'X-GCN DH 770214'); tên chứa nhiều số → gán cho từng GCN đó. Tên file không
    có số GCN hiện hành nào → không đẩy. Trả thêm [(file, lý do bỏ qua)].
    """
    bo_qua: list[tuple[str, str]] = []
    ten_ds = "; ".join(str(g.get("soPhatHanh") or g.get("giayChungNhanId")) for g in ds_gcn)
    so_gcn = {core.safe_int(g.get("giayChungNhanId")): dm.chuan_hoa(g.get("soPhatHanh")) for g in ds_gcn}
    theo_gcn: dict[int, list[tuple[str, str, int]]] = {}
    for p, ten, so, stt in file_gcn:
        if so:  # "-GCN <số>": số ghi sau hậu tố phải đúng 1 GCN hiện hành
            khop = [gid for gid, s_ in so_gcn.items() if s_ == dm.chuan_hoa(so)]
            ly_do = f"Tên file ghi GCN '{so}' nhưng GCN hiện hành của đơn là: {ten_ds}"
        else:
            khoa = dm.chuan_hoa(ten)
            khop = [gid for gid, s_ in so_gcn.items() if len(s_) >= 5 and s_ in khoa]
            ly_do = (f"Tên file không chứa số GCN hiện hành nào của đơn ({ten_ds}) — đổi tên có số GCN, "
                     f"vd '{ten} {ds_gcn[0].get('soPhatHanh')}-GCN.pdf'")
        if not khop:
            bo_qua.append((p, ly_do))
        for gid in khop:
            theo_gcn.setdefault(gid, []).append((p, ten, stt))

    kq: dict[int, tuple[str, str, int]] = {}
    for gid, ds in theo_gcn.items():
        p, ly_do = _mot_file([x[0] for x in ds])
        if p is None:
            bo_qua += [(x[0], f"GCN {ds_gcn_so(ds_gcn, gid)}: {ly_do}") for x in ds]
        else:
            kq[gid] = next(x for x in ds if x[0] == p)
            bo_qua += [(x[0], f"Đã dùng {os.path.basename(p)} cho GCN {ds_gcn_so(ds_gcn, gid)}")
                       for x in ds if x[0] != p]
    return kq, bo_qua


def ds_gcn_so(ds_gcn: list[dict], gid: int) -> str:
    return next((str(g.get("soPhatHanh")) for g in ds_gcn if core.safe_int(g.get("giayChungNhanId")) == gid), "")


# ============================ KẾ HOẠCH 1 ĐƠN ============================

def _viec(tp: dict | None, path: str = "", gcn: dict | None = None, kieu: str = "thay",
          hoso: dict | None = None) -> dict[str, Any]:
    """kieu: thay | thay_gcn | gcn (chỉ ghi nhận loại) | them (thành phần GCN mới) | tao (HSQ mới)."""
    return {"tp": tp, "hoso": tp["hoso"] if tp else hoso, "path": path,
            "da_ky": bool(path) and day.pdf_da_ky_so(path), "gcn": gcn, "kieu": kieu}


def mo_ta_viec(v: dict[str, Any]) -> str:
    so = (v.get("gcn") or {}).get("soPhatHanh")
    return {"thay": "Thay file", "thay_gcn": f"Thay file + ghi nhận GCN {so}", "gcn": f"Ghi nhận GCN {so}",
            "them": f"Thêm thành phần GCN {so}", "tao": f"Tạo HSQ mới: GCN {so}"}[v["kieu"]]


def lap_ke_hoach(detail: dict[str, Any], pdfs: list[str]) -> tuple[list[dict], list[tuple[str, dict | None, str]]]:
    """→ (việc cần làm, [(file, thành phần hoặc None, lý do bỏ qua)])."""
    ds_tp = dm.thanh_phan_cua_don(detail)
    ds_gcn = bs.gcn_hien_hanh(detail)
    hsq_list = [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]
    thuong, file_gcn = tach_file(pdfs)
    viec: list[dict] = []
    bo_qua: list[tuple[str, dict | None, str]] = []

    # ---- file "-GCN" → GCN có số nằm trong tên file
    da_dung: set[int] = set()  # stt thành phần đã nhận file -GCN
    if file_gcn and not ds_gcn:
        bo_qua += [(x[0], None, "File -GCN nhưng đơn chưa có GCN hiện hành") for x in file_gcn]
    elif file_gcn:
        gan, bq = chia_file_gcn(ds_gcn, file_gcn)
        bo_qua += [(p, None, ly_do) for p, ly_do in bq]
        for gcn in ds_gcn:
            gid = core.safe_int(gcn.get("giayChungNhanId"))
            if gid not in gan:
                bo_qua.append(("", None, f"GCN {gcn.get('soPhatHanh')}: không có file -GCN nào có số GCN này trong tên"))
                continue
            path, ten, stt = gan[gid]
            if not hsq_list:
                viec.append(_viec(None, path, gcn, "tao"))
                continue
            tp_khop = dm._khop_ten(stt, ten, ten, ds_tp)[0]
            tp = next((t for t in ds_tp if t["stt"] not in da_dung
                       and core.safe_int(t["node"].get("giayChungNhanId")) == gid), None)
            if tp is not None:
                viec.append(_viec(tp, path, gcn, "thay_gcn"))
            elif tp_khop is not None and tp_khop["stt"] not in da_dung:
                tp = tp_khop
                viec.append(_viec(tp, path, gcn, "thay_gcn"))
            else:
                viec.append(_viec(None, path, gcn, "them", tp_khop["hoso"] if tp_khop else hsq_list[0]))
            if tp is not None:
                da_dung.add(tp["stt"])

    # ---- file thường → thành phần theo mô tả
    if thuong and not ds_tp:
        bo_qua += [(p, None, "Đơn chưa có HSQ — file thường không có thành phần để thay") for p in thuong]
    theo_tp: dict[int, list[str]] = {}
    for p in thuong if ds_tp else []:
        tp, ly_do, _ = dm.khop_file(p, ds_tp)
        if tp is None:
            bo_qua.append((p, None, ly_do))
        elif tp["stt"] in da_dung:
            bo_qua.append((p, tp, "Thành phần này đã nhận file -GCN"))
        else:
            theo_tp.setdefault(tp["stt"], []).append(p)
    for stt, files in theo_tp.items():
        tp = ds_tp[stt - 1]
        p, ly_do = _mot_file(files)
        if p is None:
            bo_qua += [(f, tp, ly_do) for f in files]
            continue
        bo_qua += [(f, tp, f"Đã dùng file {os.path.basename(p)} cho thành phần này") for f in files if f != p]
        viec.append(_viec(tp, p))

    # ---- mô tả có "GCN" + số GCN hiện hành (GCN chưa có file -GCN) → ghi nhận loại, như day_hsq_theo_mota
    gcn_da_gan = {core.safe_int((v["gcn"] or {}).get("giayChungNhanId")) for v in viec if v["gcn"]}
    hsq_co_viec = {id(v["hoso"]) for v in viec if v["hoso"] is not None}
    theo_stt = {v["tp"]["stt"]: v for v in viec if v["tp"]}
    for stt_, gcn in _bo_sung_tu_mo_ta(detail, ds_tp).items():
        tp = ds_tp[stt_ - 1]
        if core.safe_int(gcn.get("giayChungNhanId")) in gcn_da_gan or id(tp["hoso"]) not in hsq_co_viec:
            continue
        if stt_ in theo_stt:
            theo_stt[stt_].update(gcn=gcn, kieu="thay_gcn")
        else:
            viec.append(_viec(tp, "", gcn, "gcn"))
    return viec, bo_qua


def _bo_sung_tu_mo_ta(detail: dict[str, Any], ds_tp: list[dict]) -> dict[int, dict]:
    kq = {}
    for hoso in {id(t["hoso"]): t["hoso"] for t in ds_tp}.values():
        tp_hsq = [t for t in ds_tp if t["hoso"] is hoso]
        for vi_tri, gcn in dm.gcn_cho_bo_sung(detail, tp_hsq).items():
            kq[next(t["stt"] for t in tp_hsq if t["vi_tri"] == vi_tri)] = gcn
    return kq


LOAI_GIAY_TO = 0


def thanh_phan_sang_giay_to(hoso: dict[str, Any], ds: list[dict], gcn_giu: set[int]) -> list[dict[str, Any]]:
    """
    HSQ có đẩy GCN lần này → mọi thành phần khác chuyển loại Giấy tờ. Giữ nguyên: thành phần nhận GCN lần này,
    thành phần GCN gắn GCN hiện hành chưa có file lần này (gcn_giu), thành phần đã là Giấy tờ.
    HSQ không đẩy GCN → [] (không đụng). → [{"hoso", "node", "vi_tri"}].
    """
    if not any(v["gcn"] for v in ds):
        return []
    nhan_gcn = {v["tp"]["vi_tri"] for v in ds if v["tp"] and v["gcn"]}
    kq = []
    for i, node in enumerate(day.cac_node(hoso)):
        loai = core.safe_int(node.get("loaiHoSoQuet"), -1)
        if i in nhan_gcn or (loai == 1 and core.safe_int(node.get("giayChungNhanId")) in gcn_giu):
            continue
        if loai == LOAI_GIAY_TO and not node.get("laGiayChungNhan"):
            continue
        kq.append({"hoso": hoso, "node": node, "vi_tri": i})
    return kq


def items_hsq(hoso: dict[str, Any], ds: list[dict],
              giay_to: list[dict] | None = None) -> list[tuple[dict, str | None]]:
    """ĐỦ thành phần hiện có (thay / ghi nhận loại / giữ nguyên / sang Giấy tờ) + thành phần GCN mới, theo thứ tự."""
    hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
    theo_vi_tri = {v["tp"]["vi_tri"]: v for v in ds if v["tp"]}
    sang_giay_to = {t["vi_tri"] for t in giay_to or []}
    items: list[tuple[dict, str | None]] = []
    for i, node in enumerate(day.cac_node(hoso)):
        v = theo_vi_tri.get(i)
        if v is None:
            info = day.info_giu_nguyen(node)
            if i in sang_giay_to:
                info.update(loaiHoSoQuet=LOAI_GIAY_TO, laGiayChungNhan=False)
            items.append((info, None))
            continue
        info = day.info_thay_file(node, hsq_id, v["da_ky"]) if v["path"] else day.info_giu_nguyen(node)
        if v["gcn"]:
            dm.gan_loai_gcn(info, v["gcn"])
        elif i in sang_giay_to:
            info.update(loaiHoSoQuet=LOAI_GIAY_TO, laGiayChungNhan=False)
        items.append((info, v["path"] or None))
    items += [(dg.info_gcn_moi(v["gcn"], v["da_ky"]), v["path"]) for v in ds if v["kieu"] == "them"]
    return items


# ============================ XỬ LÝ 1 ĐƠN ============================

def _dong(madon: int, path: str = "", tp: dict | None = None, v: dict | None = None, **kw) -> dict[str, Any]:
    v = v or {}
    tp = tp or v.get("tp")
    path = path or v.get("path") or ""
    node = (tp or {}).get("node") or {}
    gcn = v.get("gcn") or {}
    hoso = (tp or {}).get("hoso") or v.get("hoso") or {}
    loai = core.safe_int(node.get("loaiHoSoQuet"), -1)
    return {
        "madon": madon,
        "thu_muc": os.path.dirname(path) if path else "",
        "file": os.path.basename(path) if path else "",
        "hosoquetid": hoso.get("hoSoQuetId", ""),
        "vitri": tp["stt"] if tp else "",
        "mota": node.get("moTa") or (bs.mo_ta_gcn(gcn) if v.get("kieu") in ("them", "tao") else ""),
        "loaihosoquet": core.LOAI_HO_SO_QUET_OPTIONS.get(loai, node.get("loaiHoSoQuet") or "") if tp else "",
        "so_gcn": gcn.get("soPhatHanh", ""),
        "viec": mo_ta_viec(v) if v else "",
        "da_ky_so": "Có" if v.get("da_ky") else "",
        **kw,
    }


def _kiem_tra_loai_bat_buoc(hoso: dict[str, Any], ds: list[dict]) -> bool:
    gan = {v["tp"]["vi_tri"] for v in ds if v["tp"] and v["gcn"]}
    return any(v["kieu"] == "them" for v in ds) or bool(gan) or any(
        core.safe_int(n.get("loaiHoSoQuet"), -1) in dm.LOAI_BAT_BUOC for n in day.cac_node(hoso))


def _dong_giay_to(madon: int, t: dict[str, Any], **kw) -> dict[str, Any]:
    node = t["node"]
    loai = core.safe_int(node.get("loaiHoSoQuet"), -1)
    return {"madon": madon, "hosoquetid": t["hoso"].get("hoSoQuetId", ""), "vitri": t["vi_tri"] + 1,
            "mota": node.get("moTa") or "", "loaihosoquet": core.LOAI_HO_SO_QUET_OPTIONS.get(loai, ""),
            "viec": "Chuyển Giấy tờ", **kw}


def xu_ly_don(client: core.MplisClient, madon: int, pdfs: list[str], chay_thu: bool,
              giay_to: bool = False, bo_qua_nhom_1: bool = True) -> list[dict[str, Any]]:
    if not pdfs:
        return [_dong(madon, trang_thai="Bỏ qua", ghi_chu="Thư mục đơn không có file PDF")]
    detail = client.lay_chi_tiet_don(madon)
    if detail is None:
        return [_dong(madon, p, trang_thai="Lỗi", ghi_chu=dm.LOI_CHI_TIET) for p in pdfs]
    if bo_qua_nhom_1 and nhom1.da_dat(client, madon, detail):
        return [_dong(madon, p, trang_thai="Bỏ qua", ghi_chu=nhom1.LY_DO_DA_DAT) for p in pdfs]

    viec, bo_qua = lap_ke_hoach(detail, pdfs)
    rows = [_dong(madon, p, tp, trang_thai="Bỏ qua", ghi_chu=ly_do) for p, tp, ly_do in bo_qua]
    for p, _, ly_do in bo_qua:
        log(f"   ⚠ {os.path.basename(p) + ': ' if p else ''}{ly_do}")

    # ---- nhóm theo HSQ (None = tạo HSQ mới)
    nhom: list[tuple[dict | None, list[dict]]] = []
    tao = [v for v in viec if v["kieu"] == "tao"]
    if tao:
        nhom.append((None, tao))
    for h in [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]:
        ds = [v for v in viec if v["hoso"] is h and v["kieu"] != "tao"]
        if not ds:
            continue
        if not _kiem_tra_loai_bat_buoc(h, ds):
            ly_do = ("HSQ chưa có thành phần loại Giấy chứng nhận/Đơn đăng ký/Thông báo xác nhận → MPLIS sẽ từ "
                     "chối; đặt tên file GCN '<tên> <số GCN>-GCN.pdf' rồi chạy lại")
            log(f"   ⚠ HSQ {h.get('hoSoQuetId')}: {ly_do}")
            rows += [_dong(madon, v=v, trang_thai="Bỏ qua", ghi_chu=ly_do) for v in ds]
            continue
        nhom.append((h, ds))

    # GCN hiện hành không có file lần này: thành phần GCN đang gắn nó vẫn giữ loại GCN khi chuyển Giấy tờ.
    gcn_day = {core.safe_int(v["gcn"].get("giayChungNhanId")) for v in viec if v["gcn"]}
    gcn_giu = {core.safe_int(g.get("giayChungNhanId")) for g in bs.gcn_hien_hanh(detail)} - gcn_day
    da_chuyen: set[tuple[int, int]] = set()  # (id HSQ, vị trí) chuyển Giấy tờ
    for hoso, ds in nhom:
        log(f"   {'HSQ MỚI' if hoso is None else 'HSQ ' + str(hoso.get('hoSoQuetId'))}:")
        for v in ds:
            log(f"     {mo_ta_viec(v)}"
                + (f" ← {os.path.basename(v['path'])}{' (đã ký số)' if v['da_ky'] else ''}" if v["path"] else "")
                + (f"  —  {v['tp']['node'].get('moTa')}" if v["tp"] else ""))
        giay_to_hsq = thanh_phan_sang_giay_to(hoso, ds, gcn_giu) if giay_to and hoso is not None else []
        for t in giay_to_hsq:
            log(f"     Chuyển Giấy tờ  —  {t['node'].get('moTa')} "
                f"[{core.LOAI_HO_SO_QUET_OPTIONS.get(core.safe_int(t['node'].get('loaiHoSoQuet'), -1), 'chưa có loại')}]")
            da_chuyen.add((id(hoso), t["vi_tri"]))
        if giay_to and hoso is not None and not any(v["gcn"] for v in ds):
            log("     (HSQ không đẩy GCN → không chuyển Giấy tờ)")
        if chay_thu:
            rows += [_dong(madon, v=v, trang_thai="Chạy thử", ghi_chu=f"Sẽ {mo_ta_viec(v).lower()}") for v in ds]
            rows += [_dong_giay_to(madon, t, trang_thai="Chạy thử", ghi_chu="Sẽ chuyển loại Giấy tờ")
                     for t in giay_to_hsq]
            continue
        rows += _gui_va_xac_nhan(client, madon, detail, hoso, ds, giay_to_hsq)
        time.sleep(core.REQUEST_DELAY_SECONDS)

    # ---- thành phần không đụng tới
    co_viec = {v["tp"]["stt"] for v in viec if v["tp"]}
    rows += [_dong(madon, tp=tp, trang_thai="Giữ nguyên", ghi_chu="Không có file mới")
             for tp in dm.thanh_phan_cua_don(detail)
             if tp["stt"] not in co_viec and (id(tp["hoso"]), tp["vi_tri"]) not in da_chuyen]
    return rows


def _gui_va_xac_nhan(client: core.MplisClient, madon: int, detail: dict[str, Any], hoso: dict | None,
                     ds: list[dict], giay_to: list[dict] | None = None) -> list[dict[str, Any]]:
    giay_to = giay_to or []
    nodes = day.cac_node(hoso) if hoso else []
    node_id_truoc = {v["tp"]["vi_tri"]: str(v["tp"]["node"].get("nodeId") or "") for v in ds if v["tp"]}
    try:
        if hoso is None:
            tths_id = core.safe_int(core.lay_thong_tin_ho_so_id(detail))
            if not tths_id:
                raise RuntimeError("Không tìm được thongTinHoSoId của đơn để tạo HSQ mới")
            dg.gui_tao_moi(client, tths_id, [(dg.info_gcn_moi(v["gcn"], v["da_ky"]), v["path"]) for v in ds])
        else:
            bs.gui_hsq(client, hoso, items_hsq(hoso, ds, giay_to))
    except core.HetPhien:
        raise
    except Exception as exc:
        log(f"   ❌ {core.rut_gon_text(exc, 300)}")
        return ([_dong(madon, v=v, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for v in ds]
                + [_dong_giay_to(madon, t, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for t in giay_to])

    # ---- xác nhận
    detail_moi = client.lay_chi_tiet_don(madon) or {}
    if hoso is not None:
        hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
        hoso_moi = next((h for h in detail_moi.get("ListHoSoQuet") or []
                         if isinstance(h, dict) and core.safe_int(h.get("hoSoQuetId")) == hsq_id), None)
        nodes_moi = day.cac_node(hoso_moi) if hoso_moi else []
        if len(nodes_moi) < len(nodes):
            raise day.DungKhanCap(f"HSQ {hsq_id} (đơn {madon}) giảm từ {len(nodes)} còn {len(nodes_moi)} "
                                  "thành phần sau khi đẩy → DỪNG TOÀN BỘ.")
    else:
        nodes_moi = [n for h in detail_moi.get("ListHoSoQuet") or [] if isinstance(h, dict) for n in day.cac_node(h)]

    rows = []
    for v in ds:
        loi = []
        if v["tp"]:
            tp_id = core.safe_int(v["tp"]["node"].get("thanhPhanHoSoQuetId"))
            node_moi = next((n for n in nodes_moi
                             if core.safe_int(n.get("thanhPhanHoSoQuetId")) == tp_id), None) or {}
            if v["path"] and not (node_moi.get("nodeId")
                                  and str(node_moi.get("nodeId")) != node_id_truoc[v["tp"]["vi_tri"]]):
                loi.append("file chưa đổi")
        if v["gcn"]:
            gid = core.safe_int(v["gcn"].get("giayChungNhanId"))
            if not any(core.safe_int(n.get("giayChungNhanId")) == gid and core.safe_int(n.get("loaiHoSoQuet"), -1) == 1
                       for n in nodes_moi):
                loi.append("chưa có thành phần loại Giấy chứng nhận gắn GCN này")
        if loi:
            rows.append(_dong(madon, v=v, trang_thai="Lỗi",
                              ghi_chu=f"Server báo OK nhưng {', '.join(loi)} (kiểm tra trên MPLIS)"))
            log(f"   ❌ {mo_ta_viec(v)}: server báo OK nhưng {', '.join(loi)}")
        else:
            rows.append(_dong(madon, v=v, trang_thai="Đã đẩy", ghi_chu=time.strftime("%d/%m/%Y %H:%M:%S")))
            log(f"   ✅ {mo_ta_viec(v)}" + (f": {v['tp']['node'].get('moTa')}" if v["tp"] else ""))
    for t in giay_to:
        tp_id = core.safe_int(t["node"].get("thanhPhanHoSoQuetId"))
        node_moi = next((n for n in nodes_moi if core.safe_int(n.get("thanhPhanHoSoQuetId")) == tp_id), None) or {}
        if core.safe_int(node_moi.get("loaiHoSoQuet"), -1) == LOAI_GIAY_TO and not node_moi.get("laGiayChungNhan"):
            rows.append(_dong_giay_to(madon, t, trang_thai="Đã đẩy", ghi_chu=time.strftime("%d/%m/%Y %H:%M:%S")))
            log(f"   ✅ Chuyển Giấy tờ: {t['node'].get('moTa')}")
        else:
            rows.append(_dong_giay_to(madon, t, trang_thai="Lỗi",
                                      ghi_chu="Server báo OK nhưng loại chưa thành Giấy tờ (kiểm tra trên MPLIS)"))
            log(f"   ❌ Chuyển Giấy tờ: {t['node'].get('moTa')} — server báo OK nhưng loại chưa đổi")
    return rows


# ============================ EXCEL + LUỒNG CHÍNH ============================

def ghi_excel(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_tt = EXCEL_HEADERS.index("trang_thai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in EXCEL_HEADERS])
        fill = MAU_TRANG_THAI.get(str(row.get("trang_thai") or ""))
        if fill:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=fill)
    for i, h in enumerate(EXCEL_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}{max(ws.max_row, 1)}"
    wb.save(path)


def _ghi_an_toan(path: str, rows: list[dict[str, Any]]) -> str:
    try:
        ghi_excel(path, rows)
        return path
    except PermissionError:
        alt = os.path.splitext(path)[0] + "_capnhat.xlsx"
        ghi_excel(alt, rows)
        log(f"⚠ Excel đang mở, đã lưu sang: {alt}")
        return alt


# ============================ TIẾN ĐỘ ============================
# <thư mục gốc>/FILE_TIEN_DO: {"<mã đơn>": {"luc": ..., "files": [[tên tương đối, cỡ, mtime], ...]}} — đơn đã
# đẩy xong không lỗi. Chạy lại thì bỏ qua đơn này, trừ khi file trong thư mục đơn đã đổi (thêm / bớt / sửa).
FILE_TIEN_DO = "tien_do_day_hsq_gan_gcn.json"


def doc_tien_do(goc: str) -> dict[str, Any]:
    try:
        with open(os.path.join(goc, FILE_TIEN_DO), encoding="utf-8") as fp:
            data = json.load(fp)
        return data if isinstance(data, dict) else {}
    except FileNotFoundError:
        return {}
    except (OSError, ValueError) as exc:
        log(f"⚠ Không đọc được {FILE_TIEN_DO} ({exc}) — coi như chưa đơn nào xong.")
        return {}


def ghi_tien_do(goc: str, tien_do: dict[str, Any]) -> None:
    path = os.path.join(goc, FILE_TIEN_DO)
    tam = path + ".tmp"
    with open(tam, "w", encoding="utf-8") as fp:
        json.dump(tien_do, fp, ensure_ascii=False, indent=1)
    os.replace(tam, path)  # ghi file tạm rồi thay: tắt ngang không hỏng file tiến độ


def dau_van_tay(goc: str, pdfs: list[str]) -> list[list]:
    kq = []
    for p in pdfs:
        st = os.stat(p)
        kq.append([os.path.relpath(p, goc).replace(os.sep, "/"), st.st_size, int(st.st_mtime)])
    return sorted(kq)


def don_da_xong(tien_do: dict[str, Any], goc: str, madon: int, pdfs: list[str]) -> bool:
    muc = tien_do.get(str(madon))
    return isinstance(muc, dict) and muc.get("files") == dau_van_tay(goc, pdfs)


def chay(
    client: core.MplisClient,
    goc: str,
    chay_thu: bool = True,
    nen_dung: Callable[[], bool] = lambda: False,
    giay_to: bool = False,
    bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None,
    bo_qua_da_xong: bool = True,
) -> str:
    """Chạy cho mọi thư mục mã đơn trong `goc`; trả về đường dẫn Excel kết quả (ghi trong `goc`).

    Đơn đẩy xong không lỗi được ghi vào FILE_TIEN_DO; bo_qua_da_xong → lần sau bỏ qua (file trong thư mục không đổi).
    """
    theo_don, khong_nhan = dm.quet_thu_muc_goc(goc)
    tien_do = doc_tien_do(goc)
    if bo_qua_da_xong and tien_do and theo_don:
        da_xong = {m for m, pdfs in theo_don.items() if don_da_xong(tien_do, goc, m, pdfs)}
        if da_xong:
            log(f"⏭ Bỏ qua {len(da_xong)} đơn đã đẩy xong ({FILE_TIEN_DO}); còn {len(theo_don) - len(da_xong)} đơn.")
            theo_don = {m: p for m, p in theo_don.items() if m not in da_xong}
        if not theo_don:
            log("✅ Mọi đơn đều đã đẩy xong — không còn gì để chạy. (Bỏ tích 'Bỏ qua đơn đã xong' để chạy lại.)")
            return ""
    if khong_nhan:
        log(f"⚠ Bỏ qua {len(khong_nhan)} thư mục có PDF nhưng không nằm trong thư mục mã đơn nào: {', '.join(khong_nhan[:10])}"
            + (" ..." if len(khong_nhan) > 10 else ""))
    if not theo_don:
        raise RuntimeError(f"Không có thư mục mã đơn nào trong {goc}")
    excel = os.path.join(goc, f"ket_qua_day_hsq_gan_gcn_{'thu_' if chay_thu else ''}"
                              f"{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
    log(f"{'[CHẠY THỬ] ' if chay_thu else ''}{len(theo_don)} đơn | {sum(len(v) for v in theo_don.values())} file PDF"
        + (" | HSQ có đẩy GCN: thành phần khác → Giấy tờ" if giay_to else ""))

    rows: list[dict[str, Any]] = []
    loi_lien_tiep = 0
    try:
        for n, (madon, pdfs) in enumerate(theo_don.items(), 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(n - 1, len(theo_don), f"Đơn {madon} ({n}/{len(theo_don)})")
            log(f"--- [{n}/{len(theo_don)}] Đơn {madon}: {len(pdfs)} file ---")
            try:
                dong = xu_ly_don(client, madon, pdfs, chay_thu, giay_to)
                rows += dong
                if not chay_thu and dong and not any(r.get("trang_thai") == "Lỗi" for r in dong):
                    tien_do[str(madon)] = {"luc": time.strftime("%d/%m/%Y %H:%M:%S"),
                                           "files": dau_van_tay(goc, pdfs)}
                    ghi_tien_do(goc, tien_do)
                loi_lien_tiep = loi_lien_tiep + 1 if dong and dong[0].get("ghi_chu") == dm.LOI_CHI_TIET else 0
                if loi_lien_tiep >= dm.DUNG_SAU_SO_DON_LOI:
                    log(f"❌ {loi_lien_tiep} đơn liên tiếp không đọc được từ MPLIS — có thể đã hết phiên đăng nhập. "
                        "ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            except day.DungKhanCap:
                rows += [_dong(madon, p, trang_thai="Lỗi", ghi_chu="MPLIS giảm thành phần sau khi đẩy — ĐÃ DỪNG")
                         for p in pdfs]
                raise
            except Exception as exc:
                rows += [_dong(madon, p, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for p in pdfs]
                log(f"   ❌ {core.rut_gon_text(exc, 300)}")
                if isinstance(exc, core.HetPhien):
                    log(f"❌ {exc} — ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            excel = _ghi_an_toan(excel, rows)
            core.nghi_giua_don()
    finally:
        if rows:
            excel = _ghi_an_toan(excel, rows)

    dem: dict[str, int] = {}
    for r in rows:
        if r.get("viec") or r.get("file"):
            dem[r["trang_thai"]] = dem.get(r["trang_thai"], 0) + 1
    bao_tien_do(len(theo_don), len(theo_don), "Hoàn tất")
    log("=" * 50)
    log(("[CHẠY THỬ — chưa gửi gì] " if chay_thu else "✅ ") + " | ".join(f"{k}: {v}" for k, v in dem.items()))
    log(f"Excel kết quả: {excel}")
    return excel
