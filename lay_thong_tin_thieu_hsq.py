# -*- coding: utf-8 -*-
"""
Lấy thông tin còn thiếu trên MPLIS từ hồ sơ quét (HSQ) bằng AI (Gemini) → Excel để duyệt. KHÔNG ghi MPLIS.

Đầu vào: các dòng thông báo thiếu thông tin của MPLIS (dán tay, hoặc Excel / TXT bất kỳ — mọi ô có chuỗi
"Tình hình đăng ký <id> có ... không có ..." đều được đọc), ví dụ:
    Tình hình đăng ký 18566127 có thửa đất 4703930_0 có mục đích sử dụng 9205258 không có thời hạn sử dụng
    Tình hình đăng ký 19351209 có hộ gia đình 2152642_1 có cá nhân 22455172_1 có giấy tờ tùy thân 7223544 không có số giấy tờ
Dòng trùng gộp lại; dòng bị cắt cụt ('không có thời h') vẫn nhận theo tiền tố tên trường.

Mỗi đơn:
  1. Lấy chi tiết đơn → ngữ cảnh từng đối tượng thiếu (họ tên, số tờ/thửa, số phát hành...) để AI đọc đúng
     người / đúng thửa / đúng GCN.
  2. Tải HSQ (thứ tự ưu tiên file GCN trước, như tach_gcn_theo_don) về <thư mục ra>/HSQ_goc/<madon>/, đọc tới
     đâu tải tới đó: mỗi PDF lọc trang GCN bằng ảnh nhỏ (rẻ) → gửi ảnh nét các trang đó hỏi các trường còn
     thiếu. Đủ hết thì dừng, không tải / đọc tiếp.
  3. Còn trường chưa đọc được → tìm trang đơn đăng ký (file tên là đơn trước) và hỏi tiếp trên trang đơn.
  Chạy cùng lượt tải & tách HSQ: tai_tach_hsq_don.chay(..., muc_thieu=...) — dùng chung file đã tải và cache AI.
  4. Excel: mỗi trường thiếu 1 dòng — giá trị AI đọc, file/trang nguồn, độ tin cậy, trích dẫn căn cứ.
     loaiGiayChungNhanId không ghi thẳng trên giấy → AI chọn mã theo bảng LOAI_GCN, luôn "Cần kiểm tra".
     loaiBanDoDiaChinhId luôn là "Bản đồ địa chính (VN2000)" (GIA_TRI_CO_DINH) — không đọc HSQ, không gọi AI.
"""

from __future__ import annotations

import os
import re
from collections import deque
from dataclasses import dataclass, field
from typing import Any, Callable

import tach_gcn_theo_don as tg  # nạp trước: thêm src/ vào sys.path cho package tach_gcn

from tach_gcn.gemini import DEFAULT_MODEL, GeminiAnalyzer, GeminiFieldValue, Usage
from tach_gcn.pdfio import InspectedFile, inspect_pdf

import cap_nhat_don_hsq as core
import nhom1_mplis as nhom1
import tai_hsq_don as thd
import tai_tach_hsq_don as tth

log = thd.log

SO_TRANG_MOI_LAN_DOC = 8  # ảnh nét mỗi request đọc trường

# ============================ ĐỌC DÒNG THÔNG BÁO THIẾU ============================

# loại đối tượng trong thông báo → khoá id trong chi tiết đơn
LOAI_DOI_TUONG = {
    "thửa đất": "thuaDatId",
    "mục đích sử dụng": "mucDichSuDungId",
    "nguồn gốc sử dụng": "nguonGocSuDungDatId",
    "tài liệu đo đạc": "taiLieuDoDacId",
    "giấy chứng nhận": "giayChungNhanId",
    "nội dung thay đổi": "noiDungThayDoiId",
    "hộ gia đình": "hoGiaDinhId",
    "vợ chồng": "voChongId",
    "cá nhân": "caNhanId",
    "giấy tờ tùy thân": "giayToTuyThanId",
    "tổ chức": "toChucId",
    "cộng đồng dân cư": "congDongDanCuId",
}

# tên trường thiếu → (mô tả hỏi AI, có phải suy luận)
TRUONG = {
    "thời hạn sử dụng": ("thời hạn sử dụng đất của mục đích sử dụng này (vd 'Lâu dài', '50 năm', "
                         "'Đến ngày 15/10/2043')", False),
    "số giấy tờ": ("số giấy tờ tùy thân (CCCD / CMND / hộ chiếu) của người này", False),
    "ngayVaoSo": ("ngày ký cấp GCN (dòng '..., ngày ... tháng ... năm ...' ở trang có chữ ký, cạnh số vào sổ "
                  "cấp GCN), dạng DD/MM/YYYY", False),
    "loaiGiayChungNhanId": ("", True),  # mô tả dựng từ LOAI_GCN bên dưới
    "số phát hành": ("số phát hành / số phôi GCN (2 chữ cái + 6 chữ số, vd AA 123456; mẫu A4 mới: AA + 8 chữ số)", False),
    "địa chỉ": ("địa chỉ thường trú của chủ sử dụng này (không lấy địa chỉ thửa đất)", False),
    "diện tích": ("diện tích (m²)", False),
    "loaiBanDoDiaChinhId": ("loại bản đồ địa chính (giá trị cố định, không đọc HSQ)", False),
    "năm sinh": ("năm sinh của người này", False),
    "ngayXacNhan": ("ngày xác nhận của nội dung thay đổi này ở trang 'Những thay đổi sau khi cấp GCN', "
                    "dạng DD/MM/YYYY", False),
}

# Bảng mã loaiGiayChungNhanId của MPLIS: mã → (tên loại, mốc thời gian áp dụng).
LOAI_GCN = {
    "1": ("Giấy chứng nhận QSDĐ theo Luật Đất Đai 2003", "01/07/2004 - 09/12/2009"),
    "2": ("Giấy chứng nhận QSDĐ theo Luật Đất Đai 1993", "15/10/1993 - 30/06/2004"),
    "3": ("Giấy chứng nhận QSHNƠ & QSDĐƠ theo Nghị định 60/NĐ-CP", "05/07/1994 - 07/10/2006"),
    "4": ("Giấy chứng nhận QSHNƠ & QSDĐƠ theo Nghị định 90/NĐ-CP", "08/10/2006 - 09/12/2009"),
    "5": ("Giấy chứng nhận sở hữu công trình theo quy định 95", "09/08/2005 - 09/12/2009"),
    "6": ("Giấy chứng nhận QSDĐ & QSHNƠ và TSKGLVĐ theo NĐ 88/NĐ-CP", "10/12/2009 - 30/06/2014"),
    "7": ("Giấy hợp thức hóa", "trước 01/07/2004"),
    "11": ("Giấy chứng nhận QSDĐ & QSHNƠ và TSKGLVĐ theo NĐ 43/NĐ-CP", "01/07/2014 - 31/07/2024"),
    "17": ("Giấy phép Xây dựng", "mọi thời kỳ"),
    "18": ("CÁC LOẠI GCN KHÁC (bằng khoán điền thổ, văn tự mua bán...)", "trước 15/10/1993"),
    "19": ("Giấy phép mua bán, chuyển dịch nhà", "trước năm 1993"),
    "97": ("Giấy chứng nhận TT23/2025", "từ năm 2025"),
    "98": ("Giấy chứng nhận QSDĐ, QSHTSGLVĐ năm 2024", "từ 01/08/2024"),
    "99": ("Hợp đồng mua bán tài sản hình thành trong tương lai", "từ 01/08/2024"),
}
TRUONG["loaiGiayChungNhanId"] = (
    "MÃ loại GCN theo bảng sau — chọn theo tiêu đề / mẫu phôi in trên GCN VÀ ngày cấp (ngày ký) trên GCN; value "
    "CHỈ ghi mã số, evidence ghi tiêu đề phôi + ngày cấp đọc được. Bảng: "
    + "; ".join(f"{ma} = {ten} ({moc})" for ma, (ten, moc) in LOAI_GCN.items()) + " (suy luận)",
    True,
)

# Trường luôn nhận một giá trị cố định (không gọi AI) — theo quy ước của đơn vị.
GIA_TRI_CO_DINH = {"loaiBanDoDiaChinhId": "Bản đồ địa chính (VN2000)"}
NGUON_CO_DINH = "Giá trị mặc định"

_MENH_DE = re.compile(r"Tình hình đăng ký\s+(\d+)\s+(.*?)\s*không có\s*(.*)$")
_TEN_DOI_TUONG = "|".join(map(re.escape, LOAI_DOI_TUONG))
_DOI_TUONG = re.compile(r"có\s+(" + _TEN_DOI_TUONG + r")\s+(\d+)(?:_(\d+))?")
_DOI_TUONG_EXCEL = re.compile(r"(" + _TEN_DOI_TUONG + r")\s+(\d+)(?:_(\d+))?")


@dataclass
class MucThieu:
    id_don: int
    chuoi: list[tuple[str, int]]  # [(loại đối tượng, id), ...] từ ngoài vào trong
    truong: str
    ngu_canh: str = ""
    ket_qua: GeminiFieldValue | None = None
    nguon: str = ""
    loi: list[str] = field(default_factory=list)
    phien_ban: dict[int, str] = field(default_factory=dict)  # id đối tượng → version ('4703930_0' → '0')

    @property
    def khoa(self) -> tuple:
        return self.id_don, tuple(self.chuoi), self.truong

    @property
    def doi_tuong(self) -> str:
        return " › ".join(f"{loai} {so}" + (f"_{self.phien_ban[so]}" if so in self.phien_ban else "")
                          for loai, so in self.chuoi)


# đối tượng cuối chuỗi → trường hợp lệ (dùng để chọn khi tên trường bị cắt cụt mơ hồ, vd 'không có số')
TRUONG_THEO_DOI_TUONG = {
    "giấy chứng nhận": ("ngayVaoSo", "loaiGiayChungNhanId", "số phát hành"),
    "giấy tờ tùy thân": ("số giấy tờ",),
    "mục đích sử dụng": ("thời hạn sử dụng",),
    "nguồn gốc sử dụng": ("diện tích",),
    "thửa đất": ("diện tích",),
    "tài liệu đo đạc": ("loaiBanDoDiaChinhId",),
    "nội dung thay đổi": ("ngayXacNhan",),
}


def _ten_truong(text: str, doi_tuong: str = "") -> str | None:
    """'thời hạn sử dụng' / 'thời h' (cắt cụt) → tên trường chuẩn; không nhận ra / còn mơ hồ → None."""
    t = text.strip().rstrip(".;, ")
    for ten in TRUONG:
        if t.casefold().startswith(ten.casefold()):
            return ten
    if len(t) < 2:
        return None
    ung_vien = [ten for ten in TRUONG if ten.casefold().startswith(t.casefold())]
    if len(ung_vien) > 1 and doi_tuong in TRUONG_THEO_DOI_TUONG:
        ung_vien = [ten for ten in ung_vien if ten in TRUONG_THEO_DOI_TUONG[doi_tuong]]
    return ung_vien[0] if len(ung_vien) == 1 else None


def doc_thong_bao(text: str) -> tuple[list[MucThieu], list[str]]:
    """Văn bản (nhiều dòng, mệnh đề cách nhau ';') → (mục thiếu không trùng, mệnh đề không đọc được)."""
    muc: dict[tuple, MucThieu] = {}
    bo_qua: list[str] = []
    for dong in text.splitlines():
        id_truoc = None
        for menh_de in dong.split(";"):
            menh_de = menh_de.strip()
            if not menh_de:
                continue
            if id_truoc and "Tình hình đăng ký" not in menh_de and menh_de.startswith("có "):
                menh_de = f"Tình hình đăng ký {id_truoc} {menh_de}"  # mệnh đề sau dùng chung số đơn mệnh đề trước
            m = _MENH_DE.search(menh_de)
            if m:
                id_truoc = m.group(1)
            tim = _DOI_TUONG.findall(m.group(2)) if m else []
            chuoi = [(loai, int(so)) for loai, so, _ in tim]
            ten = _ten_truong(m.group(3), chuoi[-1][0] if chuoi else "hộ gia đình") if m else None
            if not m or not ten:
                if "Tình hình đăng ký" in menh_de:
                    bo_qua.append(menh_de)
                continue
            x = MucThieu(int(m.group(1)), chuoi, ten, phien_ban={int(so): ver for _, so, ver in tim if ver})
            muc.setdefault(x.khoa, x)
    return list(muc.values()), bo_qua


def doc_file_dau_vao(path: str) -> str:
    """Excel (mọi sheet, mọi ô) hoặc file văn bản → văn bản nối dòng."""
    if os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)
                             for v in r if isinstance(v, str) and "Tình hình đăng ký" in v)
        finally:
            wb.close()
    for enc in ("utf-8-sig", "utf-16", "cp1258"):
        try:
            with open(path, encoding=enc) as fp:
                return fp.read()
        except UnicodeError:
            continue
    raise RuntimeError(f"Không đọc được bảng mã file {path}")


# ============================ HỎI MPLIS ĐƠN ĐANG THIẾU GÌ ============================

def tu_doi_tuong(id_don: int, doi_tuong: str, truong: str) -> MucThieu:
    """Dựng lại mục thiếu từ cột doi_tuong của Excel ('thửa đất 4703930_0 › mục đích sử dụng 9205258')."""
    tim = _DOI_TUONG_EXCEL.findall(str(doi_tuong or ""))
    return MucThieu(id_don, [(loai, int(so)) for loai, so, _ in tim], truong,
                    phien_ban={int(so): ver for _, so, ver in tim if ver})


def thong_bao_cua(ban_ghi: list[dict]) -> str:
    """Bản ghi làm sạch → văn bản thông báo thiếu (chiTietThongTinThieu...)."""
    import bao_cao_thong_tin_thieu_mplis as bc

    return "\n".join(bc._as_source_text(x.value) for x in map(bc.extract_missing_information, ban_ghi) if x)


def lay_muc_thieu_mplis(client, id_don: int, detail: dict) -> list[MucThieu]:
    """MPLIS tự báo đơn đang thiếu gì → mục thiếu của chính đơn này. Hết phiên → core.HetPhien."""
    muc = muc_thieu_tu_ban_ghi(id_don, lay_ban_ghi_lam_sach(client, id_don, detail))
    if muc is None:
        log(f"   MPLIS: đơn {id_don} đã đáp ứng nhóm 1 (thongTinDangKyDapUngNhom1 = true) — bỏ qua đơn")
        return []
    return muc


def muc_thieu_tu_ban_ghi(id_don: int, ban_ghi: list[dict]) -> list[MucThieu] | None:
    """Response GetThongKePhanLoaiThuaDatChiTiet → mục thiếu của đơn; None = đơn đã đáp ứng nhóm 1 (bỏ qua)."""
    if TRA_CUU_TRUC_TIEP_THEO_ID:
        if any(_dap_ung_nhom_1(bg) for bg in ban_ghi):
            return None
    else:
        da_dap_ung, ban_ghi = loc_chua_dap_ung(ban_ghi)
        if da_dap_ung:
            return None
    muc, _ = doc_thong_bao(thong_bao_cua(ban_ghi))
    return [m for m in muc if m.id_don == id_don]


loc_chua_dap_ung = nhom1.loc_chua_dap_ung
_dap_ung_nhom_1 = nhom1.dap_ung_nhom_1
payload_tra_cuu_truc_tiep_theo_id = nhom1.payload_tra_cuu_truc_tiep_theo_id
payload_tra_cuu_theo_to_thua = nhom1.payload_tra_cuu_theo_to_thua
to_thua_cua_don = nhom1.to_thua_cua_don

TRA_CUU_TRUC_TIEP_THEO_ID = False


def dat_tra_cuu_truc_tiep_theo_id(bat: bool = True) -> None:
    """Bật cho bản tool riêng: lọc endpoint làm sạch bằng đúng tinhHinhDangKyId, không lọc phân loại."""
    global TRA_CUU_TRUC_TIEP_THEO_ID
    TRA_CUU_TRUC_TIEP_THEO_ID = bool(bat)


def lay_ban_ghi_lam_sach(client, id_don: int, detail: dict) -> list[dict]:
    """Bản ghi làm sạch MPLIS của 1 đơn (nhom1_mplis.lay_ban_ghi_lam_sach). Hết phiên → core.HetPhien."""
    return nhom1.lay_ban_ghi_lam_sach(client, id_don, detail, TRA_CUU_TRUC_TIEP_THEO_ID)


# ============================ NGỮ CẢNH TỪ CHI TIẾT ĐƠN ============================

_UU_TIEN_KHOA = ("hoTen", "tenToChuc", "namSinh", "soGiayTo", "soTo", "soThua", "soHieuToBanDo",
                 "soHieuThua", "soToBanDo", "soThuTuThua", "dienTich", "kyHieu", "maMucDich", "tenMucDich",
                 "soPhatHanh", "soVaoSo", "noiDung", "diaChi")


def _tim_dict(goc: Any, khoa: str, so: int) -> dict | None:
    """Duyệt theo chiều rộng: dict nông nhất có goc[khoa] == so (không có khoá đó thì bất kỳ khoá *Id)."""
    hang: deque = deque([goc])
    du_phong = None
    while hang:
        o = hang.popleft()
        if isinstance(o, dict):
            if core.safe_int(o.get(khoa)) == so:
                return o
            if du_phong is None and any(k.endswith("Id") and core.safe_int(v) == so for k, v in o.items()
                                        if not isinstance(v, (dict, list))):
                du_phong = o
            hang.extend(v for v in o.values() if isinstance(v, (dict, list)))
        elif isinstance(o, list):
            hang.extend(o)
    return du_phong


def _tom_tat(d: dict) -> str:
    gia_tri = [(k, v) for k, v in d.items()
               if not isinstance(v, (dict, list)) and v not in (None, "", 0, False)
               and not k.lower().endswith("id") and len(str(v)) <= 60 and "/Date(" not in str(v)]
    gia_tri.sort(key=lambda kv: _UU_TIEN_KHOA.index(kv[0]) if kv[0] in _UU_TIEN_KHOA else 99)
    return ", ".join(f"{k}={v}" for k, v in gia_tri[:8])


def gan_ngu_canh(detail: dict, cac_muc: list[MucThieu]) -> None:
    for muc in cac_muc:
        phan: list[str] = []
        goc: Any = detail
        for loai, so in muc.chuoi:
            d = _tim_dict(goc, LOAI_DOI_TUONG[loai], so)
            tt = _tom_tat(d) if d else ""
            phan.append(f"{loai} {so}" + (f" ({tt})" if tt else ""))
            goc = d if d else goc
        muc.ngu_canh = " › ".join(phan)


# ============================ AI ĐỌC TRƯỜNG ============================

def _mo_ta_ai(muc: MucThieu) -> str:
    mo_ta, _ = TRUONG[muc.truong]
    return f"{mo_ta} — đối tượng: {muc.ngu_canh or muc.doi_tuong}"


_DIEM = {"cao": 3, "trung bình": 2, "thấp": 1}


def doc_tren_trang(ins: InspectedFile, trang: list[int], cac_muc: list[MucThieu], analyzer: GeminiAnalyzer,
                   model: str, cache: dict, usage: Usage, loai_trang: str) -> None:
    """Hỏi AI các mục còn thiếu trên các trang (0-based) của 1 PDF; giữ giá trị tin cậy cao nhất."""
    path = os.path.basename(ins.path)
    for dau in range(0, len(trang), SO_TRANG_MOI_LAN_DOC):
        con = [m for m in cac_muc if not (m.ket_qua and m.ket_qua.value)]
        if not con:
            return
        lo = trang[dau:dau + SO_TRANG_MOI_LAN_DOC]
        hoi = [(f"f{i}", _mo_ta_ai(m)) for i, m in enumerate(con, 1)]
        khoa = (f"fields1:{ins.content_hash}:{loai_trang}:{','.join(map(str, lo))}:{model}:"
                + "|".join(d for _, d in hoi))
        cached = cache.get(khoa)
        if isinstance(cached, dict):
            kq = {k: GeminiFieldValue(**v) for k, v in cached.items()}
        else:
            kq, dung = analyzer.extract_fields(tg.render_trang(ins, lo), hoi,
                                               f"Các ảnh là trang {loai_trang}, số i là số trang thật trong file.")
            usage.add(dung)
            cache[khoa] = {k: {"value": v.value, "page": v.page, "confidence": v.confidence,
                               "evidence": v.evidence} for k, v in kq.items()}
        for (k, _), m in zip(hoi, con):
            v = kq.get(k)
            if not v or not v.value:
                continue
            if m.ket_qua is None or _DIEM.get(v.confidence, 0) > _DIEM.get(m.ket_qua.confidence, 0):
                m.ket_qua = v
                m.nguon = f"{path} tr.{v.page}" if v.page else path


# ============================ 1 ĐƠN ============================

def xu_ly_don(analyzer: GeminiAnalyzer, id_don: int, cac_muc: list[MucThieu], cache: dict, usage: Usage,
              tt: tg.TaiTruoc, model: str = DEFAULT_MODEL, nen_dung: Callable[[], bool] = lambda: False) -> None:
    """Điền ket_qua cho các mục thiếu của 1 đơn. Hết phiên → ném core.HetPhien."""
    def con_thieu() -> list[MucThieu]:
        return [m for m in cac_muc if not (m.ket_qua and m.ket_qua.value)]

    def ghi_loi(msg: str) -> None:
        for m in con_thieu():
            m.loi.append(msg)

    for m in cac_muc:
        if m.truong in GIA_TRI_CO_DINH:
            m.ket_qua = GeminiFieldValue(GIA_TRI_CO_DINH[m.truong], 0, "cao", "")
            m.nguon = NGUON_CO_DINH
            log(f"   {m.truong} ({m.doi_tuong}) → {GIA_TRI_CO_DINH[m.truong]} (mặc định)")
    if not con_thieu():
        return

    detail = tt.chi_tiet(id_don)
    if detail is None:
        ghi_loi("Không lấy được chi tiết đơn")
        log(f"❌ Đơn {id_don}: không lấy được chi tiết")
        return
    gan_ngu_canh(detail, cac_muc)
    for m in con_thieu():
        log(f"   Thiếu {m.truong}: {m.ngu_canh}")
    files = thd.liet_ke_file_hsq(detail)
    if not files:
        ghi_loi("Đơn không có HSQ")
        log("   ⚠ Đơn không có HSQ")
        return

    thu_tu = sorted(enumerate(files, 1), key=lambda x: tg.uu_tien_file(x[1]))
    da_mo: dict[int, InspectedFile] = {}

    def mo(vi_tri: int, f: dict[str, Any]) -> InspectedFile | None:
        if vi_tri in da_mo:
            return da_mo[vi_tri]
        mo_ta = f.get("moTa") or f.get("tenGiayTo") or f"file_{vi_tri}"
        try:
            path = tt.file(id_don, vi_tri, f)
        except core.HetPhien:
            raise
        except Exception as exc:
            log(f"   ❌ Tải {mo_ta}: {core.rut_gon_text(exc, 200)}")
            ghi_loi(f"Tải '{mo_ta}' lỗi: {core.rut_gon_text(exc, 120)}")
            return None
        if not path.lower().endswith(".pdf"):
            return None
        ins = inspect_pdf(path, os.path.basename(path), (), False, False)
        if ins.error:
            ghi_loi(f"'{os.path.basename(path)}' lỗi: {ins.error}")
            return None
        da_mo[vi_tri] = ins
        return ins

    # ---- trang GCN
    for n, (vi_tri, f) in enumerate(thu_tu):
        if nen_dung() or not con_thieu():
            break
        tt.dat_truoc_file(id_don, thu_tu[n:n + 2], n)
        ins = mo(vi_tri, f)
        if ins is None:
            continue
        try:
            trang = tg.loc_trang(ins, analyzer, model, cache, usage)
            if trang:
                log(f"   {os.path.basename(ins.path)}: trang GCN {tg._trang_text(trang)} → đọc {len(con_thieu())} mục")
                doc_tren_trang(ins, trang, con_thieu(), analyzer, model, cache, usage, "giấy chứng nhận")
        except Exception as exc:
            log(f"   ❌ AI đọc {os.path.basename(ins.path)}: {core.rut_gon_text(exc, 300)}")
            ghi_loi(f"AI đọc '{os.path.basename(ins.path)}' lỗi: {core.rut_gon_text(exc, 120)}")

    # ---- trang đơn đăng ký (file tên là đơn trước)
    if con_thieu() and not nen_dung():
        log(f"   Còn {len(con_thieu())} mục chưa thấy trên GCN → tìm trên đơn đăng ký")
        for vi_tri, f in sorted(thu_tu, key=lambda x: not tth.ten_la_don(x[1])):
            if nen_dung() or not con_thieu():
                break
            ins = mo(vi_tri, f)
            if ins is None:
                continue
            try:
                ban, _ = tth.tim_trang_don(ins.path, analyzer, model, cache, usage)
                trang = [i for b in ban for i in b]
                if trang:
                    log(f"   {os.path.basename(ins.path)}: trang đơn {tg._trang_text(trang)} → đọc {len(con_thieu())} mục")
                    doc_tren_trang(ins, trang, con_thieu(), analyzer, model, cache, usage, "đơn đăng ký")
            except Exception as exc:
                log(f"   ❌ AI đọc {os.path.basename(ins.path)}: {core.rut_gon_text(exc, 300)}")
                ghi_loi(f"AI đọc '{os.path.basename(ins.path)}' lỗi: {core.rut_gon_text(exc, 120)}")

    for m in cac_muc:
        if m.ket_qua and m.ket_qua.value:
            log(f"   ✅ {m.truong} ({m.doi_tuong}): {m.ket_qua.value} [{m.ket_qua.confidence}] — {m.nguon}")
        else:
            log(f"   ⚠ {m.truong} ({m.doi_tuong}): không đọc được")


# ============================ EXCEL ============================

# duyet: 'x' = cho phép tool cập nhật ghi lên MPLIS (điền sẵn cho Đã đọc / Mặc định);
# gia_tri_sua: sửa tay giá trị trước khi ghi (để trống = dùng gia_tri_ai).
EXCEL_HEADERS = ["stt", "madon", "doi_tuong", "id_doi_tuong", "truong_thieu", "ngu_canh", "gia_tri_ai",
                 "duyet", "gia_tri_sua", "do_tin_cay", "nguon", "trich_dan", "trang_thai", "ghi_chu"]
EXCEL_WIDTHS = {"stt": 6, "madon": 11, "doi_tuong": 40, "id_doi_tuong": 12, "truong_thieu": 20, "ngu_canh": 60,
                "gia_tri_ai": 30, "duyet": 7, "gia_tri_sua": 24, "do_tin_cay": 11, "nguon": 36, "trich_dan": 50, "trang_thai": 14, "ghi_chu": 50}
MAU_TRANG_THAI = {"Đã đọc": "C6EFCE", "Cần kiểm tra": "FFF2CC", "Không thấy": "F4CCCC", "Lỗi": "F4CCCC",
                  "Chưa xử lý": "EDEDED", "Mặc định": "D9EAF7"}


def dong_excel(m: MucThieu, xong: bool) -> dict[str, Any]:
    v = m.ket_qua if m.ket_qua and m.ket_qua.value else None
    suy_luan = TRUONG[m.truong][1]
    if v and m.nguon == NGUON_CO_DINH:
        trang_thai = "Mặc định"
    elif v:
        trang_thai = "Cần kiểm tra" if suy_luan or v.confidence != "cao" else "Đã đọc"
    elif not xong:
        trang_thai = "Chưa xử lý"
    else:
        trang_thai = "Lỗi" if any("lỗi" in e.lower() or "HẾT PHIÊN" in e for e in m.loi) else "Không thấy"
    ghi_chu = ["AI suy luận — không ghi thẳng trên giấy" if suy_luan and v else "", *dict.fromkeys(m.loi)]
    if v and m.truong == "loaiGiayChungNhanId":
        ma = re.sub(r"\D", "", v.value)
        if ma in LOAI_GCN:
            ghi_chu.insert(0, f"Mã {ma}: {LOAI_GCN[ma][0]} ({LOAI_GCN[ma][1]})")
        else:
            ghi_chu.insert(0, f"Mã '{v.value}' không có trong bảng loại GCN — kiểm tra tay")
    return {"madon": m.id_don, "doi_tuong": m.doi_tuong, "id_doi_tuong": m.chuoi[-1][1] if m.chuoi else "",
            "truong_thieu": m.truong, "ngu_canh": m.ngu_canh, "gia_tri_ai": v.value if v else "",
            "duyet": "x" if trang_thai in ("Đã đọc", "Mặc định") else "",
            "do_tin_cay": v.confidence if v else "", "nguon": m.nguon if v else "",
            "trich_dan": v.evidence if v else "", "trang_thai": trang_thai,
            "ghi_chu": "; ".join(filter(None, ghi_chu))}


def ghi_excel(path: str, rows: list[dict[str, Any]], them_cot: tuple[str, ...] = ()) -> None:
    """them_cot: cột thêm cuối bảng (vd trạng thái cập nhật MPLIS)."""
    headers = EXCEL_HEADERS + [c for c in them_cot if c not in EXCEL_HEADERS]
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "ThongTinThieu"
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_tt = headers.index("trang_thai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in headers])
        mau = MAU_TRANG_THAI.get(str(row.get("trang_thai") or ""))
        if mau:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=mau)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
