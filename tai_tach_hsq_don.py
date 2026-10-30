# -*- coding: utf-8 -*-
"""
Tải & tách HSQ theo đơn: tải TOÀN BỘ file HSQ của mọi đơn (kể cả đơn chưa có GCN) → AI (Gemini) tách GCN
trong HSQ (nếu có) → HSQ không có trang GCN nào thì tìm đơn đăng ký để tách.

Luồng mỗi đơn (tinhHinhDangKyId):
  1. Lấy chi tiết đơn, xếp hàng tải mọi file HSQ về <thư mục ra>/HSQ_goc/<madon>/ (thứ tự ưu tiên như
     tach_gcn_theo_don: file có "GCN" trong tên trước) — tải nền trong lúc AI đọc file trước.
  2. Tách GCN:
       a. Đơn có GCN hiện hành trên MPLIS: giống hệt tach_gcn_theo_don (gán trang theo số phát hành, đối chiếu
          chủ) → GCN/<madon>/<SỐ GCN>-GCN.pdf, GCN_can_kiem_tra/, GCN_khong_khop/. Đủ GCN đúng chủ thì AI dừng
          đọc, nhưng các file còn lại VẪN được tải về.
       b. Đơn chưa có GCN trên MPLIS: AI đọc mọi PDF tìm trang GCN (số phát hành đọc trên trang) →
          GCN_khong_khop/<madon>/<số đọc được>-GCN.pdf (tool đẩy không nhận file -GCN cho đơn chưa có GCN
          nên không để vào GCN/).
  3. Không tách được trang GCN nào → tìm đơn đăng ký:
       - File có tên (mô tả) là đơn ('Đơn đăng ký', '01_DON_DK', 'DONDANGKY'...): < 15 MB → chép nguyên file,
         KHÔNG gọi AI; ≥ 15 MB → AI tìm trang đơn rồi tách.
       - Không file nào có tên là đơn → AI đọc lần lượt các PDF; file đầu tiên có trang đơn: < 15 MB → chép
         nguyên file, ≥ 15 MB → tách các trang đơn.
       → DON/<madon>/<tên>-DON.pdf
  4. Excel: ket_qua_tach_hsq_*.xlsx (mỗi GCN / đơn 1 dòng) + ket_qua_tai_hsq_*.xlsx (mỗi file tải 1 dòng,
     đúng mẫu tai_hsq_don → dùng được cho tab Đẩy HSQ ký số).

Không ghi gì lên MPLIS.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import time
from pathlib import Path
from typing import Any, Callable

import tach_gcn_theo_don as tg  # nạp trước: thêm src/ vào sys.path cho package tach_gcn

from tach_gcn.gemini import DEFAULT_MODEL, GeminiAnalyzer, Usage
from tach_gcn.pdfio import InspectedFile, _save_selected_pages, inspect_pdf, open_pdf_document, \
    render_page_jpeg_checked
from tach_gcn.report import load_cache, save_cache

import bo_sung_gcn_hsq as bs
import cap_nhat_don_hsq as core
import nhom1_mplis as nhom1
import day_hsq_ky_so as day
import day_hsq_theo_mota as dm
import tai_hsq_don as thd

log = thd.log

THU_MUC_DON = "DON"
GIOI_HAN_KHONG_TACH_DON = 15 * 1024 * 1024  # file chứa đơn nhỏ hơn mức này → chép nguyên file
MB = 1024 * 1024
SO_TRANG_LO_DON = 16  # số ảnh nhỏ mỗi request tìm trang đơn

EXCEL_HEADERS = ["stt", "madon", "loai", "giaychungnhanid", "so_gcn", "so_vao_so", "chu_don", "file_nguon",
                 "trang", "so_ai_doc", "so_vao_so_ai", "ngay_cap_ai", "chu_ai", "cccd_ai", "dia_chi_ai",
                 "file_tach", "trang_thai", "ghi_chu"]
EXCEL_WIDTHS = {**tg.EXCEL_WIDTHS, "loai": 7, "file_tach": 55}
MAU_TRANG_THAI = {**tg.MAU_TRANG_THAI, "Đã lấy đơn": "C6EFCE", "Đã tách đơn": "C6EFCE",
                  "Không thấy đơn": "F4CCCC", "Không khớp": "FFF2CC", "Chưa có số GCN": "D9EAF7"}


# ============================ NHẬN DIỆN FILE ĐƠN THEO TÊN ============================

_CHU_DON = re.compile(r"(?<![a-z])don(?![a-z])")  # 'Đơn đăng ký', '01_DON_DK' — không bắt 'hợp đồng'
TU_DON_GHEP = ("dondangky", "dondenghi", "dondk", "donxincap", "dondangki")


def ten_la_don(f: dict[str, Any]) -> bool:
    ten = " ".join(str(f.get(k) or "") for k in tg.KHOA_TEN_FILE)
    s = core.remove_accents(ten).lower().replace("đ", "d").replace("ð", "d")
    return bool(_CHU_DON.search(s)) or any(t in dm.chuan_hoa(ten) for t in TU_DON_GHEP)


# ============================ AI TÌM TRANG ĐƠN ============================

def tim_trang_don(path: str, analyzer: GeminiAnalyzer, model: str, cache: dict,
                  usage: Usage) -> tuple[list[list[int]], InspectedFile]:
    """Trang đơn đăng ký (0-based) trong 1 PDF, gom theo từng bản đơn → (các bản, file đã đọc)."""
    ins = inspect_pdf(path, os.path.basename(path), (), False, False)
    if ins.error:
        raise RuntimeError(ins.error)
    khoa = f"don1:{ins.content_hash}:{model}"
    cached = cache.get(khoa)
    if isinstance(cached, dict) and cached.get("so_trang") == ins.page_count:
        return [list(b) for b in cached["ban"]], ins

    la_don: dict[int, tuple[bool, bool]] = {}
    doc = open_pdf_document(Path(ins.path), ins.recovered_page_xrefs)
    try:
        for dau in range(0, ins.page_count, SO_TRANG_LO_DON):
            lo = list(range(dau, min(dau + SO_TRANG_LO_DON, ins.page_count)))
            anh = [(i + 1, render_page_jpeg_checked(doc, Path(ins.path), i, tg._canh_dai(doc.load_page(i), 1000),
                                                    60, ins.recovered_page_xrefs)) for i in lo]
            truoc = ""
            if dau:
                truoc = (f"Trang ngay trước lô này (i={dau}) "
                         + ("là trang đơn đăng ký." if la_don[dau - 1][0] else "không phải trang đơn đăng ký."))
            kq, dung = analyzer.find_applications(anh, truoc, low_res=True)
            usage.add(dung)
            for i in lo:
                la_don[i] = (kq[i + 1].is_application, kq[i + 1].starts_copy)
    finally:
        doc.close()

    ban: list[list[int]] = []
    for i in range(ins.page_count):
        co, bat_dau = la_don[i]
        if not co:
            continue
        if ban and ban[-1][-1] == i - 1 and not bat_dau:
            ban[-1].append(i)
        else:
            ban.append([i])
    cache[khoa] = {"so_trang": ins.page_count, "ban": ban}
    return ban, ins


def _dong_don(id_don: int, path: str = "", trang: list[int] | None = None, **kw) -> dict[str, Any]:
    return {"madon": id_don, "loai": "Đơn", "file_nguon": os.path.basename(path),
            "trang": tg._trang_text(trang) if trang else "", **kw}


def lay_don(id_don: int, tai_ve: list[tuple[dict[str, Any], str]], analyzer: GeminiAnalyzer, model: str,
            cache: dict, usage: Usage, out_dir: str, nen_dung: Callable[[], bool]) -> list[dict[str, Any]]:
    """
    tai_ve: (file HSQ, đường dẫn đã tải) theo thứ tự HSQ. File tên là đơn ưu tiên; không có thì AI đọc lần lượt
    các PDF tới file đầu tiên có trang đơn. File chứa đơn < 15 MB → chép nguyên; lớn hơn → tách trang đơn.
    """
    thu_muc = os.path.join(out_dir, THU_MUC_DON, str(id_don))
    da_dung: set[str] = set()
    rows: list[dict[str, Any]] = []
    loi: list[str] = []

    theo_ten = [(f, p) for f, p in tai_ve if ten_la_don(f)]
    if theo_ten:
        log(f"   Tìm đơn đăng ký: {len(theo_ten)} file có tên là đơn")
        ung_vien, can_ai = theo_ten, False
    else:
        log("   Tìm đơn đăng ký: không file nào có tên là đơn → AI đọc lần lượt các PDF")
        ung_vien, can_ai = [(f, p) for f, p in tai_ve if p.lower().endswith(".pdf")], True

    for f, path in ung_vien:
        if nen_dung():
            break
        ten = os.path.basename(path)
        co = os.path.getsize(path)
        nho = co < GIOI_HAN_KHONG_TACH_DON
        if not can_ai and nho:
            dich = tg._ten_khong_trung(thu_muc, thd.ten_file_an_toan(Path(path).stem, 100) + "-DON", da_dung)
            dich = os.path.splitext(dich)[0] + os.path.splitext(path)[1]
            os.makedirs(thu_muc, exist_ok=True)
            shutil.copy2(path, dich)
            log(f"   ✅ Đơn: {ten} ({co / MB:.1f} MB < {GIOI_HAN_KHONG_TACH_DON // MB} MB) → chép nguyên file")
            rows.append(_dong_don(id_don, path, file_tach=dich, trang_thai="Đã lấy đơn",
                                  ghi_chu=f"Tên file là đơn, {co / MB:.1f} MB < {GIOI_HAN_KHONG_TACH_DON // MB} MB — chép nguyên, không tách"))
            continue
        if not path.lower().endswith(".pdf"):
            loi.append(f"{ten} ≥ {GIOI_HAN_KHONG_TACH_DON // MB} MB nhưng không phải PDF — không tách được")
            continue
        try:
            ds_ban, ins = tim_trang_don(path, analyzer, model, cache, usage)
        except Exception as exc:
            loi.append(f"AI đọc '{ten}' lỗi: {core.rut_gon_text(exc, 200)}")
            log(f"   ❌ {ten}: {core.rut_gon_text(exc, 300)}")
            continue
        if not ds_ban:
            log(f"   {ten} ({ins.page_count} trang): không có trang đơn")
            if not can_ai:
                loi.append(f"Tên '{ten}' là đơn nhưng AI không thấy trang đơn nào")
            continue
        if nho:  # AI tìm thấy đơn trong file nhỏ → chép nguyên file
            dich = tg._ten_khong_trung(thu_muc, thd.ten_file_an_toan(Path(path).stem, 100) + "-DON", da_dung)
            os.makedirs(thu_muc, exist_ok=True)
            shutil.copy2(path, dich)
            trang = [i for b in ds_ban for i in b]
            log(f"   ✅ Đơn: {ten} tr.{tg._trang_text(trang)} ({co / MB:.1f} MB < {GIOI_HAN_KHONG_TACH_DON // MB} MB) → chép nguyên file")
            rows.append(_dong_don(id_don, path, trang, file_tach=dich, trang_thai="Đã lấy đơn",
                                  ghi_chu=f"AI thấy trang đơn; {co / MB:.1f} MB < {GIOI_HAN_KHONG_TACH_DON // MB} MB — chép nguyên, không tách"))
        else:
            for trang in ds_ban:
                dich = tg._ten_khong_trung(
                    thu_muc, thd.ten_file_an_toan(f"{Path(path).stem}_tr{tg._trang_text(trang)}", 100) + "-DON", da_dung)
                try:
                    _save_selected_pages(Path(path), trang, Path(dich), ins.recovered_page_xrefs)
                except Exception as exc:
                    log(f"   ❌ Ghi PDF đơn lỗi: {exc}")
                    rows.append(_dong_don(id_don, path, trang, trang_thai="Lỗi", ghi_chu=f"Ghi PDF lỗi: {exc}"))
                    continue
                log(f"   ✅ Đơn: {ten} tr.{tg._trang_text(trang)} ({co / MB:.1f} MB) → {os.path.basename(dich)}")
                rows.append(_dong_don(id_don, path, trang, file_tach=dich, trang_thai="Đã tách đơn",
                                      ghi_chu=f"File {co / MB:.1f} MB ≥ {GIOI_HAN_KHONG_TACH_DON // MB} MB — tách trang đơn"
                                      + (f"; {len(ds_ban)} bản đơn" if len(ds_ban) > 1 else "")))
        if can_ai:
            break  # đã thấy file chứa đơn

    if not rows:
        log("   ⚠ Không tìm thấy đơn đăng ký trong HSQ")
        rows.append(_dong_don(id_don, trang_thai="Không thấy đơn",
                              ghi_chu="; ".join(["HSQ không có GCN và không tìm thấy đơn đăng ký", *loi])))
    elif loi:
        rows[-1]["ghi_chu"] = "; ".join([rows[-1]["ghi_chu"], *loi])
    return rows


# ============================ GCN — ĐƠN CHƯA CÓ GCN TRÊN MPLIS ============================

def tach_gcn_khong_co_tren_mplis(id_don: int, tai_ve: list[tuple[dict[str, Any], str]], analyzer: GeminiAnalyzer,
                                 model: str, cache: dict, usage: Usage, out_dir: str, loc_truoc: bool,
                                 nen_dung: Callable[[], bool]) -> list[dict[str, Any]]:
    """AI tìm mọi trang GCN trong các PDF (không có danh sách GCN để đối chiếu) → GCN_khong_khop/<madon>/."""
    thu_muc = os.path.join(out_dir, tg.THU_MUC_KHONG_KHOP, str(id_don))
    da_dung: set[str] = set()
    rows: list[dict[str, Any]] = []
    for f, path in tai_ve:
        if nen_dung():
            break
        if not path.lower().endswith(".pdf"):
            continue
        ten = os.path.basename(path)
        try:
            gan, ins = tg.gan_trang(path, [], analyzer, model, cache, usage, loc_truoc)
        except Exception as exc:
            log(f"   ❌ {ten}: {core.rut_gon_text(exc, 300)}")
            continue
        kho = tg.kho_trang(ins)
        nhom = [tg.bo_sung_mat(t, gan, kho)[0] if len(t) == 1 and tg.la_so_a4(gan[t[0]]["serial_seen"]) else t
                for t in tg._nhom_khong_khop(gan)]
        log(f"   {ten} ({ins.page_count} trang): "
            + ("; ".join(f"GCN tr.{tg._trang_text(t)}" for t in nhom) or "không có trang GCN"))
        for trang in nhom:
            so_doc = next((gan[i]["serial_seen"] for i in trang if gan[i]["serial_seen"]), "")
            ten_ra = f"{so_doc}-GCN" if so_doc else f"{Path(path).stem}_tr{tg._trang_text(trang)}-GCN"
            dich = tg._ten_khong_trung(thu_muc, thd.ten_file_an_toan(ten_ra), da_dung)
            try:
                _save_selected_pages(Path(path), trang, Path(dich), ins.recovered_page_xrefs)
            except Exception as exc:
                log(f"   ❌ Ghi trang GCN lỗi: {exc}")
                dich = ""
            rows.append(tg._dong(id_don, None, path, trang, gan, loai="GCN", file_gcn=dich,
                                 trang_thai="Không khớp" if dich else "Lỗi",
                                 ghi_chu="Đơn chưa có GCN trên MPLIS — trang GCN trong HSQ, kiểm tra tay"
                                 + ("" if so_doc else "; AI không đọc được số phát hành")))
    return rows


def _dong_gcn_chua_so(id_don: int, gcn_chua_so: list[dict[str, Any]],
                      rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """
    GCN đã có giayChungNhanId trên MPLIS nhưng chưa có số phát hành → mỗi GCN 1 dòng ghi các số GCN AI đọc
    được trên trang GCN không thuộc GCN có số nào của đơn (file GCN_khong_khop/), để điền số bằng tay.
    """
    if not gcn_chua_so:
        return []
    khac = [r for r in rows if r.get("trang_thai") == "Không khớp" and r.get("loai", "GCN") == "GCN"]
    so_doc = list(dict.fromkeys(s.strip() for r in khac for s in str(r.get("so_ai_doc") or "").split(",")
                                if s.strip()))
    nguon = "; ".join(f"{r.get('file_nguon')} tr.{r.get('trang')}" for r in khac if r.get("file_nguon"))
    if so_doc:
        ghi_chu = f"GCN trên MPLIS chưa có số phát hành — số đọc được trong HSQ: {', '.join(so_doc)} ({nguon})"
        if len(gcn_chua_so) > 1:
            ghi_chu += f"; đơn có {len(gcn_chua_so)} GCN chưa có số, đối chiếu tay"
    elif khac:
        ghi_chu = f"GCN trên MPLIS chưa có số phát hành — có trang GCN trong HSQ nhưng AI không đọc được số ({nguon})"
    else:
        ghi_chu = "GCN trên MPLIS chưa có số phát hành — HSQ không có trang GCN nào"
    for g in gcn_chua_so:
        log(f"   ℹ GCN id {g.get('giayChungNhanId')} chưa có số phát hành: "
            + (f"HSQ có số {', '.join(so_doc)}" if so_doc else "không đọc được số trong HSQ"))
    return [{"madon": id_don, "loai": "GCN", "giaychungnhanid": g.get("giayChungNhanId", ""), "so_gcn": "",
             "so_vao_so": g.get("soVaoSo") or "", "so_ai_doc": ", ".join(so_doc), "file_tach": "",
             "trang_thai": "Chưa có số GCN", "ghi_chu": ghi_chu} for g in gcn_chua_so]


# ============================ 1 ĐƠN ============================

def xu_ly_don(client, analyzer: GeminiAnalyzer, id_don: int, out_dir: str, cache: dict, usage: Usage,
              tt: tg.TaiTruoc, model: str = DEFAULT_MODEL, loc_truoc: bool = True, quet_het: bool = False,
              nen_dung: Callable[[], bool] = lambda: False, bo_qua_nhom_1: bool = True,
              nhom1_kq: dict[int, Any] | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """
    → (dòng Excel tách, dòng Excel tải). Hết phiên → ném core.HetPhien (ket_qua = các dòng đã có).
    bo_qua_nhom_1: đơn đã đạt nhóm 1 trên MPLIS → không tải / tách. nhom1_kq[id_don] = (đạt?, bản ghi làm sạch)
    để bước tự kiểm tra thông tin thiếu dùng lại, không tra MPLIS lần nữa.
    """
    detail = tt.chi_tiet(id_don)
    if detail is None:
        log(f"❌ Đơn {id_don}: không lấy được chi tiết")
        return ([_dong_don(id_don, trang_thai="Lỗi", ghi_chu="Không lấy được chi tiết đơn")],
                [{"madon": id_don, "trang_thai_tai": "Lỗi", "ghi_chu_tai": "Không lấy được chi tiết đơn"}])
    if bo_qua_nhom_1:
        with tt.khoa:  # không gửi chồng với luồng tải nền
            kq_nhom1 = nhom1.kiem_tra(tt.client, id_don, detail)
        if nhom1_kq is not None:
            nhom1_kq[id_don] = kq_nhom1
        if kq_nhom1[0]:
            tt.huy_don(id_don)
            return ([_dong_don(id_don, trang_thai="Bỏ qua", ghi_chu=nhom1.LY_DO_DA_DAT)],
                    [{"madon": id_don, "trang_thai_tai": "Bỏ qua", "ghi_chu_tai": nhom1.LY_DO_DA_DAT}])

    files = thd.liet_ke_file_hsq(detail)
    ds_gcn = [g for g in bs.gcn_hien_hanh(detail) if str(g.get("soPhatHanh") or "").strip()]
    gcn_chua_so = [g for g in bs.gcn_hien_hanh(detail) if not str(g.get("soPhatHanh") or "").strip()]
    log(f"Đơn {id_don}: {len(files)} file HSQ, {len(ds_gcn)} GCN hiện hành"
        + (": " + ", ".join(str(g.get("soPhatHanh")) for g in ds_gcn) if ds_gcn else "")
        + (f", {len(gcn_chua_so)} GCN chưa có số phát hành (id "
           + ", ".join(str(g.get("giayChungNhanId")) for g in gcn_chua_so) + ")" if gcn_chua_so else ""))
    if not files:
        return ([_dong_don(id_don, trang_thai="Bỏ qua", ghi_chu="Đơn không có HSQ")]
                + _dong_gcn_chua_so(id_don, gcn_chua_so, []),
                [{"madon": id_don, "trang_thai_tai": "Không có HSQ"}])
    quet_het = quet_het or bool(gcn_chua_so)  # cần đọc hết HSQ để tìm số cho GCN chưa có số

    thu_muc_hsq = os.path.join(out_dir, tg.THU_MUC_HSQ, str(id_don))
    os.makedirs(thu_muc_hsq, exist_ok=True)
    with open(os.path.join(thu_muc_hsq, "hosoquet.json"), "w", encoding="utf-8") as fp:
        json.dump(detail.get("ListHoSoQuet"), fp, ensure_ascii=False, indent=2, default=str)

    thu_tu = sorted(enumerate(files, 1), key=lambda x: tg.uu_tien_file(x[1]))
    tt.dat_truoc_file(id_don, thu_tu)  # tải nền toàn bộ, theo thứ tự ưu tiên
    rows: list[dict[str, Any]] = []

    # ---- 2a. Đơn có GCN trên MPLIS: tách như tach_gcn_theo_don (tải tới đâu đọc tới đó)
    if ds_gcn:
        rows += [{**r, "loai": "GCN"} for r in tg._tach_gcn_don(tt, analyzer, detail, ds_gcn, id_don, out_dir,
                                                                cache, usage, model, nen_dung, loc_truoc, quet_het)]

    # ---- 1. Tải nốt mọi file (file AI đã đọc lấy lại ngay từ kết quả tải nền)
    tai_ve: list[tuple[dict[str, Any], str]] = []
    dong_tai: list[dict[str, Any]] = []
    for vi_tri, f in enumerate(files, 1):
        if nen_dung():
            break
        mo_ta = f.get("moTa") or f.get("tenGiayTo") or f"file_{vi_tri}"
        try:
            path = tt.file(id_don, vi_tri, f)
        except core.HetPhien as exc:
            dong_tai.append(thd.dong_excel_tu_file(id_don, f, "", "Lỗi", f"HẾT PHIÊN — ĐÃ DỪNG: {exc}"))
            raise core.HetPhien(str(exc), ket_qua={"dong": rows, "dong_tai": dong_tai}) from exc
        except Exception as exc:
            log(f"   ❌ Tải {vi_tri}. {mo_ta}: {core.rut_gon_text(exc, 300)}")
            dong_tai.append(thd.dong_excel_tu_file(id_don, f, "", "Lỗi", core.rut_gon_text(exc, 300)))
            continue
        tai_ve.append((f, path))
        dong_tai.append(thd.dong_excel_tu_file(id_don, f, path, "Đã tải", ""))
    log(f"   ⬇ Đã tải {len(tai_ve)}/{len(files)} file → {thu_muc_hsq}")

    # ---- 2b. Đơn chưa có GCN trên MPLIS: AI tìm trang GCN trong HSQ
    if not ds_gcn:
        rows += tach_gcn_khong_co_tren_mplis(id_don, tai_ve, analyzer, model, cache, usage, out_dir, loc_truoc,
                                             nen_dung)

    # ---- 3. Không tách được trang GCN nào → tìm đơn đăng ký
    if not any(r.get("file_gcn") for r in rows) and not nen_dung():
        rows += lay_don(id_don, tai_ve, analyzer, model, cache, usage, out_dir, nen_dung)

    rows += _dong_gcn_chua_so(id_don, gcn_chua_so, rows)
    for r in rows:
        r.setdefault("file_tach", r.pop("file_gcn", ""))
    return rows, dong_tai


# ============================ EXCEL ============================

def ghi_excel(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_tt = EXCEL_HEADERS.index("trang_thai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in EXCEL_HEADERS])
        mau = MAU_TRANG_THAI.get(str(row.get("trang_thai") or ""))
        if mau:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=mau)
    for i, h in enumerate(EXCEL_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}1"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


# ============================ CHẠY NHIỀU ĐƠN ============================

def chay(client, ids: list[int], out_dir: str, api_key: str, model: str = DEFAULT_MODEL, rpm: float = 30,
         dung_file_da_tai: bool = False, nen_dung: Callable[[], bool] = lambda: False,
         bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None,
         loc_truoc: bool = True, quet_het: bool = False,
         muc_thieu: list | None = None, tu_kiem_tra_thieu: bool = False,
         bo_qua_nhom_1: bool = True) -> tuple[str, str, str]:
    """
    Xử lý danh sách đơn → (Excel tách, Excel tải, Excel thông tin thiếu hoặc ''). Hết phiên → dừng cả loạt
    (đã ghi Excel). muc_thieu: các lay_thong_tin_thieu_hsq.MucThieu — đơn của chúng cũng được tải & tách, rồi
    AI đọc các trường thiếu trên chính HSQ vừa tải (chung file tải, chung cache AI).
    tu_kiem_tra_thieu: mỗi đơn hỏi MPLIS đơn đang thiếu thông tin gì (thông báo chiTietThongTinThieu) rồi đọc
    đúng các thông tin đó — cộng thêm vào muc_thieu (bỏ trùng).
    bo_qua_nhom_1: đơn đã đạt nhóm 1 trên MPLIS → bỏ qua (không tải, không tách, không đọc thông tin thiếu).
    """
    import lay_thong_tin_thieu_hsq as ltt  # import ở đây: ltt import ngược module này

    muc_thieu = list(muc_thieu or [])
    theo_don: dict[int, list] = {}
    for m in muc_thieu:
        theo_don.setdefault(m.id_don, []).append(m)
    ids = list(dict.fromkeys([*ids, *theo_don]))
    xong: set[int] = set()
    if not api_key:
        raise RuntimeError("Chưa có Gemini API key (GEMINI_API_KEY)")
    analyzer = GeminiAnalyzer(api_key, model, rpm=rpm)
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    token = client.session.headers.get("__requestverificationtoken", "")
    downloader = thd.Downloader(None, client.session, token, thd.doc_mau_da_luu(), thd.MAU_MAC_DINH)
    cache_path = Path(out_dir) / tg.FILE_CACHE
    cache = load_cache(cache_path)
    usage = Usage()
    stamp = time.strftime("%Y%m%d_%H%M%S")
    excel = os.path.join(out_dir, f"ket_qua_tach_hsq_{stamp}.xlsx")
    excel_tai = os.path.join(out_dir, f"ket_qua_tai_hsq_{stamp}.xlsx")
    excel_thieu = os.path.join(out_dir, f"ket_qua_thong_tin_thieu_{stamp}.xlsx") if muc_thieu or tu_kiem_tra_thieu else ""
    rows: list[dict[str, Any]] = []
    rows_tai: list[dict[str, Any]] = []

    def luu() -> None:
        try:
            ghi_excel(excel, rows)
            thd.ghi_excel_ket_qua(excel_tai, rows_tai)
            if excel_thieu:
                ltt.ghi_excel(excel_thieu, [ltt.dong_excel(m, m.id_don in xong) for m in muc_thieu])
        except PermissionError:
            log("⚠ Không ghi được Excel kết quả (file đang mở?)")
        save_cache(cache_path, cache)

    log(f"Tải & tách HSQ cho {len(ids)} đơn" + (f", lấy {len(muc_thieu)} mục thông tin thiếu" if muc_thieu else "")
        + (", tự kiểm tra thông tin thiếu trên MPLIS" if tu_kiem_tra_thieu else "")
        + f" | model {model} → {out_dir}")
    nhom1_kq: dict[int, Any] = {}
    tt = tg.TaiTruoc(client, downloader, out_dir, dung_file_da_tai)
    tt.dang_ky(ids)
    try:
        for n, id_don in enumerate(ids, 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(n - 1, len(ids), f"Đơn {id_don} ({n}/{len(ids)})")
            log(f"--- [{n}/{len(ids)}] Đơn {id_don} ---")
            for id_sau in ids[n:n + tg.SO_DON_TAI_TRUOC]:
                tt.chuan_bi_don(id_sau)
            try:
                dong, dong_tai = xu_ly_don(client, analyzer, id_don, out_dir, cache, usage, tt, model, loc_truoc,
                                           quet_het, nen_dung, bo_qua_nhom_1, nhom1_kq)
                rows.extend(dong)
                rows_tai.extend(dong_tai)
                da_dat, ban_ghi = nhom1_kq.get(id_don, (False, None))
                if da_dat:
                    for m in theo_don.get(id_don, []):
                        m.loi.append(nhom1.LY_DO_DA_DAT)
                    xong.add(id_don)
                    luu()
                    continue
                if tu_kiem_tra_thieu and not nen_dung():
                    _them_muc_thieu_mplis(client, tt, id_don, muc_thieu, theo_don, ban_ghi)
                if theo_don.get(id_don) and not nen_dung():
                    log(f"   --- Lấy {len(theo_don[id_don])} mục thông tin thiếu ---")
                    ltt.xu_ly_don(analyzer, id_don, theo_don[id_don], cache, usage, tt, model, nen_dung)
            except core.HetPhien as exc:
                log(f"❌ {exc}")
                log("⛔ HẾT PHIÊN — đã dừng. Đăng nhập lại rồi chạy tiếp các đơn chưa xử lý.")
                kq = exc.ket_qua if isinstance(exc.ket_qua, dict) else {}
                rows.extend(kq.get("dong", []))
                rows.append(_dong_don(id_don, trang_thai="Lỗi", ghi_chu=f"HẾT PHIÊN — ĐÃ DỪNG: {exc}"))
                rows_tai.extend(kq.get("dong_tai", []))
                for m in theo_don.get(id_don, []):
                    m.loi.append(f"HẾT PHIÊN — ĐÃ DỪNG: {exc}")
                xong.add(id_don)
                luu()
                break
            except day.DungKhanCap:
                luu()
                raise
            except Exception as exc:
                log(f"❌ Đơn {id_don}: {core.rut_gon_text(exc, 300)}")
                rows.append(_dong_don(id_don, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)))
                for m in theo_don.get(id_don, []):
                    m.loi.append(f"Lỗi: {core.rut_gon_text(exc, 200)}")
            xong.add(id_don)
            luu()
    finally:
        tt.dong()
    bao_tien_do(len(ids), len(ids), "Xong")
    _tong_ket(rows, rows_tai, out_dir, usage, excel)
    if excel_thieu:
        dem = [ltt.dong_excel(m, m.id_don in xong)["trang_thai"] for m in muc_thieu]
        log("Thông tin thiếu: " + ", ".join(f"{t} {dem.count(t)}" for t in ltt.MAU_TRANG_THAI if dem.count(t))
            + f" → {excel_thieu}")
    return excel, excel_tai, excel_thieu


def _them_muc_thieu_mplis(client, tt: tg.TaiTruoc, id_don: int, muc_thieu: list, theo_don: dict,
                          ban_ghi: list[dict] | None = None) -> None:
    """
    Hỏi MPLIS đơn đang thiếu gì → thêm các mục chưa có vào muc_thieu / theo_don. Hết phiên → ném lên.
    ban_ghi: bản ghi làm sạch đã tra lúc kiểm tra nhóm 1 → dùng lại, không tra MPLIS lần nữa.
    """
    import lay_thong_tin_thieu_hsq as ltt

    detail = tt.chi_tiet(id_don)
    if not detail:
        return
    try:
        if ban_ghi is not None and not ltt.TRA_CUU_TRUC_TIEP_THEO_ID:
            moi = ltt.muc_thieu_tu_ban_ghi(id_don, ban_ghi) or []
        else:
            with tt.khoa:  # không gửi chồng với luồng tải nền
                moi = ltt.lay_muc_thieu_mplis(client, id_don, detail)
    except core.HetPhien:
        raise
    except Exception as exc:
        log(f"   ⚠ Không tra được thông tin thiếu trên MPLIS: {core.rut_gon_text(exc, 200)}")
        return
    co = {m.khoa for m in theo_don.get(id_don, [])}
    them = [m for m in moi if m.khoa not in co]
    muc_thieu.extend(them)
    theo_don.setdefault(id_don, []).extend(them)
    log(f"   MPLIS: đơn thiếu {len(moi)} mục" + (": " + ", ".join(dict.fromkeys(m.truong for m in moi)) if moi else
                                               " — không cần đọc thêm"))


def _tong_ket(rows: list[dict[str, Any]], rows_tai: list[dict[str, Any]], out_dir: str, usage: Usage,
              excel: str) -> None:
    def dem(tt: str) -> int:
        return sum(r.get("trang_thai") == tt for r in rows)

    log("=" * 60)
    log(f"⬇ Đã tải {sum(r.get('trang_thai_tai') == 'Đã tải' for r in rows_tai)} file HSQ → "
        f"{os.path.join(out_dir, tg.THU_MUC_HSQ)}")
    log(f"✅ Đã tách {dem('Đã tách')} GCN → {os.path.join(out_dir, tg.THU_MUC_GCN)}")
    log(f"✅ Đơn đăng ký: {dem('Đã lấy đơn')} chép nguyên, {dem('Đã tách đơn')} tách trang → "
        f"{os.path.join(out_dir, THU_MUC_DON)}")
    if dem("Không thấy"):
        log(f"⚠ {dem('Không thấy')} GCN (có trên MPLIS) không tìm thấy trong HSQ")
    if dem("Sai chủ") + dem("Chưa rõ chủ"):
        log(f"⚠ {dem('Sai chủ') + dem('Chưa rõ chủ')} GCN không có bản đúng chủ → "
            f"{os.path.join(out_dir, tg.THU_MUC_CAN_KIEM_TRA)}")
    if dem("Không khớp"):
        log(f"⚠ {dem('Không khớp')} nhóm trang GCN không thuộc GCN hiện hành → "
            f"{os.path.join(out_dir, tg.THU_MUC_KHONG_KHOP)}")
    so_nhom_1 = sum(r.get("ghi_chu") == nhom1.LY_DO_DA_DAT for r in rows)
    if so_nhom_1:
        log(f"⏭ {so_nhom_1} đơn đã đạt nhóm 1 trên MPLIS — bỏ qua, không tải")
    if dem("Không thấy đơn"):
        log(f"⚠ {dem('Không thấy đơn')} đơn không có GCN và không tìm thấy đơn đăng ký")
    log(f"Chi phí Gemini ước tính: {usage.estimated_cost_usd:.4f} USD "
        f"({usage.input_tokens} token vào, {usage.output_tokens} token ra)")
    log(f"Excel kết quả: {excel}")
