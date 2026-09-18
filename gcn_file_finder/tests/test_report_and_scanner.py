from pathlib import Path

import pytest
from openpyxl import load_workbook

from gcn_file_finder.core.excel_reader import read_gcn_excel
from gcn_file_finder.core.file_scanner import scan_source_files, validate_source_destination
from gcn_file_finder.core.report_writer import CopiedFile, write_report
from gcn_file_finder.models import FileScanResult, GCNMatch


def _excel_data(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DS"
    sheet.append(["GCN"])
    sheet.append(["AM 143443"])
    sheet.append(["CS 012345"])
    sheet.append(["sai"])
    workbook.save(path)
    return read_gcn_excel(path, "DS", "GCN")


def test_report_has_required_sheets_and_display_name(tmp_path):
    data = _excel_data(tmp_path)
    source = tmp_path / "source.pdf"
    destination = tmp_path / "AM 143443.pdf"
    source.write_bytes(b"source")
    destination.write_bytes(b"source")
    result = FileScanResult(
        source,
        "PDF",
        page_count=3,
        has_text_layer=True,
        matches=[
            GCNMatch("AM143443", "AM 143443", "PDF TEXT", page=1, raw_text="AM-143443"),
            GCNMatch("AM143443", "AM 143443", "PDF TEXT", page=3, raw_text="AM 143443"),
        ],
        status="ĐÃ TÌM THẤY",
    )
    report = write_report(
        tmp_path / "KET_QUA_TIM_GCN.xlsx",
        data,
        [result],
        [CopiedFile("AM143443", source, destination, "ĐÃ TÌM THẤY")],
        1.2,
    )
    workbook = load_workbook(report, data_only=True)
    assert workbook.sheetnames == ["KET_QUA", "FILE_DA_QUET", "TONG_HOP"]
    rows = list(workbook["KET_QUA"].iter_rows(min_row=2, values_only=True))
    assert rows[0][4] == "AM 143443"
    assert rows[0][9] == "1, 3"
    assert rows[0][14] == "AM 143443.pdf"
    assert rows[1][5] == "KHÔNG TÌM THẤY"
    assert rows[2][5] == "GCN KHÔNG HỢP LỆ"
    assert workbook["KET_QUA"].freeze_panes == "A2"
    workbook.close()


def test_scanner_filters_temp_report_and_unsupported(tmp_path):
    source = tmp_path / "source"
    destination = tmp_path / "result"
    source.mkdir()
    destination.mkdir()
    (source / "a.pdf").write_bytes(b"")
    (source / "b.JPG").write_bytes(b"")
    (source / "~$temp.pdf").write_bytes(b"")
    (source / "KET_QUA_TIM_GCN.xlsx").write_bytes(b"")
    (source / "notes.txt").write_bytes(b"")
    assert [path.name for path in scan_source_files(source, destination)] == ["a.pdf", "b.JPG"]


def test_destination_inside_source_is_rejected(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    with pytest.raises(ValueError):
        validate_source_destination(source, source / "result")
