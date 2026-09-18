"""Xuất báo cáo Excel định dạng đầy đủ bằng openpyxl."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gcn_file_finder.core.gcn_normalizer import format_gcn_display
from gcn_file_finder.models import ExcelData, FileScanResult, GCNMatch


@dataclass(frozen=True)
class CopiedFile:
    """Liên kết một GCN, file nguồn và file kết quả."""

    gcn_key: str
    source_path: Path
    destination_path: Path
    status: str


RESULT_HEADERS = [
    "STT", "Dòng Excel", "GCN gốc trong Excel", "Khóa GCN nội bộ", "GCN chuẩn hiển thị",
    "Trạng thái", "Số lượng file khớp", "Tên file nguồn", "Đường dẫn file nguồn",
    "Trang PDF tìm thấy", "Chuỗi OCR gốc", "Chuỗi sau chuẩn hóa", "Phương thức tìm thấy",
    "Độ tin cậy OCR", "Tên file kết quả", "Đường dẫn file kết quả", "Thời gian xử lý", "Ghi chú",
]
FILE_HEADERS = [
    "STT", "Tên file", "Đường dẫn", "Loại file", "Số trang", "Có text layer", "GCN tìm thấy",
    "Trang tìm thấy", "Trạng thái", "Thông báo lỗi", "Thời gian xử lý",
]


def _pages(matches: Iterable[GCNMatch]) -> str:
    return ", ".join(str(page) for page in sorted({m.page for m in matches if m.page is not None}))


def _format_sheet(sheet, headers: list[str]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        status = str(row[5].value if sheet.title == "KET_QUA" else row[8].value)
        color = None
        if status in {"ĐÃ TÌM THẤY", "ĐÃ TỒN TẠI"}:
            color = "E2F0D9"
        elif status in {"NHIỀU FILE", "CẦN KIỂM TRA"}:
            color = "FFF2CC"
        elif status in {"LỖI XỬ LÝ", "KHÔNG TÌM THẤY", "GCN KHÔNG HỢP LỆ"}:
            color = "FCE4D6"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {9, 16, 18})
            if color:
                cell.fill = PatternFill("solid", fgColor=color)
    for index, header in enumerate(headers, 1):
        values = [len(str(sheet.cell(row, index).value or "")) for row in range(1, min(sheet.max_row, 200) + 1)]
        width = min(65, max(len(header) + 2, max(values, default=8) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_report(
    path: str | Path,
    excel_data: ExcelData,
    file_results: list[FileScanResult],
    copied_files: list[CopiedFile],
    total_elapsed_seconds: float,
) -> Path:
    """Tạo hoặc thay thế báo cáo của chính phiên chạy hiện tại."""

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    result_sheet = workbook.active
    result_sheet.title = "KET_QUA"
    result_sheet.append(RESULT_HEADERS)

    by_key: dict[str, list[tuple[FileScanResult, GCNMatch]]] = defaultdict(list)
    review_by_key: dict[str, list[tuple[FileScanResult, GCNMatch]]] = defaultdict(list)
    for file_result in file_results:
        for match in file_result.matches:
            if match.key:
                by_key[match.key].append((file_result, match))
            for alternative in match.alternatives if match.needs_review else ():
                review_by_key[alternative].append((file_result, match))
    copies_by_pair = {(item.gcn_key, str(item.source_path.resolve())): item for item in copied_files}

    result_index = 1
    for excel_row in excel_data.rows:
        if not excel_row.key:
            result_sheet.append([
                result_index, excel_row.row_number, excel_row.original_value, "", "", "GCN KHÔNG HỢP LỆ",
                0, "", "", "", "", "", "", "", "", "", "", "Không đưa vào tập đối chiếu",
            ])
            result_index += 1
            continue
        grouped: dict[str, list[tuple[FileScanResult, GCNMatch]]] = defaultdict(list)
        for item in by_key.get(excel_row.key, []):
            grouped[str(item[0].source_path.resolve())].append(item)
        if not grouped and review_by_key.get(excel_row.key):
            for file_result, match in review_by_key[excel_row.key]:
                result_sheet.append([
                    result_index, excel_row.row_number, excel_row.original_value, excel_row.key, excel_row.display,
                    "CẦN KIỂM TRA", 0, file_result.source_path.name, str(file_result.source_path),
                    match.page or "", match.raw_text, " / ".join(match.alternatives), match.method,
                    match.confidence if match.confidence is not None else "", "", "", file_result.elapsed_seconds,
                    "Có nhiều phương án sửa OCR cùng khớp Excel; không sao chép.",
                ])
                result_index += 1
            continue
        if not grouped:
            result_sheet.append([
                result_index, excel_row.row_number, excel_row.original_value, excel_row.key, excel_row.display,
                "KHÔNG TÌM THẤY", 0, "", "", "", "", excel_row.key, "", "", "", "", "", "",
            ])
            result_index += 1
            continue
        file_count = len(grouped)
        for source_string, entries in grouped.items():
            file_result = entries[0][0]
            matches = [entry[1] for entry in entries]
            copied = copies_by_pair.get((excel_row.key, source_string))
            status = copied.status if copied else "LỖI XỬ LÝ"
            if file_count > 1 and status == "ĐÃ TÌM THẤY":
                status = "NHIỀU FILE"
            raw_texts = " | ".join(dict.fromkeys(match.raw_text for match in matches if match.raw_text))
            methods = " | ".join(dict.fromkeys(match.method for match in matches))
            confidences = [match.confidence for match in matches if match.confidence is not None]
            result_sheet.append([
                result_index, excel_row.row_number, excel_row.original_value, excel_row.key, excel_row.display,
                status, file_count, file_result.source_path.name, str(file_result.source_path), _pages(matches),
                raw_texts, excel_row.key, methods, min(confidences) if confidences else "",
                copied.destination_path.name if copied else "", str(copied.destination_path) if copied else "",
                file_result.elapsed_seconds, "Các dòng trùng Excel cùng tham chiếu kết quả này" if excel_row.is_duplicate else "",
            ])
            result_index += 1

    file_sheet = workbook.create_sheet("FILE_DA_QUET")
    file_sheet.append(FILE_HEADERS)
    for index, result in enumerate(file_results, 1):
        exact = [match for match in result.matches if match.key]
        file_sheet.append([
            index, result.source_path.name, str(result.source_path), result.file_type, result.page_count,
            "Có" if result.has_text_layer else "Không",
            ", ".join(format_gcn_display(key) for key in dict.fromkeys(m.key for m in exact) if key),
            _pages(exact), result.status, result.error, result.elapsed_seconds,
        ])

    summary = workbook.create_sheet("TONG_HOP")
    summary.append(["Chỉ tiêu", "Giá trị"])
    found_keys = {match.key for result in file_results for match in result.matches if match.key}
    counts = {key: len({str(result.source_path.resolve()) for result in file_results for match in result.matches if match.key == key}) for key in found_keys}
    metrics = [
        ("Tổng số dòng Excel", excel_data.total_data_rows),
        ("Tổng số GCN hợp lệ", len(excel_data.valid_gcn_keys)),
        ("Số GCN không hợp lệ", excel_data.invalid_count),
        ("Số GCN trùng", excel_data.duplicate_count),
        ("Số GCN đã tìm thấy", len(found_keys)),
        ("Số GCN chưa tìm thấy", len(excel_data.valid_gcn_keys - found_keys)),
        ("Số GCN khớp nhiều file", sum(count > 1 for count in counts.values())),
        ("Số trường hợp cần kiểm tra", sum(match.needs_review for result in file_results for match in result.matches)),
        ("Tổng số file đã quét", len(file_results)),
        ("Tổng số file lỗi", sum(result.status == "LỖI XỬ LÝ" for result in file_results)),
        ("Tổng số file đã sao chép", sum(item.status == "ĐÃ TÌM THẤY" for item in copied_files)),
        ("Tổng số file đã tồn tại", sum(item.status == "ĐÃ TỒN TẠI" for item in copied_files)),
        ("Tổng thời gian xử lý", f"{total_elapsed_seconds:.2f} giây"),
        ("Thời điểm xuất báo cáo", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
    ]
    for metric in metrics:
        summary.append(metric)

    _format_sheet(result_sheet, RESULT_HEADERS)
    _format_sheet(file_sheet, FILE_HEADERS)
    summary.freeze_panes = "A2"
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 24
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    workbook.save(target)
    workbook.close()
    return target
