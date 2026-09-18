"""Đọc GCN từ Excel bằng openpyxl và luôn giữ dữ liệu dạng chuỗi."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from gcn_file_finder.core.gcn_normalizer import format_gcn_display, normalize_gcn_key
from gcn_file_finder.models import ExcelData, GCNRow


def list_sheets(path: str | Path) -> list[str]:
    """Liệt kê sheet mà không nạp toàn bộ workbook."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def list_columns(path: str | Path, sheet_name: str) -> list[str]:
    """Liệt kê cột theo hàng tiêu đề đầu tiên có dữ liệu."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows():
            if any(cell.value not in (None, "") for cell in row):
                return [f"{get_column_letter(cell.column)} - {str(cell.value).strip() if cell.value is not None else '(trống)'}" for cell in row]
        return []
    finally:
        workbook.close()


def _header_key(value: object) -> str:
    return "".join(str(value or "").upper().split())


def find_header_and_column(sheet, requested_column: str = "GCN") -> tuple[int, int]:
    """Tìm hàng tiêu đề và chỉ số cột theo tên hoặc chữ cột."""

    requested = requested_column.split(" - ", 1)[0].strip()
    requested_header = requested_column.split(" - ", 1)[-1].strip()
    for row in sheet.iter_rows():
        nonempty = [cell for cell in row if cell.value not in (None, "")]
        if not nonempty:
            continue
        if requested.isalpha() and len(requested) <= 3 and " - " in requested_column:
            for cell in row:
                if get_column_letter(cell.column).upper() == requested.upper():
                    return cell.row, cell.column
        wanted = _header_key(requested_header if " - " in requested_column else requested_column)
        for cell in row:
            if _header_key(cell.value) == wanted:
                return cell.row, cell.column
        if wanted == "GCN":
            for cell in row:
                if _header_key(cell.value) == "GCN":
                    return cell.row, cell.column
        break
    raise ValueError(f"Không tìm thấy cột {requested_column!r} trong sheet {sheet.title!r}.")


def read_gcn_excel(path: str | Path, sheet_name: str, column: str = "GCN") -> ExcelData:
    """Đọc cột GCN, ghi nhận dòng lỗi và giá trị trùng."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet không tồn tại: {sheet_name}")
        sheet = workbook[sheet_name]
        header_row, column_index = find_header_and_column(sheet, column)
        staged: list[tuple[int, str, str | None]] = []
        for row_number in range(header_row + 1, sheet.max_row + 1):
            value = sheet.cell(row_number, column_index).value
            if value is None or str(value).strip() == "":
                continue
            # Một ô có thể chứa nhiều GCN cách nhau bằng dấu ";" (ví dụ
            # "DH 358799;DH 358800"); mỗi phần được tách và chuẩn hóa riêng.
            for part in str(value).split(";"):
                part = part.strip()
                if not part:
                    continue
                staged.append((row_number, part, normalize_gcn_key(part)))
    finally:
        workbook.close()

    counts: dict[str, int] = {}
    for _, _, key in staged:
        if key:
            counts[key] = counts.get(key, 0) + 1
    rows = [
        GCNRow(number, original, key, format_gcn_display(key) if key else None, bool(key and counts[key] > 1))
        for number, original, key in staged
    ]
    gcn_rows: dict[str, list[int]] = {}
    originals: dict[str, list[str]] = {}
    for item in rows:
        if item.key:
            gcn_rows.setdefault(item.key, []).append(item.row_number)
            originals.setdefault(item.key, []).append(item.original_value)
    return ExcelData(
        rows=rows,
        valid_gcn_keys=set(gcn_rows),
        gcn_rows=gcn_rows,
        gcn_original_values=originals,
        total_data_rows=len(rows),
        invalid_count=sum(item.key is None for item in rows),
        duplicate_count=sum(max(0, count - 1) for count in counts.values()),
    )
