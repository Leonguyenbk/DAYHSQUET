# -*- coding: utf-8 -*-
"""
Đẩy file HSQ theo mô tả, mỗi đơn 1 thư mục riêng — thay ĐÚNG thành phần có mô tả trùng tên file.

Cấu trúc thư mục:
    <thư mục gốc>/
        13456477/                         ← tên thư mục bắt đầu bằng mã đơn (tinhHinhDangKyId)
            Đơn đăng ký.pdf               ← tên file = mô tả (moTa) của thành phần HSQ cần thay
            Giấy chứng nhận CP 099228.pdf
        12913166_Nguyen Van A/            ← được, chỉ lấy dãy số đầu
            03_Trích lục thửa đất.pdf     ← tên file tải bằng TaiDayHSQ ("NN_<mô tả>") cũng được
            24232_GCN_DH 770214-GCN.pdf   ← thêm "-GCN" cuối tên + có số GCN: thay file + ghi nhận Giấy chứng nhận
            24232_GCN_BA 796153-GCN DH 770214.pdf  ← "-GCN <số phát hành>": chỉ rõ gắn GCN nào

Với mỗi đơn:
  1. Lấy chi tiết đơn → mọi thành phần HSQ (moTa, vị trí).
  2. Khớp từng PDF trong thư mục đơn với mô tả: so sánh tên đã bỏ dấu / ký tự đặc biệt / hậu tố ký số
     (_signed...), KHÔNG đoán gần đúng. Mô tả trùng nhau trong đơn → cần tiền tố "NN_" (số thứ tự như
     lúc tải bằng TaiDayHSQ) để phân biệt, không có thì bỏ qua file đó.
     Nhiều file cùng khớp 1 thành phần → ưu tiên file đã ký số nếu chỉ có 1 file ký; còn lại bỏ qua.
  3. Mỗi HSQ gửi 1 request kèm ĐỦ thành phần (day_hsq_ky_so.gui_ca_hsq): thành phần không có file mới
     giữ nguyên, thành phần có file mới → nodeId="" + file mới. Mô tả, loại HSQ, GCN giữ nguyên như lúc tra cứu.
     Riêng thành phần có mô tả chứa "GCN" (hoặc "Giấy chứng nhận") + số phát hành của đúng 1 GCN hiện hành
     của đơn → ghi nhận là Giấy chứng nhận + gắn GCN đó (nếu chưa). Sau đó HSQ vẫn chưa có thành phần loại
     Giấy chứng nhận / Đơn đăng ký / Thông báo xác nhận → không gửi (MPLIS sẽ từ chối: "Hồ sơ quét không
     đảm bảo thành phần theo quy định"), báo trong Excel.
     File đặt tên "<mô tả>-GCN" / "<mô tả>-GCN <số phát hành>" → luôn ghi nhận là Giấy chứng nhận, gắn GCN
     theo số ghi trong tên file; không ghi số thì GCN có số nằm trong mô tả / tên file. Không có số GCN
     hiện hành nào → bỏ qua file đó (không gán bừa).
  4. Xác nhận: số thành phần không giảm (giảm → DỪNG TOÀN BỘ), thành phần thay phải có nodeId mới,
     loại đã bổ sung phải đúng trên MPLIS.
Không tạo HSQ mới, không thêm thành phần: đơn chưa có HSQ / file không khớp mô tả nào → báo lại.
"""

from __future__ import annotations

import copy
import os
import re
import time
from typing import Any, Callable

import cap_nhat_don_hsq as core
import day_hsq_ky_so as day
import tai_hsq_don as thd

log = thd.log

EXCEL_HEADERS = ["stt", "madon", "thu_muc", "file", "hosoquetid", "vitri", "mota", "loaihosoquet",
                 "loai_moi", "da_ky_so", "trang_thai", "ghi_chu"]
EXCEL_WIDTHS = {"stt": 6, "madon": 12, "thu_muc": 22, "file": 45, "hosoquetid": 11, "vitri": 6, "mota": 45,
                "loaihosoquet": 18, "loai_moi": 28, "da_ky_so": 8, "trang_thai": 13, "ghi_chu": 60}
# MPLIS chỉ lưu HSQ có ít nhất 1 thành phần thuộc các loại này (Giấy chứng nhận / Đơn đăng ký / Thông báo xác nhận).
LOAI_BAT_BUOC = {1, 2, 3}
# MPLIS trả rỗng (không báo lỗi) khi hết phiên đăng nhập → bao nhiêu đơn liên tiếp như vậy thì dừng hẳn.
LOI_CHI_TIET = "Không lấy được chi tiết đơn (MPLIS trả rỗng — thường do hết phiên đăng nhập)"
DUNG_SAU_SO_DON_LOI = 3
MAU_TRANG_THAI = {"Đã thay": "C6EFCE", "Đã bổ sung loại": "C6EFCE", "Lỗi": "F4CCCC", "Bỏ qua": "FFF2CC", "Chạy thử": "D9EAF7",
                  "Giữ nguyên": "EDEDED"}


# ============================ KHỚP TÊN FILE ↔ MÔ TẢ ============================

def chuan_hoa(s: Any) -> str:
    """Bỏ dấu, chữ thường, chỉ giữ chữ + số (tên file không chứa được / : ? ... nên bỏ hết ký tự đặc biệt)."""
    s = core.remove_accents(str(s or "")).lower().replace("đ", "d")
    return re.sub(r"[^a-z0-9]+", "", s)


def tach_ten_file(path: str) -> tuple[int, str]:
    """'03_Đơn đăng ký_signed.pdf' → (3, 'Đơn đăng ký'). Không có tiền tố số → (0, tên)."""
    stem = os.path.splitext(os.path.basename(path))[0].strip()
    for hau_to in sorted((h for h in day.HAU_TO_KY_SO if h), key=len, reverse=True):
        if stem.lower().endswith(hau_to.lower()):
            stem = stem[: -len(hau_to)].strip()
            break
    m = re.match(r"^(\d{1,3})[_.\-\s]+(.+)$", stem)
    if m:
        return int(m.group(1)), m.group(2).strip()
    return 0, stem


def ma_don_tu_thu_muc(ten: str) -> int:
    m = re.match(r"^\s*(\d{5,10})(?!\d)", ten)
    return int(m.group(1)) if m else 0


def quet_thu_muc_goc(goc: str) -> tuple[dict[int, list[str]], list[str]]:
    """{mã đơn: [file PDF]} theo thư mục con + danh sách thư mục con không nhận ra mã đơn."""
    theo_don: dict[int, list[str]] = {}
    khong_nhan: list[str] = []
    for ten in sorted(os.listdir(goc)):
        thu_muc = os.path.join(goc, ten)
        if not os.path.isdir(thu_muc):
            continue
        madon = ma_don_tu_thu_muc(ten)
        if not madon:
            khong_nhan.append(ten)
            continue
        pdfs = [os.path.join(thu_muc, f) for f in sorted(os.listdir(thu_muc))
                if f.lower().endswith(".pdf") and os.path.isfile(os.path.join(thu_muc, f))]
        theo_don.setdefault(madon, []).extend(pdfs)
    return theo_don, khong_nhan


def thanh_phan_cua_don(detail: dict[str, Any]) -> list[dict[str, Any]]:
    """
    Mọi thành phần HSQ của đơn, đánh số thứ tự giống lúc TaiDayHSQ tải (tai_hsq_don.liet_ke_file_hsq:
    duyệt lần lượt các HSQ, bỏ nodeId trùng) để tiền tố "NN_" của file tải về khớp đúng thành phần.
    """
    ket_qua: list[dict[str, Any]] = []
    da_co: set[str] = set()
    for hoso in detail.get("ListHoSoQuet") or []:
        if not isinstance(hoso, dict):
            continue
        for vi_tri, node in enumerate(day.cac_node(hoso)):
            node_id = str(node.get("nodeId") or "")
            if node_id and node_id in da_co:
                continue
            da_co.add(node_id)
            ket_qua.append({
                "hoso": hoso, "node": node, "vi_tri": vi_tri, "stt": len(ket_qua) + 1,
                "khoa": chuan_hoa(node.get("moTa") or node.get("tenGiayTo")),
            })
    return ket_qua


def tach_hau_to_gcn(ten: str) -> tuple[str, str] | None:
    """'PC_A-GCN' → ('PC_A', ''); 'X-GCN DH 770214' → ('X', 'DH 770214'); không có hậu tố → None."""
    m = re.match(r"^(.+?)\s*-\s*GCN(?:[\s_]+(.+?))?\s*$", ten, re.I)
    return (m.group(1).strip(), (m.group(2) or "").strip()) if m else None


def khop_file(path: str, ds_tp: list[dict[str, Any]]) -> tuple[dict | None, str, str | None]:
    """
    Tìm thành phần cho 1 file. Trả về (thành phần, '', None) hoặc (None, lý do, None);
    file có hậu tố "-GCN [số]" → phần tử thứ 3 là số GCN ghi trong tên ('' nếu không ghi).
    """
    stt, ten = tach_ten_file(path)
    hau_to = tach_hau_to_gcn(ten)
    if hau_to is not None:
        # Bỏ hậu tố "-GCN [số]" trước: khớp kiểu "tên chứa mô tả" với tên đầy đủ sẽ làm mất đánh dấu GCN.
        tp, ly_do = _khop_ten(stt, hau_to[0], hau_to[0], ds_tp)
        if tp is not None:
            return tp, "", hau_to[1]
    tp, ly_do = _khop_ten(stt, ten, os.path.splitext(os.path.basename(path))[0], ds_tp)
    return tp, ly_do, None


def _khop_ten(stt: int, ten: str, ten_day_du: str, ds_tp: list[dict[str, Any]]) -> tuple[dict | None, str]:
    for khoa in dict.fromkeys((chuan_hoa(ten), chuan_hoa(ten_day_du))):
        if not khoa:
            continue
        # TaiDayHSQ cắt tên file ở 120 ký tự → mô tả dài chỉ cần khớp phần đầu.
        ung_vien = [tp for tp in ds_tp if tp["khoa"] == khoa or (len(khoa) >= 60 and tp["khoa"].startswith(khoa))]
        if len(ung_vien) == 1:
            return ung_vien[0], ""
        if len(ung_vien) > 1:
            theo_stt = [tp for tp in ung_vien if stt and tp["stt"] == stt]
            if len(theo_stt) == 1:
                return theo_stt[0], ""
            return None, (f"Mô tả '{ten}' trùng {len(ung_vien)} thành phần trong đơn — đặt tên file có tiền tố "
                          f"số thứ tự (vd {ung_vien[0]['stt']:02d}_{ten}.pdf) để chỉ rõ thành phần")
    # Tên file có thêm chữ trước/sau mô tả (vd 'Ly Kim Hoang 24221_3_96_GCN_DI 923410') → mô tả dài nhất nằm trong tên.
    khoa = chuan_hoa(ten)
    ung_vien = [tp for tp in ds_tp if mo_ta_nam_trong(khoa, tp["khoa"])]
    if ung_vien:
        dai_nhat = max(len(tp["khoa"]) for tp in ung_vien)
        ung_vien = [tp for tp in ung_vien if len(tp["khoa"]) == dai_nhat]
        if len(ung_vien) == 1:
            return ung_vien[0], ""
        theo_stt = [tp for tp in ung_vien if stt and tp["stt"] == stt]
        if len(theo_stt) == 1:
            return theo_stt[0], ""
        return None, (f"Tên file '{ten}' chứa mô tả '{ung_vien[0]['node'].get('moTa')}' của {len(ung_vien)} thành "
                      f"phần — đặt tên file có tiền tố số thứ tự để chỉ rõ thành phần")
    return None, f"Không có thành phần nào mô tả '{ten}'"


MO_TA_TOI_THIEU = 8  # mô tả ngắn hơn (GT, PC, GCN...) phải khớp nguyên tên file, không khớp kiểu "nằm trong"


def mo_ta_nam_trong(khoa_file: str, khoa_mo_ta: str) -> bool:
    """Tên file (đã chuẩn hoá) chứa nguyên mô tả (đã chuẩn hoá) đủ dài."""
    return len(khoa_mo_ta) >= MO_TA_TOI_THIEU and khoa_mo_ta in khoa_file


# ============================ BỔ SUNG LOẠI HSQ ============================

def gcn_cho_bo_sung(detail: dict[str, Any], tp_hsq: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    """
    {vị trí: GCN hiện hành} cho các thành phần có mô tả chứa "GCN"/"Giấy chứng nhận" + số phát hành của
    đúng 1 GCN hiện hành (vd '24221_3_96_GCN_DI 923410' ↔ GCN DI 923410) mà chưa phải loại Giấy chứng nhận
    gắn đúng GCN đó.
    """
    ds_gcn = [g for g in detail.get("ListGiayChungNhan") or []
              if isinstance(g, dict) and g.get("giayChungNhanId") and g.get("isLastest") is not False
              and len(chuan_hoa(g.get("soPhatHanh"))) >= 5]
    kq: dict[int, dict[str, Any]] = {}
    for tp in tp_hsq:
        node = tp["node"]
        mo_ta = chuan_hoa(node.get("moTa"))
        if "gcn" not in mo_ta and "giaychungnhan" not in mo_ta:
            continue
        khop = [g for g in ds_gcn if chuan_hoa(g.get("soPhatHanh")) in mo_ta]
        if len(khop) != 1:
            continue
        da_dung = (core.safe_int(node.get("loaiHoSoQuet"), -1) == 1
                   and core.safe_int(node.get("giayChungNhanId")) == core.safe_int(khop[0].get("giayChungNhanId")))
        if not da_dung:
            kq[tp["vi_tri"]] = khop[0]
    return kq


def gcn_cho_file_danh_dau(detail: dict[str, Any], tp: dict[str, Any], path: str,
                         so_ghi: str) -> tuple[dict | None, str]:
    """GCN hiện hành để gắn cho file '<mô tả>-GCN [số]'. Trả về (GCN, '') hoặc (None, lý do)."""
    ds = [g for g in detail.get("ListGiayChungNhan") or []
          if isinstance(g, dict) and g.get("giayChungNhanId") and g.get("isLastest") is not False]
    if not ds:
        return None, "File đánh dấu -GCN nhưng đơn chưa có GCN hiện hành nào"
    ten_ds = "; ".join(str(g.get("soPhatHanh") or g.get("giayChungNhanId")) for g in ds)
    if so_ghi:
        khop = [g for g in ds if chuan_hoa(g.get("soPhatHanh")) == chuan_hoa(so_ghi)]
        if len(khop) == 1:
            return khop[0], ""
        return None, f"Tên file ghi GCN '{so_ghi}' nhưng GCN hiện hành của đơn là: {ten_ds}"
    # Không ghi số sau "-GCN" → số GCN phải nằm trong mô tả / tên file (không gán bừa dù đơn chỉ 1 GCN).
    chuoi = chuan_hoa(tp["node"].get("moTa")) + chuan_hoa(os.path.basename(path))
    khop = [g for g in ds if len(chuan_hoa(g.get("soPhatHanh"))) >= 5 and chuan_hoa(g.get("soPhatHanh")) in chuoi]
    if len(khop) == 1:
        return khop[0], ""
    vd = f"{(tach_hau_to_gcn(tach_ten_file(path)[1]) or (tach_ten_file(path)[1],))[0]}-GCN {ds[0].get('soPhatHanh')}"
    return None, (f"Tên file / mô tả không chứa số GCN hiện hành nào của đơn ({ten_ds}) — ghi số vào tên file, "
                  f"vd '{vd}.pdf'")


def gan_loai_gcn(node: dict[str, Any], gcn: dict[str, Any]) -> None:
    node["loaiHoSoQuet"] = 1
    node["laGiayChungNhan"] = True
    node["giayChungNhanId"] = core.safe_int(gcn.get("giayChungNhanId"))
    node["versionGiayChungNhan"] = core.safe_int(gcn.get("version"), 11)


# ============================ XỬ LÝ 1 ĐƠN ============================

def _dong(madon: int, path: str = "", tp: dict | None = None, **kw) -> dict[str, Any]:
    node = (tp or {}).get("node") or {}
    loai = core.safe_int(node.get("loaiHoSoQuet"), -1)
    return {
        "madon": madon,
        "thu_muc": os.path.basename(os.path.dirname(path)) if path else "",
        "file": os.path.basename(path) if path else "",
        "hosoquetid": (tp or {}).get("hoso", {}).get("hoSoQuetId", "") if tp else "",
        "vitri": tp["stt"] if tp else "",
        "mota": node.get("moTa") or "",
        "loaihosoquet": core.LOAI_HO_SO_QUET_OPTIONS.get(loai, node.get("loaiHoSoQuet") or "") if tp else "",
        **kw,
    }


def xu_ly_don(client: core.MplisClient, madon: int, pdfs: list[str], chay_thu: bool) -> list[dict[str, Any]]:
    if not pdfs:
        return [_dong(madon, trang_thai="Bỏ qua", ghi_chu="Thư mục đơn không có file PDF")]

    detail = client.lay_chi_tiet_don(madon)
    if detail is None:
        return [_dong(madon, p, trang_thai="Lỗi", ghi_chu=LOI_CHI_TIET) for p in pdfs]
    ds_tp = thanh_phan_cua_don(detail)
    if not ds_tp:
        return [_dong(madon, p, trang_thai="Bỏ qua",
                      ghi_chu="Đơn chưa có HSQ — dùng tool đẩy HSQ (cap_nhat_don_hsq) để tạo mới") for p in pdfs]

    rows: list[dict[str, Any]] = []
    # ---- khớp file → thành phần
    theo_tp: dict[int, list[str]] = {}
    danh_dau_gcn: dict[str, str] = {}  # file "-GCN [số]" → số ghi trong tên
    for p in pdfs:
        tp, ly_do, so_ghi = khop_file(p, ds_tp)
        if so_ghi is not None:
            danh_dau_gcn[p] = so_ghi
        if tp is None:
            rows.append(_dong(madon, p, trang_thai="Bỏ qua", ghi_chu=ly_do))
            log(f"   ⚠ {os.path.basename(p)}: {ly_do}")
            continue
        theo_tp.setdefault(tp["stt"], []).append(p)

    # ---- mỗi thành phần đúng 1 file (nhiều file: chọn file ký số nếu chỉ 1 file đã ký)
    chon: dict[int, tuple[str, bool]] = {}
    for stt, files in theo_tp.items():
        tp = ds_tp[stt - 1]
        ky = [p for p in files if day.pdf_da_ky_so(p)]
        if len(files) == 1:
            chon[stt] = (files[0], bool(ky))
        elif len(ky) == 1:
            chon[stt] = (ky[0], True)
            for p in files:
                if p != ky[0]:
                    rows.append(_dong(madon, p, tp, trang_thai="Bỏ qua",
                                      ghi_chu=f"Đã dùng file ký số {os.path.basename(ky[0])} cho thành phần này"))
        else:
            ten = ", ".join(os.path.basename(p) for p in files)
            for p in files:
                rows.append(_dong(madon, p, tp, trang_thai="Bỏ qua",
                                  ghi_chu=f"{len(files)} file cùng khớp 1 thành phần ({ten}) — giữ lại 1 file"))
            log(f"   ⚠ '{tp['node'].get('moTa')}': {len(files)} file cùng khớp → bỏ qua")

    # ---- gửi theo từng HSQ
    for hoso in [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]:
        hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
        tp_hsq = [tp for tp in ds_tp if tp["hoso"] is hoso]
        thay = {tp["vi_tri"]: chon[tp["stt"]] for tp in tp_hsq if tp["stt"] in chon}

        # ---- file đánh dấu "-GCN": ghi nhận Giấy chứng nhận (không xác định được GCN → bỏ file đó)
        chi_dinh: dict[int, dict[str, Any]] = {}
        for tp in tp_hsq:
            i = tp["vi_tri"]
            if i not in thay or thay[i][0] not in danh_dau_gcn:
                continue
            gcn, ly_do = gcn_cho_file_danh_dau(detail, tp, thay[i][0], danh_dau_gcn[thay[i][0]])
            if gcn is None:
                rows.append(_dong(madon, thay.pop(i)[0], tp, trang_thai="Bỏ qua", ghi_chu=ly_do))
                log(f"   ⚠ '{tp['node'].get('moTa')}': {ly_do}")
            else:
                chi_dinh[i] = gcn

        if not thay:
            rows += [_dong(madon, tp=tp, trang_thai="Giữ nguyên", ghi_chu="Không có file mới") for tp in tp_hsq]
            continue

        # ---- file có "GCN" + số phát hành GCN hiện hành → ghi nhận là Giấy chứng nhận
        bo_sung = {**gcn_cho_bo_sung(detail, tp_hsq), **chi_dinh}
        co_loai_bat_buoc = bo_sung or any(core.safe_int(n.get("loaiHoSoQuet"), -1) in LOAI_BAT_BUOC
                                          for n in day.cac_node(hoso))
        if not co_loai_bat_buoc:
            ly_do = ("HSQ chưa có thành phần loại Giấy chứng nhận/Đơn đăng ký/Thông báo xác nhận và không có "
                     "mô tả nào có 'GCN' + số phát hành GCN hiện hành của đơn → MPLIS sẽ từ chối; đổi tên "
                     "file GCN thành '<mô tả>-GCN.pdf' (hoặc chọn loại HSQ trên MPLIS) rồi chạy lại")
            log(f"   ⚠ HSQ {hsq_id}: {ly_do}")
            rows += [_dong(madon, thay[tp["vi_tri"]][0], tp, trang_thai="Bỏ qua", ghi_chu=ly_do)
                     for tp in tp_hsq if tp["vi_tri"] in thay]
            rows += [_dong(madon, tp=tp, trang_thai="Giữ nguyên", ghi_chu="Không có file mới")
                     for tp in tp_hsq if tp["vi_tri"] not in thay]
            continue

        hoso_gui = copy.deepcopy(hoso)
        nodes = day.cac_node(hoso_gui)
        for i, gcn in bo_sung.items():
            gan_loai_gcn(nodes[i], gcn)
        log(f"   HSQ {hsq_id}: {len(nodes)} thành phần, thay {len(thay)} file"
            + (f", bổ sung loại GCN {len(bo_sung)}" if bo_sung else ""))
        for i, node in enumerate(nodes):
            viec = []
            if i in thay:
                viec.append(f"THAY ← {os.path.basename(thay[i][0])}{' (đã ký số)' if thay[i][1] else ''}")
            if i in bo_sung:
                viec.append(f"LOẠI → Giấy chứng nhận {bo_sung[i].get('soPhatHanh')}")
            log(f"     {i + 1}. {' | '.join(viec) or 'giữ'}  —  {node.get('moTa')}")

        co_viec = [tp for tp in tp_hsq if tp["vi_tri"] in thay or tp["vi_tri"] in bo_sung]

        def dong_viec(tp: dict, **kw) -> dict[str, Any]:
            i = tp["vi_tri"]
            path, da_ky = thay.get(i, ("", False))
            loai_moi = f"Giấy chứng nhận ({bo_sung[i].get('soPhatHanh')})" if i in bo_sung else ""
            return _dong(madon, path, tp, da_ky_so="Có" if da_ky else "", loai_moi=loai_moi, **kw)

        def viec_cua(i: int) -> str:
            return " + ".join((["thay file"] if i in thay else []) + (["bổ sung loại GCN"] if i in bo_sung else []))

        rows += [_dong(madon, tp=tp, trang_thai="Giữ nguyên", ghi_chu="Không có file mới")
                 for tp in tp_hsq if tp not in co_viec]
        if chay_thu:
            rows += [dong_viec(tp, trang_thai="Chạy thử", ghi_chu=f"Sẽ {viec_cua(tp['vi_tri'])}") for tp in co_viec]
            continue

        node_ids_truoc = [str(n.get("nodeId") or "") for n in nodes]
        try:
            day.gui_ca_hsq(client, hoso_gui, thay)
        except core.HetPhien:
            raise
        except Exception as exc:
            rows += [dong_viec(tp, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for tp in co_viec]
            log(f"   ❌ {core.rut_gon_text(exc, 300)}")
            continue

        # ---- xác nhận
        detail_moi = client.lay_chi_tiet_don(madon) or {}
        hoso_moi = next((h for h in detail_moi.get("ListHoSoQuet") or []
                         if isinstance(h, dict) and core.safe_int(h.get("hoSoQuetId")) == hsq_id), None)
        nodes_moi = day.cac_node(hoso_moi) if hoso_moi else []
        if len(nodes_moi) < len(nodes):
            rows += [dong_viec(tp, trang_thai="Lỗi", ghi_chu=f"MPLIS còn {len(nodes_moi)}/{len(nodes)} thành phần!")
                     for tp in co_viec]
            raise day.DungKhanCap(
                f"HSQ {hsq_id} (đơn {madon}) giảm từ {len(nodes)} còn {len(nodes_moi)} thành phần "
                "sau khi đẩy → DỪNG TOÀN BỘ."
            )
        ids_moi = {str(n.get("nodeId") or "") for n in nodes_moi}
        mat = [nodes[i].get("moTa") for i in range(len(nodes)) if i not in thay and node_ids_truoc[i] not in ids_moi]
        if mat:
            log(f"   ⚠ Thành phần giữ nguyên nhưng nodeId đổi: {mat}")
        for tp in co_viec:
            i = tp["vi_tri"]
            tp_id = core.safe_int(nodes[i].get("thanhPhanHoSoQuetId"))
            node_moi = next((n for n in nodes_moi if core.safe_int(n.get("thanhPhanHoSoQuetId")) == tp_id), None) or {}
            loi = []
            if i in thay and not (node_moi.get("nodeId") and str(node_moi.get("nodeId")) != node_ids_truoc[i]):
                loi.append("file chưa đổi")
            if i in bo_sung and core.safe_int(node_moi.get("loaiHoSoQuet"), -1) != 1:
                loi.append("loại chưa thành Giấy chứng nhận")
            if loi:
                rows.append(dong_viec(tp, trang_thai="Lỗi",
                                      ghi_chu=f"Server báo OK nhưng {', '.join(loi)} (kiểm tra trên MPLIS)"))
                log(f"   ❌ {nodes[i].get('moTa')}: server báo OK nhưng {', '.join(loi)}")
            else:
                rows.append(dong_viec(tp, trang_thai="Đã thay" if i in thay else "Đã bổ sung loại",
                                      ghi_chu=f"Đã {viec_cua(i)} {time.strftime('%d/%m/%Y %H:%M:%S')}"))
                log(f"   ✅ Đã {viec_cua(i)}: {nodes[i].get('moTa')}")
        time.sleep(core.REQUEST_DELAY_SECONDS)
    return rows


# ============================ EXCEL + LUỒNG CHÍNH ============================

def ghi_excel(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_tt = EXCEL_HEADERS.index("trang_thai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in EXCEL_HEADERS])
        fill = MAU_TRANG_THAI.get(str(row.get("trang_thai") or ""))
        if fill:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=fill)
    for i, h in enumerate(EXCEL_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}{max(ws.max_row, 1)}"
    wb.save(path)


def _ghi_an_toan(path: str, rows: list[dict[str, Any]]) -> str:
    try:
        ghi_excel(path, rows)
        return path
    except PermissionError:
        alt = os.path.splitext(path)[0] + "_capnhat.xlsx"
        ghi_excel(alt, rows)
        log(f"⚠ Excel đang mở, đã lưu sang: {alt}")
        return alt


def chay(
    client: core.MplisClient,
    goc: str,
    chay_thu: bool = True,
    nen_dung: Callable[[], bool] = lambda: False,
    bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None,
) -> str:
    """Chạy cho mọi thư mục mã đơn trong `goc`; trả về đường dẫn Excel kết quả (ghi trong `goc`)."""
    theo_don, khong_nhan = quet_thu_muc_goc(goc)
    if khong_nhan:
        log(f"⚠ Bỏ qua {len(khong_nhan)} thư mục không bắt đầu bằng mã đơn: {', '.join(khong_nhan[:10])}"
            + (" ..." if len(khong_nhan) > 10 else ""))
    if not theo_don:
        raise RuntimeError(f"Không có thư mục mã đơn nào trong {goc}")
    excel = os.path.join(goc, f"ket_qua_day_hsq_theo_mota_{'thu_' if chay_thu else ''}"
                              f"{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
    log(f"{'[CHẠY THỬ] ' if chay_thu else ''}{len(theo_don)} đơn | "
        f"{sum(len(v) for v in theo_don.values())} file PDF")

    rows: list[dict[str, Any]] = []
    loi_lien_tiep = 0
    try:
        for n, (madon, pdfs) in enumerate(theo_don.items(), 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(n - 1, len(theo_don), f"Đơn {madon} ({n}/{len(theo_don)})")
            log(f"--- [{n}/{len(theo_don)}] Đơn {madon}: {len(pdfs)} file ---")
            try:
                dong = xu_ly_don(client, madon, pdfs, chay_thu)
                rows += dong
                loi_lien_tiep = loi_lien_tiep + 1 if dong and dong[0].get("ghi_chu") == LOI_CHI_TIET else 0
                if loi_lien_tiep >= DUNG_SAU_SO_DON_LOI:
                    log(f"❌ {loi_lien_tiep} đơn liên tiếp không đọc được từ MPLIS — có thể đã hết phiên đăng nhập. "
                        "ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            except day.DungKhanCap:
                rows += [_dong(madon, p, trang_thai="Lỗi", ghi_chu="MPLIS giảm thành phần sau khi đẩy — ĐÃ DỪNG")
                         for p in pdfs]
                raise
            except Exception as exc:
                rows += [_dong(madon, p, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for p in pdfs]
                log(f"   ❌ {core.rut_gon_text(exc, 300)}")
                if isinstance(exc, core.HetPhien):
                    log(f"❌ {exc} — ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            excel = _ghi_an_toan(excel, rows)
            time.sleep(core.REQUEST_DELAY_SECONDS)
    finally:
        if rows:
            excel = _ghi_an_toan(excel, rows)

    dem: dict[str, int] = {}
    for r in rows:
        if r.get("file") or r.get("loai_moi"):
            dem[r["trang_thai"]] = dem.get(r["trang_thai"], 0) + 1
    bao_tien_do(len(theo_don), len(theo_don), "Hoàn tất")
    log("=" * 50)
    log(("[CHẠY THỬ — chưa gửi gì] " if chay_thu else "✅ ") + "File PDF: "
        + " | ".join(f"{k}: {v}" for k, v in dem.items()))
    log(f"Excel kết quả: {excel}")
    return excel
