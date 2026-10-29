# -*- coding: utf-8 -*-
"""
Đẩy file ĐƠN ĐĂNG KÝ theo thư mục mã đơn cho đơn CHƯA CÓ GCN + cập nhật thông tin đăng ký:
coQuyenQuanLy = True, dieuKienCapGiay = "3" (chưa đủ điều kiện cấp giấy), duDieuKienCapGiay = False.

Cấu trúc thư mục (vd thư mục DON/ do TaiTachHSQ tạo):
    <thư mục gốc>/
        13868776/                            ← tên thư mục bắt đầu bằng mã đơn (tinhHinhDangKyId)
            03_Đơn đăng ký-DON.pdf           ← tên có hậu tố "-DON" (TaiTachHSQ) ...
            Don dang ky.pdf                  ← ... hoặc tên là đơn đăng ký / đơn đề nghị
            PC_NGUYEN VAN A.pdf              ← file khác: không đụng tới

Với mỗi đơn:
  1. Lấy chi tiết đơn. Đơn đã có GCN hiện hành (kể cả GCN chưa có số phát hành) → bỏ qua.
  2. Đẩy file đơn vào HSQ, ghi nhận loại Đơn đăng ký (loaiHoSoQuet = 2):
       - HSQ đã có thành phần loại Đơn đăng ký     → không đẩy (tránh trùng), chỉ cập nhật đăng ký;
       b. thành phần có mô tả ĐÚNG BẰNG tên file (bỏ "-DON", tiền tố "NN_", đuôi ký số)
                                                  → thay file + ghi nhận loại Đơn đăng ký;
       c. không có                                → thêm thành phần mới "Đơn đăng ký" vào HSQ đầu tiên;
       d. đơn chưa có HSQ                         → tạo HSQ mới.
     Không khớp kiểu "tên file chứa mô tả": file đơn tách trang ('Hồ sơ_tr2-4-DON') sẽ đè mất cả file gốc.
     Mỗi HSQ gửi 1 request kèm ĐỦ thành phần cũ; số thành phần giảm → DỪNG TOÀN BỘ.
  3. UpdateThongTinDangKy theo mẫu RESPONes.json (body UI gửi khi lưu đơn): TinhHinhDangKy, ChuSoHuu, TaiSan,
     ListDangKyQuyen (dạng rút gọn, duDieuKienCapGiay = false), ListGiayChungNhan, ... Đã đúng sẵn → không gửi.
     Xác nhận lại trên MPLIS; số đăng ký quyền đổi sau khi gửi → DỪNG TOÀN BỘ.
"""

from __future__ import annotations

import copy
import os
import re
import time
import uuid
from typing import Any, Callable

import bo_sung_gcn_hsq as bs
import cap_nhat_don_hsq as core
import day_gcn_theo_thu_muc as dg
import day_hsq_ky_so as day
import day_hsq_theo_mota as dm
import tai_hsq_don as thd

log = thd.log

LOAI_DON_DANG_KY = 2
MO_TA_DON = "Đơn đăng ký"
DIEU_KIEN_CAP_GIAY = "3"  # chưa đủ điều kiện cấp giấy

EXCEL_HEADERS = ["stt", "madon", "viec", "file", "cach", "hosoquetid", "mota_thanh_phan", "da_ky_so",
                 "trang_thai", "ghi_chu"]
EXCEL_WIDTHS = {"stt": 6, "madon": 12, "viec": 14, "file": 45, "cach": 26, "hosoquetid": 11,
                "mota_thanh_phan": 40, "da_ky_so": 8, "trang_thai": 13, "ghi_chu": 70}
MAU_TRANG_THAI = {"Đã đẩy": "C6EFCE", "Đã cập nhật": "C6EFCE", "Lỗi": "F4CCCC", "Bỏ qua": "FFF2CC",
                  "Chạy thử": "D9EAF7", "Giữ nguyên": "EDEDED"}
CACH = {"co_san": "HSQ đã có thành phần Đơn đăng ký", "b": "Thay file + ghi nhận Đơn đăng ký",
        "c": "Thêm thành phần Đơn đăng ký mới", "d": "Tạo HSQ mới"}
VIEC_DON, VIEC_TTDK = "Đẩy đơn", "Cập nhật ĐK"

TU_DON = ("dondangky", "dondangki", "dondenghi", "dondk", "donxincap")


# ============================ TÌM FILE ĐƠN ============================

def ten_khong_hau_to_don(path: str) -> tuple[int, str, bool]:
    """'03_Đơn đăng ký-DON_signed.pdf' → (3, 'Đơn đăng ký', True)."""
    stt, ten = dm.tach_ten_file(path)
    m = re.match(r"^(.+?)\s*-\s*DON\s*$", ten, re.I)
    return (stt, m.group(1).strip(), True) if m else (stt, ten, False)


def la_file_don(path: str) -> bool:
    _, ten, hau_to = ten_khong_hau_to_don(path)
    return hau_to or any(t in dm.chuan_hoa(ten) for t in TU_DON)


# ============================ KẾ HOẠCH 1 ĐƠN ============================

def info_don_moi(da_ky: bool) -> dict[str, Any]:
    """Thành phần Đơn đăng ký tạo mới — cùng mẫu bo_sung_gcn_hsq.info_thanh_phan_moi, không gắn GCN."""
    return {
        "loaiHoSoQuet": LOAI_DON_DANG_KY,
        "laGiayToVeNguonGoc": False,
        "giayChungNhanId": "",
        "moTa": MO_TA_DON,
        "tenGiayTo": "",
        "trichYeu": "",
        "laGiayChungNhan": False,
        "daKySo": bool(da_ky),
        "__id": str(uuid.uuid4()),
        "files": None,
    }


def gan_loai_don(info: dict[str, Any]) -> None:
    info["loaiHoSoQuet"] = LOAI_DON_DANG_KY
    info["laGiayChungNhan"] = False


def lap_ke_hoach(detail: dict[str, Any], pdfs: list[str]) -> list[dict[str, Any]]:
    """Mỗi file đơn → {"path", "da_ky", "cach", "hoso", "vi_tri"} (cach: co_san / b / c / d)."""
    hsq_list = [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]
    co_don = next(((h, i) for h in hsq_list for i, n in enumerate(day.cac_node(h))
                   if core.safe_int(n.get("loaiHoSoQuet"), -1) == LOAI_DON_DANG_KY), None)
    tp = dm.thanh_phan_cua_don(detail)
    da_dung: set[int] = set()
    ke_hoach = []
    for path in pdfs:
        muc = {"path": path, "da_ky": day.pdf_da_ky_so(path), "hoso": None, "vi_tri": None}
        if co_don:
            ke_hoach.append({**muc, "cach": "co_san", "hoso": co_don[0], "vi_tri": co_don[1]})
            continue
        if not hsq_list:
            ke_hoach.append({**muc, "cach": "d"})
            continue
        stt, ten, _ = ten_khong_hau_to_don(path)
        khoa = dm.chuan_hoa(ten)
        ung_vien = [t for t in tp if t["khoa"] == khoa and id(t["node"]) not in da_dung]
        if len(ung_vien) > 1:
            ung_vien = [t for t in ung_vien if stt and t["stt"] == stt]
        if len(ung_vien) == 1:
            t = ung_vien[0]
            da_dung.add(id(t["node"]))
            ke_hoach.append({**muc, "cach": "b", "hoso": t["hoso"], "vi_tri": t["vi_tri"]})
        else:
            ke_hoach.append({**muc, "cach": "c", "hoso": hsq_list[0]})
    return ke_hoach


def items_hsq_co_san(hoso: dict[str, Any], ke_hoach_hsq: list[dict]) -> list[tuple[dict, str | None]]:
    """ĐỦ thành phần hiện có (thay hoặc giữ nguyên) + thành phần Đơn đăng ký thêm mới, theo thứ tự."""
    hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
    thay = {k["vi_tri"]: k for k in ke_hoach_hsq if k["cach"] == "b"}
    items: list[tuple[dict, str | None]] = []
    for i, node in enumerate(day.cac_node(hoso)):
        if i in thay:
            info = day.info_thay_file(node, hsq_id, thay[i]["da_ky"])
            gan_loai_don(info)
            items.append((info, thay[i]["path"]))
        else:
            items.append((day.info_giu_nguyen(node), None))
    items += [(info_don_moi(k["da_ky"]), k["path"]) for k in ke_hoach_hsq if k["cach"] == "c"]
    return items


# ============================ PAYLOAD CẬP NHẬT ĐĂNG KÝ (mẫu RESPONes.json) ============================

KHOA_TINH_HINH = (
    "maDon", "soThuTu", "dieuKienCapGiay", "ngayTiepNhan", "thoiDiemDangKyLanDau", "thoiDiemDangKy",
    "daiDienKhaiTrinhId", "coQuyenSuDung", "coQuyenQuanLy", "coQuyenSoHuu", "dongSuDung", "capGiayNguoiDaiDien",
    "dangKyLanDau", "laQuyenSoHuuKhac", "ghiChu", "CreatedUserId", "CreatedDate", "ModifiedDate",
    "tinhHinhDangKyId", "tinhId", "huyenId", "xaId", "hoSoTiepNhanId", "versionChu", "isLastest", "rootId",
    "parentId", "Id", "Title", "Name", "Path",
)
KHOA_NULL_DANG_KY_QUYEN = (
    "CaNhan", "VoChong", "HoGiaDinh", "ToChuc", "CongDong", "NhomNguoi", "ThuaDat", "MucDichSuDung", "NhaRiengLe",
    "HangMucNhaRiengLe", "CanHo", "HangMucSoHuuChung", "NhaChungCu", "CongTrinhXayDung", "CongTrinhNgam",
    "HangMucCongTrinh", "RungTrong", "CayLauNam",
)


def _chuoi_id(v: Any) -> str | None:
    return None if v in (None, "") else str(v)


def dang_ky_quyen_rut_gon(d: dict[str, Any]) -> dict[str, Any]:
    """1 phần tử ListDangKyQuyen của chi tiết đơn → dạng UI gửi khi lưu (RESPONes.json)."""
    loai = core.safe_int(d.get("loaiDoiTuong"))
    chu = core.safe_int(d.get("chuSuDungId"))
    ca_nhan = d.get("caNhanId")
    if ca_nhan in (None, "") and loai == 0:  # cá nhân: UI ghi caNhanId = chuSuDungId
        ca_nhan = chu
    return {
        "loaiDoiTuong": loai,
        "chuSuDungId": chu,
        "caNhanId": ca_nhan,
        "typeItem": core.safe_int(d.get("typeItem")),
        "itemId": core.safe_int(d.get("itemId")),
        "subItemId": core.safe_int(d.get("subItemId")),
        "dangKyQuyenSuDungId": 0,
        "tinhHinhDangKyId": core.safe_int(d.get("tinhHinhDangKyId")),
        "giayChungNhanId": core.safe_int(d.get("giayChungNhanId")),
        "loaiQuyen": core.safe_int(d.get("loaiQuyen")),
        "inGiayNguoiDaiDien": bool(d.get("inGiayNguoiDaiDien")),
        "inVoChong": bool(d.get("inVoChong")),
        "tyLe": d.get("tyLe") or 0,
        "dienTich": d.get("dienTich") or 0,
        "ngayBatDauSoHuu": d.get("ngayBatDauSoHuu"),
        "thoiHanSoHuu": d.get("thoiHanSoHuu") or "",
        "duDieuKienCapGiay": False,
        "trangThai": True,
        **{k: None for k in KHOA_NULL_DANG_KY_QUYEN},
    }


def build_payload_ttdk(detail: dict[str, Any]) -> dict[str, Any]:
    """Body UpdateThongTinDangKy theo mẫu RESPONes.json, đặt quyền quản lý + chưa đủ điều kiện cấp giấy."""
    d = core.convert_dates_recursive(copy.deepcopy(detail))
    th = d.get("TinhHinhDangKy") or {}
    tinh_hinh = {k: th.get(k) for k in KHOA_TINH_HINH}
    tinh_hinh.update({
        "dieuKienCapGiay": DIEU_KIEN_CAP_GIAY,
        "coQuyenQuanLy": True,
        "daiDienKhaiTrinhId": _chuoi_id(th.get("daiDienKhaiTrinhId")),
        "hoSoTiepNhanId": _chuoi_id(th.get("hoSoTiepNhanId")),
        "ghiChu": th.get("ghiChu") or "",
    })
    return {
        "TinhHinhDangKy": tinh_hinh,
        "ChuSoHuu": d.get("ChuSoHuu") or {},
        "TaiSan": d.get("TaiSan") or {},
        "ListDangKyQuyen": [dang_ky_quyen_rut_gon(q) for q in d.get("ListDangKyQuyen") or [] if isinstance(q, dict)],
        "ListGiayChungNhan": d.get("ListGiayChungNhan") or [],
        "ListDangKyTaiSan": d.get("ListDangKyTaiSan") or [],
        "ListHanCheQuyen": d.get("ListHanCheQuyen") or [],
        "ListNghiaVuTaiChinh": d.get("ListNghiaVuTaiChinh") or [],
        "ListHanCheThuaLienKe": d.get("ListHanCheThuaLienKe") or [],
        "GhiChuDongSuDungTrang1": d.get("GhiChuDongSuDungTrang1"),
        "GhiChuDongSuDungTrang2": d.get("GhiChuDongSuDungTrang2"),
    }


def ttdk_da_dung(detail: dict[str, Any]) -> bool:
    th = detail.get("TinhHinhDangKy") or {}
    return (th.get("coQuyenQuanLy") is True and core.safe_int(th.get("dieuKienCapGiay")) == int(DIEU_KIEN_CAP_GIAY)
            and not any(isinstance(q, dict) and q.get("duDieuKienCapGiay") is True
                        for q in detail.get("ListDangKyQuyen") or []))


# ============================ XỬ LÝ 1 ĐƠN ============================

def _dong(madon: int, k: dict | None = None, viec: str = VIEC_DON, **kw) -> dict[str, Any]:
    k = k or {}
    node = {}
    if k.get("hoso") is not None and k.get("vi_tri") is not None:
        node = day.cac_node(k["hoso"])[k["vi_tri"]]
    return {
        "madon": madon,
        "viec": viec,
        "file": os.path.basename(k["path"]) if k.get("path") else "",
        "cach": CACH.get(k.get("cach"), ""),
        "hosoquetid": (k.get("hoso") or {}).get("hoSoQuetId", "") if k.get("hoso") else "",
        "mota_thanh_phan": node.get("moTa") or (MO_TA_DON if k.get("cach") in ("c", "d") else ""),
        "da_ky_so": "Có" if k.get("da_ky") else "",
        **kw,
    }


def day_file_don(client: core.MplisClient, madon: int, detail: dict[str, Any],
                 ke_hoach: list[dict]) -> tuple[list[dict[str, Any]], bool]:
    """Gửi các file đơn theo kế hoạch → (dòng Excel, đơn đã có thành phần Đơn đăng ký trên MPLIS)."""
    rows: list[dict[str, Any]] = []
    nhom: list[tuple[dict | None, list[dict]]] = []
    tao_moi = [k for k in ke_hoach if k["cach"] == "d"]
    if tao_moi:
        nhom.append((None, tao_moi))
    for h in [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]:
        ds = [k for k in ke_hoach if k["hoso"] is h and k["cach"] in ("b", "c")]
        if ds:
            nhom.append((h, ds))

    ok = True
    for hoso, ds in nhom:
        so_tp_truoc = len(day.cac_node(hoso)) if hoso else 0
        node_id_truoc = {k["vi_tri"]: str(day.cac_node(hoso)[k["vi_tri"]].get("nodeId") or "")
                         for k in ds if k["cach"] == "b"}
        try:
            if hoso is None:
                tths_id = core.safe_int(core.lay_thong_tin_ho_so_id(detail))
                if not tths_id:
                    raise RuntimeError("Không tìm được thongTinHoSoId của đơn để tạo HSQ mới")
                dg.gui_tao_moi(client, tths_id, [(info_don_moi(k["da_ky"]), k["path"]) for k in ds])
            else:
                bs.gui_hsq(client, hoso, items_hsq_co_san(hoso, ds))
        except core.HetPhien:
            raise
        except Exception as exc:
            rows += [_dong(madon, k, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for k in ds]
            log(f"   ❌ {core.rut_gon_text(exc, 300)}")
            ok = False
            continue

        # ---- xác nhận
        detail_moi = client.lay_chi_tiet_don(madon) or {}
        if hoso is not None:
            hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
            hoso_moi = next((h for h in detail_moi.get("ListHoSoQuet") or []
                             if isinstance(h, dict) and core.safe_int(h.get("hoSoQuetId")) == hsq_id), None)
            nodes_moi = day.cac_node(hoso_moi) if hoso_moi else []
            if len(nodes_moi) < so_tp_truoc:
                rows += [_dong(madon, k, trang_thai="Lỗi", ghi_chu=f"MPLIS còn {len(nodes_moi)}/{so_tp_truoc} "
                                                                     "thành phần!") for k in ds]
                raise day.DungKhanCap(f"HSQ {hsq_id} (đơn {madon}) giảm từ {so_tp_truoc} còn {len(nodes_moi)} "
                                      "thành phần sau khi đẩy → DỪNG TOÀN BỘ.")
        else:
            nodes_moi = [n for h in detail_moi.get("ListHoSoQuet") or [] if isinstance(h, dict)
                         for n in day.cac_node(h)]
        don_nodes = [n for n in nodes_moi if core.safe_int(n.get("loaiHoSoQuet"), -1) == LOAI_DON_DANG_KY]
        for k in ds:
            loi = "" if don_nodes else "chưa có thành phần loại Đơn đăng ký"
            if not loi and k["cach"] == "b" and all(
                    str(n.get("nodeId") or "") == node_id_truoc[k["vi_tri"]] for n in don_nodes):
                loi = "file chưa đổi"
            if loi:
                ok = False
                rows.append(_dong(madon, k, trang_thai="Lỗi", ghi_chu=f"Server báo OK nhưng {loi} (kiểm tra MPLIS)"))
                log(f"   ❌ {os.path.basename(k['path'])}: server báo OK nhưng {loi}")
            else:
                rows.append(_dong(madon, k, trang_thai="Đã đẩy", ghi_chu=time.strftime("%d/%m/%Y %H:%M:%S")))
                log(f"   ✅ {os.path.basename(k['path'])}: {CACH[k['cach']].lower()}")
        time.sleep(core.REQUEST_DELAY_SECONDS)
    return rows, ok


def cap_nhat_ttdk(client: core.MplisClient, madon: int, detail: dict[str, Any]) -> dict[str, Any]:
    """Gửi UpdateThongTinDangKy rồi đọc lại xác nhận. Số đăng ký quyền đổi → DungKhanCap."""
    so_dkq_truoc = len(detail.get("ListDangKyQuyen") or [])
    try:
        client.api_update_thong_tin_dang_ky(build_payload_ttdk(detail), debug_key=str(madon))
    except core.HetPhien:
        raise
    except Exception as exc:
        log(f"   ❌ Cập nhật đăng ký: {core.rut_gon_text(exc, 300)}")
        return _dong(madon, viec=VIEC_TTDK, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300))
    detail_moi = client.lay_chi_tiet_don(madon) or {}
    so_dkq_sau = len(detail_moi.get("ListDangKyQuyen") or [])
    if so_dkq_sau != so_dkq_truoc:
        raise day.DungKhanCap(f"Đơn {madon}: số đăng ký quyền đổi từ {so_dkq_truoc} thành {so_dkq_sau} sau khi cập "
                              "nhật đăng ký → DỪNG TOÀN BỘ, kiểm tra đơn trên MPLIS.")
    if not ttdk_da_dung(detail_moi):
        th = detail_moi.get("TinhHinhDangKy") or {}
        ghi_chu = (f"Server báo OK nhưng MPLIS vẫn coQuyenQuanLy={th.get('coQuyenQuanLy')}, "
                   f"dieuKienCapGiay={th.get('dieuKienCapGiay')}")
        log(f"   ❌ {ghi_chu}")
        return _dong(madon, viec=VIEC_TTDK, trang_thai="Lỗi", ghi_chu=ghi_chu)
    log("   ✅ Đã cập nhật: có quyền quản lý, chưa đủ điều kiện cấp giấy")
    return _dong(madon, viec=VIEC_TTDK, trang_thai="Đã cập nhật",
                 ghi_chu="coQuyenQuanLy = true, dieuKienCapGiay = 3, duDieuKienCapGiay = false")


def xu_ly_don(client: core.MplisClient, madon: int, pdfs: list[str], chay_thu: bool) -> list[dict[str, Any]]:
    detail = client.lay_chi_tiet_don(madon)
    if detail is None:
        return [_dong(madon, trang_thai="Lỗi", ghi_chu=dm.LOI_CHI_TIET)]
    ds_gcn = bs.gcn_hien_hanh(detail)
    if ds_gcn:
        so = ", ".join(str(g.get("soPhatHanh") or f"id {g.get('giayChungNhanId')} (chưa có số)") for g in ds_gcn)
        return [_dong(madon, trang_thai="Bỏ qua", ghi_chu=f"Đơn đã có GCN hiện hành: {so}")]
    file_don = [p for p in pdfs if la_file_don(p)]
    if not file_don:
        return [_dong(madon, trang_thai="Bỏ qua",
                      ghi_chu="Thư mục không có file đơn (tên có hậu tố '-DON' hoặc là 'Đơn đăng ký' / 'Đơn đề nghị')")]

    ke_hoach = lap_ke_hoach(detail, file_don)
    for k in ke_hoach:
        log(f"   {os.path.basename(k['path'])}{' (đã ký số)' if k['da_ky'] else ''} | {CACH[k['cach']]}"
            + (f" '{day.cac_node(k['hoso'])[k['vi_tri']].get('moTa')}'" if k["cach"] in ("co_san", "b") else ""))
    can_ttdk = not ttdk_da_dung(detail)
    if chay_thu:
        rows = [_dong(madon, k, trang_thai="Chạy thử" if k["cach"] != "co_san" else "Giữ nguyên",
                      ghi_chu=f"Sẽ {CACH[k['cach']].lower()}" if k["cach"] != "co_san" else "Không đẩy file")
                for k in ke_hoach]
        return rows + [_dong(madon, viec=VIEC_TTDK, trang_thai="Chạy thử" if can_ttdk else "Giữ nguyên",
                             ghi_chu="Sẽ đặt coQuyenQuanLy = true, dieuKienCapGiay = 3, duDieuKienCapGiay = false"
                             if can_ttdk else "Đã đúng sẵn")]

    rows = [_dong(madon, k, trang_thai="Giữ nguyên", ghi_chu="Không đẩy file")
            for k in ke_hoach if k["cach"] == "co_san"]
    dong_day, ok = day_file_don(client, madon, detail, ke_hoach)
    rows += dong_day
    if not ok:
        return rows + [_dong(madon, viec=VIEC_TTDK, trang_thai="Bỏ qua",
                             ghi_chu="Đẩy file đơn lỗi → chưa cập nhật đăng ký")]
    if not can_ttdk:
        return rows + [_dong(madon, viec=VIEC_TTDK, trang_thai="Giữ nguyên", ghi_chu="Đã đúng sẵn")]
    detail = client.lay_chi_tiet_don(madon) or detail  # bản mới nhất sau khi đẩy HSQ
    return rows + [cap_nhat_ttdk(client, madon, detail)]


# ============================ EXCEL + LUỒNG CHÍNH ============================

def ghi_excel(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_tt = EXCEL_HEADERS.index("trang_thai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in EXCEL_HEADERS])
        fill = MAU_TRANG_THAI.get(str(row.get("trang_thai") or ""))
        if fill:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=fill)
    for i, h in enumerate(EXCEL_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}{max(ws.max_row, 1)}"
    wb.save(path)


def _ghi_an_toan(path: str, rows: list[dict[str, Any]]) -> str:
    try:
        ghi_excel(path, rows)
        return path
    except PermissionError:
        alt = os.path.splitext(path)[0] + "_capnhat.xlsx"
        ghi_excel(alt, rows)
        log(f"⚠ Excel đang mở, đã lưu sang: {alt}")
        return alt


def chay(
    client: core.MplisClient,
    goc: str,
    chay_thu: bool = True,
    nen_dung: Callable[[], bool] = lambda: False,
    bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None,
) -> str:
    """Chạy cho mọi thư mục mã đơn trong `goc`; trả về đường dẫn Excel kết quả (ghi trong `goc`)."""
    theo_don, khong_nhan = dm.quet_thu_muc_goc(goc)
    if khong_nhan:
        log(f"⚠ Bỏ qua {len(khong_nhan)} thư mục có PDF nhưng không nằm trong thư mục mã đơn nào: {', '.join(khong_nhan[:10])}"
            + (" ..." if len(khong_nhan) > 10 else ""))
    theo_don = {m: p for m, p in theo_don.items() if p}
    if not theo_don:
        raise RuntimeError(f"Không có thư mục mã đơn nào chứa file PDF trong {goc}")
    excel = os.path.join(goc, f"ket_qua_day_don_{'thu_' if chay_thu else ''}{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
    log(f"{'[CHẠY THỬ] ' if chay_thu else ''}{len(theo_don)} đơn | {sum(len(v) for v in theo_don.values())} file PDF")

    rows: list[dict[str, Any]] = []
    loi_lien_tiep = 0
    try:
        for n, (madon, pdfs) in enumerate(theo_don.items(), 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(n - 1, len(theo_don), f"Đơn {madon} ({n}/{len(theo_don)})")
            log(f"--- [{n}/{len(theo_don)}] Đơn {madon}: {len(pdfs)} file ---")
            try:
                dong = xu_ly_don(client, madon, pdfs, chay_thu)
                rows += dong
                loi_lien_tiep = loi_lien_tiep + 1 if dong and dong[0].get("ghi_chu") == dm.LOI_CHI_TIET else 0
                if loi_lien_tiep >= dm.DUNG_SAU_SO_DON_LOI:
                    log(f"❌ {loi_lien_tiep} đơn liên tiếp không đọc được từ MPLIS — có thể đã hết phiên đăng nhập. "
                        "ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            except day.DungKhanCap as exc:
                rows.append(_dong(madon, trang_thai="Lỗi", ghi_chu=f"{exc}"))
                raise
            except Exception as exc:
                rows.append(_dong(madon, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)))
                log(f"   ❌ {core.rut_gon_text(exc, 300)}")
                if isinstance(exc, core.HetPhien):
                    log(f"❌ {exc} — ĐÃ DỪNG: lấy lại session (hoặc dán token + cookie mới) rồi chạy lại.")
                    break
            excel = _ghi_an_toan(excel, rows)
            core.nghi_giua_don()
    finally:
        if rows:
            excel = _ghi_an_toan(excel, rows)

    dem: dict[str, int] = {}
    for r in rows:
        dem[r["trang_thai"]] = dem.get(r["trang_thai"], 0) + 1
    bao_tien_do(len(theo_don), len(theo_don), "Hoàn tất")
    log("=" * 50)
    log(("[CHẠY THỬ — chưa gửi gì] " if chay_thu else "✅ ") + " | ".join(f"{k}: {v}" for k, v in dem.items()))
    log(f"Excel kết quả: {excel}")
    return excel
