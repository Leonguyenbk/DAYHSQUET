# -*- coding: utf-8 -*-
"""
Tool thay thế TOÀN BỘ file hồ sơ quét trong đơn theo tên mô tả tương ứng.

Nghiệp vụ:
- 1 đơn đăng ký (tinhHinhDangKyId) có nhiều file HSQ (đơn, GCN, trích lục, nguồn gốc...).
- Người dùng có thư mục chứa các file PDF mới mà tên file tương ứng với mô tả (moTa) của các file trong đơn.
- Tool sẽ:
  1. Đọc danh sách đơn từ Excel (cột id_don/madon hoặc soto/sothua) hoặc nhập trực tiếp 1 đơn trên giao diện.
  2. Lấy chi tiết đơn và danh sách toàn bộ file HSQ hiện có (ListFileHoSoQuet).
  3. Tự động so khớp từng file trên đĩa với trường `moTa` của từng node HSQ.
  4. Đóng gói payload multipart/form-data gửi lên UpdateHoSoQuetExistFile để thay thế đồng thời toàn bộ các file đã khớp.
  5. Giữ nguyên cấu trúc các node/file không thay thế hoặc không tìm thấy file mới để bảo đảm an toàn dữ liệu trên MPLIS.
  6. Xuất báo cáo Excel chi tiết kết quả thay thế từng file.
"""

from __future__ import annotations

import copy
import json
import os
import re
import threading
import time
import unicodedata
import uuid
from contextlib import ExitStack
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager


# ============================ CẤU HÌNH API ============================

BASE_URL = "https://dla.mplis.gov.vn"

REFERER_LOGIN = f"{BASE_URL}/dc/DonDangKy/KeKhaiDangKyV2"
REFERER_API = REFERER_LOGIN

URL_TIM_DON = f"{BASE_URL}/dc/DangKyAjax/AdvancedSearchTinhHinhDangKy"
URL_CHI_TIET = f"{BASE_URL}/dc/DangKyAjax/GetThongTinDangKyByTinhHinhDangKyIds"
URL_UPDATE_HOSOQUET = f"{BASE_URL}/dc/HoSoQuetAjax/UpdateHoSoQuetExistFile"

PAGE_SIZE = 10
TIMEOUT = 180
SAVE_EVERY_ROWS = 5
REQUEST_DELAY_SECONDS = 0.2

LOAI_HO_SO_QUET_OPTIONS = {
    0: "Giấy tờ",
    1: "Giấy chứng nhận",
    2: "Đơn đăng ký",
    3: "Thông báo xác nhận đăng ký",
}

OUTPUT_HEADERS = [
    "STT",
    "Mã đơn (ID)",
    "Số tờ",
    "Số thửa",
    "Chủ sử dụng",
    "Hồ sơ quét ID",
    "STT File trong HSQ",
    "Mô tả file cũ (MPLIS)",
    "Loại HSQ",
    "File PDF mới trên đĩa",
    "Kiểu khớp",
    "Trạng thái",
    "Ghi chú",
]


# ============================ CÁC HÀM TIỆN ÍCH ============================

def safe_int(value: Any, default: int = 0) -> int:
    try:
        if value is None or value == "":
            return default
        return int(value)
    except Exception:
        return default


def rut_gon_text(value: Any, limit: int = 400) -> str:
    text = str(value or "").strip().replace("\r", " ").replace("\n", " ")
    if len(text) > limit:
        return text[:limit] + "..."
    return text


def now_iso_z() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def lay_token_tu_trang(driver: webdriver.Chrome) -> str:
    js = """
    return (
        document.querySelector('input[name="__RequestVerificationToken"]')?.value ||
        document.querySelector('input[name="__requestverificationtoken"]')?.value ||
        document.querySelector('meta[name="__RequestVerificationToken"]')?.content ||
        document.querySelector('meta[name="__requestverificationtoken"]')?.content ||
        document.querySelector('meta[name="RequestVerificationToken"]')?.content ||
        ''
    );
    """
    value = driver.execute_script(js)
    return str(value or "").strip()


def remove_accents(input_str: str) -> str:
    """Loại bỏ dấu tiếng Việt để so khớp mềm."""
    if not input_str:
        return ""
    nfkd_form = unicodedata.normalize("NFKD", input_str)
    return "".join([c for c in nfkd_form if not unicodedata.combining(c)])


def normalize_string_for_match(s: Any) -> str:
    """
    Chuẩn hoá chuỗi: bỏ dấu tiếng Việt, loại bỏ toàn bộ khoảng trắng và ký tự đặc biệt,
    chuyển thành chữ thường để so sánh không phân biệt định dạng.
    """
    if not s:
        return ""
    raw = remove_accents(str(s)).lower()
    # Loại bỏ đuôi .pdf nếu có
    if raw.endswith(".pdf"):
        raw = raw[:-4]
    # Chỉ giữ chữ và số
    return re.sub(r"[^a-z0-9]", "", raw)


def extract_numbers(s: Any) -> str:
    """Trích xuất chuỗi số liên tiếp từ chuỗi (hữu ích cho số GCN hoặc số hiệu)."""
    if not s:
        return ""
    digits = re.findall(r"\d+", str(s))
    return "".join(digits)


# ============================ BỘ SO KHỚP FILE THÔNG MINH ============================

class FileIndex:
    """
    Quét và tạo chỉ mục toàn bộ file PDF trong thư mục được chỉ định.
    Hỗ trợ tìm kiếm siêu tốc theo tên chính xác, tên chuẩn hoá và mã nhận diện.
    """
    def __init__(self, folder_path: str):
        self.folder_path = folder_path
        self.files_list: list[dict[str, Any]] = []
        self.exact_map: dict[str, str] = {}
        self.norm_map: dict[str, list[str]] = {}
        self._scan()

    def _scan(self) -> None:
        if not self.folder_path or not os.path.isdir(self.folder_path):
            return

        for root, _, files in os.walk(self.folder_path):
            for f in files:
                if f.lower().endswith(".pdf"):
                    full_p = os.path.join(root, f)
                    stem = Path(f).stem
                    norm_stem = normalize_string_for_match(stem)
                    norm_f = normalize_string_for_match(f)

                    item = {
                        "filename": f,
                        "stem": stem,
                        "full_path": full_p,
                        "norm_stem": norm_stem,
                        "norm_filename": norm_f,
                        "digits": extract_numbers(stem),
                    }
                    self.files_list.append(item)

                    # Exact map
                    self.exact_map[f.lower()] = full_p
                    self.exact_map[stem.lower()] = full_p

                    # Normalized map
                    if norm_stem:
                        self.norm_map.setdefault(norm_stem, []).append(full_p)

    def find_match(self, mo_ta: str, ten_giay_to: str = "", so_gcn: str = "") -> tuple[str | None, str]:
        """
        Tìm file PDF khớp nhất với mô tả của node HSQ.
        Trả về (full_path_or_None, reason).
        """
        if not mo_ta and not ten_giay_to and not so_gcn:
            return None, "Node không có mô tả (moTa rỗng)"

        mo_ta_clean = (mo_ta or "").strip()
        stem_target = Path(mo_ta_clean).stem if mo_ta_clean.lower().endswith(".pdf") else mo_ta_clean

        # 1. Khớp chính xác theo tên file hoặc stem
        if mo_ta_clean.lower() in self.exact_map:
            return self.exact_map[mo_ta_clean.lower()], f"Khớp chính xác tên '{mo_ta_clean}'"
        if f"{mo_ta_clean.lower()}.pdf" in self.exact_map:
            return self.exact_map[f"{mo_ta_clean.lower()}.pdf"], f"Khớp chính xác tên '{mo_ta_clean}.pdf'"
        if stem_target.lower() in self.exact_map:
            return self.exact_map[stem_target.lower()], f"Khớp chính xác stem '{stem_target}'"

        # 2. Khớp chuẩn hoá (bỏ dấu tiếng Việt, khoảng trắng, gạch nối...)
        norm_mota = normalize_string_for_match(mo_ta_clean)
        if norm_mota and norm_mota in self.norm_map:
            paths = self.norm_map[norm_mota]
            return paths[0], f"Khớp chuẩn hoá mô tả '{mo_ta_clean}' -> '{Path(paths[0]).name}'"

        # 3. Khớp nếu tên file chứa toàn bộ mô tả hoặc ngược lại
        for item in self.files_list:
            if norm_mota and len(norm_mota) >= 5:
                if norm_mota == item["norm_stem"]:
                    return item["full_path"], f"Khớp hoàn toàn chuẩn hoá '{item['filename']}'"
                if norm_mota in item["norm_stem"] or item["norm_stem"] in norm_mota:
                    len_diff = abs(len(norm_mota) - len(item["norm_stem"]))
                    if len_diff <= 10:
                        return item["full_path"], f"Khớp chứa chuỗi '{item['filename']}'"

        # 4. Khớp theo số GCN / số hiệu nếu mô tả có đề cập
        targets_code = [c for c in [so_gcn, stem_target, ten_giay_to] if c]
        for t in targets_code:
            norm_code = normalize_string_for_match(t)
            if norm_code and len(norm_code) >= 4:
                for item in self.files_list:
                    if norm_code in item["norm_stem"]:
                        return item["full_path"], f"Khớp theo mã/số '{t}' trong file '{item['filename']}'"

        # 5. Khớp theo chuỗi số (digits) nếu có số GCN rõ ràng (>= 5 chữ số)
        digits_target = extract_numbers(mo_ta_clean) or extract_numbers(so_gcn)
        if digits_target and len(digits_target) >= 5:
            matched_digit_items = [
                item for item in self.files_list if digits_target in item["digits"]
            ]
            if len(matched_digit_items) == 1:
                f_path = matched_digit_items[0]["full_path"]
                return f_path, f"Khớp duy nhất theo chuỗi số '{digits_target}' trong file '{Path(f_path).name}'"

        # 6. Kiểm tra qua ten_giay_to nếu có
        if ten_giay_to:
            norm_tgt = normalize_string_for_match(ten_giay_to)
            if norm_tgt and norm_tgt in self.norm_map:
                paths = self.norm_map[norm_tgt]
                return paths[0], f"Khớp theo tên giấy tờ '{ten_giay_to}'"

        return None, f"Không tìm thấy file trên đĩa khớp với mô tả '{mo_ta_clean}'"


# ============================ EXCEL I/O & BÓC TÁCH MÃ ĐƠN ============================

def tach_thong_tin_tu_chuoi_madon(val: Any) -> tuple[str, str]:
    """
    Tự động bóc tách mã đơn (tinhHinhDangKyId) và thông tin GCN (nếu có)
    từ chuỗi mô tả/thông báo kiểm tra hoặc lỗi.
    Ví dụ:
    "Tình hình đăng ký 13456477 có giấy chứng nhận 2502899_11 có hồ sơ quét không có dữ liệu tập tin; Chưa đồng bộ thông tin ba khối (không có hồ sơ quét)"
    -> id_don = "13456477", gcn_hint = "2502899_11"
    """
    if val is None:
        return "", ""
    s = str(val).strip()
    if not s:
        return "", ""

    # 1. Trường hợp là số nguyên hoặc float thuần túy (e.g. 13456477 hoặc 13456477.0)
    if re.fullmatch(r"\d+(\.0+)?", s):
        return str(int(float(s))), ""

    id_don = ""
    gcn_hint = ""

    s_clean = remove_accents(s).lower()

    # Pattern 1: Tình hình đăng ký <mã> / Đơn đăng ký <mã> / Mã đơn <mã> / THDK <mã>
    m_id = re.search(
        r"(?:tinh\s*hinh\s*dang\s*ky|don\s*dang\s*ky|ma\s*don|id\s*don|thdk|tinhhinhdangky)\D*?(\d{5,10})",
        s_clean,
    )
    if m_id:
        id_don = m_id.group(1)
    else:
        # Pattern 2: Tìm chữ số sau 'id' hoặc 'don'
        m_id2 = re.search(r"\b(?:id|don)\D*?(\d{5,10})\b", s_clean)
        if m_id2:
            id_don = m_id2.group(1)
        else:
            # Pattern 3: Tìm dãy số 6-10 chữ số đầu tiên trong chuỗi
            m_digits = re.search(r"\b\d{6,10}\b", s)
            if m_digits:
                id_don = m_digits.group(0)

    # 2. Tách thông tin GCN nếu có trong chuỗi (ví dụ: 'giấy chứng nhận 2502899_11')
    m_gcn = re.search(r"gi[aấ]y\s*ch[uứ]ng\s*nh[aậ]n\s*([A-Za-z0-9_\-]+)", s, re.IGNORECASE)
    if m_gcn:
        gcn_hint = m_gcn.group(1).strip()

    return id_don, gcn_hint


def chuan_hoa_gia_tri_excel(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


ALIASES_ID_DON = {
    "madon", "ma_don", "tinhhinhdangkyid", "id_don", "iddon", "id",
    "tinh_hinh_dang_ky_id", "mã đơn", "id đơn", "mã_đơn", "id_đơn"
}
ALIASES_SO_TO = {"soto", "so_to", "số tờ", "sohieutobando", "so_hieu_to_ban_do", "số_tờ"}
ALIASES_SO_THUA = {"sothua", "so_thua", "số thửa", "sothututhua", "so_thu_tu_thua", "số_thửa"}
ALIASES_SO_GCN = {"sogcn", "so_gcn", "số gcn", "sophathanh", "so_phat_hanh", "gcn", "số_gcn"}


def doc_danh_sach_don(file_path: str) -> list[dict[str, Any]]:
    """Đọc file Excel lấy danh sách các đơn cần xử lý (tự động bóc tách ID đơn nếu chứa chuỗi thông báo)."""
    extension = Path(file_path).suffix.lower()
    if extension not in {".xlsx", ".xlsm"}:
        raise ValueError("Chỉ hỗ trợ file Excel .xlsx hoặc .xlsm.")

    workbook = load_workbook(file_path, data_only=True)
    try:
        worksheet = workbook.active
        headers: dict[str, int] = {}
        for col_idx in range(1, worksheet.max_column + 1):
            val = worksheet.cell(row=1, column=col_idx).value
            if val:
                headers[str(val).strip().lower()] = col_idx

        def tim_cot(aliases: set[str]) -> int | None:
            for alias in aliases:
                if alias in headers:
                    return headers[alias]
            return None

        col_id_don = tim_cot(ALIASES_ID_DON)
        col_so_to = tim_cot(ALIASES_SO_TO)
        col_so_thua = tim_cot(ALIASES_SO_THUA)
        col_so_gcn = tim_cot(ALIASES_SO_GCN)

        if not col_id_don and not (col_so_to and col_so_thua):
            raise ValueError(
                "File Excel cần có cột Mã đơn/ID ('madon' hoặc 'tinhhinhdangkyid') "
                "HOẶC cặp cột Tờ/Thửa ('soto', 'sothua')."
            )

        rows: list[dict[str, Any]] = []
        for r_idx in range(2, worksheet.max_row + 1):
            id_don_val = worksheet.cell(row=r_idx, column=col_id_don).value if col_id_don else None
            so_to_val = worksheet.cell(row=r_idx, column=col_so_to).value if col_so_to else None
            so_thua_val = worksheet.cell(row=r_idx, column=col_so_thua).value if col_so_thua else None
            so_gcn_val = worksheet.cell(row=r_idx, column=col_so_gcn).value if col_so_gcn else None

            if not any([id_don_val, so_to_val, so_thua_val, so_gcn_val]):
                continue

            raw_id_don_str = chuan_hoa_gia_tri_excel(id_don_val)
            extracted_id, extracted_gcn = tach_thong_tin_tu_chuoi_madon(id_don_val)

            id_don_str = extracted_id if extracted_id else raw_id_don_str
            so_to_str = chuan_hoa_gia_tri_excel(so_to_val)
            so_thua_str = chuan_hoa_gia_tri_excel(so_thua_val)
            so_gcn_str = chuan_hoa_gia_tri_excel(so_gcn_val)

            # Nếu cột GCN rỗng nhưng trong chuỗi mã đơn có nhắc tới GCN, dùng luôn
            if not so_gcn_str and extracted_gcn:
                so_gcn_str = extracted_gcn

            rows.append({
                "excel_row": r_idx,
                "id_don": id_don_str,
                "soto": so_to_str,
                "sothua": so_thua_str,
                "sogcn": so_gcn_str,
            })

        if not rows:
            raise ValueError("File Excel không có dữ liệu để xử lý.")
        return rows
    finally:
        workbook.close()



class MultiFileResultWriter:
    """Ghi kết quả chi tiết từng file ra Excel có định dạng đẹp mắt."""
    def __init__(self, output_path: str):
        self.output_path = output_path
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "KetQuaThayTheHSQ"
        self.stt = 0
        self._setup()

    def _setup(self) -> None:
        self.worksheet.append(OUTPUT_HEADERS)
        header_fill = PatternFill("solid", fgColor="1B365D")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="B7B7B7")

        for cell in self.worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        widths = {
            1: 6, 2: 14, 3: 10, 4: 10, 5: 24, 6: 14, 7: 10,
            8: 35, 9: 14, 10: 38, 11: 25, 12: 16, 13: 45
        }
        for col_idx, width in widths.items():
            self.worksheet.column_dimensions[get_column_letter(col_idx)].width = width

        self.worksheet.freeze_panes = "A2"
        self.worksheet.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_HEADERS))}1"
        self.worksheet.row_dimensions[1].height = 32

    def append_row(self, data: dict[str, Any]) -> None:
        self.stt += 1
        self.worksheet.append([
            self.stt,
            data.get("id_don", ""),
            data.get("soto", ""),
            data.get("sothua", ""),
            data.get("chu_su_dung", ""),
            data.get("hosoquet_id", ""),
            data.get("file_idx", ""),
            data.get("mota_cu", ""),
            data.get("loai_hsq", ""),
            data.get("file_moi", ""),
            data.get("kieu_khop", ""),
            data.get("trang_thai", ""),
            data.get("ghi_chu", ""),
        ])

        r_idx = self.worksheet.max_row
        thin = Side(style="thin", color="D9D9D9")
        for cell in self.worksheet[r_idx]:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        trang_thai = str(data.get("trang_thai") or "")
        cell_tt = self.worksheet.cell(row=r_idx, column=12)

        if trang_thai == "Thành công":
            cell_tt.fill = PatternFill("solid", fgColor="C6EFCE")
            cell_tt.font = Font(color="006100", bold=True)
        elif trang_thai == "DRY_RUN":
            cell_tt.fill = PatternFill("solid", fgColor="D9EAF7")
            cell_tt.font = Font(color="1F4E78")
        elif trang_thai in ("Bỏ qua", "Giữ nguyên"):
            cell_tt.fill = PatternFill("solid", fgColor="FFF2CC")
            cell_tt.font = Font(color="7F6000")
        else:
            cell_tt.fill = PatternFill("solid", fgColor="F4CCCC")
            cell_tt.font = Font(color="9C0006")

    def save(self) -> None:
        parent = Path(self.output_path).resolve().parent
        parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(self.output_path)

    def close(self) -> None:
        try:
            self.workbook.close()
        except Exception:
            pass


# ============================ CORE API MPLIS ============================

class MplisClient:
    def __init__(self, log_fn):
        self.log = log_fn
        self.session: requests.Session | None = None
        self.driver: webdriver.Chrome | None = None

    def open_browser_and_fill_login(self, username: str, password: str) -> None:
        options = Options()
        options.add_argument("--start-maximized")
        try:
            service = Service(ChromeDriverManager().install())
            self.driver = webdriver.Chrome(service=service, options=options)
        except Exception as exc:
            self.log(f"Cảnh báo: ChromeDriverManager ({exc}), thử khởi chạy Chrome trực tiếp qua Selenium...")
            self.driver = webdriver.Chrome(options=options)
        self.driver.get(REFERER_LOGIN)
        time.sleep(2)

        try:
            inputs = self.driver.find_elements(By.CSS_SELECTOR, "input")
            user_box, pass_box = None, None
            for inp in inputs:
                itype = (inp.get_attribute("type") or "").lower()
                if user_box is None and itype in {"text", "email"}:
                    user_box = inp
                if pass_box is None and itype == "password":
                    pass_box = inp

            if user_box and pass_box:
                user_box.clear()
                user_box.send_keys(username)
                pass_box.clear()
                pass_box.send_keys(password)
                pass_box.send_keys(Keys.ENTER)
                self.log("Đã tự động điền form đăng nhập trên Chrome.")
            else:
                self.log("Không nhận diện được form, hãy đăng nhập tay trên Chrome.")
        except Exception as exc:
            self.log(f"Lỗi điền form đăng nhập ({exc}), hãy thao tác tay trên Chrome.")

    def build_session_from_browser(self) -> None:
        if not self.driver:
            raise RuntimeError("Chưa mở trình duyệt Chrome.")

        token = lay_token_tu_trang(self.driver)
        if not token:
            raise RuntimeError("Không lấy được token. Hãy bảo đảm đã đăng nhập thành công trang MPLIS.")

        session = requests.Session()
        user_agent = self.driver.execute_script("return navigator.userAgent;")
        session.headers.update({
            "User-Agent": user_agent,
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "Accept-Language": "vi-VN,vi;q=0.9",
            "X-Requested-With": "XMLHttpRequest",
            "Origin": BASE_URL,
            "Referer": REFERER_API,
            "__requestverificationtoken": token,
        })

        for cookie in self.driver.get_cookies():
            session.cookies.set(
                name=cookie["name"],
                value=cookie["value"],
                domain=cookie.get("domain"),
                path=cookie.get("path", "/"),
            )

        self.session = session
        self.log("✅ Đã lấy session + token thành công.")

    def build_session_from_manual(self, token: str, cookie_raw: str) -> None:
        token = token.strip()
        cookie_raw = cookie_raw.strip()
        if not token:
            raise RuntimeError("Chưa nhập token.")
        if not cookie_raw:
            raise RuntimeError("Chưa nhập cookie.")

        session = requests.Session()
        session.headers.update({
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36"
            ),
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "Accept-Language": "vi-VN,vi;q=0.9",
            "X-Requested-With": "XMLHttpRequest",
            "Origin": BASE_URL,
            "Referer": REFERER_API,
            "__requestverificationtoken": token,
        })

        cnt = 0
        for part in cookie_raw.split(";"):
            part = part.strip()
            if not part or "=" not in part:
                continue
            name, _, val = part.partition("=")
            name = name.strip()
            val = val.strip()
            if name:
                session.cookies.set(name=name, value=val)
                cnt += 1

        if cnt == 0:
            raise RuntimeError("Không đọc được cookie nào từ chuỗi đã dán.")
        self.session = session
        self.log(f"✅ Đã tạo session từ Token/Cookie thủ công ({cnt} cookie).")

    def close_browser(self) -> None:
        if self.driver:
            try:
                self.driver.quit()
            except Exception:
                pass
            self.driver = None

    def _require_session(self) -> requests.Session:
        if self.session is None:
            raise RuntimeError("Chưa khởi tạo session MPLIS (hãy lấy session trước).")
        return self.session

    def _response_json(self, response: requests.Response, api_name: str) -> Any:
        if response.status_code in {301, 302, 303, 307, 308} or response.headers.get("Location"):
            raise RuntimeError(
                f"{api_name} bị chuyển hướng (HTTP {response.status_code}). Phiên đăng nhập có thể đã hết hạn."
            )
        response.raise_for_status()
        try:
            return response.json()
        except Exception as exc:
            raise RuntimeError(
                f"{api_name} không trả JSON (HTTP {response.status_code}): {response.text[:400]}"
            ) from exc

    def tim_tat_ca_don_theo_to_thua(self, xa_id: str, so_to: str, so_thua: str) -> list[dict[str, Any]]:
        """Tra cứu đơn theo tờ/thửa."""
        session = self._require_session()
        payload = {
            "draw": "1",
            "order[0][column]": "5",
            "order[0][dir]": "desc",
            "start": "0",
            "length": "50",
            "model[xaId]": xa_id,
            "model[soThuTuThua]": so_thua,
            "model[soHieuToBanDo]": so_to,
            "model[phucHoiDuLieu]": "false",
        }
        res = session.post(URL_TIM_DON, data=payload, timeout=TIMEOUT)
        js = self._response_json(res, "AdvancedSearchTinhHinhDangKy")
        return js.get("data") or []

    def lay_chi_tiet_don(self, tinh_hinh_dang_ky_id: int) -> dict[str, Any] | None:
        session = self._require_session()
        payload = {"getHoSoQuet": True, "tinhHinhDangKyIds": [int(tinh_hinh_dang_ky_id)]}
        res = session.post(URL_CHI_TIET, json=payload, timeout=TIMEOUT)
        js = self._response_json(res, "GetThongTinDangKyByTinhHinhDangKyIds")
        vals = js.get("value") or []
        if isinstance(vals, dict):
            vals = [vals]
        if not isinstance(vals, list) or not vals:
            return None
        for it in vals:
            if isinstance(it, dict):
                cur_id = (it.get("TinhHinhDangKy") or {}).get("tinhHinhDangKyId")
                if str(cur_id) == str(tinh_hinh_dang_ky_id):
                    return it
        return vals[0] if vals and isinstance(vals[0], dict) else None

    def api_update_all_matched_hosoquet(
        self,
        hoso: dict[str, Any],
        matched_items: list[dict[str, Any]],
        cap_nhat_mota_theo_ten_file: bool = False,
    ) -> str:
        """
        Gửi request UpdateHoSoQuetExistFile thay thế đồng thời toàn bộ các file đã khớp.
        - count = len(matched_items)
        - Giữ nguyên các node không có file mới (files = None) để server giữ nguyên file cũ.
        - Node nào có file mới: mở và gửi nhị phân fileHoSoQuet_{i}.
        """
        session = self._require_session()

        ho_so_quet_id_int = safe_int(hoso.get("hoSoQuetId") or hoso.get("Title"), 0)
        hoso_core = copy.deepcopy(hoso)
        hoso_core.pop("ListFileHoSoQuet", None)
        hoso_core["hoSoQuetId"] = ho_so_quet_id_int

        data = {
            "hoSoQuet": json.dumps(hoso_core, ensure_ascii=False),
            "count": str(len(matched_items)),
            "isLuuKhoHoSoQuet": "false",
        }

        files_to_send = {}

        with ExitStack() as stack:
            for idx, item in enumerate(matched_items, start=1):
                orig_node = item["node"]
                new_pdf_path = item["matched_path"]

                node_copy = copy.deepcopy(orig_node)
                node_copy["files"] = None
                node_copy["isOldFile"] = False

                if new_pdf_path and os.path.isfile(new_pdf_path):
                    # Đính kèm file mới
                    f_handle = stack.enter_context(open(new_pdf_path, "rb"))
                    files_to_send[f"fileHoSoQuet_{idx}"] = (
                        os.path.basename(new_pdf_path),
                        f_handle,
                        "application/pdf",
                    )
                    if cap_nhat_mota_theo_ten_file:
                        node_copy["moTa"] = Path(new_pdf_path).stem

                data[f"infoHoSoQuet_{idx}"] = json.dumps(node_copy, ensure_ascii=False)

            headers = dict(session.headers)
            headers.pop("Content-Type", None)

            res = session.post(
                URL_UPDATE_HOSOQUET,
                data=data,
                files=files_to_send if files_to_send else None,
                headers=headers,
                timeout=TIMEOUT,
            )

            result = self._response_json(res, "UpdateHoSoQuetExistFile (multi-file)")
            if not isinstance(result, dict) or not result.get("success"):
                raise RuntimeError(f"Update HSQ không thành công: {rut_gon_text(result)}")

        return f"Cập nhật thành công {len(files_to_send)} file mới vào HSQ"


# ============================ LUỒNG XỬ LÝ 1 ĐƠN ============================

def xu_ly_mot_don_nhieu_file(
    client: MplisClient,
    file_index: FileIndex,
    tinh_hinh_dang_ky_id: int,
    so_to: str,
    so_thua: str,
    so_gcn_hint: str,
    cap_nhat_mota: bool,
    require_all_match: bool,
    dry_run: bool,
    log_fn,
) -> list[dict[str, Any]]:
    """
    Xử lý 1 đơn đăng ký:
    1. Lấy chi tiết đơn kèm ListHoSoQuet.
    2. Duyệt qua từng file trong HSQ và tìm file PDF trên đĩa khớp theo mô tả.
    3. Cập nhật đồng loạt các file đã khớp.
    4. Trả về danh sách kết quả chi tiết từng file để ghi Excel.
    """
    results: list[dict[str, Any]] = []

    try:
        detail = client.lay_chi_tiet_don(tinh_hinh_dang_ky_id)
    except Exception as exc:
        results.append({
            "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
            "chu_su_dung": "", "hosoquet_id": "", "file_idx": "",
            "mota_cu": "", "loai_hsq": "", "file_moi": "", "kieu_khop": "",
            "trang_thai": "Lỗi", "ghi_chu": f"Lỗi lấy chi tiết đơn: {exc}",
        })
        return results

    if not detail:
        results.append({
            "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
            "chu_su_dung": "", "hosoquet_id": "", "file_idx": "",
            "mota_cu": "", "loai_hsq": "", "file_moi": "", "kieu_khop": "",
            "trang_thai": "Lỗi", "ghi_chu": "Không tìm thấy dữ liệu chi tiết đơn trên MPLIS.",
        })
        return results

    # Thông tin thửa & chủ
    thdk = detail.get("TinhHinhDangKy") or {}
    so_to = so_to or str(thdk.get("soHieuToBanDo") or "")
    so_thua = so_thua or str(thdk.get("soThuTuThua") or "")

    chu_so_huu = detail.get("ChuSoHuu") or {}
    ca_nhans = chu_so_huu.get("CaNhans") or []
    to_chucs = chu_so_huu.get("ToChucs") or []
    chu_ten = ""
    if ca_nhans:
        chu_ten = ca_nhans[0].get("hoTen") or ""
    elif to_chucs:
        chu_ten = to_chucs[0].get("tenToChuc") or to_chucs[0].get("ten") or ""

    list_hsq = detail.get("ListHoSoQuet") or []
    if not list_hsq:
        results.append({
            "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
            "chu_su_dung": chu_ten, "hosoquet_id": "", "file_idx": "",
            "mota_cu": "", "loai_hsq": "", "file_moi": "", "kieu_khop": "",
            "trang_thai": "Bỏ qua", "ghi_chu": "Đơn không có hồ sơ quét nào (ListHoSoQuet rỗng).",
        })
        return results

    # Xử lý từng gói HSQ trong đơn (thường chỉ có 1 gói)
    for hoso in list_hsq:
        hsq_id = safe_int(hoso.get("hoSoQuetId") or hoso.get("Title"), 0)
        wrapper = hoso.get("ListFileHoSoQuet") or {}
        list_file = wrapper.get("ListFileHoSoQuet") or []

        if not list_file:
            results.append({
                "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
                "chu_su_dung": chu_ten, "hosoquet_id": str(hsq_id), "file_idx": "",
                "mota_cu": "", "loai_hsq": "", "file_moi": "", "kieu_khop": "",
                "trang_thai": "Bỏ qua", "ghi_chu": f"HSQ {hsq_id} không có file thành phần nào.",
            })
            continue

        log_fn(f"   📂 HSQ {hsq_id}: có {len(list_file)} file thành phần.")

        # So khớp từng file
        matched_items: list[dict[str, Any]] = []
        so_file_khop = 0

        for f_idx, node in enumerate(list_file, start=1):
            mo_ta = node.get("moTa") or ""
            ten_giay_to = node.get("tenGiayTo") or ""
            loai_hsq_code = safe_int(node.get("loaiHoSoQuet"), 0)
            loai_hsq_str = LOAI_HO_SO_QUET_OPTIONS.get(loai_hsq_code, str(loai_hsq_code))

            matched_path, reason = file_index.find_match(
                mo_ta=mo_ta,
                ten_giay_to=ten_giay_to,
                so_gcn=so_gcn_hint,
            )

            if matched_path:
                so_file_khop += 1
                log_fn(f"      [File {f_idx}/{len(list_file)}] '{mo_ta}' -> 🎯 {reason}")
            else:
                log_fn(f"      [File {f_idx}/{len(list_file)}] '{mo_ta}' -> ⚠️ {reason}")

            matched_items.append({
                "node": node,
                "file_idx": f_idx,
                "mo_ta": mo_ta,
                "loai_hsq_str": loai_hsq_str,
                "matched_path": matched_path,
                "reason": reason,
            })

        # Kiểm tra điều kiện bắt buộc khớp 100%
        if require_all_match and so_file_khop < len(list_file):
            msg = f"Bỏ qua: chỉ khớp {so_file_khop}/{len(list_file)} file (yêu cầu khớp đủ 100%)."
            log_fn(f"   ⛔ {msg}")
            for item in matched_items:
                results.append({
                    "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
                    "chu_su_dung": chu_ten, "hosoquet_id": str(hsq_id),
                    "file_idx": item["file_idx"], "mota_cu": item["mo_ta"],
                    "loai_hsq": item["loai_hsq_str"],
                    "file_moi": Path(item["matched_path"]).name if item["matched_path"] else "",
                    "kieu_khop": item["reason"], "trang_thai": "Bỏ qua", "ghi_chu": msg,
                })
            continue

        if so_file_khop == 0:
            msg = f"Bỏ qua: Không tìm thấy bất kỳ file nào trên đĩa khớp với {len(list_file)} file trong HSQ."
            log_fn(f"   ⚠️ {msg}")
            for item in matched_items:
                results.append({
                    "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
                    "chu_su_dung": chu_ten, "hosoquet_id": str(hsq_id),
                    "file_idx": item["file_idx"], "mota_cu": item["mo_ta"],
                    "loai_hsq": item["loai_hsq_str"], "file_moi": "",
                    "kieu_khop": item["reason"], "trang_thai": "Bỏ qua", "ghi_chu": msg,
                })
            continue

        # DRY RUN
        if dry_run:
            log_fn(f"   🧪 [DRY RUN] Giả lập thay thế {so_file_khop}/{len(list_file)} file thành công.")
            for item in matched_items:
                is_matched = bool(item["matched_path"])
                results.append({
                    "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
                    "chu_su_dung": chu_ten, "hosoquet_id": str(hsq_id),
                    "file_idx": item["file_idx"], "mota_cu": item["mo_ta"],
                    "loai_hsq": item["loai_hsq_str"],
                    "file_moi": Path(item["matched_path"]).name if is_matched else "",
                    "kieu_khop": item["reason"],
                    "trang_thai": "DRY_RUN",
                    "ghi_chu": f"DRY RUN: {'Sẽ thay file mới' if is_matched else 'Giữ nguyên file cũ'}",
                })
            continue

        # Gửi API cập nhật thật
        try:
            msg_ok = client.api_update_all_matched_hosoquet(
                hoso=hoso,
                matched_items=matched_items,
                cap_nhat_mota_theo_ten_file=cap_nhat_mota,
            )
            log_fn(f"   ✅ {msg_ok}")
            for item in matched_items:
                is_matched = bool(item["matched_path"])
                results.append({
                    "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
                    "chu_su_dung": chu_ten, "hosoquet_id": str(hsq_id),
                    "file_idx": item["file_idx"], "mota_cu": item["mo_ta"],
                    "loai_hsq": item["loai_hsq_str"],
                    "file_moi": Path(item["matched_path"]).name if is_matched else "",
                    "kieu_khop": item["reason"],
                    "trang_thai": "Thành công" if is_matched else "Giữ nguyên",
                    "ghi_chu": "Đã thay file mới vào HSQ" if is_matched else "Không có file mới, giữ nguyên file cũ",
                })
        except Exception as exc:
            err_msg = f"Lỗi Update HSQ: {rut_gon_text(exc)}"
            log_fn(f"   ❌ {err_msg}")
            for item in matched_items:
                results.append({
                    "id_don": tinh_hinh_dang_ky_id, "soto": so_to, "sothua": so_thua,
                    "chu_su_dung": chu_ten, "hosoquet_id": str(hsq_id),
                    "file_idx": item["file_idx"], "mota_cu": item["mo_ta"],
                    "loai_hsq": item["loai_hsq_str"],
                    "file_moi": Path(item["matched_path"]).name if item["matched_path"] else "",
                    "kieu_khop": item["reason"], "trang_thai": "Lỗi", "ghi_chu": err_msg,
                })

    return results


# ============================ GIAO DIỆN CUSTOM TKINTER SIÊU NGẦU ============================

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")


class App(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("⚡ MPLIS PRO - BATCH HSQ REPLACER ⚡ (Thay Thế Toàn Bộ File HSQ Theo Mô Tả)")
        self.geometry("1280x880")
        self.minsize(1120, 760)

        # Màu sắc chủ đạo (Neo-Dark / Cyber-Slate)
        self.COLOR_BG = "#0b0f19"
        self.COLOR_CARD = "#151e32"
        self.COLOR_CARD_BORDER = "#23314f"
        self.COLOR_ACCENT = "#3b82f6"
        self.COLOR_SUCCESS = "#10b981"
        self.COLOR_DANGER = "#ef4444"
        self.COLOR_WARNING = "#f59e0b"
        self.COLOR_TEXT_DIM = "#94a3b8"

        self.configure(fg_color=self.COLOR_BG)

        self.client = MplisClient(self.log)
        self.running = False
        self.stop_flag = False

        # Biến điều khiển
        self.var_token = ctk.StringVar()
        self.var_cookie = ctk.StringVar()
        self.var_folder_pdf = ctk.StringVar()
        self.var_input_excel = ctk.StringVar()
        self.var_output_excel = ctk.StringVar()
        self.var_single_id = ctk.StringVar()
        self.var_xa_id = ctk.StringVar()
        self.var_dry_run = ctk.BooleanVar(value=False)
        self.var_require_all_match = ctk.BooleanVar(value=False)
        self.var_cap_nhat_mota = ctk.BooleanVar(value=False)

        self.total_don_count = 0
        self.success_don_count = 0
        self.total_files_replaced_count = 0

        self._build_ui()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def _build_ui(self) -> None:
        # ==================== 1. TOP HEADER BANNER ====================
        header_frame = ctk.CTkFrame(self, fg_color="#101728", corner_radius=0, height=72, border_width=1, border_color="#1e293b")
        header_frame.pack(fill="x", side="top")
        header_frame.pack_propagate(False)

        left_header = ctk.CTkFrame(header_frame, fg_color="transparent")
        left_header.pack(side="left", padx=20, pady=10)

        lbl_logo = ctk.CTkLabel(
            left_header,
            text="⚡ MPLIS PRO : BATCH HSQ REPLACER",
            font=ctk.CTkFont(family="Segoe UI", size=20, weight="bold"),
            text_color="#60a5fa",
        )
        lbl_logo.pack(anchor="w")

        lbl_sub = ctk.CTkLabel(
            left_header,
            text="Hệ thống tự động so khớp mô tả và thay thế toàn bộ file hồ sơ quét MPLIS (Bảo toàn ID liên kết)",
            font=ctk.CTkFont(family="Segoe UI", size=12),
            text_color=self.COLOR_TEXT_DIM,
        )
        lbl_sub.pack(anchor="w")

        right_header = ctk.CTkFrame(header_frame, fg_color="transparent")
        right_header.pack(side="right", padx=20, pady=10)

        self.lbl_status_badge = ctk.CTkLabel(
            right_header,
            text="● SẴN SÀNG",
            font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
            text_color="#34d399",
            fg_color="#064e3b",
            corner_radius=12,
            padx=14,
            pady=4,
        )
        self.lbl_status_badge.pack(side="right", padx=(10, 0))

        opt_theme = ctk.CTkOptionMenu(
            right_header,
            values=["Dark", "Light", "System"],
            command=self._change_appearance_mode,
            width=90,
            height=28,
            fg_color="#1e293b",
            button_color="#334155",
            dropdown_fg_color="#1e293b",
        )
        opt_theme.set("Dark")
        opt_theme.pack(side="right")

        # ==================== 2. MAIN CONTAINER (2 COLUMNS) ====================
        main_container = ctk.CTkFrame(self, fg_color="transparent")
        main_container.pack(fill="both", expand=True, padx=16, pady=12)
        main_container.grid_columnconfigure(0, weight=5)  # Left panel (Controls)
        main_container.grid_columnconfigure(1, weight=6)  # Right panel (Dashboard + Logs)
        main_container.grid_rowconfigure(0, weight=1)

        # ----------------- CỘT TRÁI: CẤU HÌNH & THAO TÁC -----------------
        left_scroll = ctk.CTkScrollableFrame(
            main_container,
            fg_color="#111827",
            corner_radius=14,
            border_width=1,
            border_color=self.COLOR_CARD_BORDER,
            scrollbar_button_color="#2563eb",
        )
        left_scroll.grid(row=0, column=0, sticky="nsew", padx=(0, 8), pady=0)

        # CARD 1: XÁC THỰC MPLIS
        card_auth = self._create_card(left_scroll, "🔐 1. XÁC THỰC & PHIÊN LÀM VIỆC MPLIS")
        card_auth.pack(fill="x", padx=12, pady=(10, 8))

        self.seg_auth_mode = ctk.CTkSegmentedButton(
            card_auth,
            values=["🌐 Tự động qua Chrome", "📋 Dán Token / Cookie"],
            command=self._on_auth_mode_changed,
            selected_color="#2563eb",
            selected_hover_color="#1d4ed8",
        )
        self.seg_auth_mode.pack(fill="x", padx=12, pady=(4, 10))
        self.seg_auth_mode.set("🌐 Tự động qua Chrome")

        # Khối Chrome
        self.frame_auth_chrome = ctk.CTkFrame(card_auth, fg_color="transparent")
        self.frame_auth_chrome.pack(fill="x", padx=12, pady=2)

        row_user = ctk.CTkFrame(self.frame_auth_chrome, fg_color="transparent")
        row_user.pack(fill="x", pady=3)
        self.ent_user = ctk.CTkEntry(
            row_user, placeholder_text="Tài khoản đăng nhập MPLIS...", height=36, corner_radius=8
        )
        self.ent_user.pack(side="left", fill="x", expand=True, padx=(0, 6))
        self.ent_pass = ctk.CTkEntry(
            row_user, placeholder_text="Mật khẩu...", show="*", height=36, corner_radius=8
        )
        self.ent_pass.pack(side="right", fill="x", expand=True, padx=(6, 0))

        row_btn_chrome = ctk.CTkFrame(self.frame_auth_chrome, fg_color="transparent")
        row_btn_chrome.pack(fill="x", pady=(6, 4))
        self.btn_login = ctk.CTkButton(
            row_btn_chrome,
            text="Mở Chrome Đăng Nhập",
            command=self.mo_chrome,
            fg_color="#2563eb",
            hover_color="#1d4ed8",
            height=36,
            corner_radius=8,
            font=ctk.CTkFont(weight="bold"),
        )
        self.btn_login.pack(side="left", fill="x", expand=True, padx=(0, 6))

        self.btn_session = ctk.CTkButton(
            row_btn_chrome,
            text="✅ Đã Đăng Nhập -> Lấy Session",
            command=self.lay_session,
            state="disabled",
            fg_color="#059669",
            hover_color="#047857",
            height=36,
            corner_radius=8,
            font=ctk.CTkFont(weight="bold"),
        )
        self.btn_session.pack(side="right", fill="x", expand=True, padx=(6, 0))

        # Khối Manual Token/Cookie
        self.frame_auth_manual = ctk.CTkFrame(card_auth, fg_color="transparent")

        self.ent_token = ctk.CTkEntry(
            self.frame_auth_manual,
            textvariable=self.var_token,
            placeholder_text="Dán __RequestVerificationToken từ F12...",
            height=36,
            corner_radius=8,
        )
        self.ent_token.pack(fill="x", pady=3)

        self.ent_cookie = ctk.CTkEntry(
            self.frame_auth_manual,
            textvariable=self.var_cookie,
            placeholder_text="Dán chuỗi Cookie từ F12 Request Headers...",
            height=36,
            corner_radius=8,
        )
        self.ent_cookie.pack(fill="x", pady=3)

        self.btn_manual_session = ctk.CTkButton(
            self.frame_auth_manual,
            text="⚡ Kích Hoạt Session Nhanh",
            command=self.dung_token_cookie,
            fg_color="#059669",
            hover_color="#047857",
            height=36,
            corner_radius=8,
            font=ctk.CTkFont(weight="bold"),
        )
        self.btn_manual_session.pack(fill="x", pady=(4, 2))

        # CARD 2: THƯ MỤC FILE PDF NGUỒN
        card_folder = self._create_card(left_scroll, "📂 2. THƯ MỤC CHỨA FILE PDF THAY THẾ (BẮT BUỘC)")
        card_folder.pack(fill="x", padx=12, pady=8)

        row_folder = ctk.CTkFrame(card_folder, fg_color="transparent")
        row_folder.pack(fill="x", padx=12, pady=4)
        self.ent_folder = ctk.CTkEntry(
            row_folder,
            textvariable=self.var_folder_pdf,
            placeholder_text="Chọn đường dẫn thư mục chứa các file PDF mới...",
            height=36,
            corner_radius=8,
        )
        self.ent_folder.pack(side="left", fill="x", expand=True, padx=(0, 6))

        btn_browse_folder = ctk.CTkButton(
            row_folder,
            text="Duyệt...",
            command=self.chon_folder_pdf,
            width=85,
            height=36,
            corner_radius=8,
            fg_color="#334155",
            hover_color="#475569",
        )
        btn_browse_folder.pack(side="right")

        self.lbl_pdf_count = ctk.CTkLabel(
            card_folder,
            text="🔍 Chưa chọn thư mục PDF",
            font=ctk.CTkFont(size=12),
            text_color="#38bdf8",
        )
        self.lbl_pdf_count.pack(anchor="w", padx=12, pady=(0, 6))

        # CARD 3: NGUỒN DANH SÁCH ĐƠN
        card_data = self._create_card(left_scroll, "📋 3. NGUỒN DANH SÁCH ĐƠN CẦN XỬ LÝ")
        card_data.pack(fill="x", padx=12, pady=8)

        self.seg_data_mode = ctk.CTkSegmentedButton(
            card_data,
            values=["📑 File Excel Danh Sách", "🎯 Chạy Nhanh 1 Mã Đơn"],
            command=self._on_data_mode_changed,
            selected_color="#2563eb",
            selected_hover_color="#1d4ed8",
        )
        self.seg_data_mode.pack(fill="x", padx=12, pady=(4, 10))
        self.seg_data_mode.set("📑 File Excel Danh Sách")

        # Tab Excel
        self.frame_data_excel = ctk.CTkFrame(card_data, fg_color="transparent")
        self.frame_data_excel.pack(fill="x", padx=12, pady=2)

        lbl_excel_in = ctk.CTkLabel(self.frame_data_excel, text="File Excel đầu vào:", font=ctk.CTkFont(size=12))
        lbl_excel_in.pack(anchor="w")
        row_in = ctk.CTkFrame(self.frame_data_excel, fg_color="transparent")
        row_in.pack(fill="x", pady=(2, 6))
        self.ent_input_excel = ctk.CTkEntry(
            row_in,
            textvariable=self.var_input_excel,
            placeholder_text="Chọn file Excel có cột madon hoặc soto/sothua...",
            height=36,
            corner_radius=8,
        )
        self.ent_input_excel.pack(side="left", fill="x", expand=True, padx=(0, 6))
        btn_in = ctk.CTkButton(
            row_in, text="Chọn...", command=self.chon_file_excel, width=75, height=36, corner_radius=8, fg_color="#334155", hover_color="#475569"
        )
        btn_in.pack(side="right")

        lbl_excel_out = ctk.CTkLabel(self.frame_data_excel, text="Nơi lưu kết quả Excel:", font=ctk.CTkFont(size=12))
        lbl_excel_out.pack(anchor="w")
        row_out = ctk.CTkFrame(self.frame_data_excel, fg_color="transparent")
        row_out.pack(fill="x", pady=(2, 4))
        self.ent_output_excel = ctk.CTkEntry(
            row_out,
            textvariable=self.var_output_excel,
            placeholder_text="Đường dẫn file Excel kết quả xuất ra...",
            height=36,
            corner_radius=8,
        )
        self.ent_output_excel.pack(side="left", fill="x", expand=True, padx=(0, 6))
        btn_out = ctk.CTkButton(
            row_out, text="Lưu tại...", command=self.chon_file_output, width=75, height=36, corner_radius=8, fg_color="#334155", hover_color="#475569"
        )
        btn_out.pack(side="right")

        # Tab Single ID
        self.frame_data_single = ctk.CTkFrame(card_data, fg_color="transparent")

        row_single = ctk.CTkFrame(self.frame_data_single, fg_color="transparent")
        row_single.pack(fill="x", pady=4)
        self.ent_single_id = ctk.CTkEntry(
            row_single,
            textvariable=self.var_single_id,
            placeholder_text="Nhập mã đơn (tinhHinhDangKyId)...",
            height=36,
            corner_radius=8,
        )
        self.ent_single_id.pack(side="left", fill="x", expand=True, padx=(0, 6))

        self.ent_xa_id = ctk.CTkEntry(
            row_single,
            textvariable=self.var_xa_id,
            placeholder_text="Mã xã (xaId)...",
            width=120,
            height=36,
            corner_radius=8,
        )
        self.ent_xa_id.pack(side="right")

        # CARD 4: TUỲ CHỌN NÂNG CAO
        card_opts = self._create_card(left_scroll, "⚙️ 4. TUỲ CHỌN THAY THẾ NÂNG CAO")
        card_opts.pack(fill="x", padx=12, pady=8)

        self.switch_dry_run = ctk.CTkSwitch(
            card_opts,
            text="🧪 DRY RUN (Quét kiểm tra khớp file, KHÔNG update thật)",
            variable=self.var_dry_run,
            progress_color="#f59e0b",
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.switch_dry_run.pack(anchor="w", padx=14, pady=6)

        self.switch_require_all = ctk.CTkSwitch(
            card_opts,
            text="🛡️ Chỉ cập nhật khi khớp đủ 100% file trong đơn (bỏ qua nếu thiếu)",
            variable=self.var_require_all_match,
            progress_color="#3b82f6",
            font=ctk.CTkFont(size=13),
        )
        self.switch_require_all.pack(anchor="w", padx=14, pady=6)

        self.switch_cap_nhat_mota = ctk.CTkSwitch(
            card_opts,
            text="🏷️ Cập nhật mô tả (moTa) theo tên file PDF mới",
            variable=self.var_cap_nhat_mota,
            progress_color="#8b5cf6",
            font=ctk.CTkFont(size=13),
        )
        self.switch_cap_nhat_mota.pack(anchor="w", padx=14, pady=(6, 10))

        # ACTION BUTTONS
        frame_actions = ctk.CTkFrame(left_scroll, fg_color="transparent")
        frame_actions.pack(fill="x", padx=12, pady=(10, 16))

        self.btn_run = ctk.CTkButton(
            frame_actions,
            text="🚀 BẮT ĐẦU XỬ LÝ & THAY THẾ FILE",
            command=self.chay,
            fg_color="#059669",
            hover_color="#10b981",
            height=50,
            corner_radius=10,
            font=ctk.CTkFont(size=15, weight="bold"),
            state="disabled",
        )
        self.btn_run.pack(fill="x", pady=4)

        self.btn_stop = ctk.CTkButton(
            frame_actions,
            text="⏹ DỪNG KHẨN CẤP",
            command=self.dung,
            fg_color="#dc2626",
            hover_color="#ef4444",
            height=40,
            corner_radius=10,
            font=ctk.CTkFont(size=13, weight="bold"),
            state="disabled",
        )
        self.btn_stop.pack(fill="x", pady=4)

        # ----------------- CỘT PHẢI: DASHBOARD THỐNG KÊ & TERMINAL REALTIME -----------------
        right_panel = ctk.CTkFrame(
            main_container,
            fg_color="#111827",
            corner_radius=14,
            border_width=1,
            border_color=self.COLOR_CARD_BORDER,
        )
        right_panel.grid(row=0, column=1, sticky="nsew", padx=(8, 0), pady=0)
        right_panel.grid_rowconfigure(2, weight=1)
        right_panel.grid_columnconfigure(0, weight=1)

        # 1. KPI CARDS
        kpi_grid = ctk.CTkFrame(right_panel, fg_color="transparent")
        kpi_grid.grid(row=0, column=0, sticky="ew", padx=14, pady=(14, 10))
        kpi_grid.grid_columnconfigure((0, 1, 2), weight=1)

        self.card_kpi_total = self._create_kpi_card(kpi_grid, 0, "📋 TỔNG ĐƠN", "0", "#38bdf8")
        self.card_kpi_success = self._create_kpi_card(kpi_grid, 1, "✅ THÀNH CÔNG", "0", "#34d399")
        self.card_kpi_files = self._create_kpi_card(kpi_grid, 2, "📁 FILE ĐÃ THAY", "0", "#c084fc")

        # 2. PROGRESS SECTION
        progress_card = ctk.CTkFrame(right_panel, fg_color="#1e293b", corner_radius=10)
        progress_card.grid(row=1, column=0, sticky="ew", padx=14, pady=6)

        row_prog_top = ctk.CTkFrame(progress_card, fg_color="transparent")
        row_prog_top.pack(fill="x", padx=12, pady=(8, 4))

        self.lbl_progress_status = ctk.CTkLabel(
            row_prog_top,
            text="Chưa hoạt động (Chờ phiên làm việc)",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#e2e8f0",
        )
        self.lbl_progress_status.pack(side="left")

        self.lbl_progress_counter = ctk.CTkLabel(
            row_prog_top,
            text="0 / 0",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#60a5fa",
        )
        self.lbl_progress_counter.pack(side="right")

        self.progress_bar = ctk.CTkProgressBar(
            progress_card,
            height=14,
            corner_radius=7,
            progress_color="#2563eb",
            fg_color="#0f172a",
        )
        self.progress_bar.pack(fill="x", padx=12, pady=(2, 10))
        self.progress_bar.set(0.0)

        # 3. TERMINAL LOG CONSOLE
        console_frame = ctk.CTkFrame(right_panel, fg_color="#0f172a", corner_radius=10, border_width=1, border_color="#1e293b")
        console_frame.grid(row=2, column=0, sticky="nsew", padx=14, pady=(8, 14))
        console_frame.grid_rowconfigure(1, weight=1)
        console_frame.grid_columnconfigure(0, weight=1)

        # Console Header
        console_header = ctk.CTkFrame(console_frame, fg_color="#1e293b", corner_radius=0, height=36)
        console_header.grid(row=0, column=0, sticky="ew")
        console_header.grid_columnconfigure(0, weight=1)

        lbl_console_title = ctk.CTkLabel(
            console_header,
            text="⚡ NHẬT KÝ HOẠT ĐỘNG REALTIME",
            font=ctk.CTkFont(family="Consolas", size=12, weight="bold"),
            text_color="#38bdf8",
        )
        lbl_console_title.grid(row=0, column=0, sticky="w", padx=12)

        btn_clear = ctk.CTkButton(
            console_header,
            text="🧹 Xoá Log",
            command=self.clear_log,
            width=70,
            height=24,
            font=ctk.CTkFont(size=11),
            fg_color="#334155",
            hover_color="#475569",
        )
        btn_clear.grid(row=0, column=1, padx=6)

        # Textbox Console
        self.txt_log = ctk.CTkTextbox(
            console_frame,
            fg_color="#090d16",
            text_color="#e2e8f0",
            font=ctk.CTkFont(family="Consolas", size=12),
            wrap="word",
            corner_radius=0,
        )
        self.txt_log.grid(row=1, column=0, sticky="nsew", padx=2, pady=2)

        # Thông báo ban đầu
        self.log("Hệ thống khởi động thành công. Sẵn sàng xử lý.")

    # ---------- HELPER BUILD UI ----------
    def _create_card(self, parent, title: str) -> ctk.CTkFrame:
        card = ctk.CTkFrame(parent, fg_color=self.COLOR_CARD, corner_radius=12, border_width=1, border_color=self.COLOR_CARD_BORDER)
        lbl_title = ctk.CTkLabel(
            card,
            text=title,
            font=ctk.CTkFont(family="Segoe UI", size=13, weight="bold"),
            text_color="#e2e8f0",
        )
        lbl_title.pack(anchor="w", padx=12, pady=(10, 4))
        return card

    def _create_kpi_card(self, parent, col: int, title: str, init_val: str, val_color: str) -> ctk.CTkLabel:
        box = ctk.CTkFrame(parent, fg_color=self.COLOR_CARD, corner_radius=10, border_width=1, border_color=self.COLOR_CARD_BORDER)
        box.grid(row=0, column=col, sticky="ew", padx=4)

        lbl_t = ctk.CTkLabel(box, text=title, font=ctk.CTkFont(size=11, weight="bold"), text_color=self.COLOR_TEXT_DIM)
        lbl_t.pack(pady=(6, 0))

        lbl_v = ctk.CTkLabel(box, text=init_val, font=ctk.CTkFont(size=22, weight="bold"), text_color=val_color)
        lbl_v.pack(pady=(0, 6))
        return lbl_v

    def _change_appearance_mode(self, new_mode: str):
        ctk.set_appearance_mode(new_mode)

    def _on_auth_mode_changed(self, value: str):
        if "Chrome" in value:
            self.frame_auth_manual.pack_forget()
            self.frame_auth_chrome.pack(fill="x", padx=12, pady=2)
        else:
            self.frame_auth_chrome.pack_forget()
            self.frame_auth_manual.pack(fill="x", padx=12, pady=2)

    def _on_data_mode_changed(self, value: str):
        if "Excel" in value:
            self.frame_data_single.pack_forget()
            self.frame_data_excel.pack(fill="x", padx=12, pady=2)
        else:
            self.frame_data_excel.pack_forget()
            self.frame_data_single.pack(fill="x", padx=12, pady=2)

    # ---------- LOGGING ----------
    def log(self, msg: str) -> None:
        def _append():
            ts = datetime.now().strftime("%H:%M:%S")
            self.txt_log.insert("end", f"[{ts}] {msg}\n")
            self.txt_log.see("end")
        self.after(0, _append)

    def clear_log(self) -> None:
        self.txt_log.delete("1.0", "end")

    def set_badge_status(self, text: str, color_text: str, color_bg: str):
        self.after(0, lambda: self.lbl_status_badge.configure(text=text, text_color=color_text, fg_color=color_bg))

    def set_progress(self, current: int, total: int, status_text: str):
        def _update():
            frac = current / max(total, 1)
            self.progress_bar.set(frac)
            self.lbl_progress_counter.configure(text=f"{current} / {total} ({int(frac * 100)}%)")
            self.lbl_progress_status.configure(text=status_text)
        self.after(0, _update)

    def update_kpi(self, total: int, success: int, files: int):
        self.after(0, lambda: self.card_kpi_total.configure(text=str(total)))
        self.after(0, lambda: self.card_kpi_success.configure(text=str(success)))
        self.after(0, lambda: self.card_kpi_files.configure(text=str(files)))

    # ---------- FILE CHOOSERS ----------
    def chon_folder_pdf(self) -> None:
        folder = filedialog.askdirectory(title="Chọn thư mục chứa các file PDF mới")
        if folder:
            self.var_folder_pdf.set(folder)
            # Quét nhanh số lượng
            cnt = 0
            for _, _, files in os.walk(folder):
                for f in files:
                    if f.lower().endswith(".pdf"):
                        cnt += 1
            self.lbl_pdf_count.configure(text=f"📁 Đã quét: Tìm thấy {cnt} file PDF trong thư mục này.")

    def chon_file_excel(self) -> None:
        f = filedialog.askopenfilename(
            title="Chọn file Excel danh sách đơn",
            filetypes=[("Excel Files", "*.xlsx;*.xlsm"), ("All Files", "*.*")],
        )
        if f:
            self.var_input_excel.set(f)
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            out_p = Path(f).with_name(f"ket_qua_thay_the_hsq_{ts}.xlsx")
            self.var_output_excel.set(str(out_p))

    def chon_file_output(self) -> None:
        init_f = Path(self.var_output_excel.get() or "ket_qua_thay_the_hsq.xlsx").name
        f = filedialog.asksaveasfilename(
            title="Chọn nơi lưu file kết quả Excel",
            defaultextension=".xlsx",
            initialfile=init_f,
            filetypes=[("Excel Workbook", "*.xlsx")],
        )
        if f:
            self.var_output_excel.set(f)

    # ---------- AUTH ACTIONS ----------
    def mo_chrome(self) -> None:
        u = self.ent_user.get().strip()
        p = self.ent_pass.get()
        if not u or not p:
            messagebox.showwarning("Thiếu thông tin", "Vui lòng nhập Username và Password trước.")
            return

        self.btn_login.configure(state="disabled")
        self.set_badge_status("● ĐANG MỞ CHROME", "#fbbf24", "#78350f")

        def _work():
            try:
                self.log("Khởi động trình duyệt Chrome qua Selenium...")
                self.client.open_browser_and_fill_login(u, p)
                self.log("Chrome đã mở. Hãy hoàn tất đăng nhập/OTP rồi bấm 'Đã đăng nhập -> Lấy Session'.")
                self.after(0, lambda: self.btn_session.configure(state="normal"))
                self.set_badge_status("● CHỜ LẤY SESSION", "#38bdf8", "#0c4a6e")
            except Exception as e:
                self.log(f"❌ Lỗi mở Chrome: {e}")
                self.after(0, lambda: self.btn_login.configure(state="normal"))
                self.set_badge_status("● LỖI CHROME", "#f87171", "#7f1d1d")

        threading.Thread(target=_work, daemon=True).start()

    def lay_session(self) -> None:
        self.btn_session.configure(state="disabled")

        def _work():
            try:
                self.client.build_session_from_browser()
                self.client.close_browser()
                self.log("✅ Đã lấy session từ Chrome và đóng trình duyệt an toàn.")
                self.after(0, lambda: self.btn_run.configure(state="normal"))
                self.set_badge_status("● SESSION SẴN SÀNG", "#34d399", "#064e3b")
            except Exception as e:
                self.log(f"❌ {e}")
                self.after(0, lambda: self.btn_session.configure(state="normal"))
                self.set_badge_status("● LỖI SESSION", "#f87171", "#7f1d1d")

        threading.Thread(target=_work, daemon=True).start()

    def dung_token_cookie(self) -> None:
        t = self.var_token.get().strip()
        c = self.var_cookie.get().strip()
        if not t or not c:
            messagebox.showwarning("Thiếu thông tin", "Dán cả Token và Cookie trước.")
            return

        try:
            self.client.build_session_from_manual(t, c)
            self.btn_run.configure(state="normal")
            self.set_badge_status("● SESSION SẴN SÀNG", "#34d399", "#064e3b")
            messagebox.showinfo("Thành công", "Đã khởi tạo session từ Token/Cookie thủ công!")
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    # ---------- RUN PROCESS ----------
    def dung(self) -> None:
        self.stop_flag = True
        self.log("⏸ Yêu cầu dừng xử lý! Hệ thống sẽ hoàn tất đơn hiện tại và lưu file kết quả...")
        self.set_badge_status("● ĐANG DỪNG...", "#f87171", "#7f1d1d")

    def chay(self) -> None:
        if self.running:
            return

        folder_pdf = self.var_folder_pdf.get().strip()
        input_excel = self.var_input_excel.get().strip()
        output_excel = self.var_output_excel.get().strip()
        single_id = self.var_single_id.get().strip()
        xa_id = self.var_xa_id.get().strip()
        dry_run = self.var_dry_run.get()
        require_all = self.var_require_all_match.get()
        cap_nhat_mota = self.var_cap_nhat_mota.get()

        if not folder_pdf or not os.path.isdir(folder_pdf):
            messagebox.showwarning("Thiếu thư mục", "Vui lòng chọn thư mục chứa các file PDF mới.")
            return

        items_to_run: list[dict[str, Any]] = []

        if self.seg_data_mode.get() == "🎯 Chạy Nhanh 1 Mã Đơn":
            ext_id, ext_gcn = tach_thong_tin_tu_chuoi_madon(single_id)
            clean_id = ext_id if ext_id else single_id
            if not clean_id:
                messagebox.showwarning("Thiếu mã đơn", "Nhập mã đơn (tinhHinhDangKyId) cần xử lý.")
                return
            items_to_run.append({"excel_row": 0, "id_don": clean_id, "soto": "", "sothua": "", "sogcn": ext_gcn})
            if not output_excel:
                ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_excel = os.path.join(folder_pdf, f"ket_qua_don_{clean_id}_{ts}.xlsx")
                self.var_output_excel.set(output_excel)
        else:
            if not input_excel or not os.path.isfile(input_excel):
                messagebox.showwarning("Thiếu nguồn dữ liệu", "Vui lòng chọn File Excel danh sách đơn.")
                return
            if not output_excel:
                messagebox.showwarning("Thiếu nơi lưu", "Vui lòng chọn nơi lưu file Excel kết quả.")
                return
            try:
                items_to_run = doc_danh_sach_don(input_excel)
            except Exception as e:
                messagebox.showerror("Lỗi đọc Excel", str(e))
                return

        confirm_msg = (
            f"BẮT ĐẦU THAY THẾ TOÀN BỘ FILE HSQ THEO MÔ TẢ:\n\n"
            f"• Số đơn cần xử lý: {len(items_to_run)}\n"
            f"• Thư mục PDF nguồn: {folder_pdf}\n"
            f"• Chế độ DRY RUN: {'CÓ (Chỉ kiểm tra, KHÔNG ghi thật)' if dry_run else 'KHÔNG (Update thật)'}\n"
            f"• Bắt buộc khớp 100%: {'CÓ' if require_all else 'KHÔNG'}\n"
            f"• Cập nhật moTa theo file mới: {'CÓ' if cap_nhat_mota else 'KHÔNG'}\n\n"
            "Xác nhận tiến hành?"
        )
        if not messagebox.askyesno("Xác nhận tiến trình", confirm_msg):
            return

        self.running = True
        self.stop_flag = False
        self.btn_run.configure(state="disabled")
        self.btn_stop.configure(state="normal")
        self.set_badge_status("● ĐANG XỬ LÝ...", "#38bdf8", "#0c4a6e")

        threading.Thread(
            target=self._worker_run,
            args=(folder_pdf, items_to_run, output_excel, xa_id, dry_run, require_all, cap_nhat_mota),
            daemon=True,
        ).start()

    def _worker_run(
        self,
        folder_pdf: str,
        items: list[dict[str, Any]],
        output_excel: str,
        xa_id: str,
        dry_run: bool,
        require_all: bool,
        cap_nhat_mota: bool,
    ) -> None:
        writer = MultiFileResultWriter(output_excel)
        total_don = len(items)
        success_don = 0
        total_files_replaced = 0

        self.update_kpi(total_don, 0, 0)

        try:
            self.log("=" * 65)
            self.log("🚀 BẮT ĐẦU QUÉT THƯ MỤC VÀ TẠO CHỈ MỤC TÌM KIẾM FILE PDF...")
            file_index = FileIndex(folder_pdf)
            self.log(f"✅ Đã lập chỉ mục {len(file_index.files_list)} file PDF trong thư mục nguồn.")

            for idx, it in enumerate(items, start=1):
                if self.stop_flag:
                    self.log("⏹ Đã dừng xử lý theo yêu cầu của người dùng.")
                    break

                raw_id_don = it.get("id_don", "").strip()
                soto = it.get("soto", "").strip()
                sothua = it.get("sothua", "").strip()
                sogcn = it.get("sogcn", "").strip()

                ext_id, ext_gcn = tach_thong_tin_tu_chuoi_madon(raw_id_don)
                id_don = ext_id if ext_id else raw_id_don
                if not sogcn and ext_gcn:
                    sogcn = ext_gcn

                label = f"Mã đơn {id_don}" if id_don else f"Tờ {soto} - Thửa {sothua}"
                self.log(f"\n--- [{idx}/{total_don}] Đang xử lý: {label} ---")
                self.set_progress(idx, total_don, f"Đang xử lý: {label}")

                tinh_hinh_ids = []
                if id_don:
                    try:
                        tinh_hinh_ids.append(int(id_don))
                    except ValueError:
                        self.log(f"   ❌ Mã đơn '{raw_id_don}' không bóc tách được ID số hợp lệ.")
                        continue
                elif soto and sothua:
                    if not xa_id:
                        self.log("   ❌ Thiếu mã xã (xaId) khi tra cứu theo tờ thửa.")
                        continue
                    try:
                        found_rows = self.client.tim_tat_ca_don_theo_to_thua(xa_id, soto, sothua)
                        for r in found_rows:
                            tid = safe_int(r.get("tinhHinhDangKyId"))
                            if tid:
                                tinh_hinh_ids.append(tid)
                        if not tinh_hinh_ids:
                            self.log(f"   ⚠️ Không tìm thấy đơn nào cho tờ {soto} thửa {sothua}.")
                    except Exception as e:
                        self.log(f"   ❌ Lỗi tra cứu tờ/thửa: {e}")
                        continue

                don_has_success = False
                for tid in tinh_hinh_ids:
                    row_results = xu_ly_mot_don_nhieu_file(
                        client=self.client,
                        file_index=file_index,
                        tinh_hinh_dang_ky_id=tid,
                        so_to=soto,
                        so_thua=sothua,
                        so_gcn_hint=sogcn,
                        cap_nhat_mota=cap_nhat_mota,
                        require_all_match=require_all,
                        dry_run=dry_run,
                        log_fn=self.log,
                    )

                    for res in row_results:
                        writer.append_row(res)
                        if res.get("trang_thai") in ("Thành công", "DRY_RUN"):
                            don_has_success = True
                            if res.get("file_moi"):
                                total_files_replaced += 1

                    time.sleep(REQUEST_DELAY_SECONDS)

                if don_has_success:
                    success_don += 1

                self.update_kpi(total_don, success_don, total_files_replaced)

                if idx % SAVE_EVERY_ROWS == 0:
                    writer.save()
                    self.log(f"💾 Đã tự động lưu kết quả sau {idx} đơn.")

            writer.save()
            self.log("\n" + "=" * 65)
            self.log("🎉 HOÀN TẤT QUÁ TRÌNH XỬ LÝ!")
            self.log(f"Tổng số đơn: {total_don} | Đơn thành công: {success_don}")
            self.log(f"Tổng số file PDF đã thay thế: {total_files_replaced}")
            self.log(f"File kết quả Excel: {output_excel}")

            status_final = f"Hoàn tất: {success_don}/{total_don} đơn, {total_files_replaced} file đã thay"
            self.set_progress(total_don, total_don, status_final)
            self.set_badge_status("● HOÀN TẤT", "#34d399", "#064e3b")

            if not self.stop_flag:
                messagebox.showinfo(
                    "Hoàn tất",
                    f"Đã xử lý xong {total_don} đơn!\n"
                    f"- Số file PDF đã thay thế: {total_files_replaced}\n"
                    f"- Kết quả lưu tại: {output_excel}"
                )

        except Exception as exc:
            self.log(f"❌ Lỗi ngoài ý muốn: {exc}")
            self.set_badge_status("● GẶP LỖI", "#f87171", "#7f1d1d")
            try:
                writer.save()
            except Exception:
                pass
        finally:
            writer.close()
            self.running = False
            self.after(0, lambda: self.btn_run.configure(state="normal"))
            self.after(0, lambda: self.btn_stop.configure(state="disabled"))

    def on_close(self) -> None:
        if self.running:
            if not messagebox.askyesno("Cảnh báo", "Chương trình đang chạy. Bạn có chắc muốn dừng và thoát?"):
                return
        self.stop_flag = True
        self.client.close_browser()
        self.destroy()


if __name__ == "__main__":
    app = App()
    app.mainloop()
