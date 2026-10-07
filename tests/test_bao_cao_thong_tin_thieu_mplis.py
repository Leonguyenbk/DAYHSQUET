# -*- coding: utf-8 -*-

import json
from pathlib import Path

from openpyxl import load_workbook

from bao_cao_thong_tin_thieu_mplis import (
    build_change_history_payload,
    build_list_payload,
    build_report_row,
    extract_latest_non_dh_change,
    extract_missing_information,
    extract_giay_chung_nhan,
    extract_registration_time,
    format_registration_date_ddmmyyyy,
    parse_ma_xa_values,
    registration_responses_by_id,
    session_from_token_cookie,
    write_report,
)


def test_payload_lich_su_bien_dong_dung_network_mplis():
    assert build_change_history_payload(12882362) == {
        "tinhHinhDangKyId": 12882362,
        "soDoi": 0,
    }


def test_bien_dong_cuoi_bo_dh_va_chon_thoi_diem_moi_nhat():
    response = {
        "success": True,
        "Value": {"LichSuBienDongs": [
            {"ThongTinBienDong": {
                "maBienDong": "DH",
                "maLoaiBienDong": "DH",
                "thoiDiemBienDong": "/Date(1761551630473)/",
                "lyDoBienDong": "Điều chỉnh địa giới",
            }},
            {"ThongTinBienDong": {
                "bienDongId": 3530627,
                "loaiBienDongId": 4,
                "maBienDong": "422552.TC.001",
                "maLoaiBienDong": "TC",
                "thoiDiemBienDong": "/Date(1726135684453)/",
                "lyDoBienDong": "Thế chấp",
                "thongTinBienDong": "Nội dung thế chấp do MPLIS trả về",
            }},
            {"ThongTinBienDong": {
                "maBienDong": "422552.TA.003",
                "maLoaiBienDong": "TA",
                "thoiDiemBienDong": "/Date(1587005758000)/",
                "lyDoBienDong": None,
            }},
        ]},
    }

    result = extract_latest_non_dh_change(response)

    assert result is not None
    assert result.ma_bien_dong == "422552.TC.001"
    assert result.ma_loai_bien_dong == "TC"
    assert result.loai_bien_dong_id == 4
    assert result.ly_do_bien_dong == "Thế chấp"
    assert result.thong_tin_bien_dong == "Nội dung thế chấp do MPLIS trả về"
    assert result.thoi_diem_bien_dong == "/Date(1726135684453)/"
    assert result.source_path == "Value.LichSuBienDongs[1].ThongTinBienDong"


def test_bien_dong_khong_fallback_sang_ngay_khac():
    response = {
        "Value": {"LichSuBienDongs": [{"ThongTinBienDong": {
            "maBienDong": "422552.TC.001",
            "maLoaiBienDong": "TC",
            "thoiDiemBienDong": None,
            "ngayHopDong": "/Date(1726074000000)/",
            "CreatedDate": "2024-09-12T00:00:00Z",
            "ModifiedDate": "2024-09-13T00:00:00Z",
        }}]},
    }

    assert extract_latest_non_dh_change(response) is None


def test_bao_cao_co_ngay_va_loai_bien_dong_cuoi_tu_nguon_mplis():
    history = {
        "Value": {"LichSuBienDongs": [
            {"ThongTinBienDong": {
                "maBienDong": "DH",
                "thoiDiemBienDong": "/Date(1761551630473)/",
            }},
            {"ThongTinBienDong": {
                "maBienDong": "422552.TC.001",
                "maLoaiBienDong": "TC",
                "loaiBienDongId": 4,
                "thoiDiemBienDong": "/Date(1726135684453)/",
                "lyDoBienDong": "Thế chấp",
                "thongTinBienDong": "Chi tiết trực tiếp từ MPLIS",
            }},
        ]},
    }
    row = build_report_row(
        {"tinhHinhDangKyId": 12882362},
        ma_xa_requested="22147",
        detail_response=None,
        change_history_response=history,
    )

    assert row["ma_bien_dong_cuoi"] == "422552.TC.001"
    assert row["ma_loai_bien_dong_cuoi"] == "TC"
    assert row["loai_bien_dong_id_cuoi"] == "4"
    assert row["thoi_diem_bien_dong_cuoi"] == "/Date(1726135684453)/"
    assert row["ngay_bien_dong_cuoi_ddmmyyyy"] == "12/09/2024"
    assert row["ly_do_bien_dong_cuoi"] == "Thế chấp"
    assert row["thong_tin_bien_dong_cuoi"] == "Chi tiết trực tiếp từ MPLIS"


def test_tat_lich_su_hien_dung_trang_thai_va_de_trong_du_lieu():
    row = build_report_row(
        {"tinhHinhDangKyId": 12882362},
        ma_xa_requested="22147",
        detail_response=None,
        change_history_enabled=False,
    )

    assert row["trang_thai_lich_su_bien_dong"] == "Đã tắt"
    assert row["ma_bien_dong_cuoi"] == ""
    assert row["ngay_bien_dong_cuoi_ddmmyyyy"] == ""


def test_payload_bat_dau_tu_ma_xa_va_dung_phan_loai_7():
    payload = build_list_payload("22147", start=0, length=10, tinh_id=66)

    assert payload["traCuu[xaId]"] == "22147"
    assert payload["traCuu[soHieuToBanDo]"] == ""
    assert payload["traCuu[soThuTuThua]"] == ""
    assert payload["traCuu[tinhHinhDangKyId]"] == ""
    assert payload["traCuu[phanLoai]"] == "7"
    assert payload["traCuu[type]"] == "-1"
    assert payload["traCuu[loaiTaiSan]"] == "-1"
    assert payload["traCuu[loaiChu]"] == "-1"
    assert payload["start"] == "0"
    assert payload["length"] == "10"


def test_nhan_nhieu_ma_xa_va_bo_trung():
    assert parse_ma_xa_values("22147; 24163,22147\n25001") == [
        "22147", "24163", "25001"
    ]


def test_thong_tin_thieu_giu_nguyen_chuoi_mplis():
    source = "Thiếu thông tin đăng ký; thiếu chủ sử dụng"
    result = extract_missing_information({
        "chiTietThongTinThieu": source,
        "errorMessages": ["không được ưu tiên khi đã có mô tả trực tiếp"],
    })

    assert result is not None
    assert result.value == source
    assert result.source_field == "chiTietThongTinThieu"


def test_khong_suy_luan_thieu_tu_truong_rong():
    item = {
        "soHieuToBanDo": "12",
        "soThuTuThua": "34",
        "thongTinChu": "",
        "ListHoSoQuet": [],
    }

    assert extract_missing_information(item) is None


def test_mang_loi_duoc_giu_nguyen_duoi_dang_json_khong_tu_ghep_noi_dung():
    values = ["Thiếu hồ sơ quét", "Thiếu chủ sử dụng"]
    row = build_report_row(
        {"errorMessages": values, "tinhHinhDangKyId": 123},
        ma_xa_requested="24163",
        detail_response=None,
    )

    assert json.loads(row["thong_tin_thieu"]) == values
    assert row["thong_tin_thieu_source"] == "errorMessages"


def test_thoi_diem_dang_ky_lay_dung_field_va_giu_ca_gio():
    response = {
        "success": True,
        "value": [{
            "thongTinDangKy": {
                "TinhHinhDangKy": {
                    "CreatedDate": "2026-06-26T04:10:29.986Z",
                    "ModifiedDate": "2026-06-27T05:11:30.123Z",
                    "thoiDiemDangKyLanDau": "2016-01-01T00:00:00.000Z",
                    "thoiDiemDangKy": "2017-04-03T17:25:31.456Z",
                }
            }
        }],
    }

    result = extract_registration_time(response)

    assert result is not None
    assert result.value == "2017-04-03T17:25:31.456Z"
    assert result.source_field.endswith("TinhHinhDangKy.thoiDiemDangKy")


def test_ngay_dang_ky_ddmmyyyy_duoc_doi_sang_gio_viet_nam():
    assert format_registration_date_ddmmyyyy(
        "2017-04-03T17:25:31.456Z"
    ) == "04/04/2017"


def test_ngay_dang_ky_khong_tu_dung_ngay_khac_khi_gia_tri_rong():
    assert format_registration_date_ddmmyyyy(None) == ""
    assert format_registration_date_ddmmyyyy("") == ""


def test_thoi_diem_dang_ky_khong_fallback_created_modified_hoac_lan_dau():
    response = {
        "success": True,
        "value": [{
            "TinhHinhDangKy": {
                "CreatedDate": "2026-06-26T04:10:29.986Z",
                "ModifiedDate": "2026-06-27T05:11:30.123Z",
                "thoiDiemDangKyLanDau": "2016-01-01T00:00:00.000Z",
            }
        }],
    }

    assert extract_registration_time(response) is None


def test_thoi_diem_dang_ky_null_duoc_giu_la_null_khong_fallback():
    response = {
        "success": True,
        "value": [{
            "TinhHinhDangKy": {
                "CreatedDate": "2026-06-26T04:10:29.986Z",
                "thoiDiemDangKy": None,
            }
        }],
    }

    result = extract_registration_time(response)
    assert result is not None
    assert result.value is None


def test_bao_cao_co_truong_nguon_va_ids_de_doi_chieu():
    item = {
        "_id": "phan-loai-1",
        "xaId": 24163,
        "soHieuToBanDo": "12",
        "soThuTuThua": "34",
        "thongTinChu": "Dương Thanh Hãnh",
        "phanLoai": "Nhóm 1",
        "trangThaiLamSach": "Chưa làm sạch",
        "thuaDatId": 456,
        "tinhHinhDangKyId": 123,
        "chiTietThongTinThieu": "Thiếu hồ sơ quét",
    }
    detail = {
        "success": True,
        "value": [{
            "thongTinDangKy": {
                "TinhHinhDangKy": {
                    "tinhHinhDangKyId": 123,
                    "thoiDiemDangKy": "2020-04-15T17:00:00.000Z",
                },
                "ListGiayChungNhan": [{
                    "giayChungNhanId": 987,
                    "soPhatHanh": "AN 675030",
                    "soVaoSo": "H01973",
                    "maVach": "MV-01",
                    "ngayVaoSo": "2008-10-28T17:00:00.000Z",
                }],
            }
        }],
    }

    row = build_report_row(
        item,
        ma_xa_requested="24163",
        detail_response=detail,
    )

    assert row["thong_tin_thieu"] == "Thiếu hồ sơ quét"
    assert row["thong_tin_thieu_source"] == "chiTietThongTinThieu"
    assert row["thong_tin_chu"] == "Dương Thanh Hãnh"
    assert row["thong_tin_chu_source"] == "thongTinChu"
    assert row["thoi_diem_dang_ky"] == "2020-04-15T17:00:00.000Z"
    assert row["ngay_dang_ky_ddmmyyyy"] == "16/04/2020"
    assert row["gcn_id"] == "987"
    assert row["gcn_so_phat_hanh"] == "AN 675030"
    assert row["gcn_so_vao_so"] == "H01973"
    assert row["gcn_ma_vach"] == "MV-01"
    assert row["phan_loai_2959"] == "Nhóm 1"
    assert row["trang_thai_lam_sach"] == "Chưa làm sạch"
    assert json.loads(row["cac_id_nguon"])["tinhHinhDangKyId"] == 123


def test_excel_giu_timestamp_va_id_o_dang_text():
    row = build_report_row(
        {
            "xaId": 24163,
            "soHieuToBanDo": "001",
            "soThuTuThua": "00034",
            "tinhHinhDangKyId": 123,
            "chiTietThongTinThieu": "Thiếu hồ sơ quét",
        },
        ma_xa_requested="24163",
        detail_response={
            "success": True,
            "value": [{
                "TinhHinhDangKy": {
                    "thoiDiemDangKy": "2017-04-03T17:25:31.456Z",
                }
            }],
        },
    )
    output = Path("test_bao_cao_thong_tin_thieu_mplis_output.xlsx")
    try:
        write_report([row], output)

        workbook = load_workbook(output, data_only=False)
        sheet = workbook["KetQua"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        assert sheet.cell(2, headers["Số tờ"]).value == "001"
        assert sheet.cell(2, headers["Số thửa"]).value == "00034"
        assert sheet.cell(2, headers["Thời điểm đăng ký"]).value == "2017-04-03T17:25:31.456Z"
        assert sheet.cell(2, headers["Thời điểm đăng ký"]).number_format == "@"
        assert sheet.cell(2, headers["Ngày đăng ký (DD/MM/YYYY)"]).value == "04/04/2017"
    finally:
        if output.exists():
            output.unlink()


def test_tao_session_tu_token_cookie_copy_nguyen_header():
    session = session_from_token_cookie(
        "__RequestVerificationToken: token-abc",
        "Cookie: ASP.NET_SessionId=session-123; auth=value=co-dau-bang",
    )

    assert session.headers["__requestverificationtoken"] == "token-abc"
    assert session.cookies.get("ASP.NET_SessionId") == "session-123"
    assert session.cookies.get("auth") == "value=co-dau-bang"


def test_token_cookie_phai_co_du_noi_dung():
    try:
        session_from_token_cookie("token-abc", "")
    except Exception as exc:
        assert "Cookie" in str(exc)
    else:
        raise AssertionError("Phải báo lỗi khi thiếu Cookie")


def test_tach_response_chi_tiet_theo_lo_id():
    response = {
        "success": True,
        "value": [
            {"thongTinDangKy": {"TinhHinhDangKy": {
                "tinhHinhDangKyId": 101,
                "thoiDiemDangKy": "2020-01-01T01:02:03.000Z",
            }}},
            {"thongTinDangKy": {"TinhHinhDangKy": {
                "tinhHinhDangKyId": 102,
                "thoiDiemDangKy": "2020-01-02T04:05:06.000Z",
            }}},
        ],
    }

    mapped = registration_responses_by_id(response)

    assert set(mapped) == {"101", "102"}
    assert extract_registration_time(mapped["102"]).value == "2020-01-02T04:05:06.000Z"


def test_giu_day_du_nhieu_gcn_theo_dung_thu_tu():
    response = {
        "value": [{"thongTinDangKy": {
            "TinhHinhDangKy": {"tinhHinhDangKyId": 101},
            "ListGiayChungNhan": [
                {"giayChungNhanId": 1, "soPhatHanh": "AA 000001"},
                {"giayChungNhanId": 2, "soPhatHanh": "BB 000002"},
            ],
        }}]
    }

    gcn = extract_giay_chung_nhan(response)

    assert json.loads(gcn["giayChungNhanId"]) == [1, 2]
    assert json.loads(gcn["soPhatHanh"]) == ["AA 000001", "BB 000002"]
    assert len(json.loads(gcn["raw"])) == 2
