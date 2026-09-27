# -*- coding: utf-8 -*-
"""
Đẩy lại file hồ sơ quét (đã ký số) lên MPLIS, thay thế ĐÚNG vị trí file cũ.

QUAN TRỌNG: UpdateHoSoQuetExistFile coi danh sách infoHoSoQuet_1..count gửi lên là TOÀN BỘ
thành phần của HSQ. Gửi count=1 sẽ làm HSQ chỉ còn đúng 1 file (đã xảy ra với đơn 12830794).
→ Mỗi HSQ gửi 1 request kèm ĐỦ mọi thành phần hiện có:
    - thành phần giữ nguyên: node hiện tại, files=null (như api_update_hosoquet_metadata_only)
    - thành phần thay file : node hiện tại, nodeId="" + fileHoSoQuet_i = file ký số

Đầu vào: Excel kết quả do tai_hsq_don.py tạo ra (mỗi dòng 1 file HSQ).
Với mỗi dòng "Đã tải" chưa "Đã đẩy":
  1. Lấy lại chi tiết đơn → tìm HSQ theo hoSoQuetId, thành phần theo thanhPhanHoSoQuetId (dự phòng NId).
  2. nodeId trên MPLIS phải còn trùng với lúc tải (tránh đè file người khác vừa đổi / đẩy 2 lần).
  3. Tìm file ký số: cột file_ky_so nếu có, không thì cạnh file_tai (cùng tên hoặc *_signed...).
  4. Gửi cả HSQ (xem trên). chay_thu=True thì chỉ in kế hoạch.
  5. Kiểm tra lại: số thành phần không giảm, thành phần giữ nguyên còn nguyên nodeId,
     thành phần thay có nodeId mới. Số thành phần giảm → DỪNG TOÀN BỘ.
"""

from __future__ import annotations

import contextlib
import copy
import json
import os
import re
import time
import uuid
from typing import Any, Callable

import cap_nhat_don_hsq as core
import tai_hsq_don as thd

log = thd.log

HAU_TO_KY_SO = ["", "_signed", "-signed", "_ky", "_kyso", "_ky_so", "_daky", "_da_ky", " (signed)", "_sign"]


# ============================ FILE KÝ SỐ ============================

def pdf_da_ky_so(path: str) -> bool:
    """PDF có chữ ký số khi chứa từ điển chữ ký (/ByteRange + /Sig hoặc filter PKCS7/CAdES)."""
    try:
        with open(path, "rb") as fp:
            data = fp.read()
    except OSError:
        return False
    if b"/ByteRange" not in data:
        return False
    return any(k in data for k in (b"/Sig", b"adbe.pkcs7", b"ETSI.CAdES", b"ETSI.RFC3161"))


def tim_file_ky_so(row: dict[str, Any], thu_muc_ky_so: str = "") -> str | None:
    """
    Ưu tiên cột file_ky_so; không có thì tìm cạnh file_tai (hoặc trong thu_muc_ky_so/<madon>/):
    cùng tên (ký đè), tên + hậu tố (_signed...), hoặc file PDF cùng tiền tố "NN_".
    Trong các ứng viên, file đã có chữ ký số được ưu tiên (bản gốc chưa ký có thể vẫn nằm cạnh).
    """
    chi_dinh = str(row.get("file_ky_so") or "").strip()
    if chi_dinh:
        return chi_dinh if os.path.isfile(chi_dinh) else None

    file_tai = str(row.get("file_tai") or "").strip()
    if not file_tai:
        return None
    stem = os.path.splitext(os.path.basename(file_tai))[0]
    tien_to = re.match(r"^(\d+_)", stem)

    # (thư mục, có phải thư mục riêng của đơn không) — chỉ khớp tiền tố "NN_" trong thư mục riêng,
    # vì thư mục phẳng chứa nhiều đơn thì "01_" sẽ trùng giữa các đơn.
    thu_muc_ds: list[tuple[str, bool]] = []
    if thu_muc_ky_so:
        thu_muc_ds += [(os.path.join(thu_muc_ky_so, str(row.get("madon") or "")), True), (thu_muc_ky_so, False)]
    thu_muc_ds.append((os.path.dirname(file_tai), True))

    ung_vien: list[str] = []
    for thu_muc, rieng_don in thu_muc_ds:
        if not os.path.isdir(thu_muc):
            continue
        for hau_to in HAU_TO_KY_SO:
            p = os.path.join(thu_muc, stem + hau_to + ".pdf")
            if os.path.isfile(p) and p not in ung_vien:
                ung_vien.append(p)
        if tien_to and rieng_don:
            for f in sorted(os.listdir(thu_muc)):
                p = os.path.join(thu_muc, f)
                if f.startswith(tien_to.group(1)) and f.lower().endswith(".pdf") and p not in ung_vien:
                    ung_vien.append(p)

    return next((p for p in ung_vien if pdf_da_ky_so(p)), ung_vien[0] if ung_vien else None)


# ============================ API ĐẨY THAY THẾ ============================

def tim_hoso_va_node(detail: dict[str, Any], row: dict[str, Any]) -> tuple[dict | None, dict | None]:
    hsq_id = core.safe_int(row.get("hosoquetid"))
    tp_id = core.safe_int(row.get("thanhphanhosoquetid"))
    tp_nid = str(row.get("thanhphanhosoquetnid") or "").strip()

    for hoso in detail.get("ListHoSoQuet") or []:
        if not isinstance(hoso, dict) or core.safe_int(hoso.get("hoSoQuetId")) != hsq_id:
            continue
        wrapper = hoso.get("ListFileHoSoQuet") or {}
        files = wrapper.get("ListFileHoSoQuet") if isinstance(wrapper, dict) else wrapper
        for node in files or []:
            if not isinstance(node, dict):
                continue
            if tp_id and core.safe_int(node.get("thanhPhanHoSoQuetId")) == tp_id:
                return hoso, node
            if tp_nid and node.get("thanhPhanHoSoQuetNId") == tp_nid:
                return hoso, node
        return hoso, None
    return None, None


def ho_so_quet_payload(hoso: dict[str, Any]) -> dict[str, Any]:
    """Khối hoSoQuet cho HSQ đã có (cùng cấu trúc cap_nhat_don_hsq.api_update_hosoquet_exist_file)."""
    hsq_id = core.safe_int(hoso.get("hoSoQuetId") or hoso.get("Title"))
    return {
        "hoSoQuetId": hsq_id,
        "thongTinHoSoId": core.safe_int(hoso.get("thongTinHoSoId")),
        "tinhHinhDangKyId": core.safe_int(hoso.get("tinhHinhDangKyId")),
        "bienDongId": core.safe_int(hoso.get("bienDongId")),
        "xaId": core.safe_int(hoso.get("xaId")),
        "CreatedDate": core.dotnet_date_to_iso(hoso.get("CreatedDate")) or core.now_iso_z(),
        "ModifiedDate": core.now_iso_z(),
        "Id": str(hoso.get("Id") or uuid.uuid4()),
        "Title": str(hoso.get("Title") or hsq_id),
        "Name": hoso.get("Name"),
        "Path": hoso.get("Path"),
        "ParentPath": hoso.get("ParentPath"),
        "_id": 1,
        "TuiHoSo": None,
        "tuiHoSoId": core.safe_int(hoso.get("tuiHoSoId"), 0),
    }


def cac_node(hoso: dict[str, Any]) -> list[dict[str, Any]]:
    wrapper = hoso.get("ListFileHoSoQuet") or {}
    files = wrapper.get("ListFileHoSoQuet") if isinstance(wrapper, dict) else wrapper
    return [n for n in files or [] if isinstance(n, dict)]


def info_giu_nguyen(node: dict[str, Any]) -> dict[str, Any]:
    """Thành phần không đổi: gửi lại nguyên node, files=null → server giữ file cũ."""
    n = copy.deepcopy(node)
    n["files"] = None
    return n


def info_thay_file(node: dict[str, Any], hsq_id: int, da_ky_so: bool) -> dict[str, Any]:
    """Thành phần thay file: nguyên metadata node, bỏ nodeId để server gắn file mới."""
    n = copy.deepcopy(node)
    n.update({
        "nodeId": "",
        "deleteId": None,
        "isOldFile": False,
        "hoSoQuetId": hsq_id,
        "daKySo": bool(node.get("daKySo")) or da_ky_so,
        "files": None,
        "__id": str(uuid.uuid4()),
    })
    return n


def gui_ca_hsq(
    client: core.MplisClient,
    hoso: dict[str, Any],
    thay: dict[int, tuple[str, bool]],
) -> None:
    """
    1 request cho cả HSQ. thay: {vị trí (0-based) trong cac_node(hoso): (file ký số, đã ký?)}.
    Thành phần không có trong `thay` được gửi lại nguyên trạng.
    """
    session = client._require_session()
    nodes = cac_node(hoso)
    hsq_id = core.safe_int(hoso.get("hoSoQuetId") or hoso.get("Title"))

    data = {
        "hoSoQuet": json.dumps(ho_so_quet_payload(hoso), ensure_ascii=False),
        "count": str(len(nodes)),
        "isLuuKhoHoSoQuet": "false",
    }
    headers = dict(session.headers)
    headers.pop("Content-Type", None)

    with contextlib.ExitStack() as stack:
        files = {}
        for i, node in enumerate(nodes):
            if i in thay:
                path, da_ky = thay[i]
                data[f"infoHoSoQuet_{i + 1}"] = json.dumps(info_thay_file(node, hsq_id, da_ky), ensure_ascii=False)
                files[f"fileHoSoQuet_{i + 1}"] = (os.path.basename(path), stack.enter_context(open(path, "rb")),
                                                  "application/pdf")
            else:
                data[f"infoHoSoQuet_{i + 1}"] = json.dumps(info_giu_nguyen(node), ensure_ascii=False)
        response = session.post(core.URL_UPDATE_HOSOQUET, data=data, files=files, headers=headers,
                                timeout=core.TIMEOUT, allow_redirects=False)

    result = client._response_json(response, "UpdateHoSoQuetExistFile")
    if not isinstance(result, dict) or not result.get("success"):
        raise RuntimeError(f"Đẩy HSQ không thành công: {core.rut_gon_text(result)}")


# ============================ EXCEL ============================

def doc_excel(path: str):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    thieu = [h for h in ("madon", "hosoquetid", "thanhphanhosoquetid", "nodeid", "file_tai", "trang_thai_tai")
             if h not in headers]
    if thieu:
        wb.close()
        raise ValueError(f"Không phải Excel kết quả tải HSQ (thiếu cột: {', '.join(thieu)}).")
    # Excel cũ thiếu cột đẩy → thêm vào cuối.
    for h in ("file_ky_so", "trang_thai_day", "ghi_chu_day", "thoi_gian_day", "nodeid_moi"):
        if h not in headers:
            headers.append(h)
            ws.cell(row=1, column=len(headers), value=h)
    col = {h: i + 1 for i, h in enumerate(headers)}
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(row=r, column=c).value for h, c in col.items()}
        if row.get("madon") in (None, ""):
            continue
        row["_excel_row"] = r
        rows.append(row)
    return wb, ws, col, rows


def ghi_dong(ws, col: dict[str, int], row: dict[str, Any], **values) -> None:
    from openpyxl.styles import PatternFill

    mau = {"Đã đẩy": "C6EFCE", "Lỗi": "F4CCCC", "Bỏ qua": "FFF2CC"}
    for k, v in values.items():
        row[k] = v
        ws.cell(row=row["_excel_row"], column=col[k], value=v)
    tt = str(values.get("trang_thai_day") or "")
    fill = next((m for key, m in mau.items() if tt.startswith(key)), None)
    if fill:
        ws.cell(row=row["_excel_row"], column=col["trang_thai_day"]).fill = PatternFill("solid", fgColor=fill)


def luu_excel(wb, path: str) -> str:
    """Lưu Excel; nếu file đang mở trong Excel (bị khóa) thì lưu sang file _capnhat."""
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = os.path.splitext(path)[0] + "_capnhat.xlsx"
        wb.save(alt)
        log(f"⚠ Excel đang mở, đã lưu kết quả sang: {alt}")
        return alt


# ============================ LUỒNG CHÍNH ============================

class DungKhanCap(RuntimeError):
    """MPLIS bị giảm số thành phần sau khi đẩy → dừng toàn bộ để không lan sang đơn khác."""


def _chuan_bi_dong(r: dict, node: dict | None, thu_muc_ky_so: str, chi_file_da_ky: bool):
    """Kiểm tra 1 dòng. Trả về (file, đã ký) nếu đẩy được; ngược lại (None, (trạng thái, ghi chú, extra))."""
    if node is None:
        return None, ("Lỗi", f"Không còn thành phần {r.get('thanhphanhosoquetid')} trong HSQ {r.get('hosoquetid')}", {})
    if str(node.get("nodeId") or "") != str(r.get("nodeid") or ""):
        return None, ("Bỏ qua", "File trên MPLIS đã khác lúc tải (nodeId đổi) → không đè.",
                      {"nodeid_moi": node.get("nodeId")})
    file_ky = tim_file_ky_so(r, thu_muc_ky_so)
    if not file_ky:
        return None, ("Lỗi", "Không tìm thấy file ký số (điền cột file_ky_so hoặc để cạnh file_tai)", {})
    da_ky = pdf_da_ky_so(file_ky)
    if chi_file_da_ky and not da_ky:
        return None, ("Bỏ qua", "File chưa có chữ ký số.", {"file_ky_so": file_ky})
    return (file_ky, da_ky), None


def day_tu_excel(
    client: core.MplisClient,
    excel_path: str,
    thu_muc_ky_so: str = "",
    chi_file_da_ky: bool = True,
    chay_thu: bool = True,
    nen_dung: Callable[[], bool] = lambda: False,
    bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None,
) -> dict[str, int]:
    wb, ws, col, rows = doc_excel(excel_path)
    thong_ke = {"da_day": 0, "loi": 0, "bo_qua": 0, "chay_thu": 0}
    try:
        can_day = [
            r for r in rows
            if str(r.get("trang_thai_tai") or "") == "Đã tải"
            and not str(r.get("trang_thai_day") or "").startswith("Đã đẩy")
        ]
        da_day_truoc = sum(1 for r in rows if str(r.get("trang_thai_day") or "").startswith("Đã đẩy"))
        log(f"{'[CHẠY THỬ] ' if chay_thu else ''}Excel: {len(rows)} dòng | cần đẩy {len(can_day)} file"
            f" | đã đẩy trước đó {da_day_truoc}")

        theo_don: dict[str, list[dict]] = {}
        for r in can_day:
            theo_don.setdefault(str(r["madon"]), []).append(r)

        xong = 0
        for n, (madon, ds_don) in enumerate(theo_don.items(), 1):
            if nen_dung():
                log("⏹ Đã dừng theo yêu cầu.")
                break
            bao_tien_do(xong, len(can_day), f"Đơn {madon} ({n}/{len(theo_don)})")
            log(f"--- [{n}/{len(theo_don)}] Đơn {madon}: {len(ds_don)} file ---")
            xong += len(ds_don)

            try:
                detail = client.lay_chi_tiet_don(core.safe_int(madon))
                if detail is None:
                    raise RuntimeError("Không lấy được chi tiết đơn")
            except Exception as exc:
                for r in ds_don:
                    ghi_dong(ws, col, r, trang_thai_day="Lỗi", ghi_chu_day=core.rut_gon_text(exc, 300))
                thong_ke["loi"] += len(ds_don)
                log(f"❌ Đơn {madon}: {core.rut_gon_text(exc, 300)}")
                luu_excel(wb, excel_path)
                continue

            theo_hsq: dict[int, list[dict]] = {}
            for r in ds_don:
                theo_hsq.setdefault(core.safe_int(r.get("hosoquetid")), []).append(r)

            for hsq_id, ds in theo_hsq.items():
                hoso = next((h for h in detail.get("ListHoSoQuet") or []
                             if isinstance(h, dict) and core.safe_int(h.get("hoSoQuetId")) == hsq_id), None)
                if hoso is None:
                    for r in ds:
                        ghi_dong(ws, col, r, trang_thai_day="Lỗi", ghi_chu_day=f"Không còn HSQ {hsq_id} trong đơn")
                    thong_ke["loi"] += len(ds)
                    log(f"   ❌ Không còn HSQ {hsq_id} trong đơn")
                    continue

                nodes = cac_node(hoso)
                thay: dict[int, tuple[str, bool]] = {}
                dong_thay: dict[int, dict] = {}
                for r in ds:
                    _, node = tim_hoso_va_node(detail, r)
                    ok, loi = _chuan_bi_dong(r, node, thu_muc_ky_so, chi_file_da_ky)
                    if loi:
                        trang_thai, ghi_chu, extra = loi
                        ghi_dong(ws, col, r, trang_thai_day=trang_thai, ghi_chu_day=ghi_chu, **extra)
                        thong_ke["loi" if trang_thai == "Lỗi" else "bo_qua"] += 1
                        log(f"   {'❌' if trang_thai == 'Lỗi' else '⚠'} {r.get('mota')}: {ghi_chu}")
                        continue
                    vi_tri = next(i for i, x in enumerate(nodes) if x is node)
                    thay[vi_tri] = ok
                    dong_thay[vi_tri] = r

                if not thay:
                    continue

                log(f"   HSQ {hsq_id}: gửi {len(nodes)} thành phần, thay {len(thay)} file")
                for i, node in enumerate(nodes):
                    if i in thay:
                        log(f"     {i + 1}. THAY  {node.get('moTa')} ← {os.path.basename(thay[i][0])}")
                    else:
                        log(f"     {i + 1}. giữ   {node.get('moTa')}")

                if chay_thu:
                    for i, r in dong_thay.items():
                        ghi_dong(ws, col, r, ghi_chu_day="[Chạy thử] sẽ thay file này", file_ky_so=thay[i][0])
                    thong_ke["chay_thu"] += len(thay)
                    continue

                node_ids_truoc = {i: str(x.get("nodeId") or "") for i, x in enumerate(nodes)}
                try:
                    gui_ca_hsq(client, hoso, thay)
                except Exception as exc:
                    for r in dong_thay.values():
                        ghi_dong(ws, col, r, trang_thai_day="Lỗi", ghi_chu_day=core.rut_gon_text(exc, 300))
                    thong_ke["loi"] += len(thay)
                    log(f"   ❌ {core.rut_gon_text(exc, 300)}")
                    luu_excel(wb, excel_path)
                    continue

                # ---- xác nhận
                detail_moi = client.lay_chi_tiet_don(core.safe_int(madon)) or {}
                hoso_moi = next((h for h in detail_moi.get("ListHoSoQuet") or []
                                 if isinstance(h, dict) and core.safe_int(h.get("hoSoQuetId")) == hsq_id), None)
                nodes_moi = cac_node(hoso_moi) if hoso_moi else []
                ids_moi = {str(x.get("nodeId") or "") for x in nodes_moi}
                if len(nodes_moi) < len(nodes):
                    for i, r in dong_thay.items():
                        ghi_dong(ws, col, r, trang_thai_day="Lỗi",
                                 ghi_chu_day=f"MPLIS còn {len(nodes_moi)}/{len(nodes)} thành phần sau khi đẩy!")
                    luu_excel(wb, excel_path)
                    raise DungKhanCap(
                        f"HSQ {hsq_id} (đơn {madon}) giảm từ {len(nodes)} còn {len(nodes_moi)} thành phần "
                        "sau khi đẩy → DỪNG TOÀN BỘ. Gửi log cho người viết tool."
                    )
                mat = [nodes[i].get("moTa") for i in node_ids_truoc if i not in thay and node_ids_truoc[i] not in ids_moi]
                if mat:
                    log(f"   ⚠ Thành phần giữ nguyên nhưng nodeId đổi: {mat}")
                for i, r in dong_thay.items():
                    tp = core.safe_int(nodes[i].get("thanhPhanHoSoQuetId"))
                    node_moi = next((x for x in nodes_moi if core.safe_int(x.get("thanhPhanHoSoQuetId")) == tp), None)
                    nid_moi = str((node_moi or {}).get("nodeId") or "")
                    if nid_moi and nid_moi != node_ids_truoc[i]:
                        ghi_dong(ws, col, r, trang_thai_day="Đã đẩy", file_ky_so=thay[i][0], ghi_chu_day="",
                                 nodeid_moi=nid_moi, thoi_gian_day=time.strftime("%d/%m/%Y %H:%M:%S"))
                        thong_ke["da_day"] += 1
                        log(f"   ✅ Đã thay {r.get('mota')}")
                    else:
                        ghi_dong(ws, col, r, trang_thai_day="Lỗi", file_ky_so=thay[i][0],
                                 ghi_chu_day="Server báo OK nhưng thành phần chưa đổi file (kiểm tra trên MPLIS).")
                        thong_ke["loi"] += 1
                        log(f"   ❌ {r.get('mota')}: server báo OK nhưng file chưa đổi")
                luu_excel(wb, excel_path)
                time.sleep(core.REQUEST_DELAY_SECONDS)

        bao_tien_do(len(can_day), len(can_day), "Hoàn tất")
        log("=" * 50)
        if chay_thu:
            log(f"[CHẠY THỬ] sẽ thay {thong_ke['chay_thu']} file | bỏ qua {thong_ke['bo_qua']} | lỗi {thong_ke['loi']}"
                " — chưa gửi gì lên MPLIS.")
        else:
            log(f"✅ Đã đẩy {thong_ke['da_day']} | bỏ qua {thong_ke['bo_qua']} | lỗi {thong_ke['loi']}")
        log(f"Excel kết quả: {luu_excel(wb, excel_path)}")
        return thong_ke
    finally:
        wb.close()
