# -*- coding: utf-8 -*-
"""
Tách GCN theo đơn: lấy GCN hiện hành của đơn trên MPLIS → tải HSQ → AI (Gemini) tìm trang GCN
→ tách mỗi GCN của đơn ra 1 file riêng "<số phát hành>-GCN.pdf".

Luồng mỗi đơn (tinhHinhDangKyId):
  1. GetThongTinDangKy → GCN hiện hành (bỏ isLastest = False): soPhatHanh, soVaoSo.
     Đơn chưa có GCN: thành phần HSQ có chuỗi "DON" (có dấu hay không) trong tên → ghi nhận loại Đơn đăng ký + hồ sơ gốc
     (laGiayToVeNguonGoc) trên MPLIS. Mặc định CHẠY THỬ, --ghi-mplis mới ghi thật.
     Đơn chưa có GCN → dừng ở đây (KHÔNG tải, KHÔNG gọi AI).
  2. Đơn thường có rất nhiều file HSQ → xử lý theo thứ tự ưu tiên: tên file (mô tả) có chữ "GCN" trước
     nhất; rồi thành phần loại GCN / gắn GCN / mô tả có "giấy chứng nhận, sổ đỏ..."; rồi file chưa rõ;
     đơn đăng ký, trích lục, biên bản, CCCD... cuối. Tải từng
     file về <thư mục ra>/HSQ_goc/<madon>/ rồi đọc ngay; mọi GCN đã có bản đúng chủ → DỪNG, các file còn
     lại không tải, không gọi AI (--quet-het để xét hết).
  3. Mỗi PDF (render bằng lõi tach_gcn, mở được cả PDF lỗi cây trang):
       a. Lọc rẻ: ảnh nhỏ, media resolution thấp, prompt phân loại của tach_gcn → trang nghi là GCN
          (+ trang kẹp giữa, + tối đa 3 trang sau mỗi cụm để không lọt trang sơ đồ / trang 4).
       b. Chỉ các trang đó gửi ảnh nét KÈM danh sách số phát hành + số vào sổ của đơn → AI gán từng
          trang: thuộc GCN nào của đơn / GCN khác ("KHAC") / không phải GCN. Lô sau biết trang cuối lô
          trước thuộc GCN nào (trang sơ đồ / trang biến động không có số vẫn theo đúng GCN).
  4. Cùng 1 GCN có thể có nhiều bản (nhiều file, hoặc trước / sau đính chính trong 1 file) → Gemini đọc chủ
     hiện tại của từng bản (trang 1; đã đính chính / biến động thì lấy chủ mới nhất ở trang 4) và đối chiếu
     với chủ sử dụng trong đơn (ChuSoHuu). Chọn bản đúng chủ (trùng nhiều người hơn → nhiều trang hơn):
       có bản đúng chủ  → <thư mục ra>/GCN/<madon>/<SỐ GCN>-GCN.pdf, các bản khác ghi "Bản trùng";
       không bản nào đúng → <thư mục ra>/GCN_can_kiem_tra/<madon>/ ("Sai chủ" / "Chưa rõ chủ") để kiểm tra tay.
     Trang GCN không thuộc đơn → <thư mục ra>/GCN_khong_khop/<madon>/ để kiểm tra tay.
  5. Excel kết quả: mỗi GCN của đơn 1 dòng + mỗi bản trùng / nhóm không khớp 1 dòng.

Thư mục <thư mục ra>/GCN đúng cấu trúc của DayHSQMoTaGanGCN (day_hsq_mota_gan_gcn.py) → đẩy lên được ngay.

Cách dùng:
  set GEMINI_API_KEY=...
  python tach_gcn_theo_don.py 13456477 12913166
  python tach_gcn_theo_don.py --excel ds_don.xlsx --out ket_qua_tach_gcn
  python tach_gcn_theo_don.py 13456477 --dung-file-da-tai     # đã tải HSQ rồi thì không tải lại
  python tach_gcn_theo_don.py 13456477 --quet-het             # không dừng sớm, xét mọi file HSQ
  python tach_gcn_theo_don.py 13456477 --ghi-mplis            # ghi thật loại Đơn đăng ký + hồ sơ gốc
"""

from __future__ import annotations

import argparse
import copy
import getpass
import os
import re
import sys
import time
from pathlib import Path
from typing import Any, Callable

try:
    import tach_gcn  # noqa: F401
except ImportError:  # chưa pip install -e . → dùng thẳng src/
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "src"))

from tach_gcn.gemini import ASSIGN_OTHER, DEFAULT_MODEL, GeminiAnalyzer, Usage
from tach_gcn.pdfio import (InspectedFile, _save_selected_pages, inspect_pdf, open_pdf_document,
                            render_page_jpeg_checked)
from tach_gcn.report import cache_key, cached_gemini_result, load_cache, save_cache, serialize_gemini_result

import bo_sung_gcn_hsq as bs
import cap_nhat_don_hsq as core
import day_hsq_ky_so as day
import day_hsq_theo_mota as dm
import tai_hsq_don as thd

log = thd.log

THU_MUC_HSQ = "HSQ_goc"
THU_MUC_GCN = "GCN"
THU_MUC_KHONG_KHOP = "GCN_khong_khop"
THU_MUC_CAN_KIEM_TRA = "GCN_can_kiem_tra"
FILE_CACHE = ".tach-gcn-cache.json"
SO_TRANG_MOI_LO = 10  # số ảnh nét mỗi request gán GCN
SO_TRANG_LO_LOC = 16  # số ảnh nhỏ mỗi request lọc trang

EXCEL_HEADERS = ["stt", "madon", "giaychungnhanid", "so_gcn", "so_vao_so", "chu_don", "file_nguon", "trang",
                 "so_ai_doc", "so_vao_so_ai", "ngay_cap_ai", "chu_ai", "cccd_ai", "dia_chi_ai", "file_gcn",
                 "trang_thai", "ghi_chu"]
EXCEL_WIDTHS = {"stt": 6, "madon": 12, "giaychungnhanid": 13, "so_gcn": 14, "so_vao_so": 12, "chu_don": 30,
                "file_nguon": 40, "trang": 12, "so_ai_doc": 16, "so_vao_so_ai": 14, "ngay_cap_ai": 12,
                "chu_ai": 30, "cccd_ai": 26, "dia_chi_ai": 40, "file_gcn": 55, "trang_thai": 13,
                "ghi_chu": 70}
MAU_TRANG_THAI = {"Đã tách": "C6EFCE", "Lỗi": "F4CCCC", "Không thấy": "F4CCCC", "Sai chủ": "F4CCCC",
                  "Chưa rõ chủ": "FFF2CC", "Chưa rõ số": "FFF2CC", "Không khớp": "FFF2CC", "Bản trùng": "EDEDED",
                  "Bỏ qua": "FFF2CC", "Chạy thử": "D9EAF7", "Đã cập nhật": "C6EFCE", "Giữ nguyên": "EDEDED"}


# ============================ TẢI HSQ ============================

TU_GCN = ("gcn", "giaychungnhan", "sodo", "sohong", "bia", "gcnqsdd", "quyensudungdat")
# Mô tả các giấy tờ thường không chứa bản GCN → xử lý cuối (vẫn xử lý nếu chưa đủ GCN).
TU_KHONG_GCN = ("dondangky", "trichluc", "trichdo", "bienban", "tokhai", "thongbao", "cccd", "cancuoc",
                "chungminh", "hokhau", "phieu", "quyetdinh", "bandovitri", "giaynopthue", "bienlai", "uyquyen")


KHOA_TEN_FILE = ("moTa", "tenGiayTo", "tenFile", "fileName", "tenTapTin")
_CHU_GCN = re.compile(r"(?<![a-z])gcn(?![a-z])")  # 'GCN' đứng riêng: '24221_GCN_DH 770214', 'GCN-01'


def uu_tien_file(f: dict[str, Any]) -> int:
    """
    0: tên file (mô tả) có chữ 'GCN'; 1: dấu hiệu GCN khác (thành phần loại GCN / gắn GCN / 'giấy chứng nhận',
    'sổ đỏ'...); 2: chưa rõ (hồ sơ gốc, bản scan...); 3: giấy tờ thường không có GCN.
    """
    ten = " ".join(str(f.get(k) or "") for k in KHOA_TEN_FILE)
    if _CHU_GCN.search(core.remove_accents(ten).lower()):
        return 0
    mo_ta = dm.chuan_hoa(ten)
    if (core.safe_int(f.get("loaiHoSoQuet"), -1) == 1 or f.get("laGiayChungNhan")
            or core.safe_int(f.get("giayChungNhanId")) or any(t in mo_ta for t in TU_GCN)):
        return 1
    if any(t in mo_ta for t in TU_KHONG_GCN):
        return 3
    return 2


def _tai_ve(downloader: thd.Downloader, f: dict[str, Any], vi_tri: int, thu_muc: str,
            dung_file_da_tai: bool) -> str:
    """Tải 1 file HSQ về thu_muc/'NN_<mô tả>.<đuôi>' → đường dẫn (đuôi theo nội dung). Hết phiên → core.HetPhien."""
    mo_ta = f.get("moTa") or f.get("tenGiayTo") or f"file_{vi_tri}"
    ten = f"{vi_tri:02d}_{thd.ten_file_an_toan(mo_ta)}"
    os.makedirs(thu_muc, exist_ok=True)
    co_san = os.path.join(thu_muc, f"{ten}.pdf")
    if dung_file_da_tai and os.path.isfile(co_san):
        return co_san
    data, ten_server = downloader.tai(f)
    duong_dan = os.path.join(thu_muc, f"{ten}{thd.doan_duoi_file(data, ten_server)}")
    with open(duong_dan, "wb") as fp:
        fp.write(data)
    time.sleep(core.REQUEST_DELAY_SECONDS)
    return duong_dan


def _chi_lay_pdf(duong_dan: str) -> str | None:
    if os.path.splitext(duong_dan)[1].lower() != ".pdf":
        log(f"   ⚠ {os.path.basename(duong_dan)} không phải PDF — AI không đọc")
        return None
    return duong_dan


def tai_mot_file(downloader: thd.Downloader, f: dict[str, Any], vi_tri: int, thu_muc: str,
                 dung_file_da_tai: bool) -> str | None:
    """Tải 1 file HSQ → đường dẫn PDF (không phải PDF → None). Lỗi tải → ném lỗi; hết phiên → core.HetPhien."""
    return _chi_lay_pdf(_tai_ve(downloader, f, vi_tri, thu_muc, dung_file_da_tai))


# ============================ TẢI TRƯỚC (song song với lúc chờ AI) ============================

SO_FILE_TAI_TRUOC = 2  # đơn sắp tới: tải trước vài file đầu — đơn dừng sớm thì không phí tải cả HSQ
SO_DON_TAI_TRUOC = 2  # số đơn kế tiếp chuẩn bị trước (đơn chưa có GCN không tải file nào)


class TaiTruoc:
    """
    1 luồng nền tải TUẦN TỰ theo độ ưu tiên: file đơn đang xử lý trước, rồi chi tiết + vài file đầu của đơn kế
    tiếp → mạng chạy trong lúc luồng chính chờ AI. Mọi lệnh tải đi qua luồng nền (Downloader chỉ 1 luồng dùng);
    lệnh MPLIS của cả 2 luồng giữ self.khoa → không gửi chồng nhau trên cùng session.
    Log phát sinh trong luồng nền (thử lại, endpoint...) được giữ lại, in ra khi luồng chính dùng kết quả
    → nhật ký vẫn liền mạch theo từng đơn.
    """

    def __init__(self, client, downloader: thd.Downloader, out_dir: str, dung_file_da_tai: bool) -> None:
        import queue
        import threading

        self.client, self.downloader = client, downloader
        self.out_dir, self.dung_file_da_tai = out_dir, dung_file_da_tai
        self.khoa = threading.RLock()
        self._khoa_viec = threading.Lock()
        self._hang: queue.PriorityQueue = queue.PriorityQueue()
        self._stt = 0
        self._thu_tu_don: dict[int, int] = {}
        self._viec: dict[tuple, Any] = {}  # ("ct", id) | ("f", id, vi_tri) → Future
        self._buf: list[str] = []
        self._log_goc = thd.log
        thd.log = self._log  # Downloader gọi thd.log lúc chạy → bắt được log của luồng nền
        self._luong = threading.Thread(target=self._chay, name="tai_truoc", daemon=True)
        self._luong.start()

    # ---- luồng nền
    def _log(self, msg: str) -> None:
        import threading

        if threading.current_thread() is self._luong:
            self._buf.append(msg)
        else:
            self._log_goc(msg)

    def _chay(self) -> None:
        while True:
            _, _, fut, ham = self._hang.get()
            if fut is None:
                return
            if not fut.set_running_or_notify_cancel():
                continue
            self._buf = []
            try:
                kq = ham()
            except BaseException as exc:  # noqa: BLE001 — trả về luồng chính
                exc.log_nen = self._buf  # type: ignore[attr-defined]
                fut.set_exception(exc)
            else:
                fut.set_result((kq, self._buf))

    def _xep(self, khoa: tuple, uu_tien: tuple, ham: Callable[[], Any]):
        from concurrent.futures import Future

        with self._khoa_viec:
            fut = self._viec.get(khoa)
            if fut is not None and not fut.cancelled():
                return fut
            fut = Future()
            self._viec[khoa] = fut
            self._stt += 1
            self._hang.put((uu_tien, self._stt, fut, ham))
            return fut

    def dang_ky(self, ids: list[int]) -> None:
        """Thứ tự đơn trong loạt = độ ưu tiên tải (đơn trước tải trước)."""
        with self._khoa_viec:
            for id_don in ids:
                self._thu_tu_don.setdefault(id_don, len(self._thu_tu_don))

    def _uu_tien_don(self, id_don: int) -> int:
        with self._khoa_viec:
            return self._thu_tu_don.setdefault(id_don, len(self._thu_tu_don))

    def _ket_qua(self, fut) -> Any:
        try:
            kq, buf = fut.result()
        except BaseException as exc:
            for msg in getattr(exc, "log_nen", []):
                log(msg)
            raise
        for msg in buf:
            log(msg)
        return kq

    def _tai_chi_tiet(self, id_don: int) -> dict | None:
        with self.khoa:
            return self.client.lay_chi_tiet_don(id_don)

    def _tai_file(self, id_don: int, vi_tri: int, f: dict[str, Any]) -> str:
        thu_muc = os.path.join(self.out_dir, THU_MUC_HSQ, str(id_don))
        with self.khoa:
            return _tai_ve(self.downloader, f, vi_tri, thu_muc, self.dung_file_da_tai)

    # ---- luồng chính
    def chi_tiet(self, id_don: int) -> dict | None:
        return self._ket_qua(self._xep(("ct", id_don), (self._uu_tien_don(id_don), 0),
                                       lambda: self._tai_chi_tiet(id_don)))

    def dat_truoc_file(self, id_don: int, ds: list[tuple[int, dict[str, Any]]], bat_dau: int = 0) -> None:
        """Xếp hàng tải các file (vị trí, file); ds là đoạn thứ tự xử lý của đơn bắt đầu từ chỉ số bat_dau."""
        don = self._uu_tien_don(id_don)
        for k, (vi_tri, f) in enumerate(ds, bat_dau):
            self._xep(("f", id_don, vi_tri), (don, 1, k), lambda v=vi_tri, f_=f: self._tai_file(id_don, v, f_))

    def file(self, id_don: int, vi_tri: int, f: dict[str, Any]) -> str:
        self.dat_truoc_file(id_don, [(vi_tri, f)])
        with self._khoa_viec:
            fut = self._viec[("f", id_don, vi_tri)]
        return self._ket_qua(fut)

    def chuan_bi_don(self, id_don: int) -> None:
        """Đơn kế tiếp: tải chi tiết, đơn có GCN thì tải trước vài file đầu (theo uu_tien_file)."""
        don = self._uu_tien_don(id_don)

        def lam() -> dict | None:
            detail = self._tai_chi_tiet(id_don)
            if detail and any(str(g.get("soPhatHanh") or "").strip() for g in bs.gcn_hien_hanh(detail)):
                thu_tu = sorted(enumerate(thd.liet_ke_file_hsq(detail), 1), key=lambda x: uu_tien_file(x[1]))
                self.dat_truoc_file(id_don, thu_tu[:SO_FILE_TAI_TRUOC])
            return detail

        self._xep(("ct", id_don), (don, 0), lam)

    def huy_don(self, id_don: int) -> None:
        """Đơn xong (dừng sớm) → bỏ các file đang xếp hàng chưa tải của đơn đó."""
        with self._khoa_viec:
            for khoa, fut in self._viec.items():
                if khoa[0] == "f" and khoa[1] == id_don:
                    fut.cancel()

    def dong(self) -> None:
        with self._khoa_viec:
            for fut in self._viec.values():
                fut.cancel()
        self._hang.put(((-1,), 0, None, None))
        self._luong.join(timeout=120)
        thd.log = self._log_goc


# ============================ AI GÁN TRANG → GCN ============================

def _trang_text(indexes: list[int]) -> str:
    """[1, 2, 3, 6] (0-based) → '2-4,7'."""
    doan: list[str] = []
    for i in sorted(indexes):
        if doan and doan[-1][1] == i - 1:
            doan[-1] = (doan[-1][0], i)
        else:
            doan.append((i, i))
    return ",".join(str(a + 1) if a == b else f"{a + 1}-{b + 1}" for a, b in doan)


def _canh_dai(page, co_so: int) -> int:
    """Cạnh dài khi render: ảnh khổ ngang (scan trải 2 mặt A3) gấp đôi để mỗi mặt vẫn đủ nét."""
    w, h = float(page.rect.width), float(page.rect.height)
    return co_so * 2 if w > h * 1.2 else co_so


def render_trang(ins: InspectedFile, trang: list[int], co_so: int = 1400) -> list[tuple[int, bytes]]:
    """Như render_pdf_pages của tach_gcn (MuPDF lỗi giải mã ảnh → PDFium) nhưng ảnh khổ ngang nét gấp đôi."""
    doc = open_pdf_document(Path(ins.path), ins.recovered_page_xrefs)
    try:
        return [(i + 1, render_page_jpeg_checked(doc, Path(ins.path), i, _canh_dai(doc.load_page(i), co_so),
                                                 recovered_page_xrefs=ins.recovered_page_xrefs))
                for i in trang]
    finally:
        doc.close()


def loc_trang(ins: InspectedFile, analyzer: GeminiAnalyzer, model: str, cache: dict, usage: Usage) -> list[int]:
    """
    Lượt lọc rẻ: ảnh nhỏ + media resolution thấp, chỉ hỏi "có phải mặt GCN không" (prompt của tach_gcn).
    → trang nghi là GCN (0-based), nới thêm: trang kẹp giữa 2 trang GCN (cách ≤ 2) và tối đa 3 trang sau
    mỗi cụm (trang sơ đồ / trang 4 biến động scan mờ dễ bị lọt — trang 4 cần để đọc chủ sau đính chính).
    """
    model_loc = f"{model}:loc3"
    co_gcn: set[int] = set()
    can_gui: list[int] = []
    for i in range(ins.page_count):
        cached = cached_gemini_result(cache.get(cache_key(ins.content_hash, i + 1, model_loc), {}))
        if cached is None:
            can_gui.append(i)
        elif cached.is_certificate:
            co_gcn.add(i)
    if can_gui:
        doc = open_pdf_document(Path(ins.path), ins.recovered_page_xrefs)
        try:
            for dau in range(0, len(can_gui), SO_TRANG_LO_LOC):
                lo = can_gui[dau:dau + SO_TRANG_LO_LOC]
                anh = [(i + 1, render_page_jpeg_checked(doc, Path(ins.path), i, _canh_dai(doc.load_page(i), 900),
                                                        60, ins.recovered_page_xrefs)) for i in lo]
                kq, dung = analyzer.analyze(anh, low_res=True)
                usage.add(dung)
                for i in lo:
                    cache[cache_key(ins.content_hash, i + 1, model_loc)] = serialize_gemini_result(kq[i + 1])
                    if kq[i + 1].is_certificate:
                        co_gcn.add(i)
        finally:
            doc.close()

    chon = set(co_gcn)
    ds = sorted(co_gcn)
    for a, b in zip(ds, ds[1:]):
        if b - a <= 3:
            chon.update(range(a, b))
    for i in ds:
        if i + 1 not in co_gcn:
            chon.update(range(i + 1, min(i + 4, ins.page_count)))
    return sorted(chon)


def gan_trang(path: str, ds_so: list[tuple[str, str]], analyzer: GeminiAnalyzer, model: str, cache: dict,
              usage: Usage, loc_truoc: bool = True) -> tuple[dict[int, dict[str, str]], InspectedFile]:
    """
    Gửi trang (đã lọc, hoặc mọi trang) của 1 PDF cho Gemini kèm danh sách GCN của đơn → ({chỉ số trang
    0-based: {gcn, serial_seen, reason}}, file đã đọc (số trang, hash, recovered_page_xrefs)).
    Trang không gửi coi như không phải GCN. Cache theo nội dung file + danh sách GCN + model.
    """
    ins = inspect_pdf(path, os.path.basename(path), (), False, False)
    if ins.error:
        raise RuntimeError(ins.error)
    khoa = (f"assign4:{ins.content_hash}:{model}:{'loc' if loc_truoc else 'het'}:"
            + "|".join(f"{s}/{v}" for s, v in ds_so))
    cached = cache.get(khoa)
    if isinstance(cached, dict) and cached.get("so_trang") == ins.page_count:
        return {int(k): v for k, v in cached["trang"].items()}, ins

    ung_vien = loc_trang(ins, analyzer, model, cache, usage) if loc_truoc else list(range(ins.page_count))
    ket_qua: dict[int, dict[str, Any]] = {i: {"gcn": "", "serial_seen": "", "reason": "", "bat_dau": False}
                                          for i in range(ins.page_count)}
    ghi_chu_loc = ("Chỉ gửi các trang nghi là GCN, số i là số trang thật trong file (trang bị bỏ qua không "
                   "phải GCN). " if loc_truoc else "")
    for dau in range(0, len(ung_vien), SO_TRANG_MOI_LO):
        lo = ung_vien[dau:dau + SO_TRANG_MOI_LO]
        anh = render_trang(ins, lo)
        truoc = ghi_chu_loc
        if dau:
            j = ung_vien[dau - 1]
            t = ket_qua[j]
            thuoc = ("thuộc GCN " + t["gcn"] if t["gcn"] and t["gcn"] != ASSIGN_OTHER
                     else "là trang GCN không có trong danh sách" if t["gcn"] else "không phải trang GCN")
            truoc += f"Trang gửi ngay trước lô này (i={j + 1}) {thuoc}."
        gan, dung = analyzer.assign_certificates(anh, ds_so, truoc.strip())
        usage.add(dung)
        for i in lo:
            a = gan[i + 1]
            ket_qua[i] = {"gcn": a.gcn, "serial_seen": a.serial_seen, "reason": a.reason,
                          "bat_dau": a.starts_copy, "kem_theo": a.attached_to}
    cache[khoa] = {"so_trang": ins.page_count, "trang": {str(k): v for k, v in ket_qua.items()}}
    return ket_qua, ins


def kho_trang(ins: InspectedFile) -> list[tuple[int, int]]:
    """Kích thước (pt, làm tròn 10pt, không phân biệt xoay) từng trang — so 'cùng khổ' khi scan lệch vài pt."""
    doc = open_pdf_document(Path(ins.path), ins.recovered_page_xrefs)
    try:
        return [tuple(sorted((round(p.rect.width / 10), round(p.rect.height / 10)))) for p in doc]
    finally:
        doc.close()


def la_so_a4(so: str) -> bool:
    """Số phát hành GCN mẫu A4 mới: AA + 8 chữ số (1 tờ A4, mặt trước + mặt sau)."""
    return bool(re.fullmatch(r"AA\d{8}", re.sub(r"[^0-9A-Z]", "", str(so or "").upper())))


def bo_sung_mat(trang: list[int], gan: dict[int, dict[str, Any]], kho: list[tuple[int, int]],
                so: str = "") -> tuple[list[int], str]:
    """
    GCN luôn có mặt trước + mặt sau, thường scan cùng khổ. Bản AI chỉ thấy 1 mặt → ghép trang liền kề (sau, rồi
    trước) cùng khổ, chưa gán cho GCN nào — CHỈ khi khổ đó ít gặp trong file (vd A3 trong file toàn A4); file
    toàn một khổ thì không đoán (dễ ghép nhầm đơn / hợp đồng), chỉ trả ghi chú để kiểm tra.
    Riêng GCN mẫu A4 (số AA + 8 chữ số, `so` hoặc số AI đọc trên trang) luôn đủ 2 mặt → ghép dù file toàn A4.
    → (trang của bản, ghi chú).
    """
    if len(trang) != 1:
        return trang, ""
    i = trang[0]
    a4 = la_so_a4(so or gan[i]["serial_seen"])
    if a4 or 2 * sum(k == kho[i] for k in kho) <= len(kho):
        for j in (i + 1, i - 1):
            if 0 <= j < len(kho) and kho[j] == kho[i] and not gan[j]["gcn"] and not gan[j].get("kem_theo"):
                return sorted([i, j]), (f"GCN mẫu A4 (AA + 8 số): đã ghép trang {j + 1} làm mặt còn lại" if a4 else
                                        f"AI chỉ thấy 1 mặt; đã ghép trang {j + 1} cùng khổ làm mặt còn lại")
    return trang, "Bản GCN chỉ có 1 mặt (thiếu mặt trước/sau?) — kiểm tra lại"


def tach_ban(gan: dict[int, dict[str, str]], so: str) -> list[list[int]]:
    """
    Các trang của GCN `so` trong 1 file → từng bản (1 file có thể chứa 2 bản: trước / sau đính chính).
    Bản mới bắt đầu ở ảnh AI đánh dấu starts_copy (scan trải 2 mặt có thể lặp số phát hành ở cả 2 ảnh
    nên không tách theo số).
    """
    ban: list[list[int]] = []
    for i in sorted(gan):
        if gan[i]["gcn"] != so:
            continue
        if not ban or _mo_ban_moi(gan, i, ban[-1]):
            ban.append([i])
        else:
            ban[-1].append(i)
    return ban


def _mo_ban_moi(gan: dict[int, dict[str, Any]], i: int, ban_dang_gom: list[int]) -> bool:
    if "bat_dau" in gan[i]:
        return bool(gan[i]["bat_dau"])
    # Kết quả cũ chưa có bat_dau: trang có số khi bản đang gom đã có trang có số.
    return bool(gan[i]["serial_seen"]) and any(gan[j]["serial_seen"] for j in ban_dang_gom)


# ============================ ĐỐI CHIẾU CHỦ SỬ DỤNG ============================

def chu_cua_don(detail: dict[str, Any]) -> list[str]:
    """Họ tên / tên tổ chức trong ChuSoHuu của đơn (cá nhân, vợ chồng, hộ gia đình, tổ chức...)."""
    ten: list[str] = []

    def walk(o: Any) -> None:
        if isinstance(o, dict):
            for k in ("hoTen", "tenToChuc"):
                v = str(o.get(k) or "").strip()
                if v and v not in ten:
                    ten.append(v)
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)

    walk(detail.get("ChuSoHuu") or {})
    return ten


def _chuan_ten(s: Any) -> str:
    s = core.remove_accents(str(s or "")).lower().replace("đ", "d").replace("ð", "d")
    return " " + " ".join(re.findall(r"[a-z0-9]+", s)) + " "


def so_chu_khop(chu_ai: str, ds_chu: list[str]) -> int:
    """Số người trong đơn có tên (nguyên cụm từ) xuất hiện trong tên chủ AI đọc được trên GCN."""
    ai = _chuan_ten(chu_ai)
    return sum(1 for t in ds_chu if _chuan_ten(t).strip() and _chuan_ten(t) in ai)


def doc_chu(ban: dict[str, Any], analyzer: GeminiAnalyzer, model: str, cache: dict, usage: Usage) -> None:
    """Gemini đọc chủ hiện tại của 1 bản GCN (trang 1, có đính chính / biến động thì trang 4) → ban['chu']."""
    ins: InspectedFile = ban["ins"]
    khoa = f"owner4:{ins.content_hash}:{','.join(str(i + 1) for i in ban['trang'])}:{model}"
    cached = cache.get(khoa)
    if not (isinstance(cached, dict) and isinstance(cached.get("chu"), str)):
        anh = render_trang(ins, ban["trang"])
        info, dung = analyzer.extract_certificate(anh)
        usage.add(dung)
        cached = {"chu": info.owner_name, "ly_do": info.reason, "so": info.certificate_serial,
                  "so_vao_so": info.registration_number, "ngay_cap": info.issue_date, "cccd": info.owner_id,
                  "dia_chi": info.owner_address}
        cache[khoa] = cached
    ban["chu"], ban["ly_do_chu"] = cached["chu"], str(cached.get("ly_do") or "")
    ban["ai"] = {"so_vao_so_ai": cached.get("so_vao_so", ""), "ngay_cap_ai": cached.get("ngay_cap", ""),
                 "chu_ai": cached["chu"], "cccd_ai": cached.get("cccd", ""), "dia_chi_ai": cached.get("dia_chi", "")}
    ban["so"] = ban.get("so") or str(cached.get("so") or "")


def _nhom_khong_khop(gan: dict[int, dict[str, str]]) -> list[list[int]]:
    """Trang 'KHAC' liền nhau → 1 nhóm (1 bản GCN); trang có số phát hành khi nhóm đang gom đã có trang có số
    (GCN khác, hoặc bản khác của cùng GCN) → mở nhóm mới."""
    nhom: list[list[int]] = []
    for i in sorted(gan):
        if gan[i]["gcn"] != ASSIGN_OTHER:
            continue
        if nhom and nhom[-1][-1] == i - 1 and not _mo_ban_moi(gan, i, nhom[-1]):
            nhom[-1].append(i)
        else:
            nhom.append([i])
    return nhom


# ============================ 1 ĐƠN ============================

def _dong(id_don: int, gcn: dict | None = None, path: str = "", trang: list[int] | None = None,
          gan: dict[int, dict[str, str]] | None = None, **kw) -> dict[str, Any]:
    gcn = gcn or {}
    so_doc = list(dict.fromkeys(gan[i]["serial_seen"] for i in trang or [] if gan and gan[i]["serial_seen"]))
    return {
        "madon": id_don,
        "giaychungnhanid": gcn.get("giayChungNhanId", ""),
        "so_gcn": gcn.get("soPhatHanh", ""),
        "so_vao_so": gcn.get("soVaoSo", ""),
        "file_nguon": os.path.basename(path),
        "trang": _trang_text(trang) if trang else "",
        "so_ai_doc": ", ".join(so_doc),
        **kw,
    }


def danh_gia_ban(ban: dict[str, Any], ds_chu: list[str], analyzer: GeminiAnalyzer, model: str, cache: dict,
                 usage: Usage) -> None:
    """Đọc thông tin GCN (chủ hiện tại, CCCD, địa chỉ, ngày cấp, số vào sổ) và đếm số người trùng chủ trong
    đơn → ban['chu'/'ly_do_chu'/'khop'/'ai'/'so']."""
    ban["chu"], ban["ly_do_chu"], ban["khop"], ban["ai"] = "", "", 0, {}
    ban["so"] = next((ban["gan"][i]["serial_seen"] for i in ban["trang"] if ban["gan"][i]["serial_seen"]), "")
    try:
        doc_chu(ban, analyzer, model, cache, usage)
    except Exception as exc:
        ban["ly_do_chu"] = f"Lỗi AI đọc chủ: {core.rut_gon_text(exc, 200)}"
    ban["khop"] = so_chu_khop(ban["chu"], ds_chu)


NAM_KEM_TRICH_LUC = 2003  # GCN cấp trước năm này: trích lục / sơ đồ đính kèm đi cùng file GCN


def nam_cap(ban: dict[str, Any], gcn: dict[str, Any]) -> int | None:
    """Năm cấp GCN: ngày cấp AI đọc trên GCN, không có thì ngày vào sổ trên MPLIS."""
    for gia_tri in ((ban.get("ai") or {}).get("ngay_cap_ai"), gcn.get("ngayVaoSo")):
        text = str(gia_tri or "")
        m = re.search(r"/Date\((-?\d+)", text)
        if m:
            return time.gmtime(int(m.group(1)) / 1000).tm_year
        m = re.search(r"\b(19[5-9]\d|20\d\d)\b", text)
        if m:
            return int(m.group(1))
    return None


def kem_trich_luc(ban: dict[str, Any], so: str, gcn: dict[str, Any]) -> tuple[list[int], str]:
    """
    Trang AI đánh dấu 'trích lục / sơ đồ đính kèm GCN số <so>' trong cùng file với bản đã chọn → gộp vào file
    GCN nếu GCN cấp trước 2003. → (trang ghi ra file, ghi chú).
    """
    gan = ban["gan"]
    tl = [i for i in sorted(gan) if gan[i].get("kem_theo") == so and i not in ban["trang"]]
    if not tl:
        return ban["trang"], ""
    nam = nam_cap(ban, gcn)
    if nam is None:
        return ban["trang"], (f"Có trích lục kèm GCN ở trang {_trang_text(tl)} nhưng không rõ năm cấp — "
                              "chưa gộp, kiểm tra tay")
    if nam >= NAM_KEM_TRICH_LUC:
        return ban["trang"], ""
    return sorted(ban["trang"] + tl), f"Kèm trích lục trang {_trang_text(tl)} (GCN cấp {nam}, trước 2003)"


def chon_ban(id_don: int, gcn: dict[str, Any], so: str, ds: list[dict[str, Any]], ds_chu: list[str],
             out_dir: str) -> list[dict[str, Any]]:
    """
    Nhiều bản cùng 1 GCN (đã qua danh_gia_ban) → chọn bản có chủ hiện tại (AI đọc trang 1, đã đính chính /
    biến động thì trang 4) trùng chủ sử dụng trong đơn; trùng nhiều người hơn → nhiều trang hơn → bản đầu tiên.
    Có bản đúng chủ → GCN/<madon>/; không bản nào đúng chủ → GCN_can_kiem_tra/<madon>/ (không đẩy nhầm).
    Đơn không có chủ trên MPLIS → không đối chiếu được, lấy bản nhiều trang nhất.
    """
    dung_chu = [b for b in ds if b["khop"]]
    nguon = dung_chu or ds
    chon = max(nguon, key=lambda b: (b["khop"], len(b["trang"])))  # max giữ bản đầu tiên khi bằng nhau
    if not ds_chu:
        trang_thai, ghi_chu = "Đã tách", ["Đơn không có chủ sử dụng trên MPLIS — chưa đối chiếu chủ"]
    elif dung_chu:
        trang_thai, ghi_chu = "Đã tách", []
    else:
        trang_thai = "Sai chủ" if any(b["chu"] for b in ds) else "Chưa rõ chủ"
        ghi_chu = [f"Không bản nào có chủ trùng chủ trong đơn ({', '.join(ds_chu)}) — kiểm tra tay"]

    ten_file = f"{thd.ten_file_an_toan(so)}-GCN.pdf"
    dich = os.path.join(out_dir, THU_MUC_GCN if trang_thai == "Đã tách" else THU_MUC_CAN_KIEM_TRA,
                        str(id_don), ten_file)
    if trang_thai != "Đã tách":
        # Lần chạy trước có thể đã để file này trong GCN/ — bỏ đi để tool đẩy không đẩy nhầm.
        cu = os.path.join(out_dir, THU_MUC_GCN, str(id_don), ten_file)
        if os.path.isfile(cu):
            os.remove(cu)
    ins: InspectedFile = chon["ins"]
    trang_ghi, ghi_chu_tl = kem_trich_luc(chon, so, gcn)
    co_so = any(chon["gan"][i]["serial_seen"] for i in chon["trang"])
    vao_so_ai, vao_so_mplis = chon["ai"].get("so_vao_so_ai", ""), str(gcn.get("soVaoSo") or "")
    ghi_chu += [f"Có {len(ds) - 1} bản khác (dòng Bản trùng)" if len(ds) > 1 else "",
                "" if co_so else "AI không đọc được số phát hành trên các trang này — nên kiểm tra lại",
                f"Số vào sổ AI đọc ({vao_so_ai}) khác MPLIS ({vao_so_mplis})"
                if vao_so_ai and vao_so_mplis and dm.chuan_hoa(vao_so_ai) != dm.chuan_hoa(vao_so_mplis) else "",
                chon.get("ghi_chu_mat", ""),
                ghi_chu_tl,
                chon["ly_do_chu"]]
    try:
        _save_selected_pages(Path(chon["path"]), trang_ghi, Path(dich), ins.recovered_page_xrefs)
    except Exception as exc:
        log(f"   ❌ GCN {so}: ghi PDF lỗi {exc}")
        trang_thai, dich, ghi_chu = "Lỗi", "", [f"Ghi PDF lỗi: {exc}"]

    nhan = f"{os.path.basename(chon['path'])} tr.{_trang_text(trang_ghi)}"
    chu_text = f" | chủ: {chon['chu'] or '?'}" if ds_chu else ""
    if trang_thai == "Đã tách":
        log(f"   ✅ GCN {so}: {nhan}{chu_text} → {os.path.basename(dich)}")
    elif dich:
        log(f"   ⚠ GCN {so}: {trang_thai} — {nhan}{chu_text} → {THU_MUC_CAN_KIEM_TRA}")
    dong = [_dong(id_don, gcn, chon["path"], trang_ghi, chon["gan"], chu_don=", ".join(ds_chu),
                  **chon["ai"], file_gcn=dich, trang_thai=trang_thai, ghi_chu="; ".join(filter(None, ghi_chu)))]
    dong += [_dong(id_don, gcn, b["path"], b["trang"], b["gan"], chu_don=", ".join(ds_chu), **b["ai"],
                   trang_thai="Bản trùng",
                   ghi_chu="; ".join(filter(None, [f"Đã chọn {nhan}",
                                                   "chủ không trùng đơn" if ds_chu and not b["khop"] else "",
                                                   b["ly_do_chu"]])))
             for b in ds if b is not chon]
    return dong


def _ten_khong_trung(thu_muc: str, ten: str, da_dung: set[str]) -> str:
    goc, n = ten, 2
    while ten.casefold() in da_dung:
        ten = f"{goc}-{n}"
        n += 1
    da_dung.add(ten.casefold())
    return os.path.join(thu_muc, f"{ten}.pdf")


# ============================ GHI NHẬN THÀNH PHẦN "ĐƠN" (MỌI ĐƠN) ============================

LOAI_DON_DANG_KY = 2


def la_file_don(node: dict[str, Any]) -> bool:
    """
    Tên (mô tả) có chuỗi 'DON' ở bất kỳ đâu, có dấu hay không: 'Đơn đăng ký', '01_DON_DK', 'DONDANGKY',
    cả 'Ðơn' viết bằng chữ Ð (U+00D0, nhìn y hệt Đ).
    """
    ten = " ".join(str(node.get(k) or "") for k in KHOA_TEN_FILE)
    return "don" in dm.chuan_hoa(ten)


def cap_nhat_loai_don(client, detail: dict[str, Any], id_don: int, chay_thu: bool) -> list[dict[str, Any]]:
    """
    CHỈ đơn chưa có GCN: thành phần HSQ có chuỗi 'DON' trong tên (mô tả) → loại Đơn đăng ký + ghi nhận
    hồ sơ gốc (laGiayToVeNguonGoc). Mỗi HSQ 1 request kèm ĐỦ thành phần (thành phần khác gửi lại nguyên
    trạng, files=null — không đổi file). Sau khi gửi đọc lại: số thành phần giảm → DỪNG TOÀN BỘ
    (day.DungKhanCap); thành phần chưa đổi loại → Lỗi. Tên có 'DON' mà đang ghi nhận nhầm Giấy chứng nhận
    (đơn chưa có GCN) cũng đổi thành Đơn đăng ký.
    Luôn có ít nhất 1 dòng Excel cho đơn.
    """
    rows: list[dict[str, Any]] = []

    def dong(hoso: dict, node: dict, trang_thai: str, ghi_chu: str) -> dict[str, Any]:
        return _dong(id_don, file_nguon=node.get("moTa") or node.get("tenGiayTo") or "", trang_thai=trang_thai,
                     ghi_chu=f"HSQ {hoso.get('hoSoQuetId')}: {ghi_chu}")

    ds_hsq = [h for h in detail.get("ListHoSoQuet") or [] if isinstance(h, dict)]
    if not ds_hsq:
        log("   ⚠ Đơn chưa có GCN và không có HSQ")
        return [_dong(id_don, trang_thai="Bỏ qua", ghi_chu="Đơn chưa có GCN, không có HSQ")]
    for hoso in ds_hsq:
        hsq_id = core.safe_int(hoso.get("hoSoQuetId"))
        nodes = day.cac_node(hoso)
        viec: list[int] = []
        log(f"   HSQ {hsq_id}: {len(nodes)} thành phần")
        for i, node in enumerate(nodes):
            loai = core.LOAI_HO_SO_QUET_OPTIONS.get(core.safe_int(node.get("loaiHoSoQuet"), -1),
                                                   node.get("loaiHoSoQuet") or "chưa có loại")
            nguon_goc = " + hồ sơ gốc" if node.get("laGiayToVeNguonGoc") else ""
            nhan = f"     {i + 1}. {node.get('moTa') or node.get('tenGiayTo') or '(không mô tả)'}  [{loai}{nguon_goc}]"
            if not la_file_don(node):
                log(f"{nhan} — không có 'DON', giữ")
                continue
            if (core.safe_int(node.get("loaiHoSoQuet"), -1) == LOAI_DON_DANG_KY
                    and node.get("laGiayToVeNguonGoc") is True and not node.get("laGiayChungNhan")):
                log(f"{nhan} — đã đúng sẵn, giữ")
                rows.append(dong(hoso, node, trang_thai="Giữ nguyên", ghi_chu="Đã là Đơn đăng ký + hồ sơ gốc"))
            else:
                log(f"{nhan} — CẦN ĐỔI → Đơn đăng ký + hồ sơ gốc")
                viec.append(i)
        if not viec:
            log(f"   HSQ {hsq_id}: không có thành phần nào cần đổi → không gửi")
            continue

        loai_cu = {i: core.LOAI_HO_SO_QUET_OPTIONS.get(core.safe_int(nodes[i].get("loaiHoSoQuet"), -1),
                                                      nodes[i].get("loaiHoSoQuet") or "(chưa có loại)")
                   for i in viec}
        if chay_thu:
            log(f"   HSQ {hsq_id}: CHẠY THỬ — sẽ đổi {len(viec)} thành phần, chưa gửi MPLIS")
            rows += [dong(hoso, nodes[i], trang_thai="Chạy thử",
                          ghi_chu=f"Sẽ đổi {loai_cu[i]} → Đơn đăng ký + hồ sơ gốc") for i in viec]
            continue

        hoso_gui = copy.deepcopy(hoso)
        nodes_gui = day.cac_node(hoso_gui)
        for i in viec:
            nodes_gui[i]["loaiHoSoQuet"] = LOAI_DON_DANG_KY
            nodes_gui[i]["laGiayToVeNguonGoc"] = True
            nodes_gui[i]["laGiayChungNhan"] = False
        log(f"   HSQ {hsq_id}: gửi MPLIS {len(nodes)} thành phần (đổi {len(viec)}, không gửi file)...")
        try:
            day.gui_ca_hsq(client, hoso_gui, {})  # {}: không thay file nào, chỉ metadata
        except core.HetPhien:
            raise
        except Exception as exc:
            log(f"   ❌ HSQ {hsq_id}: {core.rut_gon_text(exc, 300)}")
            rows += [dong(hoso, nodes[i], trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)) for i in viec]
            continue

        # ---- xác nhận
        log(f"   HSQ {hsq_id}: MPLIS báo OK — đọc lại đơn để kiểm tra...")
        detail_moi = client.lay_chi_tiet_don(id_don) or {}
        hoso_moi = next((h for h in detail_moi.get("ListHoSoQuet") or []
                         if isinstance(h, dict) and core.safe_int(h.get("hoSoQuetId")) == hsq_id), None)
        nodes_moi = day.cac_node(hoso_moi) if hoso_moi else []
        if len(nodes_moi) < len(nodes):
            rows += [dong(hoso, nodes[i], trang_thai="Lỗi", ghi_chu=f"MPLIS còn {len(nodes_moi)}/{len(nodes)} thành phần!")
                     for i in viec]
            raise day.DungKhanCap(f"HSQ {hsq_id} (đơn {id_don}) giảm từ {len(nodes)} còn {len(nodes_moi)} thành "
                                  "phần sau khi cập nhật → DỪNG TOÀN BỘ.")
        for i in viec:
            tp_id = core.safe_int(nodes[i].get("thanhPhanHoSoQuetId"))
            moi = next((n for n in nodes_moi if core.safe_int(n.get("thanhPhanHoSoQuetId")) == tp_id), None) or {}
            if core.safe_int(moi.get("loaiHoSoQuet"), -1) == LOAI_DON_DANG_KY and moi.get("laGiayToVeNguonGoc"):
                rows.append(dong(hoso, nodes[i], trang_thai="Đã cập nhật",
                                 ghi_chu=f"{loai_cu[i]} → Đơn đăng ký + hồ sơ gốc {time.strftime('%d/%m/%Y %H:%M:%S')}"))
                log(f"   ✅ {nodes[i].get('moTa')}: Đơn đăng ký + hồ sơ gốc")
            else:
                rows.append(dong(hoso, nodes[i], trang_thai="Lỗi",
                                 ghi_chu="Server báo OK nhưng loại chưa đổi (kiểm tra trên MPLIS)"))
                log(f"   ❌ {nodes[i].get('moTa')}: server báo OK nhưng loại chưa đổi")
        time.sleep(core.REQUEST_DELAY_SECONDS)
    if not rows:
        log("   ⚠ Không có thành phần HSQ nào có chuỗi 'DON' trong tên")
        rows.append(_dong(id_don, trang_thai="Bỏ qua",
                          ghi_chu="Đơn chưa có GCN; không có thành phần HSQ nào có chuỗi 'DON' trong tên"))
    return rows


def xu_ly_don(client, downloader: thd.Downloader, analyzer: GeminiAnalyzer, id_don: int, out_dir: str,
              cache: dict, usage: Usage, model: str = DEFAULT_MODEL, dung_file_da_tai: bool = False,
              nen_dung: Callable[[], bool] = lambda: False, loc_truoc: bool = True,
              quet_het: bool = False, chay_thu: bool = True, tt: TaiTruoc | None = None) -> list[dict[str, Any]]:
    """
    Trả các dòng Excel của đơn. Hết phiên → ném core.HetPhien. tt: bộ tải trước dùng chung cả loạt (None → tạo
    riêng cho đơn này).
    Đơn chưa có GCN: chỉ cap_nhat_loai_don — thành phần có chuỗi 'DON' → Đơn đăng ký + hồ sơ gốc (chay_thu:
    chỉ ghi kế hoạch, không ghi MPLIS); không tải, không gọi AI.
    Đơn có GCN: chỉ tách GCN (không ghi gì lên MPLIS) — file HSQ xử lý theo uu_tien_file, tải tới đâu đọc tới
    đó; mọi GCN của đơn đã có bản đúng chủ (đơn không có chủ: đã có bản) → dừng, không tải / không gọi AI các
    file còn lại (trừ khi quet_het).
    """
    if tt is None:
        tt = TaiTruoc(client, downloader, out_dir, dung_file_da_tai)
        try:
            return xu_ly_don(client, downloader, analyzer, id_don, out_dir, cache, usage, model, dung_file_da_tai,
                             nen_dung, loc_truoc, quet_het, chay_thu, tt)
        finally:
            tt.dong()
    detail = tt.chi_tiet(id_don)
    if detail is None:
        log(f"❌ Đơn {id_don}: không lấy được chi tiết")
        return [_dong(id_don, trang_thai="Lỗi", ghi_chu="Không lấy được chi tiết đơn")]

    ds_gcn = [g for g in bs.gcn_hien_hanh(detail) if str(g.get("soPhatHanh") or "").strip()]
    log(f"Đơn {id_don}: {len(ds_gcn)} GCN hiện hành"
        + (": " + ", ".join(f"{g.get('soPhatHanh')} (vào sổ {g.get('soVaoSo') or '?'})" for g in ds_gcn)
           if ds_gcn else ""))
    if not ds_gcn:
        log("   Đơn chưa có GCN → không phân tích GCN; cập nhật loại HSQ: thành phần có chuỗi 'DON' → Đơn đăng ký"
            " + hồ sơ gốc" + (" (CHẠY THỬ)" if chay_thu else ""))
        with tt.khoa:  # luồng tải nền không gửi gì trong lúc ghi MPLIS
            return cap_nhat_loai_don(client, detail, id_don, chay_thu)
    try:
        return _tach_gcn_don(tt, analyzer, detail, ds_gcn, id_don, out_dir, cache, usage, model,
                             nen_dung, loc_truoc, quet_het)
    finally:
        tt.huy_don(id_don)


def _tach_gcn_don(tt: TaiTruoc, analyzer: GeminiAnalyzer, detail: dict[str, Any],
                  ds_gcn: list[dict[str, Any]], id_don: int, out_dir: str, cache: dict, usage: Usage, model: str,
                  nen_dung: Callable[[], bool], loc_truoc: bool, quet_het: bool) -> list[dict[str, Any]]:
    """Bước tách GCN của xu_ly_don cho đơn có GCN hiện hành."""
    files = thd.liet_ke_file_hsq(detail)
    if not files:
        log("   ⚠ Đơn không có HSQ")
        return [_dong(id_don, g, trang_thai="Không thấy", ghi_chu="Đơn không có HSQ") for g in ds_gcn]

    ds_chu = chu_cua_don(detail)
    log(f"   Chủ sử dụng trong đơn: {', '.join(ds_chu) or '(không có)'}")
    ds_so = [(str(g["soPhatHanh"]).strip(), str(g.get("soVaoSo") or "").strip()) for g in ds_gcn]
    # theo_so[số phát hành] = các bản GCN đó: {path, trang (0-based), gan, ins, chu, ly_do_chu, khop}
    theo_so: dict[str, list[dict[str, Any]]] = {}
    khong_khop: list[tuple[str, list[int], dict, InspectedFile]] = []
    loi: list[str] = []

    def du() -> bool:
        return all(any(b["khop"] or not ds_chu for b in theo_so.get(so, [])) for so, _ in ds_so)

    thu_tu = sorted(enumerate(files, 1), key=lambda x: uu_tien_file(x[1]))  # sorted ổn định: giữ thứ tự HSQ
    log(f"   {len(files)} file HSQ — thứ tự xử lý: "
        + ", ".join(f"{(f.get('moTa') or f.get('tenGiayTo') or vt)}" for vt, f in thu_tu[:8])
        + (" ..." if len(thu_tu) > 8 else ""))
    for n, (vi_tri, f) in enumerate(thu_tu):
        if nen_dung():
            break
        if not quet_het and du():
            con_lai = len(thu_tu) - n
            log(f"   ⏭ Đã đủ {len(ds_so)} GCN" + (" đúng chủ" if ds_chu else "")
                + f" — bỏ qua {con_lai} file HSQ còn lại (không tải, không gọi AI)")
            loi.append(f"Bỏ qua {con_lai} file HSQ sau khi đã đủ GCN")
            break
        mo_ta = f.get("moTa") or f.get("tenGiayTo") or f"file_{vi_tri}"
        tt.dat_truoc_file(id_don, thu_tu[n:n + 2], n)  # file kế tiếp tải trong lúc AI đọc file này
        try:
            path = _chi_lay_pdf(tt.file(id_don, vi_tri, f))
        except core.HetPhien:
            raise
        except Exception as exc:
            loi.append(f"Tải '{mo_ta}' lỗi: {core.rut_gon_text(exc, 200)}")
            log(f"   ❌ Tải {mo_ta}: {core.rut_gon_text(exc, 200)}")
            continue
        if path is None:
            continue
        ten = os.path.basename(path)
        try:
            gan, ins = gan_trang(path, ds_so, analyzer, model, cache, usage, loc_truoc)
        except Exception as exc:
            loi.append(f"AI đọc '{ten}' lỗi: {core.rut_gon_text(exc, 200)}")
            log(f"   ❌ {ten}: {core.rut_gon_text(exc, 300)}")
            continue
        tom_tat = []
        kho = kho_trang(ins)
        for so, _ in ds_so:
            for trang in tach_ban(gan, so):
                trang, ghi_chu_mat = bo_sung_mat(trang, gan, kho, so)
                ban = {"path": path, "trang": trang, "gan": gan, "ins": ins, "ghi_chu_mat": ghi_chu_mat}
                danh_gia_ban(ban, ds_chu, analyzer, model, cache, usage)
                theo_so.setdefault(so, []).append(ban)
                tom_tat.append(f"{so}: tr.{_trang_text(trang)}"
                               + (f" (chủ: {ban['chu'] or '?'}{' ✓' if ban['khop'] else ' ✗'})" if ds_chu else ""))
        for trang in _nhom_khong_khop(gan):
            if len(trang) == 1 and la_so_a4(gan[trang[0]]["serial_seen"]):
                trang, _ = bo_sung_mat(trang, gan, kho)
            khong_khop.append((path, trang, gan, ins))
            tom_tat.append(f"GCN khác: tr.{_trang_text(trang)}")
        log(f"   {ten} ({ins.page_count} trang): " + ("; ".join(tom_tat) or "không có trang GCN"))
    ghi_chu_loi = "; ".join(loi)

    dong: list[dict[str, Any]] = []
    for gcn, (so, _) in zip(ds_gcn, ds_so):
        ds = theo_so.get(so, [])
        if not ds:
            ly_do = "AI không tìm thấy trang GCN này trong HSQ"
            log(f"   ⚠ GCN {so}: {ly_do}")
            dong.append(_dong(id_don, gcn, trang_thai="Không thấy",
                              ghi_chu="; ".join(filter(None, [ly_do, ghi_chu_loi]))))
            continue
        dong += chon_ban(id_don, gcn, so, ds, ds_chu, out_dir)

    # Trang GCN không thuộc đơn → tách riêng để kiểm tra tay, không để lẫn vào thư mục GCN.
    thu_muc_kk = os.path.join(out_dir, THU_MUC_KHONG_KHOP, str(id_don))
    da_dung: set[str] = set()
    for path, trang, gan, ins in khong_khop:
        xrefs = ins.recovered_page_xrefs
        so_doc = next((gan[i]["serial_seen"] for i in trang if gan[i]["serial_seen"]), "")
        ten = f"{so_doc}-GCN" if so_doc else f"{Path(path).stem}_tr{_trang_text(trang)}-GCN"
        dich = _ten_khong_trung(thu_muc_kk, thd.ten_file_an_toan(ten), da_dung)
        try:
            _save_selected_pages(Path(path), trang, Path(dich), xrefs)
        except Exception as exc:
            log(f"   ❌ Ghi trang GCN không khớp lỗi: {exc}")
            dich = ""
        ly_do = "Trang GCN không thuộc GCN hiện hành nào của đơn (GCN cũ / thửa khác?) — " + \
                "; ".join(dict.fromkeys(gan[i]["reason"] for i in trang if gan[i]["reason"]))
        dong.append(_dong(id_don, None, path, trang, gan, file_gcn=dich, trang_thai="Không khớp",
                          ghi_chu=core.rut_gon_text(ly_do, 400)))
    return dong


# ============================ EXCEL ============================

def ghi_excel(path: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "KetQua"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_tt = EXCEL_HEADERS.index("trang_thai") + 1
    for stt, row in enumerate(rows, 1):
        ws.append([stt if h == "stt" else row.get(h, "") for h in EXCEL_HEADERS])
        mau = MAU_TRANG_THAI.get(str(row.get("trang_thai") or ""))
        if mau:
            ws.cell(row=ws.max_row, column=col_tt).fill = PatternFill("solid", fgColor=mau)
    for i, h in enumerate(EXCEL_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = EXCEL_WIDTHS.get(h, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}1"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


# ============================ CHẠY NHIỀU ĐƠN ============================

def chay(client, ids: list[int], out_dir: str, api_key: str, model: str = DEFAULT_MODEL, rpm: float = 30,
         dung_file_da_tai: bool = False, nen_dung: Callable[[], bool] = lambda: False,
         bao_tien_do: Callable[[int, int, str], None] = lambda done, total, label: None,
         analyzer: GeminiAnalyzer | None = None, loc_truoc: bool = True, quet_het: bool = False,
         chay_thu: bool = True) -> str:
    """Xử lý danh sách đơn → đường dẫn Excel kết quả. Hết phiên → dừng cả loạt (đã ghi Excel)."""
    if analyzer is None:
        if not api_key:
            raise RuntimeError("Chưa có Gemini API key (GEMINI_API_KEY)")
        analyzer = GeminiAnalyzer(api_key, model, rpm=rpm)
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    token = client.session.headers.get("__requestverificationtoken", "")
    downloader = thd.Downloader(None, client.session, token, thd.doc_mau_da_luu(), thd.MAU_MAC_DINH)
    cache_path = Path(out_dir) / FILE_CACHE
    cache = load_cache(cache_path)
    usage = Usage()
    excel = os.path.join(out_dir, f"ket_qua_tach_gcn_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
    rows: list[dict[str, Any]] = []

    def luu() -> None:
        try:
            ghi_excel(excel, rows)
        except PermissionError:
            log("⚠ Không ghi được Excel kết quả (file đang mở?)")
        save_cache(cache_path, cache)

    log(f"Tách GCN cho {len(ids)} đơn | model {model} → {out_dir}")
    tt = TaiTruoc(client, downloader, out_dir, dung_file_da_tai)
    try:
        _chay_cac_don(client, downloader, analyzer, ids, out_dir, cache, usage, model, dung_file_da_tai, nen_dung,
                      bao_tien_do, loc_truoc, quet_het, chay_thu, tt, rows, luu)
    finally:
        tt.dong()
    bao_tien_do(len(ids), len(ids), "Xong")
    _tong_ket(rows, out_dir, usage, excel)
    return excel


def _chay_cac_don(client, downloader, analyzer, ids, out_dir, cache, usage, model, dung_file_da_tai, nen_dung,
                  bao_tien_do, loc_truoc, quet_het, chay_thu, tt: TaiTruoc, rows: list, luu: Callable[[], None]) -> None:
    tt.dang_ky(ids)
    for n, id_don in enumerate(ids, 1):
        if nen_dung():
            log("⏹ Đã dừng theo yêu cầu.")
            break
        bao_tien_do(n - 1, len(ids), f"Đơn {id_don} ({n}/{len(ids)})")
        log(f"--- [{n}/{len(ids)}] Đơn {id_don} ---")
        for id_sau in ids[n:n + SO_DON_TAI_TRUOC]:
            tt.chuan_bi_don(id_sau)  # đơn kế tiếp tải nền trong lúc xử lý đơn này
        try:
            rows += xu_ly_don(client, downloader, analyzer, id_don, out_dir, cache, usage, model,
                              dung_file_da_tai, nen_dung, loc_truoc, quet_het, chay_thu, tt)
        except core.HetPhien as exc:
            log(f"❌ {exc}")
            log("⛔ HẾT PHIÊN — đã dừng. Đăng nhập lại rồi chạy tiếp các đơn chưa xử lý.")
            rows += (exc.ket_qua or {}).get("dong", [])
            rows.append(_dong(id_don, trang_thai="Lỗi", ghi_chu=f"HẾT PHIÊN — ĐÃ DỪNG: {exc}"))
            luu()
            break
        except day.DungKhanCap as exc:
            log(f"❌ {exc}")
            rows.append(_dong(id_don, trang_thai="Lỗi", ghi_chu=f"DỪNG KHẨN CẤP: {exc}"))
            luu()
            raise
        except Exception as exc:
            log(f"❌ Đơn {id_don}: {core.rut_gon_text(exc, 300)}")
            rows.append(_dong(id_don, trang_thai="Lỗi", ghi_chu=core.rut_gon_text(exc, 300)))
        luu()


def _tong_ket(rows: list[dict[str, Any]], out_dir: str, usage: Usage, excel: str) -> None:
    def dem(tt: str) -> int:
        return sum(r.get("trang_thai") == tt for r in rows)

    log("=" * 60)
    log(f"✅ Đã tách {dem('Đã tách')} GCN → {os.path.join(out_dir, THU_MUC_GCN)}")
    if dem("Chạy thử"):
        log(f"ℹ Cập nhật loại HSQ: {dem('Chạy thử')} thành phần 'DON' SẼ ghi nhận Đơn đăng ký + hồ sơ gốc "
            "(chạy thử — chưa ghi MPLIS)")
    if dem("Đã cập nhật"):
        log(f"✅ Cập nhật loại HSQ: đã ghi nhận {dem('Đã cập nhật')} thành phần Đơn đăng ký + hồ sơ gốc")
    if dem("Không thấy"):
        log(f"⚠ {dem('Không thấy')} GCN không tìm thấy trong HSQ")
    if dem("Sai chủ") + dem("Chưa rõ chủ"):
        log(f"⚠ {dem('Sai chủ') + dem('Chưa rõ chủ')} GCN không có bản nào đúng chủ trong đơn → "
            f"{os.path.join(out_dir, THU_MUC_CAN_KIEM_TRA)}")
    if dem("Chưa rõ số"):
        log(f"⚠ {dem('Chưa rõ số')} GCN đúng chủ nhưng AI không đọc được số → "
            f"{os.path.join(out_dir, THU_MUC_CAN_KIEM_TRA)}")
    if dem("Không khớp"):
        log(f"⚠ {dem('Không khớp')} nhóm trang GCN không thuộc đơn → {os.path.join(out_dir, THU_MUC_KHONG_KHOP)}")
    log(f"Chi phí Gemini ước tính: {usage.estimated_cost_usd:.4f} USD "
        f"({usage.input_tokens} token vào, {usage.output_tokens} token ra)")
    log(f"Excel kết quả: {excel}")


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="Tải HSQ của đơn, dùng AI tách trang GCN ra file <số GCN>-GCN.pdf")
    ap.add_argument("ids", nargs="*", help="tinhHinhDangKyId (hoặc chuỗi chứa mã đơn)")
    ap.add_argument("--excel", help="File Excel có cột madon (tinhHinhDangKyId)")
    ap.add_argument("--out", default="tach_gcn_theo_don", help="Thư mục kết quả (mặc định: tach_gcn_theo_don)")
    ap.add_argument("--dung-file-da-tai", action="store_true",
                    help="Đơn đã có PDF trong <out>/HSQ_goc/<madon> thì không tải lại")
    ap.add_argument("--api-key", help="Gemini API key (mặc định: biến môi trường GEMINI_API_KEY)")
    ap.add_argument("--model", default=os.environ.get("GEMINI_MODEL", DEFAULT_MODEL), help="Model Gemini")
    ap.add_argument("--rpm", type=float, default=30, help="Giới hạn request Gemini/phút")
    ap.add_argument("--quet-het", action="store_true",
                    help="Xử lý mọi file HSQ (mặc định: đủ GCN đúng chủ thì dừng, bỏ qua file còn lại)")
    ap.add_argument("--khong-loc-truoc", action="store_true",
                    help="Gửi ảnh nét mọi trang (mặc định: lọc trang GCN bằng ảnh nhỏ trước, rẻ hơn nhiều)")
    ap.add_argument("--ghi-mplis", action="store_true",
                    help="GHI THẬT loại Đơn đăng ký + hồ sơ gốc cho thành phần có chuỗi 'DON' (đơn chưa có GCN) "
                         "(mặc định chỉ chạy thử, ghi kế hoạch vào Excel)")
    ap.add_argument("--token", help="__RequestVerificationToken copy từ F12")
    ap.add_argument("--cookie", help="Chuỗi Cookie copy từ F12 (dùng cùng --token)")
    ap.add_argument("--chrome", action="store_true", help="Đăng nhập bằng Chrome, không hỏi token/cookie")
    ap.add_argument("--user", help="Username đăng nhập MPLIS")
    args = ap.parse_args()

    raw_ids = list(args.ids) + (thd.doc_id_tu_excel(args.excel) if args.excel else [])
    ids = thd.tach_danh_sach_id(raw_ids)
    if not ids:
        ap.error("Chưa có ID đơn nào (truyền ids hoặc --excel).")
    api_key = ((args.api_key or os.environ.get("GEMINI_API_KEY", "")).strip()
               or getpass.getpass("Gemini API key: ").strip())

    token, cookie = (args.token or "").strip(), (args.cookie or "").strip()
    if not args.chrome and not (token and cookie):
        print("Chọn cách đăng nhập:\n  1. Dán token + cookie (F12 > Network > Request Headers)\n  2. Mở Chrome đăng nhập")
        if (input("Chọn [1]: ").strip() or "1") == "1":
            token = token or input("__RequestVerificationToken: ").strip()
            cookie = cookie or input("Cookie: ").strip()

    client = core.MplisClient(log)
    try:
        if token and cookie:
            client.build_session_from_manual(token, cookie)
        else:
            client.open_browser_and_fill_login(args.user or input("Username: ").strip(), getpass.getpass("Password: "))
            input(">>> Đăng nhập xong trên Chrome (OTP/captcha nếu có) rồi nhấn Enter...")
            client.build_session_from_browser()
        chay(client, ids, args.out, api_key, args.model, args.rpm, args.dung_file_da_tai,
             loc_truoc=not args.khong_loc_truoc, quet_het=args.quet_het, chay_thu=not args.ghi_mplis)
    finally:
        client.close_browser()


if __name__ == "__main__":
    main()
